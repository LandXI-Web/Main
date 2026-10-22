"""엑셀 한 벌(IMPL 0-6) 검사 — ./dl/*.xlsx 를 openpyxl 로 열어 제목 줄 · 머리 줄 · 고정 · 자동 필터 · 숫자 서식 · 합계(SUBTOTAL) · 열 이름을 적는다.
python -I xlsx-check.py → ./xlsx-check.json"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
OLD_COLS = {"AI가 본 것", "믿을 만한 정도", "회차", "서비스 버전"}
out = []
for f in sorted((HERE / "dl").glob("*.xlsx")):
    wb = load_workbook(f)
    ws = wb.worksheets[0]
    head_row = 4 if ws["A1"].value and ws["A4"].value else 1
    head = [c.value for c in ws[head_row]]
    last = ws.max_row
    tot = [c.value for c in ws[last]]
    fmts = {}
    for c in ws[min(head_row + 1, last)]:
        if isinstance(c.value, (int, float)) or c.is_date:
            fmts[str(head[c.column - 1])] = c.number_format
    sheets = []
    for s in wb.worksheets:
        sheets.append({"name": s.title, "freeze": s.freeze_panes, "filter": s.auto_filter.ref, "rows": s.max_row,
                       "subtotal": any(isinstance(c.value, str) and c.value.startswith("=SUBTOTAL(109") for c in s[s.max_row])})
    out.append({"file": f.name, "title": ws["A1"].value, "sub": ws["A2"].value, "head_row": head_row, "head": head,
                "freeze": ws.freeze_panes, "filter": ws.auto_filter.ref, "num_formats": fmts, "total_row": tot[:6],
                "old_cols": sorted(OLD_COLS & {str(x) for x in head}), "sheets": sheets})
(HERE / "xlsx-check.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
sys.stdout.reconfigure(encoding="utf-8")
for o in out:
    print(o["file"], "|", o["title"], "|", o["sub"], "| 고정", o["freeze"], "| 필터", o["filter"], "| 서식", o["num_formats"], "| 합계", o["total_row"][:4], "| 옛 열", o["old_cols"])
