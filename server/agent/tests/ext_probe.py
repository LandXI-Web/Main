"""시험용 확장 도구(C2 c2-core) — 명령 바 블록 3종(chart·file·image)을 실제 경로로 한 번에 내보낸다.

운영 게이트웨이에는 붙지 않는다. 환경 변수 LX_AGENT_EXT_EXTRA=agent.tests.ext_probe 로 띄운 게이트웨이에서만 불린다(스크린샷 증거 · 테스트).
ROUTE: LX 계정이 정확히 '명령 바 블록 확인' 이라고 물을 때만 직행.
  chart = 실제 survey_stats(by=emd) 봉투(지금 지도의 시군구) → 읍면동 막대
  file  = 같은 봉투로 만든 .docx(run 폴더) → 내려받기 버튼
  image = Land-XI 자산 영상 한 장(run 폴더로 복사) → 영상 조각 + 꼬리표
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from agent.tools import Out, ToolError

TRIGGER = re.compile(r"^\s*명령\s?바\s?블록\s?확인\s*$")
SPECS = {"blocks_probe": {"description": "명령 바 블록 시험(차트·파일·영상 조각)", "properties": {"region": {"type": "string"}}, "route_only": True}}
WHY = {"blocks_probe": "명령 바 블록 시험"}
SAY = {"blocks_probe": "블록 만들기"}
IMG = Path(__file__).resolve().parents[3] / "landxi" / "v3" / "kit" / "img" / "drone.webp"


def allowed(name: str, p) -> bool:
    return getattr(p, "realm", None) == "lx"


def ROUTE(msg: str, ctx):
    if TRIGGER.match(msg or "") and allowed("blocks_probe", ctx.principal):
        return {"tool": "blocks_probe", "args": {}}
    return None


async def blocks_probe(args: dict, ctx) -> Out:
    from agent import config
    from agent.tools import registry
    region = args.get("region") or (ctx.context or {}).get("region")
    st = await registry.HANDLERS["survey_stats"]({"by": "emd", **({"region": region} if region else {})}, ctx)
    out = Out(source="시험 도구(블록 3종)")
    rows = []
    for key, meaning, e in st.envelopes:
        if key.startswith("emd_"):
            out.env(key, meaning, e)
            rows.append({"label": meaning.split(" 의심")[0].strip(), "env": key, "value": e.get("value") if isinstance(e, dict) else None})
    if not rows:
        raise ToolError("not_found", "읍면동 집계가 없습니다", 404)
    out.blocks.append({"type": "chart", "kind": "bar", "title": "읍면동별 의심 건수", "rows": [{"label": r["label"], "env": r["env"]} for r in rows[:8]]})
    d = config.run_dir(ctx.run_id)
    import docx
    doc = docx.Document()
    doc.add_heading("읍면동별 의심 건수", level=1)
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "읍면동", "의심 건수"
    for r in rows[:8]:
        c = t.add_row().cells
        c[0].text, c[1].text = r["label"], f"{r['value']:,}" if isinstance(r["value"], (int, float)) else "—"
    doc.save(d / "blocks.docx")
    out.blocks.append({"type": "file", "label": "읍면동별_의심_건수.docx", "href": config.run_href(ctx.run_id, "blocks.docx")})
    if IMG.is_file():
        shutil.copyfile(IMG, d / "tile.webp")
        out.blocks.append({"type": "image", "src": config.run_href(ctx.run_id, "tile.webp"), "caption": "드론 영상 조각", "tag": "AI 의견 · 근거 아님"})
    out.answer = "읍면동별 의심 건수를 막대로 보이고, 표 파일과 영상 조각을 붙였습니다."
    out.data = {"막대": len(rows[:8]), "파일": 1, "영상": 1 if IMG.is_file() else 0}
    return out


HANDLERS = {"blocks_probe": blocks_probe}
