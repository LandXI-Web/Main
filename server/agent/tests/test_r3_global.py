"""r3-global — 해외 말로 지도 제어(M4) · 구역 대장 × AI 대조(C4 최소) · 어긋난 구역·영어 요약(C3 최소) · 영어 서수 영상 직행.
DB·GPU·네트워크 없이 돈다(대장 저장·읽기는 가짜로 바꿔 끼운다)."""
import asyncio
import io
import re
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from agent.tools.ext import global_ as G  # noqa: E402
from agent.tools.ext import map as M  # noqa: E402
from agent.tools.ext import vlm as V  # noqa: E402
from landxi_api import global_data as GD  # noqa: E402
from landxi_api.deps import CAPS, Principal  # noqa: E402

NL = chr(10)
TAIL = NL + "(Current area: Ysyk-Ata, Autumn 2025. Answer in English.)"
KG = Principal("tenant", "manager", "kgz-agri", "u_kg", caps=CAPS[("tenant", "manager")])
KL = Principal("tenant", "manager", "kgz-land", "u_kl", caps=CAPS[("tenant", "manager")])
NW = Principal("tenant", "manager", "namwon", "u_nw", caps=CAPS[("tenant", "manager")])
ST = Principal("lx", "staff", None, "u_st", caps=CAPS[("lx", "staff")])
SIX = ["Zoom in", "Zoom out", "Turn off the imagery layer", "Turn it back on", "Go to {d}", "Tilt to 3D"]


class Ctx:
    def __init__(self, p, lang="en", context=None):
        self.principal, self.lang, self.context = p, lang, context or {}
        self.run_id, self.state, self.citations = "run_test_r3g", {}, []
        self.tokens_in = self.tokens_out = 0
        self.r = None


def run(c):
    return asyncio.run(c)


@pytest.fixture
def glob(monkeypatch):
    """두 해외 기관(구역 하나씩) · 가짜 NDVI · 가짜 '앞 답 층'."""
    yd, sk = G.find_district("Ysyk-Ata"), G.find_district("Sokuluk")

    async def fake_mine(p):
        if p.realm == "lx":
            return [{**d, "share": 1.0} for d in G.districts()]
        return {"kgz-agri": [{**yd, "share": 0.72}], "kgz-land": [{**sk, "share": 0.85}]}.get(p.tenant_id, [])

    async def fake_ndvi(d, p):
        e = {"value": 0.31, "unit": "ndvi", "basis": "measured", "as_of": "2026-09-27", "source": "S2"}
        return {"2025-10": {"env": e, "hist": None, "crop_px": 0, "res_m": 40, "job": "j"}} if d["name"] == "Ysyk-Ata" else {}

    async def fake_last(ctx):
        return None
    monkeypatch.setattr(G, "mine", fake_mine)
    monkeypatch.setattr(G, "ndvi_months", fake_ndvi)
    monkeypatch.setattr(G, "last_layer", fake_last)
    monkeypatch.setattr(G, "tenant_is_global", lambda t: t in ("kgz-agri", "kgz-land"))
    return {"kgz-agri": yd, "kgz-land": sk}


# ── M4: 지도 동작 6문 → global_map(NDVI 요약으로 새지 않음) ─────────────────────────────
@pytest.mark.parametrize("who", [KG, KL])
def test_six_map_commands_route_to_map_tool(glob, who):
    d = glob[who.tenant_id]["name"]
    c = Ctx(who)
    want = [{"op": "zoom", "delta": 1}, {"op": "zoom", "delta": -1}, {"op": "layer", "layer": "imagery", "on": False},
            {"op": "layer", "layer": "last", "on": True}, {"op": "goto", "region": d}, {"op": "view", "preset": "3d"}]
    for q, w in zip(SIX, want):
        q = q.format(d=d)
        for text in (q, q + TAIL):
            r = G.ROUTE(text, c)
            assert r == {"tool": "global_map", "args": w, "intent": "map"}, (text, r)


def test_map_command_variants_and_data_questions(glob):
    c = Ctx(KG)
    assert G.ROUTE("Hide the satellite imagery", c)["args"] == {"op": "layer", "layer": "imagery", "on": False}
    assert G.ROUTE("Show the imagery again", c)["args"] == {"op": "layer", "layer": "imagery", "on": True}
    assert G.ROUTE("Zoom in more", c)["args"] == {"op": "zoom", "delta": 2}
    assert G.ROUTE("Back to top view", c)["args"] == {"op": "view", "preset": "top"}
    assert G.ROUTE("Take me to Chuy district", c)["args"] == {"op": "goto", "region": "Chuy"}
    # 자료 질문은 지도 도구로 가지 않는다
    assert G.ROUTE("Show June NDVI for Ysyk-Ata", c)["tool"] == "global_summary"
    assert G.ROUTE("How many fields show low vegetation?", c)["tool"] == "global_parcels"
    assert G.ROUTE("Summarize built-up area change in Ysyk-Ata", c)["args"]["metric"] == "sprawl"
    assert G.ROUTE("Zoom to Sokuluk", c)["args"] == {"region": "Sokuluk"}          # 기존 계약(위치만 · 가드) 그대로


def test_domestic_english_not_intercepted(glob):
    """국내 기관의 영어 지도 말은 해외 직행이 잡지 않고 국내 지도 직행(map.py)으로 간다(회귀 0)."""
    c = Ctx(NW)
    assert G.ROUTE("Zoom in", c) is None and G.ROUTE("Turn off the imagery layer", c) is None
    assert run(M.route_map("zoom in", c)) == {"tool": "map_zoom", "args": {"delta": 1}}


def test_global_map_actions_and_answers(glob):
    c = Ctx(KG)
    o = run(G.global_map({"op": "zoom", "delta": 1}, c))
    assert o.ui_actions == [{"op": "map_zoom", "delta": 1.0}] and o.answer == "Zoomed in." and not o.envelopes
    o = run(G.global_map({"op": "layer", "layer": "imagery", "on": False}, c))
    assert o.ui_actions == [{"op": "map_layer", "layer": "imagery", "on": False}] and "imagery layer off" in o.answer
    o = run(G.global_map({"op": "layer", "layer": "last", "on": True}, c))           # 앞 답 층이 없으면 영상 층
    assert o.ui_actions[0]["layer"] == "imagery" and o.ui_actions[0]["on"] is True
    o = run(G.global_map({"op": "view", "preset": "3d"}, c))
    assert o.ui_actions[0]["op"] == "map_view" and o.ui_actions[0]["pitch"] > 0
    o = run(G.global_map({"op": "goto", "region": "Ysyk-Ata"}, c))
    assert o.ui_actions[0]["op"] == "map_region" and o.answer == "Moved the map to Ysyk-Ata."
    for o in (run(G.global_map({"op": "zoom", "delta": -1}, c)), run(G.global_map({"op": "view", "preset": "top"}, c))):
        assert not re.search(r"\d", o.answer) and not o.envelopes                       # 동작 답에 숫자 0


def test_last_layer_used_for_it(glob, monkeypatch):
    async def last(ctx):
        return "districts"
    monkeypatch.setattr(G, "last_layer", last)
    o = run(G.global_map({"op": "layer", "layer": "last", "on": True}, Ctx(KG)))
    assert o.ui_actions == [{"op": "map_layer", "layer": "districts", "on": True}]


# ── 관할 ─────────────────────────────────────────────────────────────────────────
def test_jurisdiction_guards(glob):
    c = Ctx(KG)
    o = run(G.global_map({"op": "goto", "region": "Sokuluk"}, c))                   # kgz-land 구역 — 위치만 · 값 0
    assert "outside your districts" in o.answer and o.data["status"] == "outside" and not o.envelopes
    o = run(G.global_map({"op": "goto", "region": "Namwon"}, c))                    # 국내 지명 — 지도 동작 0
    assert o.answer == "This is not your organization's data." and not o.ui_actions
    r = G.ROUTE("Show suspect parcels in Namwon-si", c)
    o = run(getattr(G, r["tool"])(r["args"], c))
    assert o.answer == "This is not your organization's data." and not o.envelopes and not o.ui_actions
    r = G.ROUTE("How many districts don't match in Sokuluk?", c)
    assert r["tool"] == "global_mismatch" and r["args"]["region"] == "Sokuluk"


# ── C4 최소: 구역 맞추기 · 열 고르기 · 어긋남 계산 ─────────────────────────────────────
def test_resolve_district_names():
    ds = G.districts()
    assert GD.resolve("Ysyk-Ata", ds)[0]["name"] == "Ysyk-Ata"
    assert GD.resolve("YSYK ATA district", ds)[0]["name"] == "Ysyk-Ata"
    assert GD.resolve("Sokuluk rayon", ds)[0]["name"] == "Sokuluk"
    assert GD.resolve("Issyk-Ata", ds) == (G.find_district("Ysyk-Ata"), "close")    # 표기 차이(음역)
    assert GD.resolve("92254566B31675215078110", ds)[1] == "code"
    assert GD.resolve("Namwon", ds) == (None, "none") and GD.resolve("", ds) == (None, "none")


def test_pick_columns_and_units():
    ds = G.districts()
    cols = ["No", "Rayon", "Declared cropland (ha)", "Note"]
    rows = [{"No": "1", "Rayon": "Ysyk-Ata", "Declared cropland (ha)": "45 000", "Note": "x"},
            {"No": "2", "Rayon": "Chuy", "Declared cropland (ha)": "30,500", "Note": ""}]
    assert GD.pick_columns(cols, rows, ds) == ("Rayon", "Declared cropland (ha)", "ha")
    assert GD.pick_columns(["district", "area_km2"], [{"district": "Sokuluk", "area_km2": "812.5"}], ds) == ("district", "area_km2", "km2")
    assert GD.num("45 000") == 45000 and GD.num("30,500") == 30500 and GD.num("12,5") == 12.5 and GD.num("n/a") is None


def test_compare_threshold_outside_unmatched():
    ds = G.districts()
    yd = G.find_district("Ysyk-Ata")
    ai = lambda did: {"value": 60000, "unit": "ha"} if did == yd["id"] else {"value": 1000, "unit": "ha"}  # noqa: E731
    rows = [{"d": "Ysyk-Ata", "v": "40000"},       # +50% → 어긋남
            {"d": "Sokuluk", "v": "85000"},        # 관할 밖 → 비교 안 함(AI 값 없음)
            {"d": "Nowhere", "v": "10"},           # 못 맞춤
            {"d": "Ysyk-Ata", "v": "55000"},       # +9.1% → 맞음
            {"d": "Ysyk-Ata", "v": ""}]            # 값 없음
    recs = GD.compare(rows, "d", "v", "ha", {yd["id"]}, 30.0, ds, ai=ai)
    assert [r["status"] for r in recs] == ["mismatch", "outside", "unmatched", "match", "no_value"]
    assert round(recs[0]["diff_pct"], 1) == 50.0 and recs[1]["ai"] is None
    s = GD.summarize(recs, 30.0)
    assert (s["rows"], s["matched"], s["compared"], s["mismatched"], s["outside"], s["unmatched"]) == (5, 4, 2, 1, 1, 1)
    assert s["match_pct"] == 80.0
    km = GD.compare([{"d": "Ysyk-Ata", "v": "600"}], "d", "v", "km2", {yd["id"]}, 30.0, ds, ai=ai)   # km² → ha
    assert km[0]["declared"] == 60000 and km[0]["status"] == "match"


def test_read_table_csv_and_xlsx():
    cols, rows = GD.read_table("district,declared_ha\nYsyk-Ata,40000\n,\nSokuluk,85000\n".encode("utf-8-sig"), "csv")
    assert cols == ["district", "declared_ha"] and len(rows) == 2
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Rayon", "Sown area, ha"])
    ws.append(["Ysyk-Ata", 40000])
    b = io.BytesIO()
    wb.save(b)
    cols, rows = GD.read_table(b.getvalue(), "xlsx")
    assert cols == ["Rayon", "Sown area, ha"] and rows[0]["Sown area, ha"] == "40000"


def test_view_is_all_envelopes():
    from landxi_api.envelope import scan
    yd = G.find_district("Ysyk-Ata")
    recs = GD.compare([{"d": "Ysyk-Ata", "v": "40000"}, {"d": "Sokuluk", "v": "1"}], "d", "v", "ha", {yd["id"]}, 30.0,
                      ai=lambda did: {"value": 60000, "unit": "ha"})
    reg = {"id": "greg_t", "filename": "t.csv", "summary": GD.summarize(recs, 30.0), "created_at": "2026-09-30T12:00:00+09:00"}
    v = GD.view(reg, recs)
    import json
    scan(json.loads(json.dumps(v, default=str)))                      # 봉투 밖 숫자 0(개발 모드 응답 검사와 같은 규칙)
    assert v["summary"]["mismatched"]["value"] == 1 and v["items"][0]["status"] == "mismatch"
    assert [f["properties"]["status"] for f in v["geojson"]["features"]] == ["mismatch"]   # 관할 밖 구역은 칠하지 않는다


# ── C3 최소: 어긋난 구역 수 · 목록 · 영어 요약(숫자는 봉투뿐) ─────────────────────────────
def _fake_register(monkeypatch, recs=None):
    yd = G.find_district("Ysyk-Ata")
    recs = recs if recs is not None else GD.compare([{"d": "Ysyk-Ata", "v": "40000"}, {"d": "Sokuluk", "v": "85000"}, {"d": "Nowhere", "v": "5"}],
                                                    "d", "v", "ha", {yd["id"]}, 30.0, ai=lambda did: {"value": 68726, "unit": "ha", "basis": "estimate",
                                                                                                       "as_of": "2026-09-30", "source": "S2 AI"})
    reg = {"id": "greg_t", "filename": "kgz-agri.csv", "summary": GD.summarize(recs, 30.0), "created_at": "2026-09-30T12:00:00+09:00"}

    async def fake_reg(ctx):
        return GD.view(reg, recs, with_geom=False)
    monkeypatch.setattr(G, "_register", fake_reg)
    return reg


def _numbers_only_in_placeholders(text: str, allow: set[str]):
    bare = re.sub(r"\{\{[a-z0-9_]+\}\}", "", text)
    for n in re.findall(r"\d+(?:\.\d+)?", bare):
        assert n in allow, (n, text)


def test_mismatch_count_answer(glob, monkeypatch):
    _fake_register(monkeypatch)
    o = run(G.global_mismatch({"list": False}, Ctx(KG)))
    env = {k: e for k, _, e in o.envelopes}
    assert env["reg_mismatched"]["value"] == 1 and env["reg_compared"]["value"] == 1 and env["reg_outside"]["value"] == 1
    assert "{{reg_mismatched}} district doesn't match out of {{reg_compared}} district compared" in o.answer
    assert next(e for k, _, e in o.envelopes if k == 'reg_mismatched')['unit'] == 'districts'
    assert "Ysyk-Ata: declared {{mis0_declared}}, AI {{mis0_ai}} ({{mis0_diff}})" in o.answer
    assert "District level, not parcels." in o.answer
    assert all(k in env for k in re.findall(r"\{\{([a-z0-9_]+)\}\}", o.answer))            # 자리표 = 이 도구 봉투
    _numbers_only_in_placeholders(o.answer, set())
    assert o.ui_actions == [{"op": "map_on", "set": "mismatch", "label": "Mismatched districts", "register_id": "greg_t"}]


def test_mismatch_no_register_and_region_guard(glob, monkeypatch):
    async def none(ctx):
        return None
    monkeypatch.setattr(G, "_register", none)
    o = run(G.global_mismatch({}, Ctx(KG)))
    assert o.data["status"] == "no_data" and "Register tab" in o.answer and not o.envelopes
    _fake_register(monkeypatch)
    o = run(G.global_mismatch({"region": "Sokuluk"}, Ctx(KG)))                      # 다른 기관 구역
    assert "outside your districts" in o.answer and not o.envelopes


def test_findings_paragraph_numbers_from_tools(glob, monkeypatch):
    _fake_register(monkeypatch)
    o = run(G.global_findings({}, Ctx(KG)))
    env = {k: e for k, _, e in o.envelopes}
    assert o.answer.startswith("Findings for Ysyk-Ata (district level, not parcels): ")
    assert NL not in o.answer and "{{reg_rows}}" in o.answer and "{{reg_mismatched}}" in o.answer
    assert any(k.startswith("ndvi_last_") for k in env) and any(k.startswith("built_change_") for k in env)
    assert all(k in env for k in re.findall(r"\{\{([a-z0-9_]+)\}\}", o.answer))
    _numbers_only_in_placeholders(o.answer, {"2025", "2017"})                              # 연도(기간 글자)만 글자로


def test_register_questions_route(glob):
    c = Ctx(KL)
    assert G.ROUTE("How many districts don't match?" + TAIL, c) == {"tool": "global_mismatch", "args": {"list": False}, "intent": "global"}
    assert G.ROUTE("Show the mismatched districts", c)["args"] == {"list": True}
    assert G.ROUTE("Summarize the findings" + TAIL, c) == {"tool": "global_findings", "args": {}, "intent": "global"}


# ── 영어 서수 영상 설명 직행(vlm.py) ────────────────────────────────────────────────────
def test_english_ordinal_goes_direct():
    c = Ctx(NW, "en")
    assert V.ROUTE("Describe the imagery of the 3rd parcel", c) == {"tool": "vlm_describe", "args": {"rank": 3}, "intent": "vlm"}
    assert V.ROUTE("Describe the imagery of the 2nd district", c)["args"] == {"rank": 2}
    assert V.ROUTE("Describe the image of parcel 4", c)["args"] == {"rank": 4}
    assert V.ref_rank("3위 필지 영상 설명해 줘") == (3, True)
    assert V.ROUTE("Zoom in", c) is None


def test_contract_shape():
    import json
    assert set(G.SPECS) == set(G.HANDLERS) and all(n in G.WHY for n in G.SPECS)
    for s in G.SPECS.values():
        json.dumps(s)
    assert G.allowed("global_map", KG) and G.allowed("global_mismatch", ST) and not G.allowed("global_map", NW)


def test_route_args_pass_registry_validation(glob):
    """직행 인자가 도구 명세(enum)를 통과한다 — 'Turn it back on'(layer='last')이 bad_request 로 모델 경로에 새지 않게."""
    from agent.tools import registry
    specs = dict(registry.SPECS)
    try:
        registry.SPECS.update(G.SPECS)
        c = Ctx(KG)
        for q in [*[s.format(d="Ysyk-Ata") for s in SIX], "Back to top view", "How many districts don't match?", "Summarize the findings",
                  "Show the mismatched districts", "Show June NDVI for Ysyk-Ata"]:
            r = G.ROUTE(q, c)
            registry.validate(r["tool"], dict(r["args"]))
    finally:
        registry.SPECS.clear()
        registry.SPECS.update(specs)


# ── 실증 must_fix 1 · 2: 지명을 붙인 확대 · 관할 밖 대장 답의 지도 동작 ─────────────────────
def test_zoom_on_named_place_routes_with_region(glob):
    for who in (KG, KL):
        c = Ctx(who)
        for q in ("Zoom in on Ak-Suu", "Zoom in on Ak-Suu" + TAIL, "zoom into Ak-Suu please"):
            r = G.ROUTE(q, c)
            assert r["tool"] == "global_map" and r["args"]["op"] == "zoom" and r["args"]["region"] == "Ak-Suu" and r["args"]["delta"] == 1, (q, r)
        assert G.ROUTE("Zoom out to Sokuluk", c)["args"] == {"op": "zoom", "delta": -1, "region": "Sokuluk"}
        assert G.ROUTE("Zoom in on Yeosu", c)["args"] == {"op": "zoom", "delta": 1, "region": "Yeosu"}
        # 지명이 아닌 말은 그냥 확대(회귀 0) · 'Zoom to {구역}'·'zoom in to 12' 계약 유지
        assert G.ROUTE("Zoom in on the map", c)["args"] == {"op": "zoom", "delta": 1}
        assert G.ROUTE("zoom in on it a bit", c)["args"] == {"op": "zoom", "delta": 1}
        assert G.ROUTE("Zoom in", c)["args"] == {"op": "zoom", "delta": 1}
        assert G.ROUTE("zoom in to 12", c)["args"] == {"op": "zoom", "zoom": 12.0}
        assert G.ROUTE("Zoom to Sokuluk", c)["tool"] == "global_summary"


def test_zoom_on_named_place_actions(glob):
    # 관할: 그 구역으로 옮긴 뒤 확대(동작 하나)
    o = run(G.global_map({"op": "zoom", "delta": 1, "region": "Sokuluk"}, Ctx(KL)))
    assert len(o.ui_actions) == 1 and o.ui_actions[0]["op"] == "map_region" and o.ui_actions[0]["name"] == "Sokuluk" and o.ui_actions[0]["zoom_delta"] == 1.0
    assert o.answer == "Moved the map to Sokuluk and zoomed in." and not o.envelopes
    # 관할 밖 해외 구역: 위치만(확대 없음) + 가드 한 줄 — 답이 말한 동작 = 보낸 동작
    for who, place in ((KL, "Ak-Suu"), (KG, "Sokuluk")):
        o = run(G.global_map({"op": "zoom", "delta": 1, "region": place}, Ctx(who)))
        assert [a["op"] for a in o.ui_actions] == ["map_region"] and "zoom_delta" not in o.ui_actions[0]
        assert "outside your districts, so the map shows its location only" in o.answer and o.data["status"] == "outside"
        assert "Zoomed" not in o.answer
    # 국내 지명·모르는 이름: 지도를 움직이지 않고 거절
    for who in (KG, KL):
        o = run(G.global_map({"op": "zoom", "delta": 1, "region": "Yeosu"}, Ctx(who)))
        assert o.ui_actions == [] and o.data["status"] == "outside" and o.answer == "This is not your organization's data."
        o = run(G.global_map({"op": "zoom", "delta": 1, "region": "Atlantis"}, Ctx(who)))
        assert o.ui_actions == [] and o.data["status"] == "not_found" and "No district named 'Atlantis'" in o.answer
    # LX 는 어느 해외 구역이든 옮겨 확대
    o = run(G.global_map({"op": "zoom", "delta": -1, "region": "Ak-Suu"}, Ctx(ST)))
    assert o.ui_actions[0]["zoom_delta"] == -1.0 and o.answer.endswith("zoomed out.")


def test_outside_register_answer_moves_map(glob, monkeypatch):
    """'How many districts don't match in Sokuluk?'(agri) — 답이 '위치만 보였다'고 말하면 위치 이동을 실제로 보낸다."""
    _fake_register(monkeypatch)
    for who, place in ((KG, "Sokuluk"), (KL, "Ysyk-Ata")):
        for q in (f"How many districts don't match in {place}?", f"Show the mismatched districts in {place}"):
            r = G.ROUTE(q, Ctx(who))
            assert r["tool"] == "global_mismatch" and r["args"]["region"] == place
            o = run(G.global_mismatch(r["args"], Ctx(who)))
            assert "the map shows its location only" in o.answer
            assert [a["op"] for a in o.ui_actions] == ["map_region"] and o.ui_actions[0]["name"] == place and not o.envelopes
    o = run(G.global_mismatch({"region": "Yeosu"}, Ctx(KG)))                      # 국내 지명: 움직이지 않음 · 동작 말도 없음
    assert o.ui_actions == [] and "map" not in o.answer
