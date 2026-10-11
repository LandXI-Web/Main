"""지도 서비스 · XI맵을 말로 — 세부 거르기 · AI 분석 결과 보고서(원칙 193 · 확인 대장 10-11 Q5 ⓑ · QA-고침-5).

  map_filter(cls?, area_min?, area_op_min?, area_max?, area_op_max?, emd?, clear?)
      '비닐하우스 단동만 펼쳐줘' · '면적 1,000㎡ 넘는 것만' · '운봉읍만' · '조건 풀어 줘' → 지금 지도에 켜진 결과 층(화면이 context.sets 로 알린 것)를
      결과 타일 속성(cls · area_m2 · emd)으로 거른다. 서버는 그 층의 실제 결과(detections — 타일과 같은 행)에 그 값이 있는지 먼저 본다:
      없는 값(예: 단동 · 연동 구분이 없는 결과)은 거르지 않고 '없다'고 답하며, 그 질문은 개선 고리(improve · '자료 없음')로 간다.
      맞는 수 = 같은 행을 같은 조건으로 센 값(봉투) · 화면은 ui_action map_filter 로 같은 조건을 바로 건다(모델 호출 0 · 1초 안 목표).
  ai_report(request)
      '보고서 초안 만들어 줘'(LX 계정) → AI 분석 결과 보고서 .docx — 읍면동별 개수 · 면적 · 지도 그림 · 영상 시점 · AI 모델 정확도.
      의심 필지 · 현장 확인 개념 없음(원칙 135). 대상 = 질문의 지역(그 지역의 결과 — 통계 답과 같은 세트 선택 · results._set)
      → 지금 화면에 켜진 층(context.sets) → 화면 지역(context.region) → 없으면 묻는다. 숫자 = 지도 서비스와 같은 행(지운 결과 제외).
ROUTE 는 모델 앞 직행(LLM 0). 지역 · 지명 고정값 0.
"""
from __future__ import annotations

import asyncio
import datetime as dt
import io
import json
import re

from .. import Out, ToolError

SPECS = {
    "map_filter": {
        "description": "지도에 켜진 AI 분석 결과 층를 분류 · 면적 · 읍면동 조건으로 거른다('비닐하우스 단동만' · '1,000㎡ 넘는 것만' · '○○면만' · '조건 풀어 줘').",
        "properties": {"cls": {"type": "string", "description": "분류(예 '비닐하우스' · '비닐하우스 단동')"},
                       "area_min": {"type": "number"}, "area_op_min": {"type": "string", "enum": [">", ">="]},
                       "area_max": {"type": "number"}, "area_op_max": {"type": "string", "enum": ["<", "<="]},
                       "emd": {"type": "string", "description": "읍면동 이름"}, "clear": {"type": "boolean"}, "request": {"type": "string"}}},
    "ai_report": {
        "description": "AI 분석 결과 보고서 초안(.docx) — 읍면동별 개수 · 면적 · 지도 그림 · 영상 시점 · AI 모델 정확도. 지역이 없으면 지금 화면의 결과 층 · 지역.",
        "properties": {"request": {"type": "string", "description": "사용자 문장 원문"}}},
}
WRITE: set[str] = set()
CONFIRM: set[str] = set()
CLIENT = {"map_filter"}
NO_LLM = {"map_filter", "ai_report"}
WHY = {"map_filter": "켜진 결과 층 거르기", "ai_report": "AI 분석 결과 보고서 .docx 만들기"}
SAY = {"map_filter": "조건으로 거르기", "ai_report": "AI 분석 결과 보고서 만들기"}
HINT = ("켜진 결과 층를 분류 · 면적 · 읍면동으로 거르는 말('○○만 보여 줘' · '1,000㎡ 넘는 것만')은 map_filter. "
        "LX 계정의 보고서 초안은 ai_report(AI 분석 결과 보고서).")


def allowed(name: str, p) -> bool:
    realm = getattr(p, "realm", None)
    if name == "ai_report":
        return realm == "lx"
    return realm in ("lx", "tenant")


# ── 말 → 조건 ─────────────────────────────────────────────────────────────────────
CLS_WORDS = [(r"비닐\s*하우스|온실|하우스", "비닐하우스"), (r"건물|건축물", "건물"), (r"경작지|농경지|논밭", "경작지"), (r"주차장", "주차장"),
             (r"차량|자동차", "차량"), (r"스티로폼", "스티로폼")]
SUB_WORDS = [(r"단동", "단동"), (r"연동|다동", "다동")]
SUB_KO = {"단동": "단동", "다동": "연동"}
UNIT = r"(㎡|m2|m²|제곱\s*미터|평방\s*미터|평|ha|헥타르)"
NUM = r"(\d[\d,]*(?:\.\d+)?)\s*(천|만)?\s*"
AREA_RX = re.compile(NUM + UNIT + r"\s*(?:(?:이|가|을|를)(?![상하]))?\s*(이상|넘는|넘은|넘게|초과|보다\s*큰|보다\s*넓은|큰|넓은|이하|미만|안\s*되는|보다\s*작은|보다\s*좁은|작은|좁은)?")
ONLY_RX = re.compile(r"만\s*(?:보여|펼쳐|펴|남겨|켜|표시|그려|띄워|줘|$)|만$|만\s*[.!?]|(?:이상|넘는|넘은|초과|이하|미만|작은|큰)\s*(?:것|거|결과|건|하우스|건물)?\s*(?:만|들)?\s*(?:보여|펼쳐|남겨|표시|그려|줘|$)")
CLEAR_RX = re.compile(r"(?:조건|거르기|필터)\S*\s*(?:다\s*)?(?:풀|없애|지워|해제|빼)|(?:모두|전부|다)\s*(?:다시\s*)?(?:보여|펼쳐|켜)|원래대로")
NOT_RX = re.compile(r"보고서|통계|차트|그래프|법령|조문|분석\s*(?:실행|돌려|해\s*줘)|내려\s*받|요약")
EMD_RX = re.compile(r"(?<![가-힣])([가-힣]{1,6}(?:읍|면|동|리|가))(?=만|에서|의|에|은|는|\s|,|$)")
EMD_SKIP = {"단동", "연동", "다동", "이동", "자동", "행동", "변동", "활동", "작동", "운동", "공동", "가동", "표면", "전면", "측면", "화면", "지면", "평면", "이면", "정면", "단면"}


def _area(m) -> tuple[float, str, str] | None:
    v = float(m.group(1).replace(",", ""))
    if m.group(2) == "천":
        v *= 1000
    elif m.group(2) == "만":
        v *= 10000
    u = re.sub(r"\s+", "", m.group(3))
    if u in ("ha", "헥타르"):
        v *= 10000
    elif u == "평":
        v *= 3.305785
    w = re.sub(r"\s+", "", m.group(4) or "")
    if w in ("이하", "미만", "안되는", "보다작은", "보다좁은", "작은", "좁은"):
        return v, "max", "<=" if w == "이하" else "<"
    return v, "min", ">=" if w == "이상" else ">"


def parse_filter(t: str) -> dict | None:
    """문장 → 거르기 인자 · 거르기 말이 아니면 None."""
    t = re.sub(r"\s+", " ", t or "").strip()
    if not t or len(t) > 100 or NOT_RX.search(t):
        return None
    if CLEAR_RX.search(t):
        return {"clear": True}
    m = AREA_RX.search(t)
    if not (ONLY_RX.search(t) or (m and m.group(4))):
        return None
    a: dict = {}
    if m:
        v, side, op = _area(m)
        a[f"area_{side}"] = round(v, 1)
        a[f"area_op_{side}"] = op
    sub = next((k for rx, k in SUB_WORDS if re.search(rx, t)), None)
    base = next((k for rx, k in CLS_WORDS if re.search(rx, t)), None)
    if sub:
        a["cls"] = f"{base or '비닐하우스'} {sub}"
    elif base:
        a["cls"] = base
    rest = AREA_RX.sub(" ", t)
    em = [x for x in EMD_RX.findall(rest) if x not in EMD_SKIP and not any(re.search(rx, x) for rx, _ in CLS_WORDS + SUB_WORDS)]
    if em:
        a["emd"] = em[0]
    if not any(k in a for k in ("cls", "emd", "area_min", "area_max")):
        return None
    a["request"] = t
    return a


REPORT_ASK = re.compile(r"보고서|리포트|report", re.I)
REPORT_DO = re.compile(r"초안|써\s*줘|써줘|작성|만들어|뽑아|draft|write", re.I)


def ROUTE(msg: str, ctx):
    t = re.sub(r"\s+", " ", msg or "").strip()
    if not t or getattr(ctx, "lang", "ko") == "en":
        return None
    p = getattr(ctx, "principal", None)
    c = getattr(ctx, "context", None) or {}
    if getattr(p, "realm", None) == "lx" and len(t) <= 120 and REPORT_ASK.search(t) and REPORT_DO.search(t) \
            and not re.search(r"법령|조문|근거\s*조", t) and not re.search(r"실태\s*조사|의심", t):
        return {"tool": "ai_report", "args": {"request": t}, "intent": "report"}
    if c.get("page") in ("lx-map", "xi-clean") or c.get("sets") is not None:
        a = parse_filter(t)
        if a:
            return {"tool": "map_filter", "args": a, "intent": "map"}
    return None


def ROUTE_FIRST(msg: str, ctx):
    return ROUTE(msg, ctx)


def _route_front():
    """파일 이름순 ROUTE 보다 앞(운영 직행 다음)에 — '○○면만'이 요약 · 숫자 직행에, '보고서 초안'이 실태조사 보고서 직행에 먹히지 않게."""
    try:
        from . import ROUTES
    except Exception:  # noqa: BLE001
        return
    ROUTES[:] = [r for r in ROUTES if r[0] not in ("mapsvc", "mapsvc_first")]
    at = 1 if ROUTES and ROUTES[0][0] == "ops" else 0
    ROUTES.insert(at, ("mapsvc_first", ROUTE_FIRST))


# ── 켜진 층 ───────────────────────────────────────────────────────────────────
_SET = re.compile(r"^results/[\w.-]+/(job_[0-9A-Z]{20,32})$")


async def _jobs_on(ctx) -> list[str]:
    """화면이 알린 켜진 결과 층(context.sets) → 이 계정이 볼 수 있는 작업 id(최대 12). 모르는 모양 · 남의 기관 것은 버린다."""
    sets = (getattr(ctx, "context", None) or {}).get("sets") or []
    ids = []
    for s in sets[:24] if isinstance(sets, list) else []:
        m = _SET.match(str(s))
        if m and m.group(1) not in ids:
            ids.append(m.group(1))
    if not ids:
        return []
    from landxi_api.deps import db
    p = ctx.principal
    async with db(realm="lx") as c:
        rows = await c.fetch("SELECT id, tenant_id FROM jobs WHERE id = ANY($1::text[]) AND kind IN ('infer','reinfer') AND state='done'", ids)
    ok = {r["id"] for r in rows if getattr(p, "realm", None) == "lx" or r["tenant_id"] in (getattr(p, "tenant_id", None), "lx")}
    return [x for x in ids if x in ok][:12]


def _nf(v: float) -> str:
    return f"{int(round(v)):,}" if abs(v - round(v)) < 0.05 else f"{v:,.1f}"


async def map_filter(args: dict, ctx) -> Out:
    from landxi_api.deps import db
    from landxi_api.envelope import env
    out = Out(source="켜진 결과 층(결과 타일과 같은 행)")
    ctx.state["talk_kind"] = None
    if args.get("clear"):
        out.ui_actions.append({"op": "map_filter", "clear": True})
        out.answer = "거르기를 풀었습니다 — 켜진 층의 결과를 모두 보여 줍니다."
        return out
    jobs = await _jobs_on(ctx)
    if not jobs:
        ctx.state["cannot"] = {"kind": "screen"}
        out.data = {"status": "none"}
        out.answer = "지금 지도에 켜진 결과 층이 없습니다. 왼쪽에서 층을 켠 뒤 다시 말씀해 주세요."
        return out
    async with db(realm="lx") as c:
        av = await c.fetch("SELECT cls, coalesce(emd, '') AS emd, count(*) AS n FROM detections WHERE job_id = ANY($1::text[]) AND edit_state <> 'deleted' GROUP BY 1, 2", jobs)
    classes = sorted({r["cls"] for r in av if r["cls"]})
    emds = sorted({r["emd"] for r in av if r["emd"]})
    miss = []
    want_cls = None
    if args.get("cls"):
        base, _, sub = str(args["cls"]).partition(" ")
        if sub:
            want_cls = [c for c in classes if c.startswith(base) and sub in c]
            if not want_cls:
                has_base = any(c.startswith(base) for c in classes)
                miss.append(f"{base} {SUB_KO.get(sub, sub)} 구분이 없습니다"
                            + (f" — AI가 '{base}' 한 가지로만 찾은 결과입니다" if has_base else ""))
        else:
            want_cls = [c for c in classes if c == base or c.startswith(base + "_")]
            if not want_cls:
                miss.append(f"{base} 결과가 없습니다")
    want_emd = None
    if args.get("emd"):
        e = str(args["emd"])
        want_emd = [x for x in emds if x == e or (len(e) >= 2 and x.startswith(e))]
        if not want_emd:
            miss.append(f"{e} 결과가 없습니다")
    if miss:
        ctx.state["cannot"] = {"kind": "nodata"}           # 개선 고리 — '자료 없음'(못 거른 조건)으로 모인다(원칙 193 자동 진화 관리)
        have = " · ".join(_ko(c) for c in classes[:4])
        out.data = {"status": "missing", "missing": miss, "classes": classes[:8]}
        out.answer = "켜진 층에는 " + " · ".join(miss) + ". 지도는 그대로 두었습니다." + (f" 지금 결과의 분류: {have}." if have else "")
        return out
    cond, a = ["job_id = ANY($1::text[])", "edit_state <> 'deleted'"], [jobs]
    if want_cls:
        a.append(want_cls)
        cond.append(f"cls = ANY(${len(a)}::text[])")
    if want_emd:
        a.append(want_emd)
        cond.append(f"emd = ANY(${len(a)}::text[])")
    if args.get("area_min") is not None:
        a.append(float(args["area_min"]))
        cond.append(f"area_m2 {'>=' if args.get('area_op_min') == '>=' else '>'} ${len(a)}")
    if args.get("area_max") is not None:
        a.append(float(args["area_max"]))
        cond.append(f"area_m2 {'<=' if args.get('area_op_max') == '<=' else '<'} ${len(a)}")
    async with db(realm="lx") as c:
        r = await c.fetchrow(f"SELECT count(*) AS n, coalesce(sum(area_m2), 0) AS m2 FROM detections WHERE {' AND '.join(cond)}", *a)
    n, m2 = int(r["n"]), float(r["m2"] or 0)
    words = []
    if want_cls:
        words.append(" · ".join(_ko(c) for c in want_cls))
    if args.get("area_min") is not None:
        words.append(f"{_nf(float(args['area_min']))}㎡ {'이상' if args.get('area_op_min') == '>=' else '넘음'}")
    if args.get("area_max") is not None:
        words.append(f"{_nf(float(args['area_max']))}㎡ {'이하' if args.get('area_op_max') == '<=' else '미만'}")
    if want_emd:
        words.append(" · ".join(want_emd))
    label = " · ".join(words)
    act = {"op": "map_filter", "label": label, "n": n}
    if want_cls:
        act["cls"] = want_cls
    if want_emd:
        act["emd"] = want_emd
    for side in ("min", "max"):
        if args.get(f"area_{side}") is not None:
            act[f"area_{side}"] = float(args[f"area_{side}"])
            act[f"area_op_{side}"] = args.get(f"area_op_{side}") or (">" if side == "min" else "<")
    out.ui_actions.append(act)
    out.env("n", f"조건({label})에 맞는 AI 분석 결과", env(n, "count", "inferred", "켜진 결과 층(지운 결과 제외)"))
    out.env("m2", f"조건({label})에 맞는 결과 면적 합", env(round(m2), "m2", "inferred", "켜진 결과 층(지운 결과 제외)"))
    out.data = {"status": "ok", "조건": label, "맞는 수": "n"}
    out.answer = (f"{label} 조건으로 걸렀습니다 — 맞는 결과 {{{{n}}}}건만 지도에 남깁니다." if n
                  else f"{label} 조건에 맞는 결과가 없습니다(0건) — 지도에는 아무것도 남지 않습니다.")
    return out


def _ko(c: str) -> str:
    from landxi_api.spaces import class_ko
    s = class_ko(c)
    return s.replace("다동", "연동")


# ── AI 분석 결과 보고서 ───────────────────────────────────────────────────────────
CLS_COLOR = {"경작지": "#0FA9A0", "비닐하우스": "#07706A", "건물": "#2B3440", "주차장": "#8F99A8"}
NOTICE = "AI 분석 결과(검수 전)는 참고자료입니다 — 처분은 담당자가 정합니다."
NEED = "어느 지역의 보고서를 만들까요? 지역 이름을 함께 말씀하시거나, 지도 서비스에서 결과 층을 켠 뒤 다시 말씀해 주세요."


async def _target_jobs(ctx, text: str) -> tuple[list[str], str]:
    """보고서 대상 작업 — ① 질문의 지역(통계 답과 같은 세트 · results._set) ② 켜진 층 ③ 화면 지역. → (작업 id 들, 어디서 골랐는지)."""
    from ... import talk
    from .. import results as RS
    groups = talk.regions_in(text)
    named = [talk.pick(ctx, g)[0] for g in groups if g]
    named = [x for x in named if x]
    cls = next((k for rx, k in CLS_WORDS if re.search(rx, text)), None)
    if named:
        s = await RS._set({"region": named[0]["sgg_cd"], **({"cls": cls} if cls else {})}, ctx)
        return [s.rsplit("/", 1)[-1]], "질문의 지역"
    on = await _jobs_on(ctx)
    if on:
        return on[:4], "켜진 층"
    if (ctx.context or {}).get("region"):
        s = await RS._set({**({"cls": cls} if cls else {})}, ctx)
        return [s.rsplit("/", 1)[-1]], "화면 지역"
    return [], ""


async def _job_data(job: str) -> dict | None:
    from landxi_api.deploys import gsd_word
    from landxi_api.deps import db
    from landxi_api.regions import region_of
    from landxi_api.release import acc_of, _acc_at
    async with db(realm="lx") as c:
        j = await c.fetchrow("SELECT j.id, j.card_id, j.model_id, j.options, j.counts, j.finished_at, c.name->>'ko' AS cname, i.name AS iname, i.year, i.epoch, "
                             "i.gsd_m, m.metrics, pr.name AS pname FROM jobs j LEFT JOIN cards c ON c.id=j.card_id LEFT JOIN imagery i ON i.id=j.imagery_id "
                             "LEFT JOIN models m ON m.id=j.model_id LEFT JOIN projects pr ON pr.id = j.options->>'project_id' WHERE j.id=$1", job)
        if not j:
            return None
        rows = await c.fetch("SELECT coalesce(nullif(emd, ''), '') AS emd, cls, count(*) AS n, coalesce(sum(area_m2), 0) AS m2 FROM detections "
                             "WHERE job_id=$1 AND edit_state <> 'deleted' GROUP BY 1, 2", job)
    o = j["options"] if isinstance(j["options"], dict) else json.loads(j["options"] or "{}")
    mt = j["metrics"] if isinstance(j["metrics"], dict) else json.loads(j["metrics"] or "{}") if j["metrics"] else {}
    iname = j["iname"]
    if isinstance(iname, str) and iname.startswith("{"):
        iname = json.loads(iname)
    iname = (iname or {}).get("ko") if isinstance(iname, dict) else (iname or "")
    yr = str(j["year"] or (str(j["epoch"])[:4] if j["epoch"] else ""))
    rg = region_of(o.get("sgg_cd")) if o.get("sgg_cd") else None
    proj = bool(o.get("project_infer"))
    if o.get("scope_full"):
        scope = f"전역 · 읍면동 {o.get('emd_total')}곳" if o.get("emd_total") else "전역"
    elif o.get("coverage") is not None:
        pc = float(o["coverage"]) * 100
        scope = "영상 있는 1% 미만" if pc < 1 else f"영상 있는 {round(pc)}%"
    else:
        scope = "배포 전 모델" if proj else ""
    return {"job": job, "service": f"프로젝트 추론 — {j['pname']}" if proj else (j["cname"] or "분석 서비스"),
            "region": " ".join(x for x in ((rg or {}).get("name") or "", o.get("range_name") or "" if proj else "") if x) or "—",
            "sgg": (rg or {}).get("sgg_cd") or o.get("sgg_cd"), "scope": scope,
            "imagery": {"name": iname or "", "year": yr, "word": " ".join(x for x in (yr, gsd_word(j["gsd_m"])) if x)},
            "finished": j["finished_at"], "acc": acc_of(mt), "acc_at": _acc_at(mt),
            "found": sum(int(v or 0) for v in (j["counts"] if isinstance(j["counts"], dict) else json.loads(j["counts"] or "{}")).values()),
            "rows": [dict(r) for r in rows]}


def _map_png(job: str, sgg: str | None, rows_geo: list, counts: dict | None = None, w_in=7.0, h_in=5.2) -> bytes | None:
    """지도 그림 — 읍면동 경계(연한 선 · 이름) + AI 분석 결과(분류 색 · 지도 서비스 범례와 같은 색). 결과가 많으면 점으로."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.collections import PatchCollection
    from matplotlib.patches import Polygon as MP
    from shapely import wkb
    from shapely.geometry import shape
    from matplotlib import font_manager
    have = {f.name for f in font_manager.fontManager.ttflist}
    plt.rcParams["font.family"] = [next((f for f in ("Malgun Gothic", "NanumGothic", "Noto Sans CJK KR", "AppleGothic") if f in have), "sans-serif")]
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(w_in, h_in), dpi=160)
    ax.set_facecolor("#F7F8FA")
    fc = None
    if sgg:
        try:
            from landxi_api.regions import emd_fc
            fc = emd_fc(sgg)
        except Exception:  # noqa: BLE001
            fc = None
    for f in (fc or {}).get("features") or []:
        try:
            g = shape(f["geometry"])
        except Exception:  # noqa: BLE001
            continue
        for poly in getattr(g, "geoms", [g]):
            x, y = poly.exterior.xy
            ax.fill(x, y, facecolor="#FFFFFF", edgecolor="#C5CBD3", linewidth=0.6, zorder=1)
        nm = (f.get("properties") or {}).get("name") or (f.get("properties") or {}).get("emd_kor_nm") or ""
        n = (counts or {}).get(nm)
        if nm and n:                                  # 결과가 있는 읍면동만 이름 · 개수(보고서 표와 같은 값)
            pt = g.representative_point()
            ax.text(pt.x, pt.y, nm + "\n" + f"{n:,}건", fontsize=6, color="#3A4651", ha="center", va="center", zorder=4,
                    bbox={"boxstyle": "round,pad=0.2", "fc": (1, 1, 1, 0.75), "ec": "none"})
    by: dict[str, list] = {}
    pts: dict[str, list] = {}
    many = len(rows_geo) > 15000
    for cls, b in rows_geo:
        try:
            g = wkb.loads(bytes(b))
        except Exception:  # noqa: BLE001
            continue
        k = next((n for n in CLS_COLOR if str(cls or "").startswith(n)), "비닐하우스")
        if many:
            c = g.centroid
            pts.setdefault(k, []).append((c.x, c.y))
            continue
        for poly in getattr(g, "geoms", [g]):
            by.setdefault(k, []).append(MP(list(poly.exterior.coords), closed=True))
    for k, ps in by.items():
        ax.add_collection(PatchCollection(ps, facecolor=CLS_COLOR[k], edgecolor=CLS_COLOR[k], linewidth=0.9, alpha=0.9, zorder=3))
    for k, ps in pts.items():
        xs, ys = zip(*ps)
        ax.scatter(xs, ys, s=0.6, c=CLS_COLOR[k], zorder=3, linewidths=0)
    ax.autoscale_view()
    if fc and fc.get("features"):
        from shapely.ops import unary_union
        try:
            u = unary_union([shape(f["geometry"]) for f in fc["features"]]).bounds
            ax.set_xlim(u[0], u[2])
            ax.set_ylim(u[1], u[3])
        except Exception:  # noqa: BLE001
            pass
    import math
    lat = sum(ax.get_ylim()) / 2
    ax.set_aspect(1 / max(0.2, math.cos(math.radians(lat))))
    ax.set_xticks([])
    ax.set_yticks([])
    for s in ax.spines.values():
        s.set_visible(False)
    from matplotlib.patches import Patch
    hs = [Patch(facecolor=CLS_COLOR[k], label=k) for k in (by or pts)]
    if hs:
        ax.legend(handles=hs, loc="lower right", fontsize=7, frameon=True, framealpha=0.9)
    fig.tight_layout(pad=0.3)
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def _emd_counts(d: dict) -> dict:
    out: dict = {}
    for r in d["rows"]:
        if r["emd"]:
            out[r["emd"]] = out.get(r["emd"], 0) + int(r["n"])
    return out


def _docx(parts: list[dict], pngs: list[bytes | None], made: str) -> tuple[bytes, str]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from survey.report import _font, _p, _shade, _tbl
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    head = parts[0]
    title = f"{head['region']} {head['service']} AI 분석 결과 보고서(초안)" if len(parts) == 1 else f"AI 분석 결과 보고서(초안) — 결과 층 {len(parts)}개"
    doc.core_properties.title = title
    doc.core_properties.author = "Land-XI"
    _font(sec.header.paragraphs[0].add_run(f"Land-XI · AI 분석 결과 · 초안"), 8, False, "6B7780")
    _p(doc, title, 16, True, "14202A", space_after=2)
    _p(doc, f"작성 {made} · AI 분석 결과(검수 전)", 9, False, "56626B")
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    c = box.rows[0].cells[0]
    _shade(c, "EEF6F6")
    c.text = ""
    _font(c.paragraphs[0].add_run(NOTICE), 10, True, "07706A")
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, (d, png) in enumerate(zip(parts, pngs), 1):
        tot_n = sum(int(r["n"]) for r in d["rows"])
        tot_m2 = sum(float(r["m2"]) for r in d["rows"])
        pre = f"{k}. " if len(parts) > 1 else ""
        _p(doc, "", space_after=2)
        _p(doc, f"{pre}개요" if not pre else f"{pre}{d['region']} · {d['service']}", 12.5, True, "14202A")
        fin = d["finished"].astimezone(dt.timezone(dt.timedelta(hours=9))).strftime("%Y.%m.%d") if d["finished"] else "—"
        _tbl(doc, ["항목", "내용"], [
            ["분석 서비스", d["service"]],
            ["지역", d["region"] + (f" · {d['scope']}" if d["scope"] else "")],
            ["영상 시점", d["imagery"]["name"] if d["imagery"]["word"] and d["imagery"]["word"] in (d["imagery"]["name"] or "")
             else " · ".join(x for x in (d["imagery"]["word"], d["imagery"]["name"]) if x) or "—"],
            ["분석한 날", fin],
            ["AI 모델 정확도", f"{d['acc']}%" + (f" (학습 끝 검증 · {d['acc_at'].replace('-', '.')})" if d.get("acc_at") else "") if d["acc"] is not None else "기록 없음"],
            ["찾은 결과", f"{tot_n:,}건 · 면적 합 {round(tot_m2):,}㎡ ({tot_m2 / 1e4:,.1f}ha)"],
        ], widths=[3.6, 13.6], size=9.5)
        if png:
            _p(doc, "", space_after=2)
            _p(doc, "지도 그림", 11, True, "14202A", space_after=2)
            doc.add_picture(io.BytesIO(png), width=Cm(16.5))
            _p(doc, "읍면동 경계 위에 AI가 찾은 결과를 칠했습니다(지도 서비스 범례와 같은 색).", 8.5, False, "6B7780")
        by_emd: dict[str, list] = {}
        for r in d["rows"]:
            e = by_emd.setdefault(r["emd"] or "읍면동 밖", [0, 0.0])
            e[0] += int(r["n"])
            e[1] += float(r["m2"])
        _p(doc, "", space_after=2)
        _p(doc, "읍면동별 개수 · 면적", 11, True, "14202A", space_after=2)
        rows = [[e, v[0], round(v[1]), f"{(v[0] / tot_n * 100 if tot_n else 0):.1f}%"] for e, v in sorted(by_emd.items(), key=lambda x: (-x[1][0], x[0]))]
        rows.append(["합계", tot_n, round(tot_m2), "100%" if tot_n else "—"])
        t = _tbl(doc, ["읍면동", "개수(건)", "면적(㎡)", "개수 비율"], rows, widths=[5.2, 3.6, 4.4, 4.0], size=9)
        for cell in t.rows[-1].cells:
            _shade(cell, "F2F4F6")
        cls_set = sorted({r["cls"] for r in d["rows"] if r["cls"]})
        if len(cls_set) > 1:
            _p(doc, "", space_after=2)
            _p(doc, "분류별", 11, True, "14202A", space_after=2)
            crow = []
            for cl in cls_set:
                n = sum(int(r["n"]) for r in d["rows"] if r["cls"] == cl)
                m2 = sum(float(r["m2"]) for r in d["rows"] if r["cls"] == cl)
                crow.append([_ko(cl), n, round(m2)])
            _tbl(doc, ["분류", "개수(건)", "면적(㎡)"], crow, widths=[6, 4, 5], size=9)
    _p(doc, "", space_after=2)
    _p(doc, "읽는 법 — 개수 · 면적은 AI가 그린 결과의 수와 넓이입니다(결과 확인 전 · 지운 결과 제외). 지도 서비스의 결과 층 · 내려받기 파일과 같은 값입니다.",
       8.5, False, "6B7780")
    buf = io.BytesIO()
    doc.save(buf)
    svc = re.sub(r"\s*분석\s*서비스$", "", head["service"]) if len(parts) == 1 else "결과층"
    name = re.sub(r"[\\/:*?\"<>|\s—]+", "_", f"{head['region']}_{svc}_AI분석보고서_{made[:10].replace('.', '')}").strip("_") + ".docx"
    return buf.getvalue(), name


async def ai_report(args: dict, ctx) -> Out:
    from landxi_api.deps import db
    from landxi_api.envelope import env
    from ... import config
    from ...report import file_href
    t = re.sub(r"\s+", " ", args.get("request") or ctx.state.get("msg") or "").strip()
    out = Out(source="AI 분석 결과(지도 서비스와 같은 행)")
    try:
        jobs, why = await _target_jobs(ctx, t)
    except ToolError as e:
        if e.status == 404:
            from ... import talk
            g = talk.regions_in(t)
            nm = (talk.pick(ctx, g[0])[0] or {}).get("full") if g else None
            out.data = {"status": "no_data", "text": "해당 지역 데이터가 없습니다", **({"next": f"XI맵에서 {nm} AI 분석을 먼저 실행하세요"} if nm else {})}
            out.answer = f"{nm or '그 지역'}은 아직 AI 분석 결과가 없습니다."
            return out
        raise
    if not jobs:
        ctx.state["cannot"] = {"kind": "other"}
        out.data = {"status": "need_region"}
        out.answer = NEED
        return out
    parts = [d for d in [await _job_data(j) for j in jobs] if d and d["rows"]]
    if not parts:
        out.data = {"status": "no_data", "text": "해당 지역 데이터가 없습니다"}
        out.answer = "고른 층에 찾은 결과가 없어(0건) 보고서를 만들지 않았습니다."
        return out
    emd_only = next((m for m in EMD_RX.findall(t) if m not in EMD_SKIP), None) if why == "질문의 지역" else None
    if emd_only:
        for d in parts:
            keep = [r for r in d["rows"] if r["emd"] == emd_only or r["emd"].startswith(emd_only)]
            if keep:
                d["rows"], d["region"] = keep, f"{d['region']} {keep[0]['emd']}"
    geos = []
    async with db(realm="lx") as c:
        for d in parts:
            emds = sorted({r["emd"] for r in d["rows"]})
            q = "SELECT cls, ST_AsBinary(ST_SimplifyPreserveTopology(geom, 0.000005)) AS b FROM detections WHERE job_id=$1 AND edit_state <> 'deleted'"
            rows = await c.fetch(q + (" AND coalesce(emd, '') = ANY($2::text[])" if emd_only else "") + " LIMIT 60000", d["job"], *([emds] if emd_only else []))
            geos.append([(r["cls"], r["b"]) for r in rows])
    now = dt.datetime.now(dt.timezone(dt.timedelta(hours=9)))
    made = now.strftime("%Y.%m.%d %H:%M")

    def build():
        pngs = [_map_png(d["job"], d["sgg"], g, _emd_counts(d)) for d, g in zip(parts, geos)]
        blob, name = _docx(parts, pngs, made)
        return blob, name, pngs
    blob, name, pngs = await asyncio.to_thread(build)
    d0 = config.ARTIFACT_DIR / ctx.run_id
    d0.mkdir(parents=True, exist_ok=True)
    (d0 / name).write_bytes(blob)
    (d0 / "draft.docx").write_bytes(blob)
    if pngs and pngs[0]:
        (d0 / "ai-report-map.png").write_bytes(pngs[0])
    tot = sum(int(r["n"]) for d in parts for r in d["rows"])
    n_emd = len({(d["job"], r["emd"]) for d in parts for r in d["rows"] if r["emd"]})
    out.env("total", "보고서의 AI 분석 결과 수(지도 서비스 결과 층과 같은 행 · 지운 결과 제외)", env(tot, "count", "inferred", "detections"))
    out.env("emds", "보고서에 실린 읍면동 수", env(n_emd, "곳", "inferred", "detections"))
    out.blocks.append({"type": "file", "label": name, "href": file_href(ctx.run_id, name)})
    if pngs and pngs[0]:
        out.blocks.append({"type": "image", "src": file_href(ctx.run_id, "ai-report-map.png"), "caption": "보고서 지도 그림", "tag": "AI 분석 결과"})
    ctx.state["artifact"] = {"docx_url": f"/api/v1/agent/runs/{ctx.run_id}/draft.docx", "href": file_href(ctx.run_id, name), "filename": name,
                             "place": parts[0]["region"]}
    out.data = {"보고서": name, "대상": [f"{d['region']} {d['service']}" for d in parts], "고른 곳": why}
    head = parts[0]
    where = f"{head['region']} {head['service']}" if len(parts) == 1 else f"켜진 층 {len(parts)}개"
    from_ = {"켜진 층": "지금 지도에 켜진 층으로 ", "화면 지역": "지금 보는 지역으로 "}.get(why, "")
    out.answer = (f"{from_}{where} AI 분석 결과 보고서 초안을 만들었습니다 — 결과 {{{{total}}}}건 · 읍면동 {{{{emds}}}}곳. "
                  "지도 그림 · 영상 시점 · AI 모델 정확도를 함께 실었습니다.")
    return out


HANDLERS = {"map_filter": map_filter, "ai_report": ai_report}
_route_front()
