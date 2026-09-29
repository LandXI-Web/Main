"""상태기계(6 전이 · 역방향 409 · client_id 멱등) → audit_log · events 표 · SSE 발행 · 보고서 docx."""
import asyncio
import io
import json
import uuid

import httpx
import pytest

from survey import db as S

# 테스트는 C 등급 R2 뒤쪽(점수 최하위권) 3건만 건드리고 끝나면 원복한다
PICK_SQL = ("SELECT id FROM survey_findings WHERE rule='R2' AND priority='C' AND state='open' AND NOT demo AND tenant_id='namwon' "
            "ORDER BY rank DESC LIMIT 3")


@pytest.fixture()
def picks():
    with S.pg() as c:
        S.lx_tx(c)
        ids = [r[0] for r in c.execute(PICK_SQL).fetchall()]
    yield ids
    with S.pg() as c:
        S.lx_tx(c)
        c.execute("DELETE FROM survey_finding_events WHERE finding_id = ANY(%s)", (ids,))
        c.execute("UPDATE survey_findings SET state='open', assignee=NULL, planned_for=NULL, reason=NULL, updated_at=NULL, updated_by=NULL, "
                  "demo=false WHERE id = ANY(%s)", (ids,))
        c.execute("DELETE FROM audit_log WHERE action='finding.state' AND subject = ANY(%s)", (ids,))
        c.commit()


def _post(api, h, fid, **body):
    body.setdefault("client_id", "pytest-" + uuid.uuid4().hex)
    return httpx.post(f"{api}/survey/findings/{fid}/state", json=body, headers=h, timeout=30)


def test_state_machine_six_transitions(api, h_nw, picks):
    a, b, c = picks
    # 1 open→assigned  2 assigned→inspected  3 inspected→closed
    r = _post(api, h_nw, a, state="assigned", assignee="토지정보과 1팀", planned_for="2026-10-05")
    assert r.status_code == 200 and r.json()["state"] == "assigned" and r.json()["assignee"] == "토지정보과 1팀"
    assert _post(api, h_nw, a, state="inspected").json()["state"] == "inspected"
    # 4 역방향 inspected→assigned 409
    rr = _post(api, h_nw, a, state="assigned")
    assert rr.status_code == 409 and rr.json()["error"]["code"] == "finding_state_invalid"
    assert rr.json()["error"]["detail"]["allowed"] == ["closed"]
    assert _post(api, h_nw, a, state="closed", verdict="violation").json()["state"] == "closed"
    # 5 open→dismissed(사유 필수)
    assert _post(api, h_nw, b, state="dismissed").status_code == 400
    assert _post(api, h_nw, b, state="dismissed", reason="농업용 창고(현장 사진)").json()["state"] == "dismissed"
    # 6 assigned→dismissed
    assert _post(api, h_nw, c, state="assigned").status_code == 200
    assert _post(api, h_nw, c, state="dismissed", reason="건축물대장 확인 — 적법").json()["state"] == "dismissed"
    # 종결 뒤 변경 금지
    assert _post(api, h_nw, a, state="dismissed", reason="x").status_code == 409
    with S.pg() as cn:
        S.lx_tx(cn)
        n_ev = cn.execute("SELECT count(*) FROM survey_finding_events WHERE finding_id = ANY(%s)", (picks,)).fetchone()[0]
        n_au = cn.execute("SELECT count(*) FROM audit_log WHERE action='finding.state' AND subject = ANY(%s)", (picks,)).fetchone()[0]
    assert n_ev == 6 and n_au == 6


def test_idempotent_client_id(api, h_nw, picks):
    fid = picks[0]
    cid = "pytest-idem-" + uuid.uuid4().hex
    r1 = _post(api, h_nw, fid, state="assigned", client_id=cid)
    r2 = _post(api, h_nw, fid, state="assigned", client_id=cid)
    assert r1.status_code == 200 and r2.status_code == 200 and r2.json().get("idempotent") is True
    r3 = _post(api, h_nw, picks[1], state="assigned", client_id=cid)
    assert r3.status_code == 409
    with S.pg() as cn:
        S.lx_tx(cn)
        assert cn.execute("SELECT count(*) FROM survey_finding_events WHERE client_id=%s", (cid,)).fetchone()[0] == 1


def test_write_guards(api, h_sales, h_gj, h_staff, picks):
    assert _post(api, h_sales, picks[0], state="assigned").status_code == 403
    assert _post(api, h_gj, picks[0], state="assigned").status_code == 404          # 타 기관: 행이 없다(RLS)
    assert httpx.post(f"{api}/survey/findings/{picks[0]}/state", json={"state": "assigned", "client_id": "x"}).status_code == 401
    r = _post(api, h_staff, picks[0], state="assigned")                               # LX 직원 = 시연 쓰기
    assert r.status_code == 200 and r.json()["basis"] == "demo" and "자동 원복" in r.json()["demo_note"]


def test_sse_published_to_tenant_and_ops(api, h_nw, picks):
    import redis as _redis
    from landxi_api import config
    rc = _redis.Redis.from_url(config.REDIS_URL, decode_responses=True)
    last_t = (rc.xrevrange("events:tenant:namwon", count=1) or [["0-0"]])[0][0]
    last_o = (rc.xrevrange("ops:events", count=1) or [["0-0"]])[0][0]
    assert _post(api, h_nw, picks[0], state="assigned").status_code == 200
    t = [json.loads(f["data"]) for _, f in rc.xrange("events:tenant:namwon", min=f"({last_t}") if f.get("event") == "finding.state"]
    o = [json.loads(f["data"]) for _, f in rc.xrange("ops:events", min=f"({last_o}") if f.get("event") == "finding.state"]
    assert any(d["id"] == picks[0] and d["from"] == "open" and d["to"] == "assigned" for d in t)
    assert any(d["id"] == picks[0] for d in o)


def test_publish_fake_bus(monkeypatch):
    from landxi_api import survey as SV
    sent = []

    class Fake:
        async def xadd(self, k, f, **kw):
            sent.append((k, f["event"], json.loads(f["data"])))

        async def expire(self, *a):
            pass

    async def fake_redis():
        return Fake()

    monkeypatch.setattr(SV, "redis", fake_redis)
    asyncio.run(SV._publish("namwon", {"id": "f_R1_x", "from": "open", "to": "assigned"}))
    assert [s[0] for s in sent] == ["events:tenant:namwon", "ops:events"] and all(s[1] == "finding.state" for s in sent)


def test_report_docx_rows_and_phrase():
    from docx import Document
    from survey import report
    blob = report.build_draft("52190450", "R1", 20, fmt="docx")
    d = Document(io.BytesIO(blob))
    top_tbl = [t for t in d.tables if t.rows[0].cells[0].text == "No"][0]
    assert len(top_tbl.rows) == 1 + 20
    txt = report.docx_text(blob)
    assert S.FIXED_PHRASE in txt and "[추정 초기값]" in txt and "⑤ 법적 근거" in txt
    assert "[법령 확인 · 2차 RAG]" not in txt and report.dev_words(txt) == [], report.dev_words(txt)
    heads = [c.text for t in d.tables for c in t.rows[0].cells]
    assert not [h for h in heads if "성명" in h or "소유자" in h]
    j = report.build_draft("52190450", "R1", 20)
    assert j["narrative_by"] == "rule" and len(j["top_items"]) == 20
    assert j["citations"][0]["label"].endswith("아영면 연속지적 필지 수")          # [1] = 요청 지역
    assert report.filename("아영면").startswith("실태조사_초안_아영면_")


def test_report_narrative_citation_check():
    from survey import report
    ok = report.build_draft("52190450", "R1", 5, narrative={"overview": ["아영면 후보는 387건이다 [2]."], "actions": ["A등급부터 배정한다 [3]."]})
    assert ok["narrative_by"] == "llm" and ok["narrative_source"].startswith("AI 가 작성한 초안")
    bad = report.build_draft("52190450", "R1", 5, narrative={"overview": ["아영면 후보는 많다. 인용이 없다."]})
    assert bad["narrative_by"] == "rule" and bad["narrative_rejected"]
    rng = report.build_draft("52190450", "R1", 5, narrative="숫자는 [99] 이다.")
    assert rng["narrative_by"] == "rule"


def test_report_cite_meaning_mismatch_dropped():
    """평균 신뢰도 문장에 '후보 필지 수' 인용 → 그 문장만 뺀다(나머지 서술은 산다)."""
    from survey import report
    j0 = report.build_draft("52190450", "R1", 5)
    by = {c["kind"]: c["n"] for c in j0["citations"]}
    narr = {"overview": [f"1위 필지의 평균 신뢰도는 높다 [{by['suspect_parcels']}].", f"후보는 {j0['citations'][1]['value']['value']:,}건이다 [{by['suspects']}]."]}
    j = report.build_draft("52190450", "R1", 5, narrative=narr)
    assert j["narrative_by"] == "llm" and len(j["narrative"]["overview"]) == 1 and "신뢰도" not in j["narrative"]["overview"][0]
    assert any("필지별 값" in x for x in j["narrative_dropped"])
    wrong = report.cite_mismatch(f"후보 필지는 {j0['citations'][3]['value']['value']:,}필지이다 [1].", [1], {c["n"]: c for c in j0["citations"]})
    assert wrong and "인용" in wrong


def _two_regions():
    """서로 다른 시도의 실태조사 시군구 두 곳(데이터에서 — 지역 고정값 없음)."""
    with S.pg() as c:
        S.lx_tx(c)
        rows = c.execute("SELECT sgg_cd, name, sido FROM survey_sgg WHERE state='done' AND findings > 0 ORDER BY findings DESC").fetchall()
    out, seen = [], set()
    for cd, nm, sido in rows:
        if sido not in seen:
            out.append((cd, nm))
            seen.add(sido)
    return out[:2]


def test_report_sgg_level_two_sido_region_cited_no_dev_words():
    from survey import report
    regs = _two_regions()
    assert len(regs) == 2
    for cd, nm in regs:
        j = report.build_draft(cd, None, 10)
        assert j["level"] == "sgg" and j["citations"][0]["label"] == f"{nm} 연속지적 필지 수"          # [1] = 요청 지역
        with S.pg() as c:
            S.lx_tx(c)
            want = c.execute("SELECT findings FROM survey_sgg WHERE sgg_cd=%s", (cd,)).fetchone()[0]
        assert j["citations"][1]["value"]["value"] == want                                            # 의심 건수 한 출처
        txt = report.docx_text(report.build_draft(cd, None, 10, fmt="docx"))
        assert report.dev_words(txt) == [], (cd, report.dev_words(txt))
        names = ["광주전남특별시", "전남광주통합특별시", "전라남도", "광주광역시", "전북특별자치도", "전라북도", "충청북도", "충청남도", "경기도",
                 "강원특별자치도", "경상북도", "경상남도", "제주특별자치도", "서울특별시", "부산광역시"]
        sidos = {x for x in names if x in txt}
        assert len(sidos) == 1, sidos                                                                  # 시도 이름 한 가지
        assert j["sido"] and j["sido"] in txt and "전남광주통합특별시" not in txt


def test_report_api_docx(api, h_nw):
    r = httpx.get(f"{api}/survey/reports/draft", params={"emd_cd": "52190450", "rule": "R1", "top": 20, "format": "docx"}, headers=h_nw, timeout=60)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    assert r.content[:2] == b"PK"
