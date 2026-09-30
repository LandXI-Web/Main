"""실태조사 보고서 초안 — 서식 `survey-emd` · python-docx · LLM 없이 닫힌다.

    build_draft(code, rule=None, top=20, narrative=None, fmt="dict"|"docx", realm="lx", tenant="") -> dict | bytes

code = 읍면동(법정동 8·10자리) 또는 시군구(5자리) — 대상은 호출자가 정한다(지역 고정값 없음).
① 개요 ② 집계표(규칙 × 등급 × 상태) ③ 의심 상위 N ④ 근거 영상(그 지역 영상의 연도·해상도) ⑤ 법적 근거(법령 원문 조문 인용) ⑥ 조치 제안.
모든 수치는 봉투(as_json) · 문서에는 사용자 말 꼬리표. 고정 문구 `AI 추론 · 검수 전 · 현장 확인 전 · 위법 판정 아님`.
인용 [n] 은 요청 지역의 사실만 가리킨다([1] = 그 지역 연속지적 필지 수). 시도 이름은 sido_label 한 가지.
narrative(에이전트 서술): 문장마다 `[n]` 인용 검사 — 없거나 범위 밖이면 서술을 버리고 정형 문장. 인용 대상과 문장 뜻이 어긋난 문장은 뺀다.
성명 열 없음(연속지적에 소유 정보 없음). 문서에 개발 정보(표 이름·코드·경로·모델 이름)를 쓰지 않는다.
"""
from __future__ import annotations

import datetime as dt
import io
import re

from . import rules as R
from .db import (AI_DRAFT, AS_OF, AUTO_DRAFT, FIXED_PHRASE, IMG23, IMG25, LEDGER, REPORT_SRC, RULE_IDS, STATE_SRC, STATES,
                 imagery_label)
from .explain import env

KST = dt.timezone(dt.timedelta(hours=9))
STATE_KO = {"open": "배정 전", "assigned": "현장조사 배정", "inspected": "현장조사 완료", "closed": "종결", "dismissed": "오탐"}
NOT_FOUND = "법령 데이터에 없습니다"
# 규칙별 근거 조문(법령 · 조 · 항) — 원문은 법령 색인에서 그대로 가져온다(색인에 없으면 '법령 데이터에 없습니다').
LAW_REFS = {
    "R1": [("농지법", "34", 1), ("건축법", "11", 1)],
    "R2": [("농지법", "10", 1), ("실태조사 요령", "Ⅱ-1", None)],
    "R3": [("국토의 계획 및 이용에 관한 법률", "56", 1), ("국토의 계획 및 이용에 관한 법률", "76", 1)],
    "R4": [("농지법", "34", 1), ("농지법", "42", 1)],
    "R5": [("산지관리법", "14", 1)],
    "R6": [("건축법", "11", 1), ("도로법", "61", 1), ("하천법", "33", 1)],
}
SENT_RE = re.compile(r"[^.!?。…]+[.!?。…]?")
# 대장 대조 규칙(L1–L3) 이름 — 문서에 코드가 남지 않게(규칙 이름만)
LEDGER_NM = {"L1": "대장 농지 위 건물", "L2": "경작 흔적 없음", "L3": "허가 필지 건물 없음"}
CODE_RE = re.compile(r"(?<![A-Za-z0-9])([RL])([1-9])(?![0-9A-Za-z])")


def rule_name(rid: str | None) -> str:
    """규칙 코드 → 사용자 말(규칙 이름). 모르는 코드는 빈 글."""
    if not rid:
        return ""
    if rid in LEDGER_NM:
        return LEDGER_NM[rid]
    try:
        return R.definitions()[rid]["name"]
    except Exception:  # noqa: BLE001
        return ""


def plain_rules(s: str) -> str:
    """글 속 규칙 코드(R1–R6 · L1–L3)를 규칙 이름으로 — 'R1–R6' → '전체 규칙', 'R1 무허가 건축 의심' → '무허가 건축 의심'."""
    if not s:
        return s
    s = re.sub(r"(?<![A-Za-z0-9])R1\s*[–~-]\s*R6(?![0-9])", "전체 규칙", s)

    def one(m):
        nm = rule_name(m.group(1) + m.group(2))
        rest = s[m.end():m.end() + len(nm) + 1].strip()
        return "" if nm and rest.startswith(nm) else (nm or "")
    s = CODE_RE.sub(one, s)
    s = re.sub(r"(?m)^[ \t]+(?=[가-힣(])", "", s)                # 줄 머리 코드를 지운 자리의 빈칸
    return re.sub(r"[ \t]{2,}", " ", s).replace("( ", "(").replace(" )", ")")


CITE_RE = re.compile(r"\[(\d+)\]")
LIST_KW = re.compile(r"신뢰도|근거\s?면적|㎡|필지 대비|점수|순위|1위|상위|지번")
LAW_KW = re.compile(r"조문|「|제\d+조|법령|법률")


class NotFound(Exception):
    pass


def _conn(realm: str, tenant: str):
    import psycopg
    from psycopg.rows import dict_row
    from landxi_api import config
    c = psycopg.connect(config.PG_DSN, autocommit=False, row_factory=dict_row)
    c.execute("SELECT set_config('app.realm', %s, true), set_config('app.tenant_id', %s, true)", (realm or "public", tenant or ""))
    return c


def sido_label(sgg_cd: str | None, sido: str | None = None) -> str | None:
    """시도 이름 한 가지 — regions.sido_label(있으면) · 없으면 기관 표기와 같은 매핑."""
    try:
        from landxi_api import regions as RG
        f = getattr(RG, "sido_label", None)
        if f:
            v = f(sgg_cd)
            if v:
                return v
    except Exception:  # noqa: BLE001
        pass
    return {"전남광주통합특별시": "광주전남특별시"}.get(sido or "", sido)


def _counts(c, sgg_cd: str) -> dict | None:
    """의심 필지 수 · 현장 확인 필요 한 출처 — survey.nation.counts_sync(첫 화면·에이전트·XI맵의 survey_counts 와 같은 식).
    그 함수가 없으면 동기 survey_counts, 그것도 없으면 같은 원천(survey_sgg)."""
    try:
        from psycopg.rows import tuple_row
        from . import nation as N
        v = N.counts_sync(c.cursor(row_factory=tuple_row), sgg_cd)
        if v.get("suspect") is not None:
            return {"suspect": int(v["suspect"]), "field_check": v.get("field_check"), "src": "survey_counts"}
    except Exception:  # noqa: BLE001
        pass
    for mod in ("landxi_api.survey", "agent.tools.survey", "survey.nation", "landxi_api.summary"):
        try:
            import importlib
            import inspect
            f = getattr(importlib.import_module(mod), "survey_counts", None)
            if f and not inspect.iscoroutinefunction(f):
                v = f(sgg_cd)
                if isinstance(v, dict) and v.get("suspect") is not None:
                    return {"suspect": int(v["suspect"]), "field_check": v.get("field_check"), "src": "survey_counts"}
        except Exception:  # noqa: BLE001
            continue
    r = c.execute("SELECT findings, by_priority, state FROM survey_sgg WHERE sgg_cd=%s", (sgg_cd,)).fetchone()
    if not r or r["findings"] is None:
        return None
    bp = r.get("by_priority") or {}
    return {"suspect": int(r["findings"]), "field_check": (bp or {}).get("A") if isinstance(bp, dict) else None, "src": "survey_sgg"}


def collect(code: str, rule: str | None = None, top: int = 20, realm: str = "lx", tenant: str = "") -> dict:
    """초안 데이터(봉투 전 원값) — RLS 는 호출자 realm/tenant 그대로. code = 읍면동(8·10자리) 또는 시군구(5자리)."""
    code = str(code or "").strip()
    level = "sgg" if re.fullmatch(r"\d{5}", code) else "emd"
    with _conn(realm, tenant) as c:
        if level == "sgg":
            sg = c.execute("SELECT sgg_cd, tenant_id, name, sido, imagery, parcels, parcels_as_of, priority_cut, parcels_src ? 'canon' AS canon "
                           "FROM survey_sgg WHERE sgg_cd=%s", (code,)).fetchone()
            if not sg:
                raise NotFound(f"{code} 실태조사 결과가 없습니다(실태조사를 만들지 않은 곳이거나 다른 기관)")
            em = c.execute("SELECT count(*) n, sum(parcels) parcels, sum(area_ha) area_ha FROM survey_emd WHERE coalesce(sgg_cd, left(emd_cd, 5))=%s",
                           (code,)).fetchone()
            e = {"emd_cd": code, "name": sg["name"], "parcels": int(sg["parcels"] or em["parcels"] or 0), "area_ha": float(em["area_ha"] or 0),
                 "sgg_cd": code, "tenant_id": sg["tenant_id"], "n_emd": int(em["n"] or 0)}
            col = "f.sgg_cd"
        else:
            e = c.execute("SELECT emd_cd, name, names, parcels, area_ha, farm_parcels, jimok_top, coalesce(sgg_cd, left(emd_cd, 5)) sgg_cd, tenant_id "
                          "FROM survey_emd WHERE emd_cd=%s", (code,)).fetchone()
            if not e:
                raise NotFound(f"읍면동 {code} 없음(실태조사를 만들지 않은 곳이거나 다른 기관)")
            sg = c.execute("SELECT sgg_cd, tenant_id, name, sido, imagery, parcels_as_of, priority_cut, parcels_src ? 'canon' AS canon "
                           "FROM survey_sgg WHERE sgg_cd=%s", (e["sgg_cd"],)).fetchone()
            col = "f.emd_cd"
        org = c.execute("SELECT name->>'ko' AS ko FROM tenants WHERE id=%s", ((sg or {}).get("tenant_id") or e["tenant_id"],)).fetchone()
        w = f"{col}=%s" + (" AND f.rule=%s" if rule else " AND f.rule = ANY(%s)")
        a = (code, rule) if rule else (code, list(RULE_IDS))
        agg = c.execute(f"SELECT f.rule, f.priority, f.state, count(*) n FROM survey_findings f WHERE {w} GROUP BY 1,2,3", a).fetchall()
        # 같은 필지는 한 줄(점수가 가장 높은 행) · 해당 규칙은 그 필지의 규칙 전부
        rows = c.execute(f"SELECT * FROM (SELECT DISTINCT ON (f.pnu) f.*, p.jibun, p.ri, "
                         f"(SELECT array_agg(DISTINCT g.rule ORDER BY g.rule) FROM survey_findings g WHERE g.pnu = f.pnu AND g.rule = ANY(%s)) AS rules_all, "
                         f"(SELECT count(*) FROM survey_timeline t, jsonb_array_elements(t.events) ev "
                         f"WHERE t.pnu = f.pnu AND ev->>'kind' = 'change') AS n_change "
                         f"FROM survey_findings f LEFT JOIN survey_parcels p USING (pnu) WHERE {w} "
                         f"ORDER BY f.pnu, f.score DESC, f.rank) x ORDER BY x.score DESC, x.rank LIMIT %s",
                         (list(RULE_IDS),) + a + (top,)).fetchall()
        sp = c.execute(f"SELECT count(DISTINCT f.pnu) FROM survey_findings f WHERE {w}", a).fetchone()["count"]
        cut = c.execute("SELECT value FROM survey_meta WHERE key='priority_cut'").fetchone()
        counts = _counts(c, e["sgg_cd"]) if (level == "sgg" and not rule) else None
    region = _region(e["sgg_cd"], sg)
    region["org"] = (org or {}).get("ko")
    return {"level": level, "emd": dict(e), "rule": rule, "top": top, "agg": [dict(x) for x in agg], "rows": [dict(x) for x in rows],
            "suspect_parcels": sp, "counts": counts,
            "cut": (cut or {}).get("value") if region["canon"] else ((sg or {}).get("priority_cut") or (cut or {}).get("value")),
            "region": region, "at": dt.datetime.now(KST)}


def _region(sgg_cd: str, sg: dict | None) -> dict:
    """시군구 이름 · 시도(한 가지) · 영상 표기(연도·해상도) — 데이터(regions · survey_sgg)에서."""
    from . import nation as N
    try:
        rg = N.region(sgg_cd)
    except Exception:
        rg = {"sgg_cd": sgg_cd, "name": (sg or {}).get("name") or sgg_cd, "sido": (sg or {}).get("sido")}
    sido = sido_label(rg["sgg_cd"], rg.get("sido") or (sg or {}).get("sido"))
    canon = bool((sg or {}).get("canon")) if sg else True
    img = [IMG23, IMG25] if canon else [x for x in [imagery_label((sg or {}).get("imagery"))] if x]
    asof = (sg or {}).get("parcels_as_of") or ""
    return {"sgg_cd": rg["sgg_cd"], "name": rg.get("name"), "sido": sido, "full": " ".join(x for x in (sido, rg.get("name")) if x),
            "canon": canon, "imagery": img, "ledger": LEDGER if canon else (f"연속지적도({asof} 기준)" if asof else "연속지적도")}


def place_of(d: dict) -> str:
    """문서가 가리키는 지역 이름 — 시군구 보고서는 '여수시', 읍면동 보고서는 '여수시 돌산읍'."""
    rg, e = d["region"], d["emd"]
    return rg["name"] if d.get("level") == "sgg" else f"{rg['name']} {e['name']}"


def _table(d: dict) -> dict:
    rules = [d["rule"]] if d["rule"] else RULE_IDS
    t = {r: {"A": 0, "B": 0, "C": 0, **{s: 0 for s in STATES}, "total": 0} for r in rules}
    for x in d["agg"]:
        if x["rule"] not in t:
            continue
        t[x["rule"]][x["priority"]] += x["n"]
        t[x["rule"]][x["state"]] += x["n"]
        t[x["rule"]]["total"] += x["n"]
    return t


# ── 법령 근거 ─────────────────────────────────────────────────────────────
def law_basis(rules: list[str]) -> list[dict]:
    """규칙 → 근거 조문(원문 그대로 · 조·항·시행일). 색인에 없으면 status = '법령 데이터에 없습니다'."""
    try:
        from agent.rag import index as LAW
    except Exception:  # noqa: BLE001
        LAW = None
    out: dict[str, dict] = {}
    for r in rules:
        for act, art, para in LAW_REFS.get(r, []):
            key = f"{act}|{art}|{para}"
            if key in out:
                out[key]["rules"].append(r)
                continue
            hit = None
            if LAW is not None:
                try:
                    hit = LAW.article(act, art, para)
                except Exception:  # noqa: BLE001
                    hit = None
            if hit:
                ref = hit["ref"]
                out[key] = {"rules": [r], "law": ref["law"], "label": ref["label"], "article": ref["article"], "para": ref["para"],
                            "effective": ref["effective"], "text": hit["text"], "status": "조문 인용"}
            else:
                shown = act if not re.fullmatch(r"\d+", art) else f"「{act}」 제{art}조" + (f" 제{para}항" if para else "")
                out[key] = {"rules": [r], "law": act, "label": shown, "article": None, "para": None, "effective": None, "text": None,
                            "status": NOT_FOUND}
    return list(out.values())


# ── 인용 ─────────────────────────────────────────────────────────────────
def citations(d: dict) -> list[dict]:
    """번호 붙은 사실(서술 문장이 [n]으로 인용) — 각 값은 봉투. [1] = 요청 지역 연속지적 필지 수."""
    e = d["emd"]
    t = _table(d)
    place = place_of(d)
    tot = sum(v["total"] for v in t.values())
    if d.get("counts") and not d["rule"]:
        tot = d["counts"]["suspect"]                    # 의심 필지 수 한 출처(survey_counts)
    a_n = sum(v["A"] for v in t.values())
    src = REPORT_SRC
    c = [
        {"n": 1, "kind": "parcels", "label": f"{place} 연속지적 필지 수", "value": env(e["parcels"], "필지", "measured", d["region"]["ledger"])},
        # 인용 라벨(봉투 뜻 · 에이전트 인용 옮기기의 열쇠)에는 규칙 코드를 둔다 — 문서·화면에 찍을 때 plain_rules 가 규칙 이름으로 바꾼다
        {"n": 2, "kind": "suspects", "label": f"{place} 실태조사 대상 후보(의심) 건수" + (f" · {d['rule']}" if d["rule"] else ""),
         "value": env(tot, "count", "inferred", src, "검수 전")},
        {"n": 3, "kind": "grade_a", "label": f"그중 A등급({d['region']['name']} 전체 점수 상위 5%)", "value": env(a_n, "count", "inferred", src, "검수 전")},
        {"n": 4, "kind": "suspect_parcels", "label": f"{place} 후보 필지 수", "value": env(d["suspect_parcels"], "필지", "inferred", src, "검수 전")},
        {"n": 5, "kind": "open", "label": f"{place} 현장조사 배정 전 건수", "value": env(sum(v["open"] for v in t.values()), "count", "recorded", STATE_SRC)},
    ]
    k = 6
    for r in t:
        if t[r]["total"]:
            c.append({"n": k, "kind": f"rule:{r}", "label": f"{place} {r} {R.definitions()[r]['name']} 건수",
                      "value": env(t[r]["total"], "count", "inferred", src, "검수 전")})
            k += 1
    c.append({"n": k, "kind": "list", "label": f"{place} 의심 상위 {len(d['rows'])}건 표(③ · 필지별 근거면적 · 신뢰도)",
              "value": env(len(d["rows"]), "count", "inferred", src, "검수 전")})
    k += 1
    for lb in d.get("_law") or []:
        if lb["status"] != NOT_FOUND:
            c.append({"n": k, "kind": "law", "label": lb["label"], "value": None, "law_ref": {"법령": lb["law"], "조": lb["article"], "항": lb["para"],
                                                                                             "시행일": lb["effective"]}})
            k += 1
    return c


def _nums(s: str) -> set[int]:
    return {int(x.replace(",", "")) for x in re.findall(r"\d{1,3}(?:,\d{3})+|\d+", s)}


def cite_mismatch(sentence: str, refs: list[int], by_n: dict[int, dict]) -> str | None:
    """문장과 인용 대상이 맞는지 — 어긋나면 이유, 맞으면 None.
    ① 필지별 값(신뢰도·근거면적·순위) 문장은 목록·필지 인용만 ② 조문 문장은 법령 인용만
    ③ 문장의 숫자가 인용한 사실 값과 하나도 안 맞는데 다른 사실 값과 맞으면 잘못 짚은 인용."""
    kinds = {(by_n.get(n) or {}).get("kind") for n in refs}
    body = CITE_RE.sub("", sentence)
    if LIST_KW.search(body) and not (kinds & {"list", "parcel"}):
        return "필지별 값 문장에 집계 인용"
    if LAW_KW.search(body) and "law" not in kinds and re.search(r"제\d+조", body):
        return "조문 문장에 법령 인용 없음"
    nums = _nums(body)
    if nums:
        vals = {n: (by_n[n].get("value") or {}).get("value") for n in by_n}
        mine = {vals.get(n) for n in refs}
        if not (nums & {v for v in mine if isinstance(v, int)}):
            other = [n for n, v in vals.items() if isinstance(v, int) and v in nums and n not in refs and v > 9]
            if other:
                return f"숫자가 인용 [{other[0]}] 의 값"
    return None


def check_narrative(narrative, n_cites: int, cites: list[dict] | None = None) -> tuple[dict | None, list[str]]:
    """문장마다 [n] 인용(1..n_cites) — 통과하면 정규화한 dict, 아니면 (None, 사유).
    cites 를 주면 인용-문장 뜻 검사도 한다(어긋난 문장은 빼고 사유에 '뺌:' 으로 남긴다 · 서술 전체는 살린다)."""
    if narrative is None:
        return None, ["서술 없음 — 정형 문장"]
    if isinstance(narrative, str):
        narrative = {"overview": [narrative]}
    elif isinstance(narrative, list):
        narrative = {"overview": narrative}
    by_n = {c["n"]: c for c in cites or []}
    out, bad, dropped = {}, [], []
    for sec, paras in narrative.items():
        if sec not in ("overview", "findings", "actions"):
            bad.append(f"알 수 없는 절 {sec}")
            continue
        paras = [paras] if isinstance(paras, str) else list(paras or [])
        keep = []
        for para in paras:
            ok_sents = []
            for s in (m.group(0).strip() for m in SENT_RE.finditer(para or "")):
                if not s:
                    continue
                refs = [int(x) for x in CITE_RE.findall(s)]
                if not refs:
                    bad.append(f"인용 없음: {s[:40]}")
                elif any(r < 1 or r > n_cites for r in refs):
                    bad.append(f"인용 번호 범위 밖: {s[:40]}")
                elif by_n and (why := cite_mismatch(s, refs, by_n)):
                    dropped.append(f"뺌: {why}: {s[:40]}")
                    continue
                ok_sents.append(s)
            if ok_sents:
                keep.append(" ".join(ok_sents))
        out[sec] = keep
    if bad:
        return None, bad
    return out, dropped


def _default_narrative(d: dict, cites: list[dict]) -> dict:
    v = {c["n"]: (c["value"] or {}).get("value") for c in cites}
    by_kind = {c["kind"]: c["n"] for c in cites}
    place = place_of(d)
    rules = [c for c in cites if c["kind"].startswith("rule:")]
    rules_txt = ", ".join(f"{c['label'].replace(place + ' ', '').split(' 건수')[0]} {c['value']['value']:,}건 [{c['n']}]" for c in rules)
    img = (d["region"]["imagery"] or ["영상"])[0]
    law_n = next((c["n"] for c in cites if c["kind"] == "law"), None)
    out = {
        "overview": [f"{place}의 연속지적 {v[1]:,}필지를 {img} AI 분석 결과와 대조했습니다 [1].",
                     f"실태조사 대상 후보(의심)는 {v[2]:,}건이며 후보 필지는 {v[4]:,}필지입니다 [2][4]."],
        "findings": ([f"규칙별로는 {rules_txt}입니다."] if rules else [])
                    + [f"점수 상위 5%인 A등급은 {v[3]:,}건으로 현장 확인 우선 대상입니다 [3].",
                       f"필지별 근거면적과 신뢰도는 ③ 의심 상위 목록에 실었습니다 [{by_kind['list']}]."],
        "actions": [f"현장조사 배정 전 {v[5]:,}건 가운데 A등급부터 현장조사 담당을 배정합니다 [5][3].",
                    "현장 확인 전에 건축물대장과 허가 대장을 먼저 대조해 오탐을 줄입니다 [2]."],
    }
    if law_n:
        out["actions"].append(f"관련 조문은 ⑤ 법적 근거에 원문 그대로 실었습니다 [{law_n}].")
    return out


def as_json(d: dict, narrative=None) -> dict:
    e = d["emd"]
    t = _table(d)
    rules = [d["rule"]] if d["rule"] else [r for r in RULE_IDS if t.get(r, {}).get("total")] or RULE_IDS
    d["_law"] = law_basis(rules)
    cites = citations(d)
    narr, why = check_narrative(narrative, len(cites), cites)
    if narr:                                              # 서술 속 규칙 코드도 규칙 이름으로(사용자 말)
        narr = {k: [plain_rules(x) for x in v] for k, v in narr.items()}
    th = R.default_thresholds()
    rg = d["region"]
    place = place_of(d)
    tops = d["rows"]
    n_multi = sum(1 for x in tops if x["n_change"])
    return {
        "template": "survey-emd", "level": d.get("level", "emd"), "emd_cd": e["emd_cd"], "emd": e["name"], "rule": d["rule"], "limit": d["top"],
        "org": rg.get("org"), "sgg_cd": rg["sgg_cd"], "sgg": rg["name"], "sido": rg.get("sido"), "region_full": rg.get("full"), "place": place,
        "title": f"{place} 실태조사 대상 후보 보고서(초안)", "fixed": FIXED_PHRASE,
        "overview": {"org": rg.get("org"), "sgg": rg["name"], "emd": e["name"], "parcels": env(e["parcels"], "필지", "measured", rg["ledger"]),
                     "area_ha": env(e["area_ha"], "ha", "measured", rg["ledger"]),
                     "imagery": rg["imagery"], "ledger": rg["ledger"],
                     "rules": [{"id": r, "name": R.definitions()[r]["name"], "condition": R.condition_text(r, th)} for r in ([d["rule"]] if d["rule"] else RULE_IDS)],
                     "thresholds_note": "판단 기준값은 모두 [추정 초기값]입니다 — 법령 기준이 아니며 현장조사 결과로 보정합니다"},
        "table": [{"rule": r, "name": R.definitions()[r]["name"],
                   **{k: env(v[k], "count", "inferred", REPORT_SRC, "검수 전") for k in ("A", "B", "C", "total")},
                   "by_state": {s: env(v[s], "count", "recorded", STATE_SRC) for s in STATES}} for r, v in t.items()],
        "top_items": [{"order": i + 1, "id": x["id"], "pnu": x["pnu"], "addr": x["addr"], "jibun": x.get("jibun"), "jimok": x["jimok"],
                       "rule": x["rule"], "rules": [rule_name(r) for r in (x.get("rules_all") or [x["rule"]]) if rule_name(r)] or [rule_name(x["rule"])],
                       "priority": x["priority"], "state": x["state"],
                       "rank": env(x["rank"], "count", "estimate", REPORT_SRC),
                       "evid_m2": env(round(x["evid_m2"], 1), "m2", "inferred", REPORT_SRC, "검수 전"),
                       "evid_pct": env(round(100 * x["evid_m2"] / max(x["parcel_m2"] or 1, 1), 1), "%", "inferred", "근거면적 ÷ 필지면적", "검수 전"),
                       "conf": env(None if x["conf"] is None else round(x["conf"], 3), "ratio", "inferred", REPORT_SRC, "검수 전"),
                       "imagery": "여러 시점 비교 가능" if x["n_change"] else (imagery_label(x.get("img_date")) or "한 시점"),
                       "img_date": imagery_label(x.get("img_date"))} for i, x in enumerate(tops)],
        "imagery_note": {"multi": env(n_multi, "count", "recorded", "필지 이력"), "single": env(len(tops) - n_multi, "count", "recorded", "필지 이력"),
                         "canon": rg["canon"], "imagery": rg["imagery"],
                         "note": ("드론 여러 시점 영상은 일부 구역만 덮습니다 — 그 밖은 두 시점(항공 · 드론) 비교입니다" if rg["canon"]
                                  else "AI 분석 영상 한 시점 — 시점 비교(변화)는 두 번째 영상 등록 후")},
        "law": [{"rules": x["rules"], "rule_names": [rule_name(r) for r in x["rules"]], "label": x["label"], "text": x["text"], "effective": x["effective"], "status": x["status"],
                 "law_ref": {"법령": x["law"], "조": x["article"], "항": x["para"], "시행일": x["effective"]} if x["article"] else None}
                for x in d["_law"]],
        "actions": {"field_targets": env(sum(v["A"] for v in t.values()), "count", "inferred", REPORT_SRC, "A등급 우선 · 검수 전"),
                    "checklist": ["건축물대장 대조(대조 전)", "농지·산지전용 허가 대장 대조(기관 대장 올리기 후)",
                                  "현장 사진·측량 확인", "결과 입력: 현장조사 완료 → 종결 또는 오탐(사유 기록)"],
                    "fixed": FIXED_PHRASE},
        "citations": cites,
        "narrative": narr or _default_narrative(d, cites),
        "narrative_by": "llm" if narr else "rule",
        "narrative_source": AI_DRAFT if narr else AUTO_DRAFT,
        "narrative_rejected": why if (narrative is not None and not narr) else [],
        "narrative_dropped": why if narr else [],
        "as_of": AS_OF, "generated": d["at"].isoformat(timespec="seconds"), "source": REPORT_SRC,
        "filename": filename(place.replace(" ", "_"), d["at"]),
    }


def filename(emd_name: str, at: dt.datetime | None = None) -> str:
    at = at or dt.datetime.now(KST)
    return f"실태조사_초안_{emd_name}_{at:%Y%m%d}.docx"


# ─────────────────────────── docx ───────────────────────────
TAG = {"inferred": "AI 추론 · 검수 전", "recorded": "기록", "measured": "실측", "estimate": "추정"}


def _font(run, size=None, bold=None, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name = "맑은 고딕"
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), "맑은 고딕")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _p(doc, text="", size=10, bold=False, color=None, align=None, space_after=4):
    from docx.shared import Pt
    p = doc.add_paragraph()
    if text:
        _font(p.add_run(text), size, bold, color)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def _shade(cell, hex_fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _tbl(doc, header: list[str], rows: list[list], widths=None, size=8.5, head_fill="E8ECEF"):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), size, True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(c, head_fill)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            s = "" if v is None else (f"{v:,}" if isinstance(v, int) else (f"{v:,.1f}" if isinstance(v, float) else str(v)))
            _font(cells[i].paragraphs[0].add_run(s), size)
            if isinstance(v, (int, float)):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def render_docx(d: dict, narrative=None) -> tuple[bytes, str]:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    j = as_json(d, narrative)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10)
    from docx.oxml.ns import qn
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    doc.core_properties.title = j["title"]
    doc.core_properties.author = j.get("org") or "Land-XI"
    doc.core_properties.comments = ""
    # 머리
    hdr = sec.header.paragraphs[0]
    _font(hdr.add_run(" · ".join(x for x in ["Land-XI 실태조사", j.get("org"), j["place"], "초안"] if x)), 8, False, "6B7780")
    _p(doc, j["title"], 17, True, "14202A", space_after=2)
    _p(doc, f"작성 {j['generated'][:16].replace('T', ' ')} · 기관 {j.get('org') or '—'} · 대상 {' '.join(x for x in (j.get('sido'), j['place']) if x)}"
            + (f" · 규칙 {rule_name(j['rule'])}" if j["rule"] else " · 전체 규칙") + f" · 상위 {j['limit']}건", 9, False, "56626B")
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    c = box.rows[0].cells[0]
    _shade(c, "FFF4E5")
    c.text = ""
    _font(c.paragraphs[0].add_run(FIXED_PHRASE), 11, True, "9A3412")
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(c.add_paragraph().add_run("현장조사 대상 후보 목록입니다. 판단 기준값은 [추정 초기값]이며 법령 기준이 아닙니다. "
                                    "수치 꼬리표: AI 추론 · 검수 전 / 기록 / 실측 / 추정."), 8.5, False, "7C2D12")
    # ① 개요
    _p(doc, "", space_after=2)
    _p(doc, "① 개요", 12.5, True, "14202A")
    o = j["overview"]
    _tbl(doc, ["항목", "내용", "꼬리표"], [
        ["기관", o.get("org") or "—", ""],
        ["대상", " ".join(x for x in (j.get("sido"), j["place"]) if x), ""],
        ["필지 수", f"{o['parcels']['value']:,}필지 · {o['area_ha']['value']:,.1f} ha", TAG["measured"]],
        ["영상", " / ".join(o["imagery"]) or "—", "기록"],
        ["대장", o["ledger"] + " (연속지적도 표기 지목 · 토지대장 원본과 다를 수 있음)", "기록"],
        ["규칙", "; ".join(f"{r['name']}: {r['condition']}" for r in o["rules"]), TAG["estimate"]],
        ["기준값", o["thresholds_note"], TAG["estimate"]],
    ], widths=[3.0, 11.6, 2.8])
    for s in j["narrative"].get("overview", []):
        _p(doc, plain_rules(s), 9.5, space_after=2)
    # ② 집계표
    _p(doc, "", space_after=2)
    _p(doc, "② 집계표 — 규칙 × 등급 × 상태", 12.5, True, "14202A")
    rows = []
    for r in j["table"]:
        rows.append([r["name"], r["A"]["value"], r["B"]["value"], r["C"]["value"], r["total"]["value"],
                     *[r["by_state"][s]["value"] for s in STATES]])
    tot = ["계"] + [sum(x[i] for x in rows) for i in range(1, 10)]
    _tbl(doc, ["규칙", "A", "B", "C", "계", *[STATE_KO[s] for s in STATES]], rows + [tot],
         widths=[4.6, 1.2, 1.2, 1.4, 1.4, 1.5, 1.4, 1.5, 1.2, 1.2])
    _p(doc, f"꼬리표: 등급·건수 = AI 추론 · 검수 전 / 상태 = 기록. 등급 A = {j['sgg']} 전체 점수 상위 5%, B = 다음 20%.",
       8, False, "56626B")
    for s in j["narrative"].get("findings", []):
        _p(doc, plain_rules(s), 9.5, space_after=2)
    # ③ 의심 상위 N
    _p(doc, "", space_after=2)
    _p(doc, f"③ 의심 상위 {len(j['top_items'])}건 (점수순)", 12.5, True, "14202A")
    strip = [x for x in ((j.get("region_full") or "") + " ", (j.get("sgg") or "") + " ") if x.strip()]

    def short(a):
        a = a or ""
        for s in strip:
            a = a.replace(s, "")
        return re.sub(r"^\S+(특별시|광역시|특별자치도|특별자치시|도)\s", "", a)
    trs = [[x["order"], x["pnu"], short(x["addr"]), x["jimok"], " · ".join(x["rules"]), x["priority"],
            x["evid_m2"]["value"], f"{x['evid_pct']['value']:.0f}%", "-" if x["conf"]["value"] is None else f"{x['conf']['value']:.2f}",
            STATE_KO.get(x["state"], x["state"])] for x in j["top_items"]]
    _tbl(doc, ["No", "PNU", "소재지", "지목", "해당 규칙", "등급", "근거면적㎡", "필지 대비", "신뢰도", "상태"], trs,
         widths=[0.8, 3.2, 3.2, 0.9, 2.6, 0.8, 1.6, 1.3, 1.1, 1.3], size=7.5)
    _p(doc, "같은 필지가 여러 규칙에 걸리면 점수가 가장 높은 한 줄로 적고, '해당 규칙' 칸에 규칙을 모두 적었습니다.", 8, False, "56626B")
    _p(doc, "꼬리표: 근거면적·비율·신뢰도 = AI 추론 · 검수 전. 소유자 성명은 싣지 않습니다(연속지적도에 없음).", 8, False, "56626B")
    # ④ 근거 영상
    _p(doc, "", space_after=2)
    _p(doc, "④ 근거 영상", 12.5, True, "14202A")
    im = j["imagery_note"]
    _tbl(doc, ["구분", "건수(상위 목록 안)", "영상"], [
        ["여러 시점 비교 가능", im["multi"]["value"], " / ".join(im.get("imagery") or []) or "—"],
        ["한 시점", im["single"]["value"], (im.get("imagery") or ["—"])[0]],
    ], widths=[4.0, 3.0, 10.4])
    _p(doc, im["note"], 8, False, "56626B")
    # ⑤ 법적 근거 — 법령 원문 조문 인용
    _p(doc, "", space_after=2)
    _p(doc, "⑤ 법적 근거", 12.5, True, "14202A")
    _tbl(doc, ["규칙", "조문", "시행일", "원문"],
         [[" · ".join(x["rule_names"]), x["label"].split(" · 시행")[0], x["effective"] or "—", x["text"] if x["text"] else x["status"]] for x in j["law"]],
         widths=[2.8, 3.8, 1.8, 9.0], size=7.5)
    _p(doc, "조문은 국가법령정보센터 공개 원문 그대로입니다. 위반 여부는 현장 확인 후 담당자가 판단합니다.", 8, False, "56626B")
    # ⑥ 조치 제안
    _p(doc, "", space_after=2)
    _p(doc, "⑥ 조치 제안", 12.5, True, "14202A")
    a = j["actions"]
    _p(doc, f"현장조사 우선 대상: A등급 {a['field_targets']['value']:,}건 (AI 추론 · 검수 전)", 10, True)
    for s in a["checklist"]:
        _p(doc, "□ " + s, 9.5, space_after=1)
    for s in j["narrative"].get("actions", []):
        _p(doc, plain_rules(s), 9.5, space_after=2)
    _p(doc, a["fixed"], 10.5, True, "9A3412", space_after=6)
    # 인용
    _p(doc, "인용", 10, True, "14202A")
    for c in j["citations"]:
        v = c.get("value")
        if v is None:
            _p(doc, f"[{c['n']}] {c['label']} — 법령 원문", 8, False, "56626B", space_after=0)
            continue
        unit = {"count": "건", "필지": "필지"}.get(v["unit"], " " + v["unit"])
        _p(doc, plain_rules(f"[{c['n']}] {c['label']}: {v['value']:,}{unit} — {TAG.get(v['basis'], v['basis'])} · {v['source']}"), 8, False, "56626B", space_after=0)
    _p(doc, f"서술: {j['narrative_source']}", 8, False, "56626B")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), j["filename"]


def build_draft(emd_cd: str, rule: str | None = None, top: int = 20, narrative=None, fmt: str = "dict",
                realm: str = "lx", tenant: str = ""):
    """에이전트·화면이 부르는 한 함수 — fmt 'dict' → 봉투 JSON · 'docx' → bytes. realm/tenant 는 호출자의 RLS."""
    d = collect(emd_cd, rule, top, realm, tenant)
    if fmt == "docx":
        return render_docx(d, narrative)[0]
    return as_json(d, narrative)


# ─────────────────────────── 공문 초안(현장 조사 협조 요청) ───────────────────────────
LETTER_DAYS = 30            # 기간 기본값(초안 · 담당자가 정함)


def letter_filename(place: str, at: dt.datetime | None = None) -> str:
    at = at or dt.datetime.now(KST)
    return f"현장조사_협조공문_초안_{place.replace(' ', '_')}_{at:%Y%m%d}.docx"


def collect_letter(code: str, realm: str = "lx", tenant: str = "") -> dict:
    """공문 초안 자료 — 대상 필지 = '현장 확인 필요'(전체 규칙 · 우선순위 A · 배정 전·배정 · 필지 단위) = 첫 화면 같은 이름의 값과 같은 식.
    시군구는 한 출처(counts_sync field_check), 읍면동은 같은 조건을 그 읍면동에서 센 값."""
    code = str(code or "").strip()
    level = "sgg" if re.fullmatch(r"\d{5}", code) else "emd"
    with _conn(realm, tenant) as c:
        if level == "sgg":
            sg = c.execute("SELECT sgg_cd, tenant_id, name, sido, imagery, parcels_as_of, priority_cut, parcels_src ? 'canon' AS canon "
                           "FROM survey_sgg WHERE sgg_cd=%s", (code,)).fetchone()
            if not sg:
                raise NotFound(f"{code} 실태조사 결과가 없습니다")
            e = {"emd_cd": code, "name": sg["name"], "sgg_cd": code, "tenant_id": sg["tenant_id"]}
            col = "f.sgg_cd"
        else:
            e = c.execute("SELECT emd_cd, name, coalesce(sgg_cd, left(emd_cd, 5)) sgg_cd, tenant_id FROM survey_emd WHERE emd_cd=%s", (code,)).fetchone()
            if not e:
                raise NotFound(f"읍면동 {code} 없음")
            sg = c.execute("SELECT sgg_cd, tenant_id, name, sido, imagery, parcels_as_of, priority_cut, parcels_src ? 'canon' AS canon "
                           "FROM survey_sgg WHERE sgg_cd=%s", (e["sgg_cd"],)).fetchone()
            col = "f.emd_cd"
        org = c.execute("SELECT name->>'ko' AS ko FROM tenants WHERE id=%s", ((sg or {}).get("tenant_id") or e["tenant_id"],)).fetchone()
        rows = c.execute(f"SELECT * FROM (SELECT DISTINCT ON (f.pnu) f.pnu, f.addr, f.jimok, f.score, f.rank, f.state, p.jibun, "
                         f"(SELECT array_agg(DISTINCT g.rule ORDER BY g.rule) FROM survey_findings g WHERE g.pnu = f.pnu AND g.rule = ANY(%s) "
                         f"AND g.priority = 'A' AND g.state IN ('open','assigned')) AS rules_all "
                         f"FROM survey_findings f LEFT JOIN survey_parcels p USING (pnu) "
                         f"WHERE {col} = %s AND f.rule = ANY(%s) AND f.priority = 'A' AND f.state IN ('open','assigned') "
                         f"ORDER BY f.pnu, f.score DESC) x ORDER BY x.score DESC, x.rank",
                         (list(RULE_IDS), code, list(RULE_IDS))).fetchall()
        counts = _counts(c, e["sgg_cd"]) if level == "sgg" else None
    region = _region(e["sgg_cd"], sg)
    region["org"] = (org or {}).get("ko")
    n = counts["field_check"] if counts and counts.get("field_check") is not None else len(rows)
    return {"level": level, "emd": dict(e), "region": region, "rows": [dict(x) for x in rows], "targets": int(n),
            "suspect": counts["suspect"] if counts else None, "at": dt.datetime.now(KST)}


def letter_json(d: dict) -> dict:
    rg = d["region"]
    place = rg["name"] if d["level"] == "sgg" else f"{rg['name']} {d['emd']['name']}"
    at = d["at"]
    end = at + dt.timedelta(days=LETTER_DAYS)
    return {"template": "field-letter", "place": place, "org": rg.get("org"), "sido": rg.get("sido"),
            "title": f"{place} 농지 등 실태조사 현장 확인 협조 요청(초안)",
            "to": f"{place} 관계 부서장(농지·건축·산지 담당)",
            "targets": env(d["targets"], "필지", "inferred", REPORT_SRC, "현장 확인 필요 · 검수 전"),
            "suspect": env(d["suspect"], "count", "inferred", REPORT_SRC, "검수 전") if d.get("suspect") is not None else None,
            "period": {"from": f"{at:%Y. %m. %d.}", "to": f"{end:%Y. %m. %d.}", "days": LETTER_DAYS},
            "items": [{"no": i + 1, "pnu": x["pnu"], "addr": x["addr"], "jimok": x.get("jimok"),
                       "rules": [rule_name(r) for r in (x.get("rules_all") or []) if rule_name(r)]} for i, x in enumerate(d["rows"])],
            "fixed": FIXED_PHRASE, "generated": at.isoformat(timespec="seconds"), "filename": letter_filename(place, at)}


def render_letter(d: dict) -> tuple[bytes, str, dict]:
    """한 장 공문 초안(제목·수신·참조·본문(목적·대상 필지 수·기간·협조 사항)·붙임(대상 필지 목록)) + 붙임 목록 쪽."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    j = letter_json(d)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    doc.core_properties.title = j["title"]
    doc.core_properties.author = j.get("org") or "Land-XI"
    doc.core_properties.comments = ""
    org = j.get("org") or ""
    n = j["targets"]["value"]
    _p(doc, org or " ", 18, True, "14202A", WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    _p(doc, f"수신  {j['to']}", 11, space_after=2)
    _p(doc, "참조  현장 확인 담당", 11, space_after=2)
    _p(doc, "(경유)", 11, space_after=6)
    _p(doc, f"제목  {j['title']}", 12, True, "14202A", space_after=10)
    _p(doc, "1. 관련: 농지 등 실태조사(AI 영상 분석 결과와 연속지적도 대조 · 검수 전).", 11, space_after=4)
    _p(doc, f"2. 위 관련으로 {j['place']}의 실태조사 대상 후보 가운데 현장 확인이 필요한 필지에 대하여 다음과 같이 현장 조사 협조를 요청합니다.",
       11, space_after=4)
    _p(doc, "  가. 목적: AI 분석 결과의 현장 확인과 조치 여부 판단", 11, space_after=2)
    tail = f" — 의심 필지 {j['suspect']['value']:,}건 가운데 우선 확인 대상" if j.get("suspect") else ""
    _p(doc, f"  나. 대상: 현장 확인 필요 {n:,}필지(붙임 목록){tail}", 11, space_after=2)
    _p(doc, f"  다. 기간: {j['period']['from']} ~ {j['period']['to']}({j['period']['days']}일 · 담당자가 조정)", 11, space_after=2)
    _p(doc, "  라. 협조 사항", 11, space_after=2)
    for k, t in enumerate(["대상 필지 현장 확인과 사진 기록", "건축물대장 · 농지전용 · 산지전용 허가 대장 대조",
                           "결과 입력: Land-XI '할 일'에서 위반 / 대장과 같음 / 불명확 중 하나로 판정"], 1):
        _p(doc, f"     {k}) {t}", 11, space_after=1)
    _p(doc, f"3. {j['fixed']}. AI 분석 결과는 참고자료이며 위법 여부는 현장 확인 후 담당자가 판단합니다.", 11, space_after=8)
    _p(doc, "붙임  현장 확인 대상 필지 목록 1부.  끝.", 11, space_after=24)
    sign = (org + "장") if org and not org.endswith("장") else org
    _p(doc, f"{sign}  (직인 생략 · 초안)", 14, True, "14202A", WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    _p(doc, "기안자            검토자            결재권자", 10, False, "56626B", space_after=2)
    _p(doc, f"시행  (문서번호 부여 전)  ({j['period']['from']})", 10, False, "56626B", space_after=2)
    _p(doc, AI_DRAFT, 9, False, "9A3412", space_after=0)
    # 붙임 — 새 쪽
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _p(doc, f"붙임  현장 확인 대상 필지 목록 — {j['place']} {n:,}필지", 12.5, True, "14202A", space_after=4)
    _p(doc, "같은 필지가 여러 규칙에 걸리면 한 줄로 적고 '해당 규칙' 칸에 모두 적었습니다. AI 추론 · 검수 전.", 8.5, False, "56626B", space_after=4)
    strip = [x for x in ((j.get("sido") or "") + " ", (d["region"].get("name") or "") + " ") if x.strip()]

    def short(a):
        a = a or ""
        for s_ in strip:
            a = a.replace(s_, "")
        return re.sub(r"^\S+(특별시|광역시|특별자치도|특별자치시|도)\s", "", a)
    _tbl(doc, ["No", "소재지", "지목", "해당 규칙", "PNU"],
         [[x["no"], short(x["addr"]), x["jimok"] or "—", " · ".join(x["rules"]) or "—", x["pnu"]] for x in j["items"]],
         widths=[1.0, 5.6, 1.2, 5.2, 4.0], size=8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), j["filename"], j


def docx_text(blob: bytes) -> str:
    """docx 안 글자 전부(본문 · 표 · 머리) — 개발 문구 검사용."""
    from docx import Document
    d = Document(io.BytesIO(blob))
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    for s in d.sections:
        parts += [p.text for p in s.header.paragraphs]
    return "\n".join(parts)


DEV_WORDS = re.compile(r"PostGIS|V-World|API\s?키|건축HUB|inspected|closed|dismissed|2차|RAG|\bllm\b|LLM|survey_|emd_cd|덕과면|AOI|도엽|"
                       r"\.py\b|/api|README|A0[1-4]\b|이 문서는 $|EPSG|run_|시연|관제|\b[RL]\d\b", re.M)


def dev_words(text: str) -> list[str]:
    return sorted({m.group(0) for m in DEV_WORDS.finditer(text)})


if __name__ == "__main__":   # python -m survey.report <법정동 8자리|시군구 5자리> R1 20 out.docx
    import sys
    emd, rl, n, out = sys.argv[1], (sys.argv[2] if len(sys.argv) > 2 and sys.argv[2] != "-" else None), \
        int(sys.argv[3]) if len(sys.argv) > 3 else 20, sys.argv[4] if len(sys.argv) > 4 else None
    blob, name = render_docx(collect(emd, rl, n))
    p = out or name
    open(p, "wb").write(blob)
    print(p, len(blob), "bytes")
