"""구현 14차 b · 바퀴 3(확인 대장 10-11 Q5 ⓑ · Q6 ⓑ · Q7 · QA-고침-5 · 원칙 59 · 135 · 192 · 193). 실서버(:8700) · GPU 0 · vLLM 0(직행만).

  지도 서비스 층 내려받기   GET /me/analyses/{작업}/formats · /download?fmt=geojson|shp|parcels(동의 한 줄 · 기록 · 외부 API 1차와 같은 열 · 같은 함수)
  말로 거르기(map_filter)   켜진 층의 실제 결과 속성만 — 없는 값(단동 · 연동)은 거르지 않고 '없다' + 개선 고리(자료 없음)로 · 맞는 수 = 같은 행 같은 조건
  AI 분석 결과 보고서        읍면동별 개수 · 면적 합 = 지도 서비스 결과 수(같은 행) · 지역이 없으면 켜진 층 → 화면 지역 → 묻기 · '의심 필지' · '현장 확인' 말 0
내려받기 시험은 ?test=1(기록 줄에 '(시험)') · 말로 하는 시험은 run_test… 번호(개선 고리가 모으지 않는다).
"""
import asyncio
import io
import json
import zipfile

import httpx
import psycopg
import pytest

from conftest import B, H
from landxi_api import config

NAMWON_GH = "job_01M4AT60R3DZ1Q0VKGC1R71CPF"     # 비닐하우스 분석서비스 · 남원시 전역(LX 직원 test@lx.or.kr 이 돌린 결과)


def pg():
    return psycopg.connect(config.PG_ADMIN_DSN, autocommit=True)


def _item(tok, job):
    j = httpx.get(B + "/me/analyses", headers=H(tok), timeout=60).json()
    return next((it for g in j["groups"] for it in g["items"] if it["job"] == job), None)


@pytest.fixture(scope="module")
def staff(tok):
    it = _item(tok["staff"], NAMWON_GH)
    if not it:
        pytest.skip("남원 비닐하우스 전역 결과가 이 계정 지도 서비스에 없음")
    return tok["staff"], it


def test_formats_and_consent(staff, tok):
    t, it = staff
    f = httpx.get(B + f"/me/analyses/{NAMWON_GH}/formats", headers=H(t), timeout=30).json()
    assert {x["fmt"]: x["ok"] for x in f["items"]} == {"geojson": True, "shp": True, "parcels": True}
    c = httpx.get(B + "/me/downloads/consent", headers=H(t), timeout=30).json()
    assert "기록이 남습니다" in c["line"] and len(c["line"]) < 80               # 동의 한 줄
    assert httpx.post(B + "/me/downloads/consent", headers=H(t), json={}, timeout=30).status_code == 400
    assert httpx.post(B + "/me/downloads/consent", headers=H(t), json={"agree": True}, timeout=30).json()["done"] is True
    # 남의 결과 · 기관 계정 · 모르는 작업
    assert httpx.get(B + "/me/analyses/job_NOPE/formats", headers=H(t), timeout=30).status_code == 404
    assert httpx.get(B + f"/me/analyses/{NAMWON_GH}/download", headers=H(tok["namwon"]), timeout=30).status_code == 403


def test_download_geojson_shp_xlsx_same_values(staff):
    t, it = staff
    found = it["found"]["value"]
    with pg() as c:
        before = c.execute("SELECT count(*) FROM space_log WHERE tenant_id='lx' AND kind='download'").fetchone()[0]
    g = httpx.get(B + f"/me/analyses/{NAMWON_GH}/download", params={"fmt": "geojson", "test": 1}, headers=H(t), timeout=120)
    assert g.status_code == 200 and g.headers["content-type"].startswith("application/geo+json")
    fc = g.json()
    assert len(fc["features"]) == found                                        # 지도 서비스 결과 수와 같다
    from landxi_api.ext_api import _props
    keys = set(_props({"fid": "a", "id": 1, "cls": "x", "cls_en": "x", "conf": 0.5, "area_m2": 1, "emd": "", "pnu": None, "edit_state": "raw"}, "", "", None))
    assert set(fc["features"][0]["properties"]) == keys                        # 외부 API 1차와 같은 열(같은 함수)
    by_emd = {}
    for ft in fc["features"]:
        by_emd[ft["properties"]["emd"]] = by_emd.get(ft["properties"]["emd"], 0) + 1
    assert by_emd.get("운봉읍") and sum(by_emd.values()) == found
    s = httpx.get(B + f"/me/analyses/{NAMWON_GH}/download", params={"fmt": "shp", "test": 1}, headers=H(t), timeout=120)
    assert s.status_code == 200
    z = zipfile.ZipFile(io.BytesIO(s.content))
    assert {n.rsplit(".", 1)[-1] for n in z.namelist() if "." in n} >= {"shp", "shx", "dbf", "prj", "cpg"}
    import shapefile
    stem = next(n for n in z.namelist() if n.endswith(".shp"))[:-4]
    r = shapefile.Reader(shp=io.BytesIO(z.read(stem + ".shp")), shx=io.BytesIO(z.read(stem + ".shx")), dbf=io.BytesIO(z.read(stem + ".dbf")), encoding="utf-8")
    assert len(r) == found
    x = httpx.get(B + f"/me/analyses/{NAMWON_GH}/download", params={"fmt": "parcels", "test": 1}, headers=H(t), timeout=120)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content), read_only=True)
    assert wb.sheetnames[:3] == ["필지", "열 설명", "안내"]
    info = {r[0]: r[1] for r in wb["안내"].iter_rows(values_only=True)}
    assert info["서비스"] == "비닐하우스 분석서비스" and info["지역"] == "남원시" and "결과 설명서" not in info
    with pg() as c:
        rows = c.execute("SELECT line, detail FROM space_log WHERE tenant_id='lx' AND kind='download' ORDER BY id DESC LIMIT 3").fetchall()
        after = c.execute("SELECT count(*) FROM space_log WHERE tenant_id='lx' AND kind='download'").fetchone()[0]
    assert after - before == 3 and all("(시험)" in r[0] and r[1]["test"] for r in rows)
    lg = httpx.get(B + "/me/downloads", headers=H(t), timeout=30).json()["items"]
    assert lg and "내려받기" in lg[0]["line"]


def arun(coro):
    """한 번 돌릴 때마다 이 루프의 새 풀(앞선 asyncio.run 이 남긴 닫힌 루프의 풀은 버린다)."""
    from landxi_api import deps

    async def go():
        deps._pool = deps._pool_sys = deps._redis = None
        try:
            return await coro
        finally:
            await deps.close() if hasattr(deps, "close") else None
    return asyncio.run(go())


def _ctx(context, msg):
    from agent import runner
    from agent.tools import registry
    from landxi_api.deps import CAPS, Principal
    registry.ensure_ext()
    p = Principal("lx", "staff", None, "u_mail_test", caps=CAPS[("lx", "staff")])
    c = runner.Ctx(run_id="run_testimpl14b", principal=p, token=None, context=context)
    c.lang = "ko"
    c.state["msg"] = msg
    return c


@pytest.mark.parametrize("q,want", [
    ("비닐하우스 단동만 펼쳐줘", {"cls": "비닐하우스 단동"}), ("비닐하우스 연동만 보여 줘", {"cls": "비닐하우스 다동"}),
    ("면적 1,000㎡ 넘는 것만", {"area_min": 1000.0, "area_op_min": ">"}), ("300평 이하만", {"area_max": 991.7, "area_op_max": "<="}),
    ("운봉읍만", {"emd": "운봉읍"}), ("조건 풀어 줘", {"clear": True}),
    ("운봉읍 비닐하우스 통계 내 줘", None), ("남원시 비닐하우스 보여 줘", None), ("보고서 초안 만들어 줘", None)])
def test_parse_filter(q, want):
    from agent.tools.ext.mapsvc import parse_filter
    got = parse_filter(q)
    if want is None:
        assert got is None
    else:
        assert got and all(got.get(k) == v for k, v in want.items()), got


def test_map_filter_only_real_values(live):
    from agent.tools.ext import mapsvc as M
    ctxv = {"page": "lx-map", "sets": [f"results/lx/{NAMWON_GH}"]}
    one = arun(M.map_filter({"cls": "비닐하우스 단동"}, c := _ctx(ctxv, "비닐하우스 단동만 펼쳐줘")))
    assert not one.ui_actions and "단동 구분이 없습니다" in one.answer and c.state["cannot"] == {"kind": "nodata"}
    big = arun(M.map_filter({"area_min": 1000.0, "area_op_min": ">"}, _ctx(ctxv, "면적 1,000㎡ 넘는 것만")))
    with pg() as c2:
        n = c2.execute("SELECT count(*) FROM detections WHERE job_id=%s AND edit_state<>'deleted' AND area_m2 > 1000", (NAMWON_GH,)).fetchone()[0]
    a = big.ui_actions[0]
    assert a["op"] == "map_filter" and a["n"] == n and a["area_min"] == 1000.0 and "1,000㎡ 넘음" in a["label"]
    emd = arun(M.map_filter({"emd": "운봉읍"}, _ctx(ctxv, "운봉읍만")))
    assert emd.ui_actions[0]["emd"] == ["운봉읍"]
    nope = arun(M.map_filter({"emd": "없는면"}, _ctx(ctxv, "없는면만")))
    assert not nope.ui_actions and "없는면 결과가 없습니다" in nope.answer
    none = arun(M.map_filter({"emd": "운봉읍"}, _ctx({"page": "lx-map", "sets": []}, "운봉읍만")))
    assert not none.ui_actions and "켜진 결과 층이 없습니다" in none.answer


def test_improve_classifies_filter_miss():
    from landxi_api.improve import classify, screen_of
    assert classify("agent.done", {"answer_md": "x", "cannot": {"kind": "nodata"}}) == [("blocked", "nodata")]
    assert screen_of({"page": "lx-map"}) == "지도 서비스"


def test_ai_report_same_numbers_and_words(staff):
    t, it = staff
    from agent.tools.ext import mapsvc as M
    from survey.report import docx_text
    c = _ctx({"page": "lx-map", "sets": [f"results/lx/{NAMWON_GH}"]}, "보고서 초안 만들어 줘")
    out = arun(M.ai_report({"request": "보고서 초안 만들어 줘"}, c))
    tot = next(e for k, _, e in out.envelopes if k == "total")
    assert tot["value"] == it["found"]["value"]                                # 지도 서비스 결과 수와 같은 값(숫자 한 출처)
    f = next(b for b in out.blocks if b["type"] == "file")
    from agent import config as AC
    d = AC.ARTIFACT_DIR / c.run_id
    txt = docx_text((d / "draft.docx").read_bytes())
    for w in ("읍면동별 개수 · 면적", "AI 모델 정확도", "영상 시점", "지도 그림", "합계"):
        assert w in txt, w
    for bad in ("의심", "현장 확인", "현장조사", "실태조사", "job_", "GPU"):
        assert bad not in txt, bad
    assert f["label"].endswith(".docx") and (d / f["label"]).exists()
    # 지역이 없고 켜진 층도 없으면 묻는다 · 질문의 지역이 있으면 그 지역(통계 답과 같은 세트)
    ask = arun(M.ai_report({"request": "보고서 초안 만들어 줘"}, _ctx({"page": "lx-console"}, "보고서 초안 만들어 줘")))
    assert "어느 지역" in ask.answer and not ask.blocks


def test_route_order():
    from agent import runner
    hit = arun(runner.ext_route(_ctx({"page": "lx-map", "sets": []}, "보고서 초안 만들어 줘"), "보고서 초안 만들어 줘"))
    assert hit["tool"] == "ai_report"
    hit = arun(runner.ext_route(_ctx({"page": "lx-map", "sets": []}, "운봉읍만 보여 줘"), "운봉읍만 보여 줘"))
    assert hit["tool"] == "map_filter"
    hit = arun(runner.ext_route(_ctx({"page": "lx-console"}, "운봉읍만 보여 줘"), "운봉읍만 보여 줘"))
    assert not hit or hit["tool"] != "map_filter"                            # 지도 · 결과 층이 없는 화면에서는 거르기로 가지 않는다
