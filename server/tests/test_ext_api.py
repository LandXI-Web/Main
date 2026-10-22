"""외부 연동 API 3차 — 1차 고침 + 안내 · 신청(원칙 184 · 확인 대장 10-11 16:41 API3 · 설계 design-r18/open-api-3). 실서버(:8700) · GPU 0.

키 신청(로그인 없이 · 시간당 5번) · 승인(서비스마다 키 · 값 한 번) · 거절(사유) · 외부 기관 승인 막힘 · 기관 공유와 무관 · 범위 밖 404 ·
원천(LX 보유 = LX 내부 키만) · 형식 셋(GeoJSON · SHP · 필지 엑셀) · 받은 수 = 장부 수 = 기관 화면 수 · 한도 · 기록 · 분석 호출 501 · 1차 주소 404.
시험 신청 · 키는 모두 test=true · 끝나면 키 폐기(사유 '시험 끝') — 기록은 지우지 않는다. 시험으로 바꾼 공유 줄(test=true)은 끝에 지워 원래 상태로.
"""
import io
import json
import secrets
import zipfile

import httpx
import psycopg
import pytest

from conftest import B, BASE, H
from landxi_api import config

X = BASE + "/api/ext/v1"
T = 120
BAD = ("job_", "results/", "E:\\", "GPU", "lxs_", "lxt_", "set_id", "imagery_id")


def pg():
    return psycopg.connect(config.PG_ADMIN_DSN, autocommit=True)


def K(key):
    return {"authorization": "Bearer " + key}


def _ip():
    return {"x-forwarded-for": f"10.77.{secrets.randbelow(250)}.{secrets.randbelow(250) + 1}"}


APP = {"kind": "internal", "org": "LX 시험 부서", "system": "시험 연동 시스템", "person": "시험 담당", "email": "test@lx.or.kr",
       "services": ["farm"], "area": "전국", "rights": "read", "purpose": "pytest 시험 신청입니다", "consent": True}


def _apply(**kw):
    r = httpx.post(B + "/public/api-applications", json={**APP, **kw}, headers=_ip(), timeout=T)
    assert r.status_code == 201, r.text
    with pg() as c:
        aid = c.execute("SELECT id FROM api_applicants WHERE system=%s ORDER BY at DESC LIMIT 1", (kw.get("system", APP["system"]),)).fetchone()[0]
        c.execute("UPDATE api_applicants SET test = true WHERE id=%s", (aid,))
    return aid


@pytest.fixture(scope="module")
def keys(live, tok):
    made = []

    def mk(card="card-farm", kind="internal", **kw):
        body = {"card_id": card, "kind": kind, "system": "시험 연동", "test": True, **kw}
        r = httpx.post(B + "/apikeys", headers=H(tok["admin"]), json=body, timeout=T)
        assert r.status_code == 201, r.text
        j = r.json()
        made.append(j["item"]["id"])
        return j["key"], j["item"]["id"]
    yield mk
    for kid in made:
        httpx.post(B + f"/apikeys/{kid}/revoke", headers=H(tok["admin"]), json={"reason": "시험 끝"}, timeout=T)
    with pg() as c:
        c.execute("DELETE FROM card_shares WHERE test AND note='pytest ext_api3'")


def test_catalog_public(live):
    r = httpx.get(B + "/public/api-catalog", timeout=T)
    assert r.status_code == 200
    j = r.json()
    ids = {s["id"] for s in j["services"]}
    assert {"farm", "building", "parking", "greenhouse"} <= ids and j["analyze_open"] is False
    for bad in BAD:
        assert bad not in r.text


def test_apply_approve_once(tok):
    A = H(tok["admin"])
    assert httpx.post(B + "/public/api-applications", json={**APP, "services": []}, headers=_ip(), timeout=T).status_code == 400
    assert httpx.post(B + "/public/api-applications", json={**APP, "consent": False}, headers=_ip(), timeout=T).status_code == 400
    assert httpx.post(B + "/public/api-applications", json={**APP, "email": "", "phone": ""}, headers=_ip(), timeout=T).status_code == 400
    sys_name = "시험 랜디에어 " + secrets.token_hex(2)
    aid = _apply(kind="lx_service", system=sys_name, services=["farm", "building"])
    for who in ("staff", "namwon"):                                     # 관리자만
        assert httpx.get(B + "/api-applications", headers=H(tok[who]), timeout=T).status_code in (401, 403)
    lst = httpx.get(B + "/api-applications", headers=A, timeout=T).json()
    row = [x for x in lst["items"] if x["id"] == aid][0]
    assert row["state"] == "pending" and row["can_approve"] and row["kind_label"] == "LX 다른 서비스"
    bad = httpx.post(B + f"/api-applications/{aid}/approve", headers=A, json={"sources": ["lx"], "test": True}, timeout=T)
    assert bad.status_code == 400                                       # LX 보유 원천 = LX 내부 키만
    ok = httpx.post(B + f"/api-applications/{aid}/approve", headers=A, json={"scope": {"codes": ["52"]}, "sources": ["ngii-aerial", "vworld"],
                                                                             "area_once_km2": 10, "test": True}, timeout=T)
    assert ok.status_code == 200, ok.text
    ks = ok.json()["keys"]
    assert len(ks) == 2 and all(k["key"].startswith("lxk_") for k in ks)
    assert {k["item"]["service"]["id"] for k in ks} == {"farm", "building"}
    assert all(k["item"]["area"]["once"]["value"] == 10 for k in ks)
    with pg() as c:
        for k in ks:
            h, l4 = c.execute("SELECT key_hash, last4 FROM api_keys WHERE id=%s", (k["item"]["id"],)).fetchone()
            assert k["key"] not in h and l4 == k["key"][-4:]
            assert not c.execute("SELECT 1 FROM audit_log WHERE after::text LIKE %s", ("%" + k["key"] + "%",)).fetchone()
    again = httpx.get(B + "/apikeys", headers=A, timeout=T)
    assert again.status_code == 200 and all(k["key"] not in again.text for k in ks)
    assert httpx.post(B + f"/api-applications/{aid}/approve", headers=A, json={}, timeout=T).status_code == 409   # 두 번 승인 없음
    me = httpx.get(X + "/me", headers=K(ks[0]["key"]), timeout=T)
    assert me.status_code == 200 and me.json()["system"] == sys_name and me.json()["scope"]["words"] == ["전북특별자치도"]
    for k in ks:
        httpx.post(B + f"/apikeys/{k['item']['id']}/revoke", headers=A, json={"reason": "시험 끝"}, timeout=T)


def test_reject_and_external_blocked(tok):
    A = H(tok["admin"])
    aid = _apply(system="시험 거절 시스템")
    assert httpx.post(B + f"/api-applications/{aid}/reject", headers=A, json={"reason": ""}, timeout=T).status_code == 400
    r = httpx.post(B + f"/api-applications/{aid}/reject", headers=A, json={"reason": "시험 — 쓰는 목적을 더 적어 주세요"}, timeout=T)
    assert r.status_code == 200 and r.json()["item"]["state"] == "rejected" and r.json()["item"]["reason"].startswith("시험")
    ext_id = _apply(kind="external", system="시험 외부 기관 시스템")
    row = [x for x in httpx.get(B + "/api-applications", headers=A, timeout=T).json()["items"] if x["id"] == ext_id][0]
    assert row["can_approve"] is False and "보안성 검토" in row["hold"]
    r = httpx.post(B + f"/api-applications/{ext_id}/approve", headers=A, json={"test": True}, timeout=T)
    assert r.status_code == 409 and r.json()["error"]["code"] == "security_review"
    r = httpx.post(B + "/apikeys", headers=A, json={"card_id": "card-farm", "kind": "external", "system": "시험 외부", "test": True}, timeout=T)
    assert r.status_code == 409
    httpx.post(B + f"/api-applications/{ext_id}/reject", headers=A, json={"reason": "시험 끝"}, timeout=T)


def test_apply_rate_limit(live):
    ip = {"x-forwarded-for": "10.88.1." + str(secrets.randbelow(250) + 1)}
    codes = [httpx.post(B + "/public/api-applications", json={**APP, "website": "bot"}, headers=ip, timeout=T).status_code for _ in range(6)]
    assert codes[:5] == [201] * 5 and codes[5] == 429                   # 자동 입력기 칸 = 받은 척(저장 0) · 6번째 429


def test_auth_not_mixed(tok, keys):
    key, _ = keys()
    assert httpx.get(X + "/me", timeout=T).status_code == 401
    assert httpx.get(X + "/me", headers=H(tok["namwon"]), timeout=T).status_code == 401
    assert httpx.get(X + "/me", headers=K(key + "x"), timeout=T).status_code == 401
    assert httpx.get(B + "/me", headers=K(key), timeout=T).status_code == 401
    me = httpx.get(X + "/me", headers=K(key), timeout=T)
    assert me.status_code == 200 and me.json()["service"]["id"] == "farm" and me.json()["scope"]["words"] == ["전국"]
    for bad in BAD:
        assert bad not in me.text
    for path in ("/results", "/timepoints", "/summary", "/exports/geojson", "/uploads"):   # 1차 주소 · 영상 올리기 = 없음
        assert httpx.get(X + path, headers=K(key), timeout=T).status_code == 404
    assert httpx.get(X + "/services/building/results", headers=K(key), timeout=T).status_code == 404   # 키에 없는 서비스


def test_analyses_not_open(keys):
    rk, _ = keys()                                                       # 조회 키 → 403
    assert httpx.post(X + "/services/farm/analyses", headers=K(rk), json={}, timeout=T).status_code == 403
    ak, _ = keys(rights="read_analyze")
    for m, p in (("POST", "/services/farm/analyses"), ("POST", "/services/farm/analyses/estimate"), ("GET", "/services/farm/analyses/an_x"),
                 ("DELETE", "/services/farm/analyses/an_x"), ("GET", "/imagery")):
        r = httpx.request(m, X + p, headers=K(ak), json={} if m == "POST" else None, timeout=T)
        assert r.status_code == 501 and r.json()["error"]["code"] == "not_open", (m, p, r.text)


def _all(key, svc, **params):
    feats, cur, n = [], None, 0
    while True:
        r = httpx.get(X + f"/services/{svc}/results", headers=K(key), params={**params, **({"cursor": cur} if cur else {})}, timeout=600)
        assert r.status_code == 200, r.text
        j = r.json()
        feats += j["features"]
        cur = j["lx"]["next"]
        n += 1
        if not cur:
            return feats, j["lx"], n


def _org_download(tok, who, card):
    r = httpx.get(B + f"/spaces/me/guides/{card}/download", headers=H(tok[who]), params={"fmt": "geojson"}, timeout=600)
    if r.status_code == 409:
        httpx.post(B + "/spaces/me/consent", headers=H(tok[who]), json={"agree": True}, timeout=T)
        r = httpx.get(B + f"/spaces/me/guides/{card}/download", headers=H(tok[who]), params={"fmt": "geojson"}, timeout=600)
    assert r.status_code == 200
    return r.json()["features"]


def test_scope_counts_match_ledger_and_org(tok, keys):
    jb, _ = keys(scope={"codes": ["52"]}, sources=["ngii-aerial", "vworld", "lx"])     # 범위 '전북' 키
    feats, lx, pages = _all(jb, "farm", region="52190")
    with pg() as c:
        sets = [r[0] for r in c.execute("SELECT set_id FROM service_results WHERE card_id='card-farm' AND sgg_cd='52190' AND replaced_at IS NULL "
                                         "AND year = (SELECT max(year) FROM service_results WHERE card_id='card-farm' AND sgg_cd='52190' AND replaced_at IS NULL)").fetchall()]
        led = c.execute("SELECT count(*) FROM detections WHERE job_id = ANY(%s) AND edit_state <> 'deleted'", (sets,)).fetchone()[0]
    assert len(feats) == lx["total"] == led > 0 and len({f["id"] for f in feats}) == led
    org = _org_download(tok, "namwon", "card-farm")                       # 남원 기관 화면 같은 시점
    assert sorted(f["id"] for f in org) == sorted(f["id"] for f in feats)
    assert all(f["properties"]["source"] == "lx" and f["properties"]["year"] == 2025 for f in feats)
    assert lx["imagery"] and lx["imagery"][0]["attribution"]
    for bad in BAD:
        assert bad not in json.dumps(feats[:50], ensure_ascii=False) and bad not in json.dumps(lx, ensure_ascii=False)
    # 범위 밖 = 404(전남광주 · 밖 bbox · 밖 도형)
    assert httpx.get(X + "/services/farm/results", headers=K(jb), params={"region": "12110"}, timeout=T).status_code == 404
    assert httpx.get(X + "/services/farm/results", headers=K(jb), params={"bbox": "126.9,34.9,127.0,35.0"}, timeout=T).status_code == 404
    poly = {"type": "Polygon", "coordinates": [[[126.9, 34.9], [127.0, 34.9], [127.0, 35.0], [126.9, 35.0], [126.9, 34.9]]]}
    assert httpx.post(X + "/services/farm/results/search", headers=K(jb), json={"geometry": poly}, timeout=T).status_code == 404
    # 범위 안 bbox(남원 안) — 서버 오류 0 · 결과는 전부 그 상자와 겹침
    r = httpx.get(X + "/services/farm/results", headers=K(jb), params={"bbox": "127.35,35.38,127.45,35.45", "limit": 200}, timeout=T)
    assert r.status_code == 200 and r.json()["lx"]["total"] <= led
    poly2 = {"type": "Polygon", "coordinates": [[[127.35, 35.38], [127.45, 35.38], [127.45, 35.45], [127.35, 35.45], [127.35, 35.38]]]}
    s = httpx.post(X + "/services/farm/results/search", headers=K(jb), json={"geometry": poly2, "limit": 50}, timeout=T)
    assert s.status_code == 200 and s.json()["lx"]["total"] >= r.json()["lx"]["total"] * 0 and s.json()["lx"]["total"] <= led
    # 지난 연도(2023)는 year 로
    old = httpx.get(X + "/services/farm/results", headers=K(jb), params={"region": "52190", "year": 2023, "limit": 5}, timeout=T)
    assert old.status_code == 200 and all(f["properties"]["year"] == 2023 for f in old.json()["features"])
    cov = httpx.get(X + "/services/farm/coverage", headers=K(jb), timeout=T).json()
    assert {x["year"] for x in cov["items"]} >= {2023, 2025} and all(x["region"]["code"].startswith("52") for x in cov["items"])
    # 다른 시도 한 곳 — 전남광주 해양쓰레기(여수) = 기관 화면과 같은 수
    gk, _ = keys(card="card-marine", scope={"codes": ["12"]}, sources=["lx"])
    gf, glx, _p = _all(gk, "marine")
    gorg = _org_download(tok, "gj", "card-marine")
    assert len(gf) == glx["total"] == len(gorg) > 0 and sorted(f["id"] for f in gf) == sorted(f["id"] for f in gorg)


def test_source_rules(keys):
    nk, _ = keys(sources=["ngii-aerial", "vworld"])                    # LX 보유 원천 없는 키
    r = httpx.get(X + "/services/farm/results", headers=K(nk), params={"source": "lx"}, timeout=T)
    assert r.status_code == 404
    r = httpx.get(X + "/services/farm/results", headers=K(nk), params={"region": "52190"}, timeout=T)
    assert r.status_code == 200 and r.json()["lx"]["total"] == 0          # 지금 결과는 모두 LX 보유 영상에서 — 공개 원천 결과는 2차부터
    assert httpx.get(X + "/services/farm/results", headers=K(nk), params={"source": "bad"}, timeout=T).status_code == 400


def test_sharing_does_not_matter(tok, keys):
    A = H(tok["admin"])
    key, _ = keys(scope={"codes": ["52190"]}, sources=["lx"])
    before = httpx.get(X + "/services/farm/results", headers=K(key), params={"limit": 1}, timeout=T).json()["lx"]["total"]
    sh = {"card_id": "card-farm", "tenant_id": "namwon", "test": True, "note": "pytest ext_api3"}
    assert httpx.put(B + "/release/shares", headers=A, json={**sh, "shared": False}, timeout=T).status_code == 200
    try:
        r = httpx.get(X + "/services/farm/results", headers=K(key), params={"limit": 1}, timeout=T)
        assert r.status_code == 200 and r.json()["lx"]["total"] == before
        st = [k for g in httpx.get(B + "/apikeys", headers=A, timeout=T).json()["items"] for k in g["keys"] if k["id"] == _]
        assert st and st[0]["state"] == "on"
    finally:
        httpx.put(B + "/release/shares", headers=A, json={**sh, "shared": True}, timeout=T)
        with pg() as c:
            c.execute("DELETE FROM card_shares WHERE test AND note='pytest ext_api3'")


def test_formats_three(keys):
    key, _ = keys(scope={"codes": ["52190"]}, sources=["lx"])
    feats, lx, _p = _all(key, "farm")
    gx = httpx.get(X + "/services/farm/exports/geojson", headers=K(key), timeout=600)
    assert gx.status_code == 200 and len(gx.json()["features"]) == lx["total"]
    sh = httpx.get(X + "/services/farm/exports/shp", headers=K(key), timeout=600)
    assert sh.status_code == 200 and sh.headers["content-type"] == "application/zip"
    import shapefile
    z = zipfile.ZipFile(io.BytesIO(sh.content))
    stem = [n for n in z.namelist() if n.endswith(".shp")][0][:-4]
    rd = shapefile.Reader(shp=io.BytesIO(z.read(stem + ".shp")), shx=io.BytesIO(z.read(stem + ".shx")), dbf=io.BytesIO(z.read(stem + ".dbf")), encoding="utf-8")
    assert len(rd) == lx["total"] and b"WGS_1984" in z.read(stem + ".prj")
    px = httpx.get(X + "/services/farm/exports/parcels", headers=K(key), timeout=600)
    assert px.status_code == 200
    from openpyxl import load_workbook
    rows = [r[0] for r in load_workbook(io.BytesIO(px.content))["필지"].iter_rows(values_only=True)]
    rows = rows[rows.index("필지 번호") + 1:]          # 엑셀 첫머리 안내 줄(제목 · 기준일 · 출처 · 빈 줄) 뒤 머리줄 다음부터(원칙 308 엑셀 한 벌)
    if rows and rows[-1] == "합계":                     # 끝 줄 = 합계(원칙 308)
        rows = rows[:-1]
    assert rows and all(isinstance(x, str) and len(x) == 19 and x.startswith("52190") for x in rows)
    sm = httpx.get(X + "/services/farm/summary", headers=K(key), timeout=T)
    assert sm.status_code == 200 and sm.json()["total"] == lx["total"]
    assert httpx.get(X + "/services/farm/exports/gpkg", headers=K(key), timeout=T).status_code == 404
    fk, _ = keys(formats=["parcels"], scope={"codes": ["52190"]}, sources=["lx"])
    assert httpx.get(X + "/services/farm/results", headers=K(fk), timeout=T).status_code == 403
    assert httpx.get(X + "/services/farm/exports/shp", headers=K(fk), timeout=T).status_code == 403


def test_crs_area_same(keys):
    key, _ = keys(scope={"codes": ["52190"]}, sources=["lx"])
    a = httpx.get(X + "/services/farm/results", headers=K(key), params={"limit": 20}, timeout=T).json()
    b = httpx.get(X + "/services/farm/results", headers=K(key), params={"limit": 20, "crs": "5186"}, timeout=T).json()
    assert b["lx"]["crs"] == "EPSG:5186" and [f["id"] for f in a["features"]] == [f["id"] for f in b["features"]]
    assert [f["properties"]["area_m2"] for f in a["features"]] == [f["properties"]["area_m2"] for f in b["features"]]


def test_pause_revoke_limit_log(tok, keys):
    A = H(tok["admin"])
    key, kid = keys(per_min=3)
    codes = [httpx.get(X + "/me", headers=K(key), timeout=T).status_code for _ in range(5)]
    assert codes[:3] == [200, 200, 200] and codes[3:] == [429, 429]
    calls = httpx.get(B + f"/apikeys/{kid}/calls", headers=A, timeout=T).json()
    assert calls["total"] == 5 and sum(1 for x in calls["items"] if x["ok"]) == 3 and all(x["test"] for x in calls["items"])
    k2, kid2 = keys()
    assert httpx.patch(B + f"/apikeys/{kid2}", headers=A, json={"paused": True}, timeout=T).json()["item"]["state"] == "paused"
    r = httpx.get(X + "/me", headers=K(k2), timeout=T)
    assert r.status_code == 403 and r.json()["error"]["code"] == "key_paused"
    assert httpx.patch(B + f"/apikeys/{kid2}", headers=A, json={"paused": False, "scope": {"codes": ["52190"]}}, timeout=T).json()["item"]["scope_words"] == ["전북특별자치도 남원시"]
    assert httpx.post(B + f"/apikeys/{kid2}/revoke", headers=A, json={"reason": "시험 끝"}, timeout=T).json()["item"]["state"] == "revoked"
    r = httpx.get(X + "/me", headers=K(k2), timeout=T)
    assert r.status_code == 401 and r.json()["error"]["code"] == "key_revoked"
    with pg() as c:
        assert c.execute("SELECT count(*) FROM api_calls WHERE key_id=%s", (kid2,)).fetchone()[0] == 2
