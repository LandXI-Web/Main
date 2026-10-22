"""구현 18차 · 바퀴 5(확인 대장 10-11 18:52 — gov-보고서 · gov-비교 이대로 구현 · QA-Q2 ⓐ · QA-Q3 ⓐ). 실서버(:8700) · GPU 0 · vLLM 0.

  기관 AI 분석 결과 보고서   GET /spaces/me/report(.hwpx · /map.png) — 합계 = 대시보드 큰 숫자(결과 설명서) = 필지 목록 total · 읍면동 합 = 합계 ·
                             규칙 · 소견 · 조치 제안 · '의심' · '현장 확인' 말 0 · 올린 대장이 있으면 '우리 대장과 다른 필지'
  행정정보와 비교            GET /spaces/me/ledger/compare(.xlsx) — 결과 세트 필지 번호 × 대장 필지 번호(실태조사 규칙 없음) ·
                             다른 필지 = 두 줄 합 · AI 결과 있음 + AI 분석 전 = 대장 필지 · 다른 기관 대장 0 · 올리기(지번만) → 맞춤 → 결과(시험 뒤 지움)
  공개 실적                  GET /public/record — 시군구 수 = 점 수 = 실태조사 집계 끝난 시군구 · 서비스 = 서비스 결과 장부 · 운영 기관 = 공개된 결과 설명서
"""
import io
import time
import zipfile

import httpx
import psycopg
import pytest

from conftest import B, H, drop_account, temp_account
from landxi_api import config

CARD = "card-farm"
BAD = ("의심", "현장 확인", "소견", "조치 제안", "우선순위", "규칙", "판정")


def pg():
    return psycopg.connect(config.PG_ADMIN_DSN, autocommit=True)


def v(e):
    return e["value"] if isinstance(e, dict) else e


@pytest.fixture(scope="module")
def pledged(tok):
    """개인정보가 든 파일(대장이 든 엑셀 · 보고서)은 이용 약속 뒤 ok=1 로 받는다(10-18 영농-q3 · impl-45). 시험이 적은 약속은 끝에 지운다."""
    r = httpx.post(B + "/spaces/me/pledge", json={"purpose": "시험 — 대장과 AI 결과 확인", "until": time.strftime("%Y-%m-%d"), "agree": True},
                   headers=H(tok["namwon"]), timeout=60)
    assert r.status_code == 200, r.text[:200]
    yield r.json()["active"]["id"]
    with pg() as c:
        c.execute("DELETE FROM download_pledges WHERE id=%s", (r.json()["active"]["id"],))


@pytest.fixture(scope="module")
def rep(tok):
    r = httpx.get(B + "/spaces/me/report", params={"card": CARD}, headers=H(tok["namwon"]), timeout=180)
    if r.status_code == 404:
        pytest.skip("남원 경작·휴경 결과 설명서 없음")
    assert r.status_code == 200, r.text[:300]
    return r.json()


def test_report_same_numbers(rep, tok):
    t = tok["namwon"]
    g = httpx.get(B + f"/spaces/me/guides/{CARD}", headers=H(t), timeout=60).json()
    total = v(rep["total"])
    assert total == g["body"]["what"]["total"]["value"]                      # 대시보드 큰 숫자(결과 설명서)와 같은 값
    pj = httpx.get(B + f"/spaces/me/guides/{CARD}/parcels", params={"limit": 1}, headers=H(t), timeout=120).json()
    assert total == pj["total"]["value"]                                       # 필지 목록 total 과 같은 값
    assert sum(e["n"] for e in rep["emd_rows"]) == total                       # 읍면동 합 = 합계
    assert sum(v(c["n"]) for c in rep["classes"]) == total                     # 분류별 합 = 합계
    assert round(sum(v(e["m2"]) for e in rep["emd_rows"]) / 1e4, 1) == v(rep["area_ha"])
    assert round(sum(v(e["ha"]) for e in rep["emd_rows"]), 1) == v(rep["area_ha"])          # 읍면동 ha 합 = 합계(GPT5-3)
    assert v(rep["acc"]) is not None and 0 < v(rep["acc"]) <= 100
    assert rep["services"] and all(s["card"] for s in rep["services"])
    words = str(rep)
    assert not any(w in words for w in ("의심", "현장 확인", "소견", "조치 제안"))


def test_report_emd_and_hwpx(rep, tok, pledged):
    t = tok["namwon"]
    emd = rep["emd_rows"][0]["emd"]
    e = httpx.get(B + "/spaces/me/report", params={"card": CARD, "emd": emd}, headers=H(t), timeout=120).json()
    assert v(e["total"]) == rep["emd_rows"][0]["n"] and e["emd"] == emd
    assert httpx.get(B + "/spaces/me/report", params={"card": CARD, "emd": "없는읍"}, headers=H(t), timeout=60).status_code == 404
    png = httpx.get(B + "/spaces/me/report/map.png", params={"card": CARD}, headers=H(t), timeout=180)
    assert png.status_code == 200 and png.content[:4] == b"\x89PNG"
    d = httpx.get(B + "/spaces/me/report.hwpx", params={"card": CARD, "ok": 1}, headers=H(t), timeout=180)
    assert d.status_code == 200 and d.headers["content-type"].startswith("application/hwp+zip")
    from landxi_api import hwpx as HX
    assert HX.check(d.content) == [] and "BinData/image1.png" in zipfile.ZipFile(io.BytesIO(d.content)).namelist()   # 지도 그림 1장
    assert ".hwpx" in d.headers["content-disposition"]
    xml = HX.text_of(d.content)
    assert httpx.get(B + "/spaces/me/report.docx", params={"card": CARD}, headers=H(t), timeout=60).status_code == 404   # .docx 는 더 내지 않음
    assert "읍면동별 개수 · 면적" in xml and f"{v(rep['total']):,}" in xml
    assert not any(w in xml for w in BAD)
    if rep.get("ledger"):
        assert "우리 대장과 다른 필지" in xml
    # 다른 기관 · 받지 않은 서비스 · LX 계정
    assert httpx.get(B + "/spaces/me/report", params={"card": CARD}, headers=H(tok["gj"]), timeout=60).status_code == 404
    assert httpx.get(B + "/spaces/me/report", params={"card": CARD}, headers=H(tok["staff"]), timeout=60).status_code in (401, 403)


def test_compare_facts(tok, pledged):
    t = tok["namwon"]
    c = httpx.get(B + "/spaces/me/ledger/compare", params={"card": CARD, "geom": 1}, headers=H(t), timeout=180).json()
    if not c.get("ledger"):
        pytest.skip("남원 올린 대장 없음")
    n = v(c["ledger"]["parcels"])
    assert v(c["with_ai"]) + v(c["no_ai"]) == n
    assert v(c["diff"]) == sum(v(x["n"]) for x in c["lines"]) == len(c["items"])
    assert [x["key"] for x in c["lines"]] == ["farm_off", "off_farm"]          # 경작·휴경 서비스 = 두 줄
    assert all(x["k"] == "diff" and x["why"] in ("farm_off", "off_farm") for x in c["items"])
    assert len(c["geojson"]["features"]) <= n
    words = str({k: c[k] for k in ("lines", "items", "no_ai_label")})
    assert not any(w in words for w in ("의심", "확인 필요", "판정", "우선", "규칙"))
    # 보고서 ④ 칸과 같은 값
    r = httpx.get(B + "/spaces/me/report", params={"card": CARD}, headers=H(t), timeout=180).json()
    assert r["ledger"] and v(r["ledger"]["diff"]) == v(c["diff"]) and v(r["ledger"]["with_ai"]) == v(c["with_ai"])
    assert httpx.get(B + "/spaces/me/ledger/compare.xlsx", params={"card": CARD}, headers=H(t), timeout=180).status_code == 409   # 받을 때마다 확인
    x = httpx.get(B + "/spaces/me/ledger/compare.xlsx", params={"card": CARD, "ok": 1}, headers=H(t), timeout=180)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content), read_only=True)["대장 대조"]   # 엑셀 한 벌(10-22) — 4행 머리 · 5행부터 자료 · 마지막 합계
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True)]
    assert rows[-1][0] == "합계" and len(rows) - 1 == n
    # 다른 기관 계정은 이 대장을 못 본다
    assert httpx.get(B + "/spaces/me/ledger/compare", params={"card": "card-marine", "import_id": c["ledger"]["import_id"]},
                     headers=H(tok["gj"]), timeout=60).status_code == 404
    # 보고서에 넣기
    ok = httpx.post(B + "/spaces/me/ledger/compare/report", json={"card": CARD, "import_id": c["ledger"]["import_id"]}, headers=H(t), timeout=60)
    assert ok.status_code == 200 and ok.json()["import_id"] == c["ledger"]["import_id"]
    assert httpx.post(B + "/spaces/me/ledger/compare/report", json={"card": "card-nope"}, headers=H(t), timeout=60).status_code == 404


def test_upload_by_jibun_then_compare(tok):
    """올리기(지번만 · 필지 번호 없음) → 서버 반입 · 맞춤 → 결과. 시험 대장은 끝나면 지우고 기관 최근 대장을 되돌린다."""
    t0 = tok["namwon"]
    base = httpx.get(B + "/spaces/me/ledger/compare", params={"card": CARD}, headers=H(t0), timeout=180).json()
    if not base.get("ledger") or not base["items"]:
        pytest.skip("견줄 대장 · 다른 필지 없음")
    with pg() as c:
        prev = c.execute("SELECT id FROM ledger_imports WHERE tenant_id='namwon' AND kind='farm_ledger' AND latest").fetchone()
        addr = {r[0]: r[1] for r in c.execute("SELECT pnu, addr FROM survey_parcels WHERE pnu = ANY(%s)",
                                               ([x["pnu"] for x in base["items"][:2]],)).fetchall()}
    uid, tt = temp_account("tenant", "u_namwon_t18", "t18@namwon.test", "manager", "namwon")
    iid = None
    try:
        rows = ["PNU,지번,상태"] + [f",{addr[p]},답" for p in addr]
        fd = {"kind": (None, "farm_ledger"), "file": ("시험_농지대장.csv", ("﻿" + "\r\n".join(rows)).encode("utf-8"), "text/csv")}
        j = httpx.post(B + "/t/namwon/survey/registry/import", files=fd, headers=H(tt), timeout=120).json()
        iid = j["import_id"]
        httpx.post(B + f"/t/namwon/survey/registry/{iid}/confirm", json={"mapping": {"pnu": "PNU", "jibun": "지번", "status": "상태"}}, headers=H(tt), timeout=60)
        for _ in range(120):
            d = httpx.get(B + f"/t/namwon/survey/registry/{iid}", headers=H(tt), timeout=60).json()
            if d["state"] in ("matched", "failed"):
                break
            time.sleep(1)
        assert d["state"] == "matched"
        c2 = httpx.get(B + "/spaces/me/ledger/compare", params={"card": CARD, "import_id": iid}, headers=H(tt), timeout=120).json()
        assert v(c2["ledger"]["parcels"]) == len(addr) and v(c2["diff"]) == len(addr)       # 지번으로 이은 두 필지 = 둘 다 다른 필지
        assert {x["pnu"] for x in c2["items"]} == set(addr)
    finally:
        with pg() as c:
            if iid:
                c.execute("DELETE FROM survey_findings WHERE import_id=%s", (iid,))
                c.execute("DELETE FROM registry_snapshots WHERE import_id=%s", (iid,))
                c.execute("DELETE FROM ledger_recent WHERE import_id=%s", (iid,))
                c.execute("DELETE FROM ledger_imports WHERE id=%s", (iid,))
            if prev:
                c.execute("UPDATE ledger_imports SET latest=true WHERE id=%s", (prev[0],))
        drop_account("tenant", uid)


def test_public_record():
    r = httpx.get(B + "/public/record", timeout=60)
    assert r.status_code == 200
    j = r.json()
    with pg() as c:
        done = {x[0] for x in c.execute("SELECT sgg_cd FROM survey_sgg WHERE state='done'").fetchall()}
        # XI맵 전국 결과와 같은 줄(GPT5-1) · 서비스 × 시군구 보이기 규칙(원칙 293 · 316 — 보령 곤포처럼 뺀 결과는 세지 않는다 · 10-22 바퀴 22)
        from landxi_api.svc_rules import SQL_NOT_EXCLUDED
        svc = c.execute("SELECT count(DISTINCT s.card_id) FROM service_results s WHERE s.replaced_at IS NULL AND s.origin IN ('deploy', 'job')" + SQL_NOT_EXCLUDED).fetchone()[0]
    assert v(j["sgg"]) == len(j["places"]) == len(done)                        # 실태조사 집계가 끝난 시군구와 같은 수
    assert {p["sgg_cd"] for p in j["places"]} == done
    assert v(j["services"]["n"]) == len(j["services"]["names"]) == svc
    assert v(j["orgs"]["n"]) == len(j["orgs"]["items"]) >= 1
    assert len(j["work"]["names"]) == v(j["work"]["n"]) >= len(j["services"]["names"])
    pub = httpx.get(B + "/regions", params={"public": 1}, timeout=60).json()
    codes = {x["sgg_cd"] for x in pub["items"]} | {x.get("prev_cd") for x in pub["items"]}
    assert {p["sgg_cd"] for p in j["places"]} <= codes                         # 점은 모두 /regions 의 시군구


def test_screens_do_not_call_survey():
    from pathlib import Path
    root = Path(config.SERVER_ROOT).parent / "landxi" / "v3"
    for f in ("gov-fusion/svc-app.js", "gov-report/svc-report.js"):
        s = (root / f).read_text(encoding="utf-8")
        assert "api('/survey" not in s and "api(`/survey" not in s, f
    assert '<script type="module" src="svc-app.js">' in (root / "gov-fusion/index.html").read_text(encoding="utf-8")
