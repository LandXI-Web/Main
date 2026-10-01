"""구현 3차 · 기관 분기 공간 1단 — 결과 설명서 자동 생성 · 관할 밖 내려받기 0 · 부서 사용자 범위 · LX 관리자 기관 공간 목록.

확인 대장 13차 분기-2(1차 결과는 '결과 설명서'에서 확인) · 분기-3 ⓒ의 1단(모든 기관의 가벼운 칸) · 8차 API-형식 ⓐ(GeoJSON · 필지 엑셀 · 요약) ·
원칙 38(부서는 자기 서비스만) · 39(계정 범위 밖 0 — 서버가 자른다) · 59(내려받기는 동의 · 범위 · 기록) · 72(LX 관리자는 보고 돕는다).
실행: python -m pytest server/tests/test_impl3_spaces.py -q   (게이트웨이 :8700 이 떠 있어야 한다 — 판 생성 논리 시험은 게이트웨이 없이도 돈다)
운영 DB 를 쓰므로 시험이 만든 것(임시 부서 사용자 · 배정 · 동의 · 내려받기 기록 · 시험 판)은 끝나면 지운다. 지역 고정값 대신 두 기관(남원 · 광주전남)을 같이 본다.
"""
import asyncio
import datetime as dt
import io
import json

import httpx
import psycopg
import pytest

from landxi_api import config

from conftest import B, H, drop_account, temp_account

SIX = ("what", "where", "when", "format", "trust", "version")
TENANTS = {"namwon": "namwon", "gj": "gwangju-jeonnam"}


def _sql(q, args=(), many=False):
    with psycopg.connect(config.PG_ADMIN_DSN, autocommit=True) as c:
        cur = c.execute(q, args)
        if cur.description is None:
            return None
        return cur.fetchall() if many else cur.fetchone()


def _open_cards(tok, tenant):
    j = httpx.get(B + f"/brand/{tenant}", headers=H(tok), timeout=60).json()
    return [s["card"] for s in j["services"] if s.get("open")]


@pytest.fixture(scope="module")
def since(live):
    return dt.datetime.now(dt.timezone.utc)


@pytest.fixture(scope="module", autouse=True)
def cleanup(live, since):
    """시험이 남긴 동의 · 내려받기 기록을 지운다(기관 관리자 계정의 원래 동의는 시험 전 상태로)."""
    before = {r[0]: r[1] for r in (_sql("SELECT tenant_id, at FROM space_consents WHERE user_id IN ('u_namwon_mail_lxadmin','u_gwangju-jeonnam_mail_lxadmin')",
                                        many=True) or [])}
    yield
    _sql("DELETE FROM space_log WHERE kind IN ('download','consent','assign') AND at >= %s", (since,))
    for t, uid in (("namwon", "u_namwon_mail_lxadmin"), ("gwangju-jeonnam", "u_gwangju-jeonnam_mail_lxadmin")):
        if t not in before:
            _sql("DELETE FROM space_consents WHERE tenant_id=%s AND user_id=%s", (t, uid))


def _consent(tok):
    r = httpx.post(B + "/spaces/me/consent", json={"agree": True}, headers=H(tok), timeout=30)
    assert r.status_code == 200


# ── 결과 설명서 — 받은 1차 서비스마다 저절로 · 여섯 칸 ───────────────────────────
@pytest.mark.parametrize("key", ["namwon", "gj"])
def test_every_received_service_has_guide(tok, key):
    t = TENANTS[key]
    me = httpx.get(B + "/spaces/me", headers=H(tok[key]), timeout=300)
    assert me.status_code == 200
    j = me.json()
    cards = [s["card"] for s in j["services"]]
    assert cards == _open_cards(tok[key], t)                      # 기관 서비스 선택과 같은 판정
    assert cards, "받은 서비스가 하나도 없다"
    for s in j["services"]:
        assert s["edition"]["value"] >= 1                         # 사람이 쓰지 않아도 판이 있다
        g = httpx.get(B + f"/spaces/me/guides/{s['card']}", headers=H(tok[key]), timeout=120).json()
        assert all(k in g["body"] for k in SIX)
        assert g["change"]                                        # 바뀐 점 한 줄
        assert g["history"][0]["edition"]["value"] >= g["edition"]["value"]
        assert "sets" not in g["body"]                            # 내부 결과 세트 이름은 내지 않는다
        assert len(g["body"]["format"]["files"]) == 3             # GeoJSON · 필지 엑셀 · 요약(8차 API-형식 ⓐ)
        assert "API" not in json.dumps(g, ensure_ascii=False)
        w = g["body"]["what"]
        if w["kind"] != "biz":                                    # 업무 결과가 아니면 개수 0(분석 칸 도형 조각 수 노출 금지 · 사용자 규칙 2)
            assert w["total"]["value"] is None and not any("n" in c for c in w["classes"])
            assert not g["body"]["trust"]["checks"] and g["body"]["where"]["outside_emd"]["value"] is None


def test_guide_numbers_match_summary(tok):
    """숫자 한 출처 — 설명서 결과 수(관할 안) = 대표 수치 요약의 AI 탐지(같은 서비스 · 같은 기관)."""
    for key, card in (("namwon", "card-farm"), ("gj", "card-marine")):
        g = httpx.get(B + f"/spaces/me/guides/{card}", headers=H(tok[key]), timeout=120).json()
        s = httpx.get(B + f"/summary?card={card}", headers=H(tok[key]), timeout=120).json()
        det = sum(i["metrics"]["detected"]["value"] or 0 for i in s["items"])
        assert g["body"]["what"]["total"]["value"] == det
        assert sum(c["n"]["value"] for c in g["body"]["what"]["classes"]) == det


def test_new_round_makes_new_edition_and_notice(live):
    """판 생성 논리(게이트웨이 없이 · 운영 배포본을 건드리지 않는다): 처음 = 지난 공개분 채우기(알림 없음) → 새 회차 = 새 판 + 알림 한 줄 →
    되돌리면 그 판으로(새 판 없음)."""
    from landxi_api import spaces as SP
    t, card = "namwon", "card-space-selftest"
    rows0 = [{"id": "dp-space-selftest", "card_version_id": "card-farm@2.1", "prev_card_version_id": "card-farm@2.0",
              "snapshot_current": "results/namwon/dp-nw-farm-25@2.1", "snapshot_prev": "results/namwon/dp-nw-farm-25@2.0",
              "stage": "ga", "year": 2025, "scale": {"as_of": "2026-06-08"}, "sgg_cd": "52190", "region_name": None}]
    state = {"rows": rows0}

    async def fake_received(tenant):
        return [{"card": card, "name": "시험 서비스", "open": True, "status": "운영"}] if tenant == t else []

    async def fake_rows(tenant, c):
        return [dict(r) for r in state["rows"]] if (tenant, c) == (t, card) else []

    old = (SP.received, SP._live_rows)
    SP.received, SP._live_rows = fake_received, fake_rows
    try:
        _sql("DELETE FROM space_guides WHERE tenant_id=%s AND card_id=%s", (t, card))
        _sql("DELETE FROM space_log WHERE tenant_id=%s AND card_id=%s", (t, card))

        async def run():
            from landxi_api import deps
            deps._pool = None                                    # 이 시험의 이벤트 루프에서 새로
            await SP._sync_one(t)
            first = _sql("SELECT edition, backfill, change FROM space_guides WHERE tenant_id=%s AND card_id=%s ORDER BY edition", (t, card), many=True)
            # 새 회차(서비스 버전 · 결과 회차가 바뀜) → 새 판 + 알림(지난 공개분 아님)
            state["rows"] = [{**rows0[0], "card_version_id": "card-farm@2.0", "prev_card_version_id": None,
                              "snapshot_current": "results/namwon/dp-nw-change@1.0", "snapshot_prev": None, "scale": {"as_of": "2026-06-08"}}]
            await SP._sync_one(t)
            second = _sql("SELECT edition, backfill, change FROM space_guides WHERE tenant_id=%s AND card_id=%s ORDER BY edition", (t, card), many=True)
            # 되돌리기(처음 회차로) → 새 판 없음 · 지금 판 = 그 판
            state["rows"] = rows0
            await SP._sync_one(t)
            third = _sql("SELECT count(*) FROM space_guides WHERE tenant_id=%s AND card_id=%s", (t, card))[0]
            async with deps.db(realm="lx") as c:
                cur = await SP._current(c, t, card, rows0)
            await deps.close()
            return first, second, third, cur

        first, second, third, cur = asyncio.run(run())
        assert [r[0] for r in first] == [1, 2] and all(r[1] for r in first)        # 직전 회차(2.0) + 지금 회차(2.1) — 지난 공개분
        # 직전 회차 = 분석 칸마다 잘린 도형 조각(셈 단위 polygons) — 개수를 어디에도 싣지 않는다(카드의 '업무 결과만'과 같은 기준)
        assert first[0][2] == "처음 공개 — 2023년 항공영상"
        b1 = _sql("SELECT body FROM space_guides WHERE tenant_id=%s AND card_id=%s AND edition=1", (t, card))[0]
        assert b1["what"]["kind"] == "fragments" and b1["what"]["total"]["value"] is None
        assert not any("n" in c or "area" in c for c in b1["what"]["classes"]) and not b1["trust"]["checks"]
        assert "결과 단위가 필지로 바뀜 · 2,098필지" in first[1][2] and "129,420" not in first[1][2]
        import re as _re
        assert not _re.search(r"\d{2,3},\d{3}건", first[1][2])                    # 조각 수(예: 129,420건)가 바뀐 점에 없다
        assert len(second) == 3 and second[2][1] is False                          # 새 회차 = 3판 · 새 알림
        assert "서비스 버전" in second[2][2] and "→" in second[2][2]
        notes = _sql("SELECT line, backfill FROM space_log WHERE tenant_id=%s AND card_id=%s AND kind='guide' ORDER BY id", (t, card), many=True)
        assert len(notes) == 3 and notes[-1][1] is False and "3판" in notes[-1][0]
        assert third == 3 and cur["edition"] == 2                                   # 되돌리면 그 판(새 판 없음)
        f = config.DATA_ROOT / "tenants" / t / "space" / "guides" / card / "3.json"
        assert f.exists() and json.loads(f.read_text(encoding="utf-8"))["edition"] == 3   # 저장 폴더에도 판 한 벌
    finally:
        SP.received, SP._live_rows = old
        _sql("DELETE FROM space_guides WHERE tenant_id=%s AND card_id=%s", (t, card))
        _sql("DELETE FROM space_log WHERE tenant_id=%s AND card_id=%s", (t, card))
        import shutil
        shutil.rmtree(config.DATA_ROOT / "tenants" / t / "space" / "guides" / card, ignore_errors=True)


def test_light_space_cells_exist(live):
    """가벼운 칸 — 기관마다 공간 줄 · 자료 칸(스키마) · 저장 폴더."""
    for t in ("namwon", "gwangju-jeonnam"):
        sp = _sql("SELECT mode, db_schema, folder FROM spaces WHERE tenant_id=%s", (t,))
        assert sp and sp[0] == "light"
        assert _sql("SELECT 1 FROM information_schema.schemata WHERE schema_name=%s", (sp[1],))
        assert (config.DATA_ROOT / sp[2]).is_dir()


# ── 관할 밖 내려받기 0 ─────────────────────────────────────────────────────────
def test_scope_clip_points():
    """모양으로 판정하는 결과(읍면동 · 필지 번호 없음 — 바다 등): 관할 안만 남는다."""
    from landxi_api import spaces as SP
    pts = [{"id": 1, "x": 127.39, "y": 35.41},    # 남원시 안
           {"id": 2, "x": 127.70, "y": 34.73},    # 여수 앞바다(광주전남 연안)
           {"id": 3, "x": 126.98, "y": 37.56}]    # 서울(둘 다 관할 밖)
    assert SP._keep_points("namwon", pts) == [1]
    assert SP._keep_points("gwangju-jeonnam", pts) == [2]
    assert SP._scope_codes("namwon") == ["52190"]
    assert "12130" in SP._scope_codes("gwangju-jeonnam") and "52190" not in SP._scope_codes("gwangju-jeonnam")


def test_other_tenant_guide_and_download_are_absent(tok):
    """남원은 광주전남의 해양쓰레기를 · 광주전남은 남원의 영농을 열 수도 받을 수도 없다(없는 서비스 = 404)."""
    _consent(tok["namwon"])
    _consent(tok["gj"])
    for key, card in (("namwon", "card-marine"), ("gj", "card-farm"), ("gj", "card-living")):
        assert httpx.get(B + f"/spaces/me/guides/{card}", headers=H(tok[key]), timeout=60).status_code == 404
        for fmt in ("geojson", "parcels", "summary"):
            r = httpx.get(B + f"/spaces/me/guides/{card}/download?fmt={fmt}", headers=H(tok[key]), timeout=60)
            assert r.status_code == 404
    # LX 계정 · 게스트는 기관 공간이 없다
    assert httpx.get(B + "/spaces/me", headers=H(tok["admin"]), timeout=30).status_code == 403
    assert httpx.get(B + "/spaces/me", timeout=30).status_code == 401


def test_geojson_only_inside_scope(tok):
    _consent(tok["namwon"])
    _consent(tok["gj"])
    for key, card, code in (("namwon", "card-farm", "52190"), ("gj", "card-marine", None)):
        g = httpx.get(B + f"/spaces/me/guides/{card}", headers=H(tok[key]), timeout=120).json()
        r = httpx.get(B + f"/spaces/me/guides/{card}/download?fmt=geojson", headers=H(tok[key]), timeout=300)
        assert r.status_code == 200 and "attachment" in r.headers.get("content-disposition", "")
        fc = r.json()
        assert fc["type"] == "FeatureCollection" and fc["lx"]["안내"]
        assert len(fc["features"]) == g["body"]["what"]["total"]["value"]
        props = [f["properties"] for f in fc["features"]]
        assert not any(k in p for p in props for k in ("job_id", "shard_id", "chip_edge"))      # 내부 값 0
        if code:
            assert all((p["parcel"] or code).startswith(code) for p in props)
    # 남원 결과 번호는 광주전남 파일에 하나도 없다(서로의 결과 0)
    nw = {f["id"] for f in httpx.get(B + "/spaces/me/guides/card-farm/download?fmt=geojson", headers=H(tok["namwon"]), timeout=300).json()["features"]}
    gj = {f["id"] for f in httpx.get(B + "/spaces/me/guides/card-marine/download?fmt=geojson", headers=H(tok["gj"]), timeout=300).json()["features"]}
    assert nw and gj and not (nw & gj)


def test_parcel_excel_and_summary(tok):
    from openpyxl import load_workbook
    _consent(tok["namwon"])
    g = httpx.get(B + "/spaces/me/guides/card-farm", headers=H(tok["namwon"]), timeout=120).json()
    files = {f["fmt"]: f for f in g["body"]["format"]["files"]}
    assert files["parcels"]["ok"]
    r = httpx.get(B + "/spaces/me/guides/card-farm/download?fmt=parcels", headers=H(tok["namwon"]), timeout=300)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["필지"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == files["parcels"]["n"]["value"]
    assert all(isinstance(x[0], str) and len(x[0]) == 19 and x[0].startswith("52190") for x in rows)   # 필지 번호 = 글자 · 관할 안
    assert ws["A2"].number_format == "@"
    assert {"열 설명", "안내"} <= set(wb.sheetnames)
    # 바다 결과(필지 없음) — 필지 엑셀은 이유 한 줄과 함께 막힌다
    gm = httpx.get(B + "/spaces/me/guides/card-marine", headers=H(tok["gj"]), timeout=120).json()
    pf = {f["fmt"]: f for f in gm["body"]["format"]["files"]}["parcels"]
    assert pf["ok"] is False and pf["why"]
    _consent(tok["gj"])
    assert httpx.get(B + "/spaces/me/guides/card-marine/download?fmt=parcels", headers=H(tok["gj"]), timeout=60).status_code == 404
    # 요약 — 설명서와 같은 숫자
    s = httpx.get(B + "/spaces/me/guides/card-farm/download?fmt=summary", headers=H(tok["namwon"]), timeout=120).json()
    assert s["결과"]["값"] == g["body"]["what"]["total"]["value"] and s["안내"]
    # 내려받은 기록 — 기관 관리자가 본다
    d = httpx.get(B + "/spaces/me/downloads", headers=H(tok["namwon"]), timeout=30).json()
    assert any("필지 엑셀" in x["line"] for x in d["items"])


# ── 부서 사용자 — 정해 준 서비스만 · 동의 먼저 ─────────────────────────────────
@pytest.fixture()
def viewer(live):
    uid, tok = temp_account("tenant", "u_namwon_space_viewer_t", "space-viewer-t@example.go.kr", "viewer", "namwon", "시험 부서 사용자")
    yield uid, tok
    _sql("DELETE FROM space_assign WHERE user_id=%s", (uid,))
    _sql("DELETE FROM space_consents WHERE user_id=%s", (uid,))
    _sql("DELETE FROM space_reads WHERE user_id=%s", (uid,))
    _sql("DELETE FROM space_log WHERE actor=%s OR detail->>'user' = %s", (uid, uid))
    drop_account("tenant", uid)


def test_viewer_sees_only_assigned(tok, viewer):
    uid, vt = viewer
    me = httpx.get(B + "/spaces/me", headers=H(vt), timeout=120).json()
    assert me["services"] == [] and me["notices"] == [] and me["can_assign"] is False
    assert httpx.get(B + "/spaces/me/guides/card-farm", headers=H(vt), timeout=60).status_code == 404
    assert httpx.put(B + f"/spaces/me/assign/{uid}", json={"cards": ["card-farm"]}, headers=H(vt), timeout=30).status_code == 403
    # 기관 관리자가 영농 하나만 정해 준다
    r = httpx.put(B + f"/spaces/me/assign/{uid}", json={"cards": ["card-farm"]}, headers=H(tok["namwon"]), timeout=60)
    assert r.status_code == 200 and r.json()["cards"] == ["card-farm"]
    me = httpx.get(B + "/spaces/me", headers=H(vt), timeout=120).json()
    assert [s["card"] for s in me["services"]] == ["card-farm"]
    assert all(n["card"] == "card-farm" for n in me["notices"])
    assert httpx.get(B + "/spaces/me/guides/card-farm", headers=H(vt), timeout=60).status_code == 200
    assert httpx.get(B + "/spaces/me/guides/card-road", headers=H(vt), timeout=60).status_code == 404
    # 동의 전 내려받기 = 409 → 동의 뒤 200
    assert httpx.get(B + "/spaces/me/guides/card-farm/download?fmt=summary", headers=H(vt), timeout=60).status_code == 409
    httpx.post(B + "/spaces/me/consent", json={"agree": True}, headers=H(vt), timeout=30)
    assert httpx.get(B + "/spaces/me/guides/card-farm/download?fmt=summary", headers=H(vt), timeout=60).status_code == 200
    # 받지 않은 서비스 · 다른 기관 사람은 정할 수 없다 · 부서 사용자는 기록을 못 본다
    assert httpx.put(B + f"/spaces/me/assign/{uid}", json={"cards": ["card-marine"]}, headers=H(tok["namwon"]), timeout=30).status_code == 400
    assert httpx.put(B + f"/spaces/me/assign/{uid}", json={"cards": []}, headers=H(tok["gj"]), timeout=30).status_code == 404
    assert httpx.get(B + "/spaces/me/downloads", headers=H(vt), timeout=30).status_code == 403
    lst = httpx.get(B + "/spaces/me/assign", headers=H(tok["namwon"]), timeout=30).json()
    assert any(u["id"] == uid and u["cards"] == ["card-farm"] for u in lst["users"])


# ── LX 관리자 — 기관 공간 목록(보기만) ──────────────────────────────────────────
def test_admin_space_list(tok):
    r = httpx.get(B + "/spaces", headers=H(tok["admin"]), timeout=300)
    assert r.status_code == 200
    rows = {x["tenant"]: x for x in r.json()["items"]}
    assert {"namwon", "gwangju-jeonnam"} <= set(rows)
    for key, t in TENANTS.items():
        x = rows[t]
        assert x["services"]["value"] == len(_open_cards(tok[key], t))
        assert x["mode_word"] == "가벼운 칸" and x["latest"]["edition"]["value"] >= 1 and x["updated"]
        u = httpx.get(B + "/ops/tenants", headers=H(tok["admin"]), timeout=120).json()
        st = next(i for i in u["items"] if i["tenant_id"] == t)["dims"]["storage_gb"]["used"]["value"]
        assert abs(x["storage"]["value"] - st) < 0.01                   # '저장' = 사용 현황 탭과 같은 값(한 출처)
    d = httpx.get(B + "/spaces/namwon", headers=H(tok["admin"]), timeout=120).json()
    assert d["services"] and d["log"]
    # 기관 · LX 직원 세션은 목록을 못 본다 · 없는 기관 = 404
    assert httpx.get(B + "/spaces", headers=H(tok["namwon"]), timeout=30).status_code == 403
    assert httpx.get(B + "/spaces", headers=H(tok["staff"]), timeout=30).status_code == 403
    assert httpx.get(B + "/spaces/lx-demo", headers=H(tok["admin"]), timeout=30).status_code == 404
