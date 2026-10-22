"""기관 AI 분석 결과 보고서 · 행정정보와 비교(서비스 기준) — 확인 대장 10-11 18:52 'gov-보고서' · 'gov-비교' 이대로 구현 · 원칙 135 · 143 · 168 · 193.
메인 실적 숫자(QA-Q2 ⓐ · QA-Q3 ⓐ)도 여기서 센다(공개 · 로그인 없음).

  [기관 세션 · 관할 · 공유받은 서비스 범위 — 공간(spaces)과 같은 판정]
  GET  /api/v1/spaces/me/report?card=&emd=             AI 분석 결과 보고서 — 개요(합계 · 분류별 · AI 모델 정확도 · 영상 · 분석한 날) · 읍면동별 개수 · 면적 ·
                                                       (올린 대장이 있으면) 우리 대장과 다른 필지. 규칙 입력 없음 · 소견 · 조치 제안 없음.
  GET  /api/v1/spaces/me/report/map.png?card=&emd=     보고서 지도 그림(읍면동 경계 + 결과 — LX 보고서와 같은 그림 함수)
  GET  /api/v1/spaces/me/report.hwpx?card=&emd=        같은 내용의 한글 문서(.hwpx · 공용 서식 hwpx.py · 10-14 원칙 250)
  GET  /api/v1/spaces/me/ledger/compare?card=&import_id=&geom=   올린 대장 × 고른 서비스 결과(필지 번호로 맞춤) — 사실 그대로:
                                                       대장과 AI 결과가 다른 필지 · 어떻게 다른지 두 줄 · AI 결과 있는 필지 · AI 분석 전(영상 범위 밖) · 표
  GET  /api/v1/spaces/me/ledger/compare.xlsx?card=&import_id=    같은 행 엑셀(대장 필지 전부 · 다른 필지 먼저)
  POST /api/v1/spaces/me/ledger/compare/report {card, import_id}  '보고서에 넣기' — 이 대장을 이 사람의 보고서 대장으로(ledger_recent 기억)
  [공개]
  GET  /api/v1/public/record                           지금까지 분석 결과를 낸 시군구 · 서비스 · 운영 기관 · 행정 업무 서비스 이름(메인 장면)

숫자 한 출처: 보고서 합계 = 결과 설명서(space_guides) 의 결과 세트 × 관할(_scope_codes · _keep_ids) — 대시보드 큰 숫자 · 필지 목록과 같은 행.
대조는 실태조사 규칙(L-*)을 쓰지 않는다 — 결과 세트 detections.pnu × 대장 registry_snapshots.pnu.
"""
from __future__ import annotations

import asyncio
import datetime as dt
import io
import json
import re
import time
from urllib.parse import quote

from fastapi import APIRouter, Request
from fastapi.responses import Response
from starlette.concurrency import run_in_threadpool

from .deps import ApiError, Principal, audit, db, principal
from .envelope import KST, env, now_iso, split_round

router = APIRouter()
CARD_RE = re.compile(r"^card-[a-z0-9-]{2,40}$")
IMP_RE = re.compile(r"^imp_[0-9A-Z]{10,40}$")
SRC = "결과 설명서의 결과 세트(관할 안 · 지운 결과 제외)"
SRC_CMP = "올린 대장 필지 × 서비스 결과(필지 번호로 맞춤)"
# 지목 28가지(공간정보관리법 시행령) — 대장 값이 지목이면 그 값, 아니면 연속지적 지목
JIMOK = {"전", "답", "과수원", "목장용지", "임야", "광천지", "염전", "대", "공장용지", "학교용지", "주차장", "주유소용지", "창고용지", "도로", "철도용지",
         "제방", "하천", "구거", "유지", "양어장", "수도용지", "공원", "체육용지", "유원지", "종교용지", "사적지", "묘지", "잡종지"}
FARM = {"전", "답", "과수원"}
NONFARM_USE = ("건물", "건축물", "주차장")        # 농지 위에 있으면 대장과 다른 쓰임(비닐하우스 · 곤포사일리지는 농지 위가 보통이라 넣지 않는다)


def _me(request: Request) -> Principal:
    from .spaces import _me as me
    return me(request)


def _card(card: str) -> str:
    if not CARD_RE.match(card or ""):
        raise ApiError("not_found", "없는 서비스입니다")
    return card


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _short_svc(name: str) -> str:
    return re.sub(r"\s*(분석\s*서비스|서비스)$", "", str(name or "")).strip()


def _ymd(s) -> str | None:
    s = str(s or "")[:10]
    return s.replace("-", ".") if re.match(r"\d{4}-\d{2}-\d{2}", s) else None


def _addr_short(addr: str | None) -> str:
    """'전북특별자치도 남원시 주생면 도산리 986-6' → '도산리 986-6'(리 · 동 + 지번 — 표 한 칸)."""
    w = str(addr or "").split()
    return " ".join(w[-2:]) if len(w) >= 2 else (w[0] if w else "")


# ── 서비스 범위(공간과 같은 판정) ───────────────────────────────────────────
async def _svc(p: Principal, card: str):
    from .spaces import _guide_or_404, _only, _scope_codes
    svc, g, rows = await _guide_or_404(p, _card(card))
    body = g["body"] or {}
    only = await _only(p, card)
    return svc, g, rows, body, (body.get("sets") or []), _scope_codes(p.tenant_id, only), only


async def _services(p: Principal) -> list[dict]:
    """보고서 · 비교에서 고를 수 있는 서비스 — 기관에 열린(부서 배정 · 공유 범위 그대로) 서비스 가운데 결과가 있는 것(결과 설명서 순서)."""
    from .spaces import _current, _visible, received, sync
    await sync(p.tenant_id)
    out = []
    async with db(p) as c:
        for s in await _visible(p, await received(p.tenant_id)):
            g = await _current(c, p.tenant_id, s["card"])
            if not g:
                continue
            b = g["body"] or {}
            tot = (b.get("what", {}).get("total") or {})
            if not b.get("sets") or tot.get("value") in (None, 0):
                continue
            rnd = (b.get("when") or {}).get("rounds") or []
            out.append({"card": s["card"], "name": s["name"], "shot": next((r.get("shot") for r in rnd if r.get("shot")), None),
                        "unit": tot.get("unit") or "건", "parcel": bool((b.get("where") or {}).get("parcel"))})
    return out


async def _acc(rows: list[dict]) -> tuple[int | None, str | None]:
    """AI 모델 정확도 — 이 서비스 공개 배포본의 서비스 판(card_versions) 모델 기록(학습 끝 검증 · 분석하기 카드와 같은 함수 acc_of)."""
    from .release import _acc_at, acc_of
    cvs = sorted({r.get("card_version_id") for r in rows if r.get("card_version_id")}, reverse=True)
    if not cvs:
        return None, None
    async with db(realm="lx") as c:
        vers = await c.fetch("SELECT id, model_ids FROM card_versions WHERE id = ANY($1::text[])", cvs)
        mids = [m for v in sorted(vers, key=lambda x: x["id"], reverse=True) for m in (v["model_ids"] or [])]
        if not mids:
            return None, None
        ms = {r["id"]: r["metrics"] for r in await c.fetch("SELECT id, metrics FROM models WHERE id = ANY($1::text[])", mids)}
    for m in mids:
        mt = ms.get(m)
        mt = json.loads(mt) if isinstance(mt, str) else (mt or {})
        a = acc_of(mt)
        if a is not None:
            return a, _acc_at(mt)
    return None, None


async def _rows(p: Principal, sets: list[str], codes: list[str] | None, emd: str | None) -> tuple[list[dict], list | None]:
    """관할 안 결과 — (읍면동, 분류)별 개수 · 면적 + 범위(bbox). 결과 설명서 · 내려받기와 같은 거르기(_scope_sql · _keep_ids)."""
    from .spaces import NO_EMD, _keep_ids, _scope_sql
    if not sets:
        return [], None
    keep = await _keep_ids(p.tenant_id, sets) if codes is not None else []
    cond, args = _scope_sql(codes, 2)
    a = [sets] + args + ([keep] if codes is not None else [])
    ew = ""
    if emd:
        a.append(emd)
        ew = f" AND d.emd = ${len(a)}"
    from .result_scope import only_cls, only_sql
    await only_cls()
    ew += "".join(" AND " + x for x in only_sql(sets, a, "d"))         # 서비스 분류만 보이는 결과(옛 판으로 되돌린 결과 등)
    async with db(realm="lx") as c:
        await c.execute("SET LOCAL jit = off")
        rows = await c.fetch("SELECT coalesce(nullif(d.emd, ''), '') AS emd, d.cls, count(*) AS n, coalesce(sum(d.area_m2), 0) AS m2 FROM detections d "
                             f"WHERE d.job_id = ANY($1::text[]) AND d.edit_state <> 'deleted' AND {cond}{ew} GROUP BY 1, 2", *a)
        bb = await c.fetchrow("SELECT ST_XMin(e) x0, ST_YMin(e) y0, ST_XMax(e) x1, ST_YMax(e) y1 FROM (SELECT ST_Extent(d.geom) e FROM detections d "
                              f"WHERE d.job_id = ANY($1::text[]) AND d.edit_state <> 'deleted' AND {cond}{ew}) s", *a)
    out = [{"emd": "" if r["emd"] in NO_EMD else r["emd"], "cls": r["cls"], "n": int(r["n"]), "m2": float(r["m2"])} for r in rows]
    box = [round(float(bb[k]), 6) for k in ("x0", "y0", "x1", "y1")] if bb and bb["x0"] is not None else None
    return out, box


async def _geo(p: Principal, sets: list[str], codes: list[str] | None, emd: str | None) -> list:
    from .spaces import _keep_ids, _scope_sql
    keep = await _keep_ids(p.tenant_id, sets) if codes is not None else []
    cond, args = _scope_sql(codes, 2)
    a = [sets] + args + ([keep] if codes is not None else [])
    ew = ""
    if emd:
        a.append(emd)
        ew = f" AND d.emd = ${len(a)}"
    from .result_scope import only_cls, only_sql
    await only_cls()
    ew += "".join(" AND " + x for x in only_sql(sets, a, "d"))
    async with db(realm="lx") as c:
        rows = await c.fetch("SELECT d.cls, ST_AsBinary(ST_SimplifyPreserveTopology(d.geom, 0.000005)) AS b FROM detections d "
                             f"WHERE d.job_id = ANY($1::text[]) AND d.edit_state <> 'deleted' AND {cond}{ew} LIMIT 60000", *a)
    return [(r["cls"], r["b"]) for r in rows]


def _sgg_of(body: dict, rows: list[dict]) -> tuple[str | None, str]:
    from .regions import region_of
    for r in rows:
        if r.get("sgg_cd"):
            rg = region_of(r["sgg_cd"])
            if rg:
                return rg["sgg_cd"], rg.get("name") or ""
    places = (body.get("where") or {}).get("places") or []
    return None, (places[0] if places else "")


# ── 대장 × 서비스 결과 ─────────────────────────────────────────────────────
async def _import_of(p: Principal, import_id: str | None) -> dict | None:
    """대장 한 건 — 주어진 것(이 기관 것만) 또는 이 사람이 마지막으로 연(보고서에 넣은) 대장 → 이 사람이 올린 최근 → 기관 최근(공용 대장)."""
    async with db(realm="lx") as c:
        if import_id:
            if not IMP_RE.match(import_id):
                raise ApiError("not_found", "올린 대장이 없습니다")
            r = await c.fetchrow("SELECT * FROM ledger_imports WHERE id=$1 AND tenant_id=$2", import_id, p.tenant_id)
            if not r:
                raise ApiError("not_found", "올린 대장이 없습니다")
            return dict(r)
        if p.user_id and await c.fetchval("SELECT to_regclass('ledger_recent') IS NOT NULL"):
            r = await c.fetchrow("SELECT i.* FROM ledger_recent l JOIN ledger_imports i ON i.id=l.import_id AND i.tenant_id=l.tenant_id "
                                 "WHERE l.tenant_id=$1 AND l.user_id=$2 AND i.state='matched' AND i.kind <> 'feedcrop' ORDER BY (l.sgg_cd='*') DESC, l.at DESC LIMIT 1",
                                 p.tenant_id, p.user_id)
            if r:
                return dict(r)
        # 사료작물 재배 신청 대장(feedcrop)은 세 시점 견주기 화면이 따로 연다(10-18 영농-q1) — 서비스 하나와 맞추는 이 비교 · 보고서의 '최근 대장'에서는 뺀다
        r = await c.fetchrow("SELECT * FROM ledger_imports WHERE tenant_id=$1 AND state='matched' AND kind <> 'feedcrop' AND created_by=$2 ORDER BY created_at DESC LIMIT 1",
                             p.tenant_id, p.user_id) if p.user_id else None
        r = r or await c.fetchrow("SELECT * FROM ledger_imports WHERE tenant_id=$1 AND state='matched' AND kind <> 'feedcrop' ORDER BY latest DESC, created_at DESC LIMIT 1",
                                  p.tenant_id)
    return dict(r) if r else None


def _mode(body: dict) -> str:
    names = {str(x.get("name") or "") for x in (body.get("what") or {}).get("classes") or []}
    if names & {"경작지", "비경작지"}:
        return "farm"
    if any(n.startswith(NONFARM_USE) for n in names):
        return "use"
    return "none"


async def compare(p: Principal, card: str, import_id: str | None = None, emd: str | None = None, geom: bool = False) -> dict:
    from .ledger import KINDS, _uploader
    from .spaces import class_ko
    svc, g, rows, body, sets, codes, _only = await _svc(p, card)
    imp = await _import_of(p, import_id)
    if not imp:
        return {"service": {"card": card, "name": svc["name"]}, "ledger": None, "as_of": now_iso()}
    unit = ((body.get("what") or {}).get("total") or {}).get("unit") or "건"
    per_parcel = unit == "필지"
    mode = _mode(body)
    async with db(realm="lx") as c:
        await c.execute("SET LOCAL jit = off")
        led = {x["pnu"]: (x["payload"] or {}) for x in await c.fetch(
            "SELECT DISTINCT ON (pnu) pnu, payload FROM registry_snapshots WHERE import_id=$1 AND tenant_id=$2 AND pnu IS NOT NULL ORDER BY pnu, row_no",
            imp["id"], p.tenant_id)}
        loose = int(await c.fetchval("SELECT count(*) FROM registry_snapshots WHERE import_id=$1 AND tenant_id=$2 AND pnu IS NULL",
                                     imp["id"], p.tenant_id) or 0)
        pn = list(led)
        attrs = {a["pnu"]: dict(a) for a in await c.fetch("SELECT pnu, addr, emd, jimok_nm, area_m2 FROM survey_parcels WHERE pnu = ANY($1::text[])", pn)} if pn else {}
        if emd:                                                  # 읍면동 보고서 — 그 읍면동의 대장 필지만
            pn = [k for k in pn if (attrs.get(k) or {}).get("emd") == emd]
        ai: dict[str, dict] = {}
        if pn and sets:
            sc = " AND left(d.pnu, 5) = ANY($3::text[])" if codes is not None else ""
            for r in await c.fetch("SELECT d.pnu, d.cls, d.cls_en, count(*) n, coalesce(sum(d.area_m2), 0) m2 FROM detections d "
                                   f"WHERE d.job_id = ANY($1::text[]) AND d.edit_state <> 'deleted' AND d.pnu = ANY($2::text[]){sc} GROUP BY 1, 2, 3",
                                   *([sets, pn, codes] if codes is not None else [sets, pn])):
                e = ai.setdefault(r["pnu"], {"by": {}, "m2": 0.0, "n": 0})
                k = class_ko(r["cls"], r["cls_en"])
                e["by"][k] = e["by"].get(k, 0.0) + float(r["m2"])
                e["m2"] += float(r["m2"])
                e["n"] += int(r["n"])
        gj = {}
        if geom and pn:
            gj = {r["pnu"]: r["g"] for r in await c.fetch(
                "SELECT pnu, ST_AsGeoJSON(ST_SimplifyPreserveTopology(geom, 0.000004), 6) g FROM survey_parcels WHERE pnu = ANY($1::text[])", pn)}
        up = await _uploader(c, imp)
    lines = []
    pay = imp["kind"] == "direct_pay"      # 공익직불 신청 대장(10-22 제안-공익직불대장) — 신청한 필지는 경작해야 한다: AI 가 비경작지 · 휴경으로 본 신청 필지가 점검 대상
    if mode == "farm" and pay:
        lines = [{"key": "pay_off", "label": "직불 신청 필지인데 AI는 비경작지", "short": "신청했는데 비경작지"}]
    elif mode == "farm":
        lines = [{"key": "farm_off", "label": "대장은 농지(전·답)인데 AI는 비경작지", "short": "농지인데 비경작지"},
                 {"key": "off_farm", "label": "대장은 농지가 아닌데 AI는 경작지", "short": "농지 아닌데 경작지"}]
    elif mode == "use":
        seen = sorted({k for e in ai.values() for k in e["by"] if k.startswith(NONFARM_USE)})
        lines = [{"key": "use:" + k, "label": f"대장은 농지(전·답)인데 AI는 {k}", "short": f"농지인데 {k}"} for k in seen] or \
                [{"key": "use:-", "label": "대장은 농지(전·답)인데 AI는 건물 · 주차장", "short": "농지인데 건물 · 주차장"}]
    cnt = {x["key"]: 0 for x in lines}
    items = []
    for k in pn:
        a = attrs.get(k) or {}
        lv = str((led.get(k) or {}).get("status") or "").strip()
        jm = lv if lv in JIMOK else (a.get("jimok_nm") or "")
        e = ai.get(k)
        top = max(e["by"].items(), key=lambda kv: kv[1])[0] if e and e["by"] else None
        why = None
        if e:
            if mode == "farm" and pay:
                if top in ("비경작지", "휴경", "휴경지"):
                    why = "pay_off"
            elif mode == "farm":
                if jm in FARM and top == "비경작지":
                    why = "farm_off"
                elif jm and jm not in FARM and top == "경작지":
                    why = "off_farm"
            elif mode == "use" and jm in FARM:
                hit = sorted(x for x in e["by"] if x.startswith(NONFARM_USE))
                if hit:
                    why = "use:" + hit[0] if ("use:" + hit[0]) in cnt else "use:-"
        if why:
            cnt[why] = cnt.get(why, 0) + 1
        items.append({"pnu": k, "where": _addr_short(a.get("addr")) or k, "emd": a.get("emd") or "", "ledger": jm or lv or "—",
                      "ai": (" · ".join(f"{x} {round(m):,}㎡" for x, m in sorted(e["by"].items(), key=lambda kv: -kv[1]))) if e else None,
                      "ai_cls": top, "ai_m2": round(e["m2"], 1) if e else None, "parcel_m2": round(float(a["area_m2"]), 1) if a.get("area_m2") else None,
                      "area": {"ai": env(round(e["m2"], 1) if e else None, "m2", "measured", SRC_CMP),
                               "parcel": env(round(float(a["area_m2"]), 1) if a.get("area_m2") else None, "m2", "recorded", "연속지적 필지 면적")},
                      "k": "diff" if why else ("ai" if e else "none"), "why": why,
                      "_by": {x: round(m, 1) for x, m in e["by"].items()} if e else {}})
    diff = [x for x in items if x["k"] == "diff"]
    diff.sort(key=lambda x: (-(x["ai_m2"] or 0), x["where"]))
    has_ai = sum(1 for x in items if x["k"] != "none")
    none = len(items) - has_ai
    kind = KINDS.get(imp["kind"]) or "대장"
    as_of = (imp["confirmed_at"] or imp["created_at"]).astimezone(KST).isoformat(timespec="seconds") if (imp["confirmed_at"] or imp["created_at"]) else now_iso()
    out = {
        "service": {"card": card, "name": svc["name"], "shot": next((r.get("shot") for r in (body.get("when") or {}).get("rounds") or [] if r.get("shot")), None)},
        "ledger": {"import_id": imp["id"], "name": re.sub(r"\.(xlsx|xls)\.csv$", r".\1", imp["filename"] or "") or kind, "kind": kind, "kind_code": imp["kind"], "rows": env(int(imp["rows"] or 0), "count", "recorded", "올린 대장 행", as_of=as_of),
                   "parcels": env(len(items), "필지", "measured", "올린 대장 × 연속지적(필지 번호)", as_of=as_of),
                   "loose": env(loose, "행", "measured", "필지와 잇지 못한 대장 행", as_of=as_of) if loose else None,
                   "uploader": up, "date": _ymd(as_of), "base_date": _ymd(imp["base_date"].isoformat()) if imp.get("base_date") else None},
        "lines": [{**x, "n": env(cnt.get(x["key"], 0), "필지", "inferred", SRC_CMP)} for x in lines],
        "diff": env(len(diff) if lines else None, "필지", "inferred", SRC_CMP, None if lines else "이 서비스는 대장과 견줄 분류가 없습니다"),
        "with_ai": env(has_ai, "필지", "inferred", SRC_CMP),
        "no_ai": env(none, "필지", "inferred", SRC_CMP),
        "no_ai_label": "영상 범위 밖 · AI 분석 전" if per_parcel else "AI 결과가 없는 필지",
        "items": diff, "emd": emd or None, "as_of": now_iso(),
    }
    if geom:
        out["geojson"] = {"type": "FeatureCollection", "features": [
            {"type": "Feature", "id": x["pnu"], "geometry": json.loads(gj[x["pnu"]]),
             "properties": {"pnu": x["pnu"], "k": x["k"], "where": x["where"], "ledger": x["ledger"], "ai": x["ai"] or "", "parcel_m2": x["parcel_m2"]}}
            for x in items if gj.get(x["pnu"])]}
    out["_all"] = items
    return out


def _clean(items: list[dict]) -> list[dict]:
    """응답용 — 숫자는 봉투(area)로만(속 계산용 날 숫자 ai_m2 · parcel_m2 는 뺀다)."""
    return [{k: v for k, v in x.items() if k not in ("ai_m2", "parcel_m2", "_by")} for x in items]


def _public(o: dict) -> dict:
    out = {k: v for k, v in o.items() if not k.startswith("_")}
    if "items" in out:
        out["items"] = _clean(out["items"])
    return out


# ── 보고서 ───────────────────────────────────────────────────────────────
async def report_data(p: Principal, card: str, emd: str | None = None) -> dict:
    from .spaces import class_ko
    svc, g, rows, body, sets, codes, _only = await _svc(p, card)
    emd = (emd or "").strip() or None
    all_rows, all_box = await _rows(p, sets, codes, None)
    emds = sorted({r["emd"] for r in all_rows if r["emd"]})
    if emd and emd not in emds:
        raise ApiError("not_found", "그 읍면동에는 이 서비스 결과가 없습니다")
    rs, box = (all_rows, all_box) if not emd else await _rows(p, sets, codes, emd)
    unit = ((body.get("what") or {}).get("total") or {}).get("unit") or "건"
    acc, acc_at = await _acc(rows)
    rnd = (body.get("when") or {}).get("rounds") or []
    shot = next((r.get("shot") for r in rnd if r.get("shot")), None)
    analyzed = next((r.get("analyzed") for r in rnd if r.get("analyzed")), None)
    sgg, place = _sgg_of(body, rows)
    tot_n = sum(r["n"] for r in rs)
    tot_m2 = sum(r["m2"] for r in rs)
    cls = {}
    for r in rs:
        k = class_ko(r["cls"])
        e = cls.setdefault(k, [0, 0.0])
        e[0] += r["n"]
        e[1] += r["m2"]
    tot_ha = round(tot_m2 / 1e4, 1)                              # 합계 = 반올림 전 면적의 합을 한 번 반올림
    cls_l = sorted(cls.items(), key=lambda kv: -kv[1][0])
    cls_ha = split_round([v[1] / 1e4 for _, v in cls_l], 1, tot_ha)   # 분류 칸도 합이 합계와 같게(최대 나머지 · GPT5-3)
    from .farm_admin import bale_kg, is_bale, yield_env
    kg = await bale_kg(p.tenant_id) if any(is_bale(k) for k, _ in cls_l) else None       # 곤포사일리지 생산량 = 개수 × 기관이 넣은 1개당 무게(10-18 영농-q2 ⓑ)
    classes = [{"name": k, "n": env(v[0], unit, "inferred", SRC), "ha": env(cls_ha[i], "ha", "measured", SRC),
                **({"yield": yield_env(v[0], kg)} if kg and is_bale(k) else {})}
               for i, (k, v) in enumerate(cls_l)]
    by: dict[str, dict] = {}
    for r in rs:
        e = by.setdefault(r["emd"] or "읍면동 밖", {"emd": r["emd"] or "읍면동 밖", "n": 0, "m2": 0.0, "by": {}})
        e["n"] += r["n"]
        e["m2"] += r["m2"]
        k = class_ko(r["cls"])
        e["by"][k] = e["by"].get(k, 0) + r["n"]
    emd_l = sorted(by.values(), key=lambda x: (-x["n"], x["emd"]))
    emd_ha = split_round([e["m2"] / 1e4 for e in emd_l], 1, tot_ha)   # 읍면동 칸 합 = 합계(GPT5-3)
    emd_rows = [{"emd": e["emd"], "n": e["n"], "ha": env(emd_ha[i], "ha", "measured", SRC), "m2": env(round(e["m2"], 1), "m2", "measured", SRC),
                 "by": [{"cls": k, "n": v} for k, v in e["by"].items()]} for i, e in enumerate(emd_l)]
    ld = None
    try:
        c = await compare(p, card, None, emd)
        if c.get("ledger") and c["ledger"]["parcels"]["value"]:
            ld = _public(c)
            ld.pop("items", None)
            ld["top"] = _clean(c["items"][:5])
    except ApiError:
        ld = None
    region = f"{place} {emd}".strip() if emd else place
    return {
        "service": {"card": card, "name": svc["name"]}, "region": region or "—", "sgg": sgg, "emd": emd, "emds": emds,
        "unit": unit, "made": dt.datetime.now(KST).strftime("%Y.%m.%d"),
        "imagery": shot, "analyzed": _ymd(analyzed),
        "total": env(tot_n, unit, "inferred", SRC), "area_ha": env(tot_ha, "ha", "measured", SRC),
        "classes": classes,
        "acc": env(acc, "%", "measured", "모델 기록 · 학습에 쓰지 않은 영상으로 잰 값(분석하기 카드와 같은 출처)", as_of=acc_at or None),
        "emd_rows": emd_rows, "emd_n": env(len([e for e in emd_rows if e["emd"] != "읍면동 밖"]), "count", "inferred", SRC),
        "bbox": box, "ledger": ld, "services": await _services(p), "as_of": now_iso(),
    }


_PNG: dict[tuple, tuple[float, bytes]] = {}


def _png_file(key):
    """그린 그림의 디스크 사본 — 게이트웨이를 다시 켜도 첫 사람이 다시 그리기를 기다리지 않게(바퀴 15 재검수: 다시 켠 뒤 보고서 탭 처음 7.4초)."""
    import hashlib
    from . import config
    return config.DATA_ROOT / "cache" / "report-map" / (hashlib.sha1(repr(key).encode()).hexdigest()[:20] + ".png")


async def _png(p: Principal, card: str, emd: str | None, d: dict | None = None) -> bytes | None:
    svc, g, rows, body, sets, codes, only = await _svc(p, card)
    key = (p.tenant_id, card, emd or "", g["edition"], tuple(only or []))
    hit = _PNG.get(key)
    if hit is None:
        try:
            f = _png_file(key)
            hit = _PNG[key] = (f.stat().st_mtime, f.read_bytes())
        except OSError:
            hit = None
    if hit and time.time() - hit[0] < 600:
        return hit[1]
    if hit and not d:            # 10분 지난 그림 — 화면에는 바로 주고 뒤에서 새로 그린다(보고서 탭 처음 열 때 3~5초 기다리던 것 · 바퀴 14 재검수). 한글 문서(d 있음)는 늘 새로
        if key not in _PNG_BUSY:
            _PNG_BUSY.add(key)

            async def _again():
                try:
                    await _png_make(p, card, emd, None, key, sets, codes)
                except Exception:            # 다음 요청이 다시 그린다
                    pass
                finally:
                    _PNG_BUSY.discard(key)
            t = asyncio.create_task(_again())
            _PNG_TASKS.add(t)
            t.add_done_callback(_PNG_TASKS.discard)
        return hit[1]
    return await _png_make(p, card, emd, d, key, sets, codes)


_PNG_BUSY: set = set()
_PNG_TASKS: set = set()


async def _png_make(p: Principal, card: str, emd: str | None, d: dict | None, key, sets, codes) -> bytes | None:
    from agent.tools.ext.mapsvc import _map_png
    d = d or await report_data(p, card, emd)
    geo = await _geo(p, sets, codes, emd)
    counts = {e["emd"]: e["n"] for e in d["emd_rows"] if e["emd"] != "읍면동 밖"}
    png = await asyncio.to_thread(_map_png, "", d["sgg"], geo, counts)
    _PNG[key] = (time.time(), png)
    try:
        f = _png_file(key)
        f.parent.mkdir(parents=True, exist_ok=True)
        tmp = f.with_suffix(".tmp")
        tmp.write_bytes(png)
        tmp.replace(f)
    except OSError:
        pass
    while len(_PNG) > 48:
        _PNG.pop(next(iter(_PNG)))
    return png


def _hwpx(d: dict, png: bytes | None) -> tuple[bytes, str]:
    """AI 분석 결과 보고서 — 한글 문서(공용 서식 hwpx.py · 10-14 원칙 250). 숫자 = report_data(화면 보고서 탭과 같은 값)."""
    from . import hwpx as HX
    name, region, unit = d["service"]["name"], d["region"], d["unit"]
    now = dt.datetime.now(KST)
    made_t = now.strftime("%Y. %m. %d. %H:%M")
    doc = HX.Doc()
    doc.title("AI 분석 결과 보고서", f"{name} · {region}")
    doc.overview([("지역", region), ("분석 서비스", name), ("영상", d["imagery"] or "—"), ("분석한 날", d["analyzed"] or "—"),
                  ("AI 모델 정확도", f"{d['acc']['value']}%" if d["acc"]["value"] is not None else "기록 없음"), ("만든 날짜", made_t)])
    doc.notice("AI 분석 결과는 참고자료입니다 — 처분은 담당자가 정합니다.")
    doc.heading("결과 개요")
    rows = [["항목", "개수", "면적(ha)"], ["AI 분석 결과 합계", f"{d['total']['value']:,}{unit}", f"{d['area_ha']['value']:,.1f}"]]
    for c in d["classes"]:
        y = (c.get("yield") or {}).get("value")
        rows.append([c["name"] + (f" (생산량 {y:,.0f}kg)" if y else ""), f"{c['n']['value']:,}{unit}", f"{c['ha']['value']:,.1f}"])
    doc.table(rows, [44, 28, 28], aligns=["l", "r", "r"])
    if png:
        doc.heading("지도")
        doc.figure(png, f"{region} {_short_svc(name)} AI 분석 결과", "읍면동 경계 위에 AI 분석 결과를 칠했습니다", max_w=int(HX.TEXT_W * 0.82))
    doc.heading("읍면동별 개수 · 면적")
    cls_names = [c["name"] for c in d["classes"]]
    many = len(cls_names) > 1
    head = ["읍면동", *(cls_names if many else []), "계", "면적(ha)"]
    body = [[e["emd"], *([f"{({b['cls']: b['n'] for b in e['by']}).get(k, 0):,}" for k in cls_names] if many else []), f"{e['n']:,}", f"{e['ha']['value']:,.1f}"]
            for e in d["emd_rows"]]
    body.append(["합계", *([f"{c['n']['value']:,}" for c in d["classes"]] if many else []), f"{d['total']['value']:,}", f"{d['area_ha']['value']:,.1f}"])
    doc.table([head] + body, [3] + [2] * (len(head) - 1), aligns=["l"] + ["r"] * (len(head) - 1), total_row=True, small=len(head) > 6,
              caption=f"읍면동별 개수 · 면적 (단위: {unit} · ha)")
    ld = d.get("ledger")
    if ld:
        doc.heading("우리 대장과 다른 필지")
        lr = [["항목", "필지"], ["올린 대장", f"{ld['ledger']['name']} · {ld['ledger']['parcels']['value']:,}필지"],
              ["AI 결과가 있는 필지", f"{ld['with_ai']['value']:,}필지"]]
        for x in ld["lines"]:
            lr.append([x["label"], f"{x['n']['value']:,}필지"])
        lr.append([ld["no_ai_label"], f"{ld['no_ai']['value']:,}필지"])
        doc.table(lr, [60, 40], aligns=["l", "r"])
        if ld.get("top"):
            doc.table([["지번", "대장", "AI 분석", "필지 면적(㎡)"]] +
                      [[x["where"], x["ledger"], x["ai"] or "—", f"{x['area']['parcel']['value']:,.0f}" if x["area"]["parcel"]["value"] else "—"] for x in ld["top"]],
                      [22, 14, 44, 20], aligns=["l", "l", "l", "r"], caption="대장과 다른 필지(면적이 큰 순 5곳)")
    doc.source("Land-XI AI 분석 결과 — 개수 · 면적은 서비스 대시보드 · 필지 목록 · 내려받기 파일과 같은 값입니다(결과 확인 전 · 지운 결과 제외).", made_t)
    fname = HX.safe_name(f"{region}_{_short_svc(name)}_AI분석결과보고서_{d['made'].replace('.', '')}")
    return doc.package(f"AI 분석 결과 보고서 — {name} · {region}", now), fname


@router.get("/spaces/me/report")
async def gov_report(request: Request, card: str, emd: str | None = None):
    p = _me(request)
    return await report_data(p, card, emd)


@router.get("/spaces/me/report/services")
async def gov_report_services(request: Request):
    p = _me(request)
    return {"items": await _services(p), "as_of": now_iso()}


@router.get("/spaces/me/report/map.png")
async def gov_report_map(request: Request, card: str, emd: str | None = None):
    p = _me(request)
    png = await _png(p, card, (emd or "").strip() or None)
    if not png:
        raise ApiError("not_found", "지도 그림을 만들지 못했습니다")
    return Response(png, media_type="image/png", headers={"cache-control": "private, max-age=300"})


@router.get("/spaces/me/report.hwpx")
async def gov_report_hwpx(request: Request, card: str, emd: str | None = None, ok: str | None = None):
    from .farm_admin import need_personal
    from .xreport import hwpx_response
    p = _me(request)
    emd = (emd or "").strip() or None
    d = await report_data(p, card, emd)
    pl = await need_personal(p, ok) if d.get("ledger") else None     # '우리 대장과 다른 필지'(지번 · 대장 값)가 든 문서 = 개인정보가 든 파일(10-18 영농-q3)
    png = await _png(p, card, emd, d)
    blob, name = await run_in_threadpool(_hwpx, d, png)
    async with db(p) as c:
        await c.execute("INSERT INTO space_log(tenant_id, kind, card_id, line, actor, realm, detail) VALUES ($1,'download',$2,$3,$4,'tenant',$5)",
                        p.tenant_id, card, f"{p.name or '담당자'} — {d['service']['name']} AI 분석 결과 보고서 내려받기({d['region']})", p.user_id,
                        {"fmt": "report", "emd": emd, "rows": d["total"]["value"], **({"pledge": pl["id"], "personal": True} if pl else {})})
        await audit(c, p, "space.report", f"{p.tenant_id}/{card}", None, {"emd": emd})
    return hwpx_response(blob, name)


@router.get("/spaces/me/ledger/compare")
async def gov_compare(request: Request, card: str, import_id: str | None = None, geom: int = 0):
    p = _me(request)
    c = await compare(p, card, import_id, None, bool(geom))
    c["services"] = [s for s in await _services(p) if s["parcel"]]
    return _public(c)


@router.get("/spaces/me/ledger/compare.xlsx")
async def gov_compare_xlsx(request: Request, card: str, import_id: str | None = None, ok: str | None = None):
    from .farm_admin import need_personal
    p = _me(request)
    pl = await need_personal(p, ok)                  # 대장 필지마다 지번 · 대장 값이 든 파일 = 개인정보가 든 파일(10-18 영농-q3 — 이용 약속 · 받을 때마다 확인)
    c = await compare(p, card, import_id)
    if not c.get("ledger"):
        raise ApiError("not_found", "올린 대장이 없습니다")
    why = {x["key"]: x["label"] for x in c["lines"]}
    rows = sorted(c["_all"], key=lambda x: ({"diff": 0, "ai": 1, "none": 2}[x["k"]], -(x["ai_m2"] or 0), x["where"]))

    from .farm_more import AREA, _book, _org_short, _save, _sheet, cap_line
    org = await _org_short(p.tenant_id)

    def build() -> bytes:
        wb = _book()
        # 종류별 AI 결과 면적 칸(10-18 원칙 286 — 운영 중인 영농관리 화면의 'AI 분석 면적' 종류별 칸) · 넓은 종류부터 · 대장 행마다 같은 필지의 AI 결과
        tot: dict[str, float] = {}
        for x in rows:
            for k, m in (x.get("_by") or {}).items():
                tot[k] = tot.get(k, 0.0) + m
        kinds = [k for k, _ in sorted(tot.items(), key=lambda kv: -kv[1])]
        head = ["필지 번호", "지번", "읍면동", "대장", "AI 분석 결과", "AI 결과 넓이(㎡)", "필지 넓이(㎡)", "대조 결과", "AI 결과 유무"] + [f"AI {k} 넓이(㎡)" for k in kinds]
        data = [[x["pnu"], x["where"], x["emd"], x["ledger"], x["ai"] or "", x["ai_m2"], x["parcel_m2"],
                 why.get(x["why"], "") if x["k"] == "diff" else ("AI 결과 있음" if x["k"] == "ai" else c["no_ai_label"]),
                 "유" if x["k"] != "none" else "무"] + [(x.get("_by") or {}).get(k, 0) if x["k"] != "none" else None for k in kinds] for x in rows]
        L = c["ledger"]
        cap = cap_line(org, c["service"]["name"], f"대장 대조 · {L['name']}", str(L.get("base_date") or now_iso()[:10])[:10],
                       "올린 대장 필지 × AI 분석 결과(필지 번호로 맞춤)")
        num = {5: AREA, 6: AREA, **{9 + i: AREA for i in range(len(kinds))}}
        _sheet(wb, "대장 대조", head, data, num=num, text_cols=(0,), widths=[22, 20, 12, 10, 26, 16, 14, 34, 12] + [18] * len(kinds), cap=cap)
        i = wb.create_sheet("안내")
        for k, v in [("서비스", c["service"]["name"]), ("올린 대장", L["name"]), ("대장 필지", L["parcels"]["value"]),
                     ("대장과 AI 결과가 다른 필지", c["diff"]["value"]), ("AI 결과가 있는 필지", c["with_ai"]["value"]),
                     (c["no_ai_label"], c["no_ai"]["value"]), ("개인정보 이용 확인", f"{pl['purpose']} · {pl['until']}까지"), ("만든 때", now_iso()),
                     ("안내", "AI 분석 결과는 참고자료입니다 — 처분은 담당자가 정합니다.")]:
            i.append([k, v])
        i.column_dimensions["A"].width = 24
        i.column_dimensions["B"].width = 60
        return _save(wb)
    blob = await run_in_threadpool(build)
    name = re.sub(r"[\\/:*?\"<>|\s—·]+", "_", f"{org}_{_short_svc(c['service']['name'])}_대장대조_{dt.datetime.now(KST).strftime('%Y%m%d')}").strip("_") + ".xlsx"
    async with db(p) as cn:
        await cn.execute("INSERT INTO space_log(tenant_id, kind, card_id, line, actor, realm, detail) VALUES ($1,'download',$2,$3,$4,'tenant',$5)",
                         p.tenant_id, card, f"{p.name or '담당자'} — {c['service']['name']} 대장과 AI 결과 내려받기", p.user_id,
                         {"fmt": "ledger_compare", "rows": len(rows), "import_id": c["ledger"]["import_id"], "pledge": pl["id"], "personal": True})
    return Response(blob, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"content-disposition": f"attachment; filename=\"compare.xlsx\"; filename*=UTF-8''{quote(name)}", "cache-control": "no-store"})


@router.post("/spaces/me/ledger/compare/report")
async def gov_compare_to_report(request: Request, body: dict | None = None):
    """'보고서에 넣기' — 이 대장을 이 사람의 보고서 대장으로 기억(ledger_recent '*') · 보고서 '우리 대장과 다른 필지' 칸이 이 대장으로."""
    from .ledger import remember
    p = _me(request)
    b = body or {}
    card = _card(str(b.get("card") or ""))
    await _svc(p, card)                                          # 받은 서비스인가(없으면 404)
    imp = await _import_of(p, str(b.get("import_id") or "") or None)
    if not imp:
        raise ApiError("not_found", "올린 대장이 없습니다")
    async with db(realm="lx") as c:
        await remember(c, p.tenant_id, p.user_id, imp["id"], imp["kind"], [])
    return {"ok": True, "import_id": imp["id"], "card": card}


# ── 공개 실적(메인 QA-Q2 ⓐ · QA-Q3 ⓐ) ─────────────────────────────────────
_REC: dict = {"t": 0.0, "v": None}


async def result_sggs() -> list[str]:
    """AI 분석 결과가 있는 시군구 — 시군구 전역 분석(scope sgg)이 끝난 작업(시험 · 시연용 제외 · 결과 장부 준비됨).
    /regions/{시군구}/results ②와 같은 거르기(그 시군구 결과 레이어가 열리는 곳)."""
    async with db(realm="lx") as c:
        rows = await c.fetch("SELECT DISTINCT options->>'sgg_cd' AS cd, card_id FROM jobs WHERE kind='infer' AND state='done' AND NOT demo "
                             "AND NOT coalesce(test,false) AND snapshot_ready AND options->>'scope'='sgg' AND options->>'sgg_cd' IS NOT NULL")
    from .svc_rules import excluded, hit
    ex = await excluded()                                        # 서비스 × 시군구 보이기 규칙(원칙 293) — 시군구 결과 층과 같은 거르기
    return sorted({r["cd"] for r in rows if r["cd"] and not hit(ex, r["card_id"], r["cd"])})


async def record() -> dict:
    if _REC["v"] and time.time() - _REC["t"] < 60:
        return _REC["v"]
    from .cards import _name as card_name
    from .regions import region_of
    from .spaces import _short_org
    from .svc_rules import SQL_NOT_EXCLUDED
    sgg = await result_sggs()
    places = []
    for cd in sgg:
        r = region_of(cd)
        if r and r.get("center"):
            places.append({"sgg_cd": r["sgg_cd"], "name": r.get("name"), "center": r["center"]})
    async with db(realm="lx") as c:
        sv = await c.fetch("SELECT s.card_id, c.name, count(*) n FROM service_results s JOIN cards c ON c.id = s.card_id "
                           "WHERE s.replaced_at IS NULL AND s.origin IN ('deploy', 'job')" + SQL_NOT_EXCLUDED + " GROUP BY 1, 2 ORDER BY 3 DESC, 1")   # XI맵 전국 결과(ximap.build_board)와 같은 줄 — '서비스 n가지' 한 출처(GPT5-1)
        cards = await c.fetch("SELECT c.id, c.name, c.scope, (SELECT count(*) FROM service_results s WHERE s.card_id=c.id AND s.replaced_at IS NULL" + SQL_NOT_EXCLUDED + ") n, "
                              "EXISTS(SELECT 1 FROM deploys d WHERE d.card_id=c.id AND NOT coalesce(d.test,false) AND d.stage <> 'draft') dp FROM cards c")
        og = await c.fetch("SELECT DISTINCT ON (g.tenant_id) g.tenant_id, t.name, d.sgg_cd FROM space_guides g JOIN tenants t ON t.id=g.tenant_id "
                           "LEFT JOIN deploys d ON d.tenant_id=g.tenant_id AND d.card_id=g.card_id AND NOT coalesce(d.test,false) "
                           "WHERE g.published_at IS NOT NULL AND (g.body->'what'->'total'->>'value') IS NOT NULL "
                           "AND (g.body->'what'->'total'->>'value')::numeric > 0 ORDER BY g.tenant_id, d.sgg_cd NULLS LAST")
    services = [_short_svc(card_name(r["name"])) for r in sv]
    # 행정 업무 서비스 = 국내 서비스 가운데 결과를 냈거나 기관에 배포된 것(결과 · 배포가 없는 서비스는 넣지 않는다)
    work = [r for r in cards if (r["scope"] or "local") != "global" and (r["n"] or r["dp"])]
    work.sort(key=lambda r: (-int(r["n"] or 0), 0 if r["dp"] else 1, card_name(r["name"])))
    orgs = []
    for r in og:
        nm = r["name"]
        nm = (nm.get("ko") if isinstance(nm, dict) else json.loads(nm).get("ko") if isinstance(nm, str) and nm.startswith("{") else nm) or r["tenant_id"]
        rg = region_of(r["sgg_cd"]) if r["sgg_cd"] else None
        orgs.append({"name": _short_org(nm), "sgg_cd": (rg or {}).get("sgg_cd"), "center": (rg or {}).get("center")})
    src = "AI 분석 결과 장부(시군구 전역 분석 · 서비스 결과 · 공개된 결과 설명서)"
    v = {"sgg": env(len(sgg), "곳", "recorded", src), "places": places,
         "services": {"n": env(len(services), "가지", "recorded", src), "names": services},
         "orgs": {"n": env(len(orgs), "곳", "recorded", src), "items": orgs},
         "work": {"n": env(len(work), "가지", "recorded", "서비스 목록(국내 · 결과 또는 배포가 있는 서비스)"), "names": [_short_svc(card_name(r["name"])) for r in work]},
         "as_of": now_iso()}
    _REC.update(t=time.time(), v=v)
    return v


@router.get("/public/record")
async def public_record(request: Request):
    principal(request)
    return await record()


_NATP: dict = {"t": 0.0, "v": None}


@router.get("/public/national")
async def public_national(request: Request):
    """게스트 메인 ① 전국 결과 지도의 숫자(10-22 원칙 317 '무엇 · 얼마 · 왜') — 로그인 없이 보이는 전국 합계만.
    값 = LX 직원 XI맵과 같은 장부(national_cells · 화면 기본 해 · 끝난 칸)와 같은 셈(같은 업무의 서비스 결과가 덮은 읍면동을 뺀 수 net · 없으면 찾은 수).
    종류마다 수 · 넓이(㎢) 합계와 기준 영상 · 끝난 비율만 낸다 — 칸 · 읍면동 · 좌표 · 모델 · 작업 정보 없음. 1분 보관."""
    principal(request)
    if _NATP["v"] and time.time() - _NATP["t"] < 60:
        return _NATP["v"]
    from . import national as N
    async with db(realm="lx") as c:
        st, _ = await N._settings(c)
        year = int(st["year"])
        sums = await c.fetch(
            "SELECT asset, k AS cls, sum((v->>'n')::bigint) AS n, sum((v->>'km2')::numeric) AS km2 FROM national_cells, jsonb_each(counts) AS e(k, v) "
            "WHERE year=$1 AND NOT test AND state='done' GROUP BY asset, k", year)
        prog = await c.fetch("SELECT asset, coalesce(sum(target_km2) FILTER (WHERE state<>'skip'),0) AS km2, "
                             "coalesce(sum(target_km2) FILTER (WHERE state='done'),0) AS done_km2 FROM national_cells WHERE year=$1 AND NOT test GROUP BY asset", year)
    src = f"Land-XI 전국 AI 분석(국토지리정보원 항공 정사영상 {year})"
    out = {"year": year, "imagery": f"국토지리정보원 항공 정사영상 {year}", "assets": {}, "as_of": now_iso()}
    for a in N.ASSETS:
        cnt = {r["cls"]: {"n": env(int(r["n"] or 0), "count", "measured", src), "km2": env(round(float(r["km2"] or 0), 2), "km2", "measured", src)}
               for r in sums if r["asset"] == a and r["cls"] in _NAT_CLS[a]}
        await N._net(cnt, a, None)
        pr = next((x for x in prog if x["asset"] == a), None)
        pct = round(100.0 * float(pr["done_km2"]) / float(pr["km2"]), 1) if pr and float(pr["km2"]) else 0.0
        out["assets"][a] = {"pct": env(min(100.0, pct), "%", "measured", "끝난 분석 범위 ÷ 전체 분석 범위"), "counts": cnt}
    _NATP.update(t=time.time(), v=out)
    return out


_NAT_CLS = {"farm": ("경작지", "비닐하우스"), "river": ("건물",)}   # XI맵과 같은 규칙 — 자산마다 한 분류씩(겹쳐 세지 않게 · xi-clean/national.js NCLS)


_SAFE = re.compile(r"[\\/:*?\"<>|\s]+")


@router.get("/spaces/me/report/bundle.zip")
async def gov_report_bundle(request: Request, cards: str | None = None, emd: str | None = None, import_id: str | None = None):
    """보고에 붙일 첨부 한 번에(영농-7) — 한글 문서(.hwpx) + 엑셀(.xlsx) + 지도 그림(.png) 한 묶음.
    같은 고르기의 multi.hwpx · multi.xlsx 와 같은 빌더(같은 숫자 · 같은 관할) · 발급 목록에는 한 줄(묶음 1회 = 1회)."""
    import zipfile
    from . import farm_more as FM
    p = _me(request)
    d = await FM.multi_data(p, cards, emd, import_id, with_avail=False)
    pngs: list[tuple[str, bytes | None]] = []
    for s in d["picked"][:4]:
        try:
            pngs.append((s["short"], await _png(p, s["card"], d["emd"])))
        except ApiError:
            pngs.append((s["short"], None))
    hw, hname = await run_in_threadpool(FM._hwpx_multi, d, pngs)
    xl = await run_in_threadpool(FM._xlsx_multi, d)
    base = _SAFE.sub("_", f"{d['region']}_{'_'.join(s['short'] for s in d['picked'])}_AI분석결과_{FM._today()}")

    def _zip() -> bytes:
        bio = io.BytesIO()
        with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr(f"{base}_보고서.hwpx", hw)
            z.writestr(f"{base}_표.xlsx", xl)
            ok = [(n, b) for n, b in pngs if b]
            for i, (n, b) in enumerate(ok):
                suf = "" if len(ok) == 1 else "_" + _SAFE.sub("_", n)
                z.writestr(f"{base}_지도{suf}.png", b)
        return bio.getvalue()

    blob = await run_in_threadpool(_zip)
    await FM._issue(p, d, (import_id or "").strip() or None, "zip")
    name = f"{base}_첨부묶음.zip"
    return Response(blob, media_type="application/zip",
                    headers={"content-disposition": f"attachment; filename=\"bundle.zip\"; filename*=UTF-8''{quote(name)}", "cache-control": "no-store",
                             "access-control-expose-headers": "Content-Disposition"})


# ── 대장 올리기 양식(영농-8) ────────────────────────────────────────────────
# 머리 줄만(값 0줄) — 열 이름은 올리기 화면의 열 자동 인식(gov-fusion/ledger.js guessColumns)이 그대로 맞추는 말.
# 성명 · 연락처 · 경영체 이름 열은 양식에 두지 않는다(읽지 않음).
_TPL_BASE = [("소재지", "지번 주소 — 읍면동 · 리 · 지번까지(예: ○○읍 ○○리 123-4)", True),
             ("필지고유번호", "19자리 필지고유번호(PNU) — 있으면 지번보다 먼저 맞춥니다", False)]
_TPL = {
    "farm_ledger": ("농지대장", _TPL_BASE + [("지목", "전 · 답 · 과수원 등 대장의 지목", False), ("기준일자", "대장 기준일(예: 2025-06-30)", False)]),
    "feedcrop": ("사료작물 재배 신청 대장", _TPL_BASE + [("신청 작물", "신청한 작물 이름(예: 옥수수 · 수단그라스)", True), ("신청 면적(㎡)", "신청 면적 숫자만", False),
                                                       ("신청일자", "신청한 날(예: 2025-04-15)", False)]),
    "direct_pay": ("공익직불 신청 대장", _TPL_BASE + [("신청 면적(㎡)", "직불금을 신청한 면적 숫자만", False), ("품목", "신청 품목(예: 벼 · 밭작물)", False),
                                                  ("신청일자", "신청한 날(예: 2025-05-31)", False)]),
    "greenhouse": ("시설 신고 대장", _TPL_BASE + [("시설 구분", "시설 종류(예: 비닐하우스 · 유리온실)", False), ("신고일자", "신고한 날", False)]),
    "dev_permit": ("개발행위허가 대장", _TPL_BASE + [("허가 종류", "허가 종류", False), ("허가일자", "허가한 날", False)]),
    "public_asset": ("공유재산 대장", _TPL_BASE + [("용도", "재산 용도", False), ("기준일자", "대장 기준일", False)]),
    "river_permit": ("점용허가 대장", _TPL_BASE + [("허가 종류", "점용 허가 종류", False), ("허가일자", "허가한 날", False)]),
}


def _template_xlsx(kind: str) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    label, cols = _TPL[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = "대장"
    for i, (name, _d, _r) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF2F7")
        ws.column_dimensions[c.column_letter].width = max(16, len(name) * 3)
    ws.freeze_panes = "A2"
    g = wb.create_sheet("설명")
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 64
    g.column_dimensions["C"].width = 10
    rows = [(f"{label} 올리기 양식", "", ""), ("", "", ""), ("열 이름", "채우는 내용", "필요")]
    rows += [(n, d, "권장" if r else "선택") for n, d, r in cols]
    rows += [("", "", ""),
             ("안내", "첫 시트 '대장'의 첫 줄을 머리로 읽습니다. 값은 둘째 줄부터 채워 주세요.", ""),
             ("", "필지고유번호와 소재지 중 하나만 있어도 필지를 맞춥니다.", ""),
             ("", "성명 · 연락처 · 경영체 이름 열은 올려도 읽지 않습니다. 양식에도 없습니다.", ""),
             ("", "농림사업정보시스템 등에서 내려받은 엑셀을 그대로 올려도 열을 자동으로 맞춥니다. 이 양식은 선택입니다.", "")]
    for r in rows:
        g.append(list(r))
    g["A1"].font = Font(bold=True, size=14)
    for c in g[3]:
        c.font = Font(bold=True)
    for row in g.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"{label}_올리기양식.xlsx"


@router.get("/spaces/me/ledger/template/{kind}.xlsx")
async def ledger_template(request: Request, kind: str):
    _me(request)
    if kind not in _TPL:
        raise ApiError("not_found", "그 대장 종류의 양식이 없습니다")
    blob, name = await run_in_threadpool(_template_xlsx, kind)
    return Response(blob, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"content-disposition": f"attachment; filename=\"template.xlsx\"; filename*=UTF-8''{quote(name)}", "cache-control": "no-store",
                             "access-control-expose-headers": "Content-Disposition"})
