"""영농관리 바퀴 19(10-18 15:52 now 답 · 원칙 293) — 기관 세션 · 관할 · 받은 서비스 범위(공간과 같은 판정).

  [영농-보고서 '시안 A 고르면 바로 만들어지는 보고서']
  GET  /api/v1/spaces/me/report/multi?cards=a,b&emd=&import_id=      고른 결과(서비스 여럿)로 바로 채운 보고서 — 개요 · 시점 비교(둘 이상) · 종류별 ·
                                                                    읍면동별 · (대장을 넣으면) 대장과 견준 결과. 적어 넣는 칸 없음.
  GET  /api/v1/spaces/me/report/multi.hwpx?…                         같은 내용 한글 문서 → 발급 목록에 한 줄(같은 사람 · 같은 고르기면 횟수만 +1)
  GET  /api/v1/spaces/me/report/multi.xlsx?…                         같은 내용 엑셀(시트마다 표 · 숫자는 숫자)
  GET  /api/v1/spaces/me/report/issues                               발급 목록(발급 일시 · 넣은 결과 · 요청자 · 내려받은 횟수 · 다시 받기 고르기)
  보고서의 대장 칸은 건수만(필지별 지번 · 대장 값 없음) — 개인정보가 든 파일이 아니라 바로 받는다(PERSONAL_RULE 그대로).

  [CSV-이름 '진짜 엑셀(.xlsx)로 바꾸기']
  GET  /api/v1/spaces/me/guides/{card}/class-stats.xlsx              통계 종류별 · 읍면동별(화면 표와 같은 서버 값 — space_class_stats)
  GET  /api/v1/history.xlsx?card=&kind=                              이력(화면과 같은 줄 — gov_history.history · 확인 기록 줄 뺌)

  [제안-이용약속기록화면 넣음]
  GET  /api/v1/spaces/me/pledges?from=&to=   (.xlsx)                기관 관리자 — 자기 기관 이용 약속 기록(누가 · 언제 · 목적 · 기간 · 받은 파일 종류 · 횟수)
  GET  /api/v1/admin/pledges?from=&to=&tenant=   (.xlsx)            LX 관리자 — 모든 기관(보기만)
숫자 한 출처: 서비스별 합계 = gov_report._rows(보고서 · 대시보드 큰 숫자 · 필지 목록과 같은 행). 어느 지역이든 같은 길(지역 이름 · 코드를 코드에 두지 않는다).
"""
from __future__ import annotations

import datetime as dt
import io
import json
import re
import secrets
from urllib.parse import quote

from fastapi import APIRouter, Request
from fastapi.responses import Response
from starlette.concurrency import run_in_threadpool

from .deps import ApiError, Principal, db, principal, require
from .envelope import KST, env, now_iso, split_round

router = APIRouter()
CARD_RE = re.compile(r"^card-[a-z0-9-]{2,40}$")
IMP_RE = re.compile(r"^imp_[0-9A-Z]{10,40}$")
SRC = "결과 설명서의 결과 세트(관할 안 · 지운 결과 제외)"
SRC_LED = "올린 대장 필지 × 고른 서비스 결과(필지 번호로 맞춤)"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# 내려받은 파일 종류(이력 · 이용 약속 기록 · 한 이름표)
FMT_KO = {"geojson": "지도 파일", "parcels": "필지 엑셀", "summary": "요약", "report": "보고서 한글 문서", "report_multi": "보고서 한글 문서",
          "report_xlsx": "보고서 엑셀", "stats": "통계 엑셀", "history": "이력 엑셀", "feedcrop": "사료작물 대장 대조 엑셀",
          "ledger_compare": "대장과 AI 결과 엑셀", "pledges": "다운로드 이력 엑셀"}


def _me(request: Request) -> Principal:
    from .spaces import _me as me
    return me(request)


# ── 공용: 엑셀 ─────────────────────────────────────────────────────────────
def _book():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    return wb


AREA, COUNT, RATIO, DATE, DATETIME = "#,##0.0", "#,##0", '0.0"%"', "yyyy-mm-dd", "yyyy-mm-dd hh:mm"
SUMMABLE = {"#,##0", "#,##0.0", "#,##0.00"}     # 합계(SUBTOTAL 109)를 내는 숫자 서식 — 넓이 · 개수 · 금액. 비율 · 날짜 · 점수는 더하지 않는다


def cap_line(org: str = "", svc: str = "", scope: str = "", base: str = "", src: str = "") -> tuple[str, str]:
    """엑셀 첫 시트 1 ~ 2행 — 제목(기관 · 서비스 · 범위) · 기준일 · 출처(10-22 엑셀 한 벌 · FARM-ADMIN §6-1)."""
    t = " · ".join(x for x in (org, svc, scope) if x)
    b = " · ".join(x for x in ((f"기준일 {base}" if base else ""), (f"출처 {src}" if src else "")) if x)
    return t, b


def _sheet(wb, title: str, head: list[str], rows: list[list], *, num: dict[int, str] | None = None, widths: list[int] | None = None,
           total_last: bool = False, text_cols: tuple[int, ...] = (), cap: tuple[str, str] | None = None, total: bool | None = None):
    """엑셀 한 벌(10-22 now 21 '다시' · FARM-ADMIN §6-1) — 모든 기관 엑셀(이력 · 필지 · 통계 · 대장 대조 · 보고서 · 다운로드 이력)이 이 함수로.
    cap = (제목 줄, 기준일 · 출처 줄) → 1 ~ 2행 · 3행 빈 줄 · 4행 머리(굵게 · 회백) · A5 고정. cap 이 없으면 1행 머리 · A2 고정.
    num = {열 번호: 서식}(넓이 AREA · 개수 COUNT · 비율 RATIO · 날짜 DATE) · 자동 필터(머리 ~ 마지막 자료 줄).
    합계: total_last=True 면 rows 의 마지막 줄이 합계 줄 — 더할 수 있는 열(SUMMABLE)이 자료 줄 합과 같으면 SUBTOTAL(109) 수식으로(거르면 보이는 줄만 더함),
          다르면(예: 읍면동 밖 결과가 있어 서버 합계가 더 큼) 서버 값 그대로 둔다(숫자 한 출처). total=True(또는 None 이고 더할 열이 있음)면 '합계' 줄을 붙인다."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(title[:31])
    h0 = 1
    if cap:
        ws.append([cap[0]])
        ws.append([cap[1]] if cap[1] else [])
        ws.append([])
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"].font = Font(color="4E535C")
        h0 = 4
    ws.append(head)
    for c in ws[h0]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F4F6")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    data = rows[:-1] if total_last and rows else list(rows)
    tot = rows[-1] if total_last and rows else None
    for r in data:
        ws.append(r)
    r0, r1 = h0 + 1, h0 + len(data)                 # 자료 줄 범위
    num = dict(num or {})
    for i, fmt in num.items():
        if i >= len(head) or not data:
            continue
        for (cell,) in ws.iter_rows(min_row=r0, max_row=r1, min_col=i + 1, max_col=i + 1):
            if isinstance(cell.value, (int, float, dt.date, dt.datetime)):
                cell.number_format = fmt
    for i in text_cols:
        for (cell,) in ws.iter_rows(min_row=r0, max_row=r1, min_col=i + 1, max_col=i + 1):
            cell.number_format = "@"
    sums = [i for i, f in num.items() if f in SUMMABLE and i < len(head) and i > 0]
    if total is None:
        total = bool(sums) and not total_last and len(data) > 1
    if data and (tot is not None or (total and sums)):
        row = list(tot) if tot is not None else (["합계"] + [None] * (len(head) - 1))
        row += [None] * (len(head) - len(row))
        for i in sums:
            col = get_column_letter(i + 1)
            vals = [r[i] for r in data if i < len(r) and isinstance(r[i], (int, float))]
            server = row[i]
            same = server is None or not isinstance(server, (int, float)) or abs(sum(vals) - float(server)) <= max(0.51, abs(float(server)) * 1e-6)
            if vals and same:
                row[i] = f"=SUBTOTAL(109,{col}{r0}:{col}{r1})"
        ws.append(row)
        for j, c in enumerate(ws[ws.max_row]):
            c.font = Font(bold=True)
            if j in num and num[j] in SUMMABLE | {RATIO}:
                c.number_format = num[j]
    if data:
        ws.auto_filter.ref = f"A{h0}:{get_column_letter(len(head))}{r1}"
    ws.freeze_panes = f"A{h0 + 1}"
    for i, h in enumerate(head):
        w = (widths or [])[i] if widths and i < len(widths) else max(10, min(30, len(str(h)) * 2 + 2))
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    if cap:
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 10, 14)
    return ws


def _info(wb, pairs: list[tuple[str, object]]):
    ws = wb.create_sheet("안내")
    for k, v in pairs:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70
    return ws


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_resp(blob: bytes, name: str) -> Response:
    name = re.sub(r"[\\/:*?\"<>|\s—]+", "_", name.replace("·", "")).strip("_") + ".xlsx"
    return Response(blob, media_type=XLSX, headers={"content-disposition": f"attachment; filename=\"landxi.xlsx\"; filename*=UTF-8''{quote(name)}",
                                                    "cache-control": "no-store", "access-control-expose-headers": "Content-Disposition"})


def _today() -> str:
    return dt.datetime.now(KST).strftime("%Y%m%d")


async def _org_short(tenant: str) -> str:
    from .spaces import _short_org
    async with db(realm="lx") as c:
        n = await c.fetchval("SELECT name FROM tenants WHERE id=$1", tenant)
    n = json.loads(n) if isinstance(n, str) and n.startswith("{") else n
    return _short_org((n or {}).get("ko") if isinstance(n, dict) else (n or ""))


async def _log(p: Principal, card: str | None, line: str, detail: dict):
    async with db(p) as c:
        await c.execute("INSERT INTO space_log(tenant_id, kind, card_id, line, actor, realm, detail) VALUES ($1,'download',$2,$3,$4,'tenant',$5)",
                        p.tenant_id, card, line, p.user_id, detail)


# ── 보고서(고르면 바로 만들어지는) ───────────────────────────────────────────
async def _avail(p: Principal) -> list[dict]:
    """고를 수 있는 결과 — 기관에 열린 서비스 가운데 결과가 있는 것(gov_report._services 와 같은 판정) + 영상 · 분석한 날."""
    from .spaces import _current, _visible, received, sync
    await sync(p.tenant_id)
    out = []
    async with db(p) as c:
        for s in await _visible(p, await received(p.tenant_id)):
            g = await _current(c, p.tenant_id, s["card"])
            if not g:
                continue
            b = g["body"] or {}
            tot = (b.get("what") or {}).get("total") or {}
            if not b.get("sets") or tot.get("value") in (None, 0):
                continue
            rnd = (b.get("when") or {}).get("rounds") or []
            out.append({"card": s["card"], "name": s["name"], "shot": next((r.get("shot") for r in rnd if r.get("shot")), None),
                        "analyzed": next((str(r.get("analyzed"))[:10] for r in rnd if r.get("analyzed")), None),
                        "unit": tot.get("unit") or "건", "parcel": bool((b.get("where") or {}).get("parcel"))})
    return out


async def _one(p: Principal, card: str, emd: str | None) -> dict:
    """서비스 하나 — 합계 · 종류별 · 읍면동별(gov_report.report_data 와 같은 행 · 같은 반올림)."""
    from .gov_report import _acc, _rows, _sgg_of, _short_svc, _svc, _ymd
    from .spaces import class_ko
    svc, g, rows, body, sets, codes, _only = await _svc(p, card)
    all_rows, _box = await _rows(p, sets, codes, None)
    emds = sorted({r["emd"] for r in all_rows if r["emd"]})
    rs = [r for r in all_rows if r["emd"] == emd] if emd else all_rows
    unit = ((body.get("what") or {}).get("total") or {}).get("unit") or "건"
    acc, acc_at = await _acc(rows)
    rnd = (body.get("when") or {}).get("rounds") or []
    sgg, place = _sgg_of(body, rows)
    tot_n = sum(r["n"] for r in rs)
    tot_ha = round(sum(r["m2"] for r in rs) / 1e4, 1)
    cls: dict[str, list] = {}
    for r in rs:
        e = cls.setdefault(class_ko(r["cls"]), [0, 0.0])
        e[0] += r["n"]
        e[1] += r["m2"]
    cls_l = sorted(cls.items(), key=lambda kv: -kv[1][1])
    cls_ha = split_round([v[1] / 1e4 for _, v in cls_l], 1, tot_ha)
    from .farm_admin import bale_kg, is_bale, yield_env
    kg = await bale_kg(p.tenant_id) if any(is_bale(k) for k, _ in cls_l) else None
    classes = [{"name": k, "n": env(v[0], unit, "inferred", SRC), "ha": env(cls_ha[i], "ha", "measured", SRC),
                **({"yield": yield_env(v[0], kg)} if kg and is_bale(k) else {}), "bale": is_bale(k)} for i, (k, v) in enumerate(cls_l)]
    by: dict[str, dict] = {}
    for r in rs:
        e = by.setdefault(r["emd"] or "읍면동 밖", {"emd": r["emd"] or "읍면동 밖", "n": 0, "m2": 0.0, "by": {}})
        e["n"] += r["n"]
        e["m2"] += r["m2"]
        k = class_ko(r["cls"])
        e["by"][k] = e["by"].get(k, 0) + r["n"]
    emd_l = sorted(by.values(), key=lambda x: (x["emd"] == "읍면동 밖", -x["n"], x["emd"]))
    emd_ha = split_round([e["m2"] / 1e4 for e in emd_l], 1, tot_ha)
    return {"card": card, "name": svc["name"], "short": _short_svc(svc["name"]), "unit": unit, "sgg": sgg, "place": place, "emds": emds,
            "shot": next((r.get("shot") for r in rnd if r.get("shot")), None), "analyzed": _ymd(next((r.get("analyzed") for r in rnd if r.get("analyzed")), None)),
            "total": env(tot_n, unit, "inferred", SRC), "area_ha": env(tot_ha, "ha", "measured", SRC), "classes": classes,
            "emd_rows": [{"emd": e["emd"], "n": e["n"], "ha": env(emd_ha[i], "ha", "measured", SRC), "by": [{"cls": k, "n": v} for k, v in e["by"].items()]}
                         for i, e in enumerate(emd_l)],
            "acc": env(acc, "%", "measured", "모델 기록 · 학습에 쓰지 않은 영상으로 잰 값", as_of=acc_at or None),
            "_sets": sets, "_codes": codes}


async def _ledger_cmp(p: Principal, import_id: str, emd: str | None, picks: list[dict]) -> dict | None:
    """대장 × 고른 결과 — 필지마다 고른 서비스마다 AI 결과 있음 · 없음 → 모두 있음 · 하나 이상 없음(건수만 · 필지 목록 없음).
    '찾지 못함'과 같은 뜻의 종류(비경작지 · 휴경)는 재배 흔적으로 세지 않는다(사료작물 견주기와 같은 판정)."""
    from .farm_admin import NEG
    from .ledger import KINDS
    from .spaces import class_ko
    if not IMP_RE.match(import_id or ""):
        raise ApiError("not_found", "올린 대장이 없습니다")
    async with db(realm="lx") as c:
        await c.execute("SET LOCAL jit = off")
        imp = await c.fetchrow("SELECT id, kind, name, filename, base_date, created_at, confirmed_at FROM ledger_imports WHERE id=$1 AND tenant_id=$2 AND state='matched'",
                               import_id, p.tenant_id)
        if not imp:
            raise ApiError("not_found", "올린 대장이 없습니다")
        feed = imp["kind"] == "feedcrop"     # 사료작물 재배 신청 대장 = 사료작물 견주기와 같은 판정(비경작지 · 휴경은 재배 흔적이 아님) · 그 밖 = 행정정보와 비교와 같은 판정
        pn = [r["pnu"] for r in await c.fetch("SELECT DISTINCT pnu FROM registry_snapshots WHERE import_id=$1 AND tenant_id=$2 AND pnu IS NOT NULL", imp["id"], p.tenant_id)]
        if emd and pn:
            pn = [r["pnu"] for r in await c.fetch("SELECT pnu FROM survey_parcels WHERE pnu = ANY($1::text[]) AND emd = $2", pn, emd)]
        found: dict[str, set] = {}
        for s in picks:
            if not pn or not s["_sets"]:
                found[s["card"]] = set()
                continue
            codes = s["_codes"]
            sc = " AND left(d.pnu, 5) = ANY($3::text[])" if codes is not None else ""
            rows = await c.fetch("SELECT DISTINCT d.pnu, d.cls, d.cls_en FROM detections d WHERE d.job_id = ANY($1::text[]) AND d.edit_state <> 'deleted' "
                                 f"AND d.pnu = ANY($2::text[]){sc}", *([s["_sets"], pn, codes] if codes is not None else [s["_sets"], pn]))
            found[s["card"]] = {r["pnu"] for r in rows if not feed or class_ko(r["cls"], r["cls_en"]) not in NEG}
    n = len(pn)
    allf = sum(1 for k in pn if all(k in found[s["card"]] for s in picks))
    nm = re.sub(r"\.(xlsx|xls)\.csv$", r".\1", imp["name"] or imp["filename"] or "") or (KINDS.get(imp["kind"]) or "대장")
    out = {"import_id": imp["id"], "name": nm, "kind": KINDS.get(imp["kind"]) or "대장",
           "base_date": imp["base_date"].isoformat() if imp["base_date"] else None,
           "parcels": env(n, "필지", "measured", "올린 대장 × 연속지적(필지 번호)"),
           "feed": feed, "all": env(allf, "필지", "inferred", SRC_LED), "some": env(n - allf, "필지", "inferred", SRC_LED),
           "labels": ({"all": "맞음 · 고른 시점 모두 AI 결과 있음", "some": "다름 · AI 결과가 없는 시점이 있음", "found": "AI 결과 있음", "missing": "AI 결과 없음"} if feed else
                      {"all": "고른 결과 모두 있음", "some": "하나 이상 없음", "found": "AI 결과가 있는 필지", "missing": "AI 결과가 없는 필지"}),
           "per": [{"card": s["card"], "short": s["short"], "found": env(len(found[s["card"]]), "필지", "inferred", SRC_LED),
                    "missing": env(n - len(found[s["card"]]), "필지", "inferred", SRC_LED)} for s in picks],
           "lines": []}
    if len(picks) == 1:                       # 서비스 하나 — '대장과 다른 필지'(지목 × AI 분류) 건수도(행정정보와 비교와 같은 판정)
        from .gov_report import compare
        try:
            if not feed:
                cm = await compare(p, picks[0]["card"], imp["id"], emd)
                out["lines"] = [{"label": x["label"], "short": x.get("short") or x["label"], "n": x["n"]} for x in cm.get("lines") or []]
                out["labels"]["missing"] = cm.get("no_ai_label") or out["labels"]["missing"]
        except ApiError:
            pass
    return out


# ── 해 바뀜 비교(영농-5 · 필지 전이표) ─────────────────────────────────────────
NONE_CLS = "결과 없음"


async def _parcel_cls(c, s: dict, only: list[str] | None) -> dict[str, tuple[str, float]]:
    """결과 하나 → 필지마다 (주된 종류, 넓이) — 필지 번호가 붙은 결과(필지 목록과 같은 행)가 있으면 그것, 없으면 필지와 겹친 저장 넓이(필지 엑셀과 같은 값)."""
    from .spaces import class_ko
    sets, codes = s["_sets"], s["_codes"]
    if not sets:
        return {}
    args = [sets] + ([codes] if codes is not None else []) + ([only] if only is not None else [])
    sc = (" AND left(d.pnu, 5) = ANY($2::text[])" if codes is not None else "") + (f" AND d.pnu = ANY(${len(args)}::text[])" if only is not None else "")
    rows = await c.fetch("SELECT d.pnu, d.cls, d.cls_en, coalesce(sum(d.area_m2), 0) m2 FROM detections d WHERE d.job_id = ANY($1::text[]) "
                         f"AND d.edit_state <> 'deleted' AND d.pnu IS NOT NULL{sc} GROUP BY 1, 2, 3", *args)
    by: dict[str, dict] = {}
    for r in rows:
        k = class_ko(r["cls"], r["cls_en"])
        e = by.setdefault(r["pnu"], {})
        e[k] = e.get(k, 0.0) + float(r["m2"])
    if not by and await c.fetchval("SELECT to_regclass('public.survey_parcel_ai') IS NOT NULL"):
        sc2 = sc.replace("d.pnu", "x.pnu")
        for r in await c.fetch("SELECT x.pnu, x.cls_ko, sum(x.hit1_m2) m2 FROM survey_parcel_ai x WHERE x.job_id = ANY($1::text[]) AND x.n1 > 0"
                               f"{sc2} GROUP BY 1, 2", *args):
            e = by.setdefault(r["pnu"], {})
            e[r["cls_ko"]] = e.get(r["cls_ko"], 0.0) + float(r["m2"] or 0)
    out = {}
    for k, e in by.items():
        top = max(e.items(), key=lambda kv: kv[1])
        out[k] = (top[0], sum(e.values()))
    return out


async def _transitions(p: Principal, picks: list[dict], emd: str | None) -> dict | None:
    """고른 결과 가운데 같은 종류 이름을 쓰는 두 결과(가장 이른 것 · 가장 늦은 것)를 필지 번호로 맞춰 '이전 종류 → 이후 종류' 표.
    대상 = 나중 결과의 필지(합계 = 나중 결과의 필지 수) · 이전 결과에 없던 필지는 '결과 없음'에서 온 줄.
    두 결과가 없거나 필지로 이어지지 않으면 None(화면 · 문서에 블록 없음 — 지어내지 않는다)."""
    if len(picks) < 2:
        return None
    order = sorted(picks, key=lambda s: (s.get("analyzed") or "", s.get("shot") or ""))
    async with db(realm="lx") as c:
        await c.execute("SET LOCAL jit = off")
        cls_of: dict[str, dict] = {}
        z, Zm = None, {}
        for s in reversed(order):                                  # 나중 결과(필지로 이어지는 것)부터
            Zm = await _parcel_cls(c, s, None)
            if Zm:
                z = s
                break
        if not z:
            return None
        cls_of[z["card"]] = Zm
        for s in order:
            if s is not z and (s.get("analyzed") or "", s.get("shot") or "") <= (z.get("analyzed") or "", z.get("shot") or ""):
                cls_of[s["card"]] = await _parcel_cls(c, s, list(Zm))
        cands = [s for s in order if cls_of.get(s["card"]) and s is not z] + [z]
        if len(cands) < 2:
            return None
        zc = {v[0] for v in cls_of[z["card"]].values()}
        a = next((s for s in cands[:-1] if {v[0] for v in cls_of[s["card"]].values()} & zc), None)
        if not a:
            return None
        A, Z = cls_of[a["card"]], cls_of[z["card"]]
        keys = set(Z)
        attrs = {r["pnu"]: dict(r) for r in await c.fetch("SELECT pnu, addr, emd FROM survey_parcels WHERE pnu = ANY($1::text[])", list(keys))} if keys else {}
    if emd:
        keys = {k for k in keys if (attrs.get(k) or {}).get("emd") == emd}
    if not (set(A) & set(Z) & keys):
        return None
    cell: dict[tuple, list] = {}
    moved = []
    for k in keys:
        fa = A[k][0] if k in A else NONE_CLS
        fz = Z[k][0] if k in Z else NONE_CLS
        m2 = (Z.get(k) or A.get(k))[1]
        e = cell.setdefault((fa, fz), [0, 0.0])
        e[0] += 1
        e[1] += m2
        if fa != fz and k in A and k in Z:
            at = attrs.get(k) or {}
            from .gov_report import _addr_short
            moved.append({"pnu": k, "where": _addr_short(at.get("addr")) or k, "emd": at.get("emd") or "", "from": fa, "to": fz,
                          "area": env(round(m2, 1), "m2", "measured", SRC)})
    moved.sort(key=lambda x: -x["area"]["value"])
    rows = sorted(cell.items(), key=lambda kv: (kv[0][0] == kv[0][1], kv[0][0] == NONE_CLS or kv[0][1] == NONE_CLS, -kv[1][0]))
    ch = sum(v[0] for (fa, fz), v in cell.items() if fa != fz and NONE_CLS not in (fa, fz))
    src = "두 결과 × 필지 번호(한 필지 = 넓이가 큰 종류)"
    return {"from": {"card": a["card"], "short": a["short"], "shot": a["shot"], "analyzed": a["analyzed"]},
            "to": {"card": z["card"], "short": z["short"], "shot": z["shot"], "analyzed": z["analyzed"]},
            "rows": [{"from": fa, "to": fz, "n": env(v[0], "필지", "inferred", src), "ha": env(round(v[1] / 1e4, 2), "ha", "measured", src),
                      "changed": fa != fz} for (fa, fz), v in rows],
            "changed": env(ch, "필지", "inferred", src, "두 결과 모두 있는 필지 가운데 종류가 바뀐 필지"),
            "parcels": env(len(keys), "필지", "inferred", src), "top": moved[:10], "none_label": NONE_CLS}


def _cards_of(cards: str | None) -> list[str]:
    out = []
    for c in str(cards or "").split(","):
        c = c.strip()
        if c and CARD_RE.match(c) and c not in out:
            out.append(c)
    if not out:
        raise ApiError("bad_request", "넣을 결과를 하나 이상 골라 주세요", {"field": "cards"})
    return out[:8]


async def multi_data(p: Principal, cards: str | None, emd: str | None, import_id: str | None, with_avail: bool = True) -> dict:
    cl = _cards_of(cards)
    emd = (emd or "").strip() or None
    avail = await _avail(p) if with_avail else []
    picks = []
    for c in cl:
        try:
            picks.append(await _one(p, c, emd))
        except ApiError:
            continue                          # 받지 않은 · 결과 없는 서비스는 조용히 뺀다(있다는 것도 알리지 않는다)
    if not picks:
        raise ApiError("not_found", "고른 결과를 불러오지 못했습니다")
    emds = sorted({e for s in picks for e in s["emds"]})
    if emd and emd not in emds:
        raise ApiError("not_found", "그 읍면동에는 고른 결과가 없습니다")
    ld = await _ledger_cmp(p, import_id, emd, picks) if (import_id or "").strip() else None
    place = next((s["place"] for s in picks if s["place"]), "")
    shots = []
    for s in picks:
        if s["shot"] and s["shot"] not in shots:
            shots.append(s["shot"])
    days = sorted({s["analyzed"] for s in picks if s["analyzed"]})
    return {"region": f"{place} {emd}".strip() if emd else (place or "—"), "place": place, "emd": emd, "emds": emds,
            "made": dt.datetime.now(KST).strftime("%Y.%m.%d"), "shots": shots, "days": days,
            "bale": any(c["bale"] for s in picks for c in s["classes"]),
            "picked": picks, "ledger": ld, "transitions": await _transitions(p, picks, emd), "services": avail, "as_of": now_iso()}


def _pub(d: dict) -> dict:
    return {**d, "picked": [{k: v for k, v in s.items() if not k.startswith("_")} for s in d["picked"]]}


def _title(d: dict) -> str:
    return " · ".join(s["short"] for s in d["picked"]) + (f" · {d['emd']}" if d.get("emd") else "")


@router.get("/spaces/me/report/multi")
async def report_multi(request: Request, cards: str | None = None, emd: str | None = None, import_id: str | None = None):
    p = _me(request)
    return _pub(await multi_data(p, cards, emd, import_id))


@router.get("/spaces/me/report/transitions")
async def report_transitions(request: Request, cards: str | None = None, emd: str | None = None):
    """필지 전이표만(영농-5) — 보고서 탭이 고른 결과가 둘 이상일 때 부른다. 같은 종류 이름을 쓰는 두 결과가 없으면 transitions = null(블록 없음)."""
    p = _me(request)
    emd = (emd or "").strip() or None
    picks = []
    for c in _cards_of(cards):
        try:
            picks.append(await _one(p, c, emd))
        except ApiError:
            continue
    return {"transitions": await _transitions(p, picks, emd) if len(picks) >= 2 else None, "as_of": now_iso()}


@router.get("/spaces/me/report/pick")
async def report_pick(request: Request):
    """고를 것만(처음 열 때 · 결과를 아직 고르지 않았을 때) — 결과 목록 · 발급 목록 · 행정정보 목록을 한 번에."""
    p = _me(request)
    from .farm_admin import ledgers
    led = await ledgers(request)
    return {"services": await _avail(p), "issues": await _issues(p), "ledgers": [{"label": g["label"], "items": [
        {"import_id": x["import_id"], "name": x["name"], "date": x["base_date"] or x["uploaded"]} for x in g["items"] if not x.get("state")]} for g in led["groups"]],
        "as_of": now_iso()}


async def _issue(p: Principal, d: dict, import_id: str | None, fmt: str):
    cards = [s["card"] for s in d["picked"]]
    key = ",".join(sorted(cards)) + "|" + (d.get("emd") or "") + "|" + (import_id or "")
    async with db(p) as c:
        await c.execute(
            "INSERT INTO report_issues(id, tenant_id, user_id, pick_key, cards, emd, import_id, title, fmt) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9) "
            "ON CONFLICT (tenant_id, user_id, pick_key) DO UPDATE SET downloads = report_issues.downloads + 1, last_at = now(), fmt = EXCLUDED.fmt, title = EXCLUDED.title",
            "ri_" + secrets.token_hex(8), p.tenant_id, p.user_id or "", key, cards, d.get("emd"), import_id or None, _title(d), fmt)
        await c.execute("INSERT INTO space_log(tenant_id, kind, card_id, line, actor, realm, detail) VALUES ($1,'download',$2,$3,$4,'tenant',$5)",
                        p.tenant_id, cards[0], f"{p.name or '담당자'} — AI 분석 결과 보고서 내려받기({_title(d)})", p.user_id,
                        {"fmt": "report_multi" if fmt == "hwpx" else "report_xlsx", "cards": cards, "emd": d.get("emd"), "import_id": import_id or None})


async def _issues(p: Principal) -> list[dict]:
    async with db(p) as c:
        rows = await c.fetch("SELECT r.id, r.cards, r.emd, r.import_id, r.title, r.fmt, r.downloads, r.at, r.last_at, u.name, u.dept, u.position "
                             "FROM report_issues r LEFT JOIN tenant_users u ON u.id = r.user_id WHERE r.tenant_id=$1 ORDER BY r.last_at DESC LIMIT 50", p.tenant_id)
    return [{"id": r["id"], "cards": list(r["cards"] or []), "emd": r["emd"], "import_id": r["import_id"], "title": r["title"], "fmt": r["fmt"],
             "downloads": env(int(r["downloads"]), "회", "recorded", "보고서 발급 목록"),
             "at": r["at"].astimezone(KST).isoformat(timespec="minutes"), "last_at": r["last_at"].astimezone(KST).isoformat(timespec="minutes"),
             "who": " · ".join(x for x in ((r["dept"] or "").strip(), " ".join(y for y in ((r["name"] or "").strip(), (r["position"] or "").strip()) if y)) if x) or "기관 담당자"}
            for r in rows]


@router.get("/spaces/me/report/issues")
async def report_issues(request: Request):
    p = _me(request)
    return {"items": await _issues(p), "as_of": now_iso()}


def _hwpx_multi(d: dict, pngs: list[tuple[str, bytes | None]]) -> tuple[bytes, str]:
    from . import hwpx as HX
    now = dt.datetime.now(KST)
    made_t = now.strftime("%Y. %m. %d. %H:%M")
    P = d["picked"]
    one = len(P) == 1
    doc = HX.Doc()
    doc.title("AI 분석 결과 보고서", f"{d['region']} · {_title(d)}")
    ov = [("지역", d["region"]), ("넣은 결과", " · ".join(s["name"] for s in P)), ("영상", " · ".join(d["shots"]) or "—"),
          ("분석한 날", " · ".join(x for x in d["days"]) or "—"), ("만든 날짜", made_t)]
    if one and P[0]["acc"]["value"] is not None:
        ov.insert(4, ("AI 모델 정확도", f"{P[0]['acc']['value']}%"))
    if d.get("ledger"):
        ov.append(("대장", f"{d['ledger']['name']} · {d['ledger']['parcels']['value']:,}필지"))
    doc.overview(ov)
    doc.notice("AI 분석 결과는 참고자료입니다 — 처분은 담당자가 정합니다.")
    if not one:
        doc.heading("시점 비교")
        head = ["결과", "영상", "분석한 날", "AI 분석 결과", "넓이(ha)"] + (["곤포사일리지"] if d["bale"] else [])
        rows = [head]
        for s in P:
            b = next((c for c in s["classes"] if c["bale"]), None)
            y = ((b or {}).get("yield") or {}).get("value")
            rows.append([s["short"], s["shot"] or "—", s["analyzed"] or "—", f"{s['total']['value']:,}{s['unit']}", f"{s['area_ha']['value']:,.1f}"]
                        + ([(f"{b['n']['value']:,}개" + (f" · {y:,.0f}kg" if y else "")) if b else "—"] if d["bale"] else []))
        doc.table(rows, [20, 18, 14, 18, 14] + ([16] if d["bale"] else []), aligns=["l", "l", "l", "r", "r"] + (["r"] if d["bale"] else []))
    doc.heading("종류별")
    if one:
        s = P[0]
        rows = [["종류", "개수", "넓이(ha)"]]
        for c in s["classes"]:
            y = (c.get("yield") or {}).get("value")
            rows.append([c["name"] + (f" (생산량 {y:,.0f}kg)" if y else ""), f"{c['n']['value']:,}{s['unit']}", f"{c['ha']['value']:,.1f}"])
        rows.append(["합계", f"{s['total']['value']:,}{s['unit']}", f"{s['area_ha']['value']:,.1f}"])
        doc.table(rows, [44, 28, 28], aligns=["l", "r", "r"], total_row=True)
    else:
        names = []
        for s in P:
            for c in s["classes"]:
                if c["name"] not in names:
                    names.append(c["name"])
        ha = {s["card"]: {c["name"]: c["ha"]["value"] for c in s["classes"]} for s in P}
        two = len(P) == 2
        rows = [["종류", *[f"{s['short']}(ha)" for s in P], *(["차이(ha)"] if two else [])]]
        for n in names:
            vals = [ha[s["card"]].get(n) for s in P]
            rows.append([n, *[f"{v:,.1f}" if v is not None else "—" for v in vals],
                         *([f"{(vals[1] or 0) - (vals[0] or 0):+,.1f}"] if two else [])])
        tv = [s["area_ha"]["value"] for s in P]
        rows.append(["합계", *[f"{v:,.1f}" for v in tv], *([f"{tv[1] - tv[0]:+,.1f}"] if two else [])])
        doc.table(rows, [3] + [2] * (len(rows[0]) - 1), aligns=["l"] + ["r"] * (len(rows[0]) - 1), total_row=True, small=len(rows[0]) > 6)
    tr = d.get("transitions")
    if tr:
        doc.heading("필지 전이")
        rows = [[f"{tr['from']['short']}", f"{tr['to']['short']}", "필지", "넓이(ha)"]]
        rows += [[x["from"], x["to"], f"{x['n']['value']:,}", f"{x['ha']['value']:,.2f}"] for x in tr["rows"]]
        rows.append(["합계", "", f"{tr['parcels']['value']:,}", f"{sum(x['ha']['value'] for x in tr['rows']):,.2f}"])
        doc.table(rows, [30, 30, 20, 20], aligns=["l", "l", "r", "r"], total_row=True,
                  caption=f"필지 번호로 맞춤 · 한 필지 = 넓이가 큰 종류 · 종류가 바뀐 필지 {tr['changed']['value']:,}필지")
        if tr["top"]:
            rows = [["소재지", "읍면동", "바뀜", "넓이(㎡)"]] + [[x["where"], x["emd"], f"{x['from']} → {x['to']}", f"{x['area']['value']:,.0f}"] for x in tr["top"]]
            doc.table(rows, [40, 16, 28, 16], aligns=["l", "l", "l", "r"], caption="바뀐 필지(넓은 순 10)")
    for name, png in pngs:
        if png:
            doc.heading(f"지도 — {name}")
            doc.figure(png, f"{d['region']} {name} AI 분석 결과", "읍면동 경계 위에 AI 분석 결과를 칠했습니다", max_w=int(HX.TEXT_W * 0.82))
    doc.heading("읍면동별")
    if one:
        s = P[0]
        cls_names = [c["name"] for c in s["classes"]]
        many = len(cls_names) > 1
        head = ["읍면동", *(cls_names if many else []), "계", "넓이(ha)"]
        body = [[e["emd"], *([f"{({b['cls']: b['n'] for b in e['by']}).get(k, 0):,}" for k in cls_names] if many else []), f"{e['n']:,}", f"{e['ha']['value']:,.1f}"]
                for e in s["emd_rows"]]
        body.append(["합계", *([f"{c['n']['value']:,}" for c in s["classes"]] if many else []), f"{s['total']['value']:,}", f"{s['area_ha']['value']:,.1f}"])
        doc.table([head] + body, [3] + [2] * (len(head) - 1), aligns=["l"] + ["r"] * (len(head) - 1), total_row=True, small=len(head) > 6,
                  caption=f"읍면동별 개수 · 넓이 (단위: {s['unit']} · ha)")
    else:
        emds = []
        for s in P:
            for e in s["emd_rows"]:
                if e["emd"] not in emds:
                    emds.append(e["emd"])
        m = {s["card"]: {e["emd"]: e["n"] for e in s["emd_rows"]} for s in P}
        head = ["읍면동", *[f"{s['short']}({s['unit']})" for s in P]]
        body = [[x, *[f"{m[s['card']].get(x, 0):,}" for s in P]] for x in emds]
        body.append(["합계", *[f"{s['total']['value']:,}" for s in P]])
        doc.table([head] + body, [3] + [2] * len(P), aligns=["l"] + ["r"] * len(P), total_row=True, small=len(head) > 6, caption="읍면동별 AI 분석 결과")
    ld = d.get("ledger")
    if ld:
        doc.heading("대장과 견준 결과")
        rows = [["항목", "필지"], ["올린 대장", f"{ld['name']} · {ld['parcels']['value']:,}필지"]]
        if one:
            rows.append([ld["labels"]["found"], f"{ld['per'][0]['found']['value']:,}필지"])
            rows.append([ld["labels"]["missing"], f"{ld['per'][0]['missing']['value']:,}필지"])
            rows += [[x["label"], f"{x['n']['value']:,}필지"] for x in ld["lines"]]
        else:
            rows.append([ld["labels"]["all"], f"{ld['all']['value']:,}필지"])
            rows.append([ld["labels"]["some"], f"{ld['some']['value']:,}필지"])
            rows += [[f"{x['short']} AI 결과 없음", f"{x['missing']['value']:,}필지"] for x in ld["per"]]
        doc.table(rows, [60, 40], aligns=["l", "r"])
        doc.small("필지 목록(지번)은 '행정정보와 비교'에서 봅니다.")
    doc.source("Land-XI AI 분석 결과 — 개수 · 넓이는 서비스 대시보드 · 필지 목록 · 내려받기 파일과 같은 값입니다(결과 확인 전 · 지운 결과 제외).", made_t)
    fname = HX.safe_name(f"{d['region']}_{'_'.join(s['short'] for s in P)}_AI분석결과보고서_{d['made'].replace('.', '')}")
    return doc.package(f"AI 분석 결과 보고서 — {d['region']} · {_title(d)}", now), fname


@router.get("/spaces/me/report/multi.hwpx")
async def report_multi_hwpx(request: Request, cards: str | None = None, emd: str | None = None, import_id: str | None = None):
    from .gov_report import _png
    from .xreport import hwpx_response
    p = _me(request)
    d = await multi_data(p, cards, emd, import_id, with_avail=False)
    pngs = []
    for s in d["picked"][:4]:                         # 지도 그림은 넷까지(문서가 너무 길어지지 않게)
        try:
            pngs.append((s["short"], await _png(p, s["card"], d["emd"])))
        except ApiError:
            pngs.append((s["short"], None))
    blob, name = await run_in_threadpool(_hwpx_multi, d, pngs)
    await _issue(p, d, (import_id or "").strip() or None, "hwpx")
    return hwpx_response(blob, name)


def _xlsx_multi(d: dict) -> bytes:
    wb = _book()
    P = d["picked"]
    one = len(P) == 1
    ov = [["지역", d["region"]], ["넣은 결과", " · ".join(s["name"] for s in P)], ["영상", " · ".join(d["shots"]) or "—"],
          ["분석한 날", " · ".join(d["days"]) or "—"], ["만든 날", d["made"]]]
    if d.get("ledger"):
        ov.append(["대장", f"{d['ledger']['name']} · {d['ledger']['parcels']['value']:,}필지"])
    ov.append(["안내", "AI 분석 결과는 참고자료입니다 — 처분은 담당자가 정합니다."])
    cap = cap_line(d.get("org") or "", " · ".join(s["short"] for s in P), d["region"], d["made"].replace(".", "-"), "AI 분석 결과(관할 안 · 지운 결과 제외)")
    _sheet(wb, "개요", ["항목", "내용"], ov, widths=[16, 70], cap=cap)
    head = ["결과", "영상", "분석한 날", "AI 분석 결과", "단위", "넓이(ha)"] + (["곤포사일리지(개)", "생산량(kg)"] if d["bale"] else [])
    rows = []
    for s in P:
        b = next((c for c in s["classes"] if c["bale"]), None)
        rows.append([s["name"], s["shot"] or "", s["analyzed"] or "", s["total"]["value"], s["unit"], s["area_ha"]["value"]]
                    + ([b["n"]["value"] if b else None, ((b or {}).get("yield") or {}).get("value")] if d["bale"] else []))
    _sheet(wb, "시점 비교" if not one else "결과", head, rows, num={3: COUNT, 5: AREA, 6: COUNT, 7: COUNT}, widths=[26, 18, 12, 14, 8, 12, 16, 14], cap=cap, total=False)
    names = []
    for s in P:
        for c in s["classes"]:
            if c["name"] not in names:
                names.append(c["name"])
    cm = {s["card"]: {c["name"]: c for c in s["classes"]} for s in P}
    head = ["종류"] + [x for s in P for x in (f"{s['short']} 개수({s['unit']})", f"{s['short']} 넓이(ha)")]
    rows = [[n] + [v for s in P for v in ((cm[s["card"]].get(n) or {}).get("n", {}).get("value"), (cm[s["card"]].get(n) or {}).get("ha", {}).get("value"))] for n in names]
    rows.append(["합계"] + [v for s in P for v in (s["total"]["value"], s["area_ha"]["value"])])
    _sheet(wb, "종류별", head, rows, num={i: (COUNT if i % 2 else AREA) for i in range(1, len(head))}, total_last=True, cap=cap)
    emds = []
    for s in P:
        for e in s["emd_rows"]:
            if e["emd"] not in emds:
                emds.append(e["emd"])
    em = {s["card"]: {e["emd"]: e for e in s["emd_rows"]} for s in P}
    head = ["읍면동"] + [x for s in P for x in (f"{s['short']} 개수({s['unit']})", f"{s['short']} 넓이(ha)")]
    rows = [[x] + [v for s in P for v in ((em[s["card"]].get(x) or {}).get("n"), ((em[s["card"]].get(x) or {}).get("ha") or {}).get("value"))] for x in emds]
    rows.append(["합계"] + [v for s in P for v in (s["total"]["value"], s["area_ha"]["value"])])
    _sheet(wb, "읍면동별", head, rows, num={i: (COUNT if i % 2 else AREA) for i in range(1, len(head))}, total_last=True, cap=cap)
    tr = d.get("transitions")
    if tr:
        rows = [[x["from"], x["to"], x["n"]["value"], x["ha"]["value"]] for x in tr["rows"]]
        rows.append(["합계", "", tr["parcels"]["value"], round(sum(x["ha"]["value"] for x in tr["rows"]), 2)])
        _sheet(wb, "필지 전이", [tr["from"]["short"], tr["to"]["short"], "필지", "넓이(ha)"], rows, num={2: COUNT, 3: "#,##0.00"}, widths=[20, 20, 12, 12],
               total_last=True, cap=cap)
        if tr["top"]:
            _sheet(wb, "바뀐 필지", ["필지 번호", "소재지", "읍면동", "이전", "이후", "넓이(㎡)"],
                   [[x["pnu"], x["where"], x["emd"], x["from"], x["to"], x["area"]["value"]] for x in tr["top"]], num={5: AREA}, text_cols=(0,),
                   widths=[22, 30, 12, 14, 14, 12], cap=cap)
    ld = d.get("ledger")
    if ld:
        rows = [["올린 대장 필지", ld["parcels"]["value"]]]
        if one:
            rows += [[ld["labels"]["found"], ld["per"][0]["found"]["value"]], [ld["labels"]["missing"], ld["per"][0]["missing"]["value"]]]
            rows += [[x["label"], x["n"]["value"]] for x in ld["lines"]]
        else:
            rows += [[ld["labels"]["all"], ld["all"]["value"]], [ld["labels"]["some"], ld["some"]["value"]]]
            rows += [[f"{x['short']} AI 결과 없음", x["missing"]["value"]] for x in ld["per"]]
        _sheet(wb, "대장 대조", ["항목", "필지"], rows, num={1: COUNT}, widths=[40, 14], cap=cap, total=False)
    return _save(wb)


@router.get("/spaces/me/report/multi.xlsx")
async def report_multi_xlsx(request: Request, cards: str | None = None, emd: str | None = None, import_id: str | None = None):
    p = _me(request)
    d = await multi_data(p, cards, emd, import_id, with_avail=False)
    blob = await run_in_threadpool(_xlsx_multi, d)
    await _issue(p, d, (import_id or "").strip() or None, "xlsx")
    org = await _org_short(p.tenant_id) if p.tenant_id else ""
    return _xlsx_resp(blob, f"{org}_{'_'.join(s['short'] for s in d['picked'])}_보고서{('_' + d['emd']) if d.get('emd') else ''}_{_today()}")


# ── 통계 · 이력 엑셀 ─────────────────────────────────────────────────────────
@router.get("/spaces/me/guides/{card}/class-stats.xlsx")
async def class_stats_xlsx(card: str, request: Request):
    from .gov_report import _short_svc
    from .spaces import space_class_stats
    p = _me(request)
    j = await space_class_stats(card, request)
    classes, emds = j.get("classes") or [], j.get("emds") or []
    if not classes:
        raise ApiError("not_found", "이 서비스는 종류별 통계가 없습니다")
    org = await _org_short(p.tenant_id)
    cnt = bool(j.get("count"))
    nw = "필지 수" if j.get("unit") == "필지" else "개수"
    v = lambda e: (round(float(e["value"]), 1) if e and e.get("value") is not None else None)   # noqa: E731
    yk = {c["name"] for c in classes if c.get("yield")}

    def build() -> bytes:
        wb = _book()
        tv = float((j.get("total") or {}).get("area", {}).get("value") or 0) or 1.0
        head = ["종류", "넓이(㎡)", "비율"] + ([nw] if cnt else []) + (["생산량(kg)"] if yk else [])
        rows = [[c["name"], v(c.get("area")), round(float((c.get("area") or {}).get("value") or 0) / tv * 100, 1)] + ([v(c.get("n"))] if cnt else [])
                + ([v(c.get("yield"))] if yk else []) for c in classes]
        rows.append(["합계", v((j.get("total") or {}).get("area")), 100.0] + ([v((j.get("total") or {}).get("n"))] if cnt else []) + ([None] if yk else []))
        cap = cap_line(org, j["service"]["name"], "관할 전체", str(j.get("as_of") or "")[:10], "이 서비스 AI 분석 결과(관할 안 · 지운 결과 제외)")
        _sheet(wb, "종류별", head, rows, num={1: AREA, 2: RATIO, 3: COUNT, 4: COUNT}, widths=[22, 16, 10, 12, 14], total_last=True, cap=cap)
        names = [c["name"] for c in classes]
        head = ["읍면동", "합계 넓이(㎡)"] + ([f"합계 {nw}"] if cnt else [])
        for n in names:
            head += [f"{n} 넓이(㎡)"] + ([f"{n} {nw}"] if cnt else []) + ([f"{n} 생산량(kg)"] if n in yk else [])
        rows = []
        for e in emds + [{"emd": f"{org} 전체".strip(), "area": (j.get("total") or {}).get("area"), "n": (j.get("total") or {}).get("n"),
                          "by": {c["name"]: c for c in classes}}]:
            r = [e["emd"], v(e.get("area"))] + ([v(e.get("n"))] if cnt else [])
            for n in names:
                b = (e.get("by") or {}).get(n) or {}
                r += [v(b.get("area"))] + ([v(b.get("n"))] if cnt else []) + ([v(b.get("yield"))] if n in yk else [])
            rows.append(r)
        fm = {}
        for i, hd in enumerate(head[1:], start=1):
            fm[i] = AREA if "㎡" in hd else COUNT
        _sheet(wb, "읍면동별", head, rows, num=fm, total_last=True, cap=cap)
        _info(wb, [("기관", org), ("서비스", j["service"]["name"]), ("기준", str(j.get("as_of") or "")[:16].replace("T", " ")),
                   ("출처", "이 서비스 AI 분석 결과(관할 안 · 지운 결과 제외) — 화면 '종류별 · 읍면동별'과 같은 값"),
                   ("안내", "AI 분석 결과는 참고자료입니다 — 처분은 담당자가 정합니다.")]
                  + ([("넓이만 보이는 까닭", "이 서비스는 분석 칸마다 나눈 결과라 개수를 싣지 않습니다.")] if j.get("note") else []))
        return _save(wb)
    blob = await run_in_threadpool(build)
    await _log(p, card, f"{p.name or '담당자'} — {j['service']['name']} 종류별 · 읍면동별 통계 엑셀 내려받기", {"fmt": "stats", "rows": len(emds)})
    return _xlsx_resp(blob, f"{org}_{_short_svc(j['service']['name'])}_통계_{_today()}")


@router.get("/history.xlsx")
async def history_xlsx(request: Request, card: str, kind: str | None = None):
    from .gov_history import history
    from .gov_report import _short_svc
    j = await history(request, card, 500)
    p = principal(request)
    items = [x for x in j.get("items") or [] if x.get("kind") != "check" and (not kind or kind == "all" or x.get("kind") == kind)]
    name = card
    try:
        from .gov_report import _svc
        name = (await _svc(p, card))[0]["name"]
    except ApiError:
        pass
    org = await _org_short(p.tenant_id) if p.tenant_id else ""

    def build() -> bytes:
        wb = _book()
        rows = []
        for x in items:
            at = None
            try:
                at = dt.datetime.fromisoformat(str(x.get("at"))).astimezone(KST).replace(tzinfo=None)
            except (TypeError, ValueError):
                pass
            rows.append([at, x.get("kind_ko") or "", x.get("title") or "", x.get("sub") or "", x.get("who") or ""])
        cap = cap_line(org, name, "이력" + ("" if not kind or kind == "all" else f" · {items[0]['kind_ko'] if items else kind}"), dt.datetime.now(KST).strftime("%Y-%m-%d"),
                       "결과 공개 · 대장 · 메모 · 요청 기록")
        _sheet(wb, "이력", ["날짜 · 시각", "종류", "내용", "덧붙임", "누가"], rows, widths=[18, 12, 50, 30, 18], num={0: DATETIME}, cap=cap)
        _info(wb, [("기관", org), ("서비스", name), ("고른 종류", "전체" if not kind or kind == "all" else (items[0]["kind_ko"] if items else kind)),
                   ("줄 수", len(items)), ("만든 때", now_iso()[:16].replace("T", " "))])
        return _save(wb)
    blob = await run_in_threadpool(build)
    if p.realm == "tenant":
        await _log(p, card, f"{p.name or '담당자'} — {name} 이력 엑셀 내려받기", {"fmt": "history", "rows": len(items)})
    return _xlsx_resp(blob, f"{org}_{_short_svc(name)}_이력_{_today()}")


# ── 이용 약속 기록(보기만) ───────────────────────────────────────────────────
def _day(s: str | None, field: str) -> dt.date | None:
    s = str(s or "").strip()[:10]
    if not s:
        return None
    try:
        return dt.date.fromisoformat(s)
    except ValueError:
        raise ApiError("bad_request", "기간은 날짜로 골라 주세요", {"field": field})


async def _pledges(realm_p: Principal | None, tenant: str | None, d0: dt.date | None, d1: dt.date | None) -> list[dict]:
    from .spaces import _short_org
    cond, args = ["TRUE"], []
    if tenant:
        args.append(tenant)
        cond.append(f"pl.tenant_id = ${len(args)}")
    if d0:
        args.append(d0)
        cond.append(f"(pl.at AT TIME ZONE 'Asia/Seoul')::date >= ${len(args)}")
    if d1:
        args.append(d1)
        cond.append(f"(pl.at AT TIME ZONE 'Asia/Seoul')::date <= ${len(args)}")
    sql = ("SELECT pl.id, pl.tenant_id, pl.purpose, pl.until_date, pl.at, u.name, u.dept, u.position, t.name AS tname, "
           "(SELECT count(*) FROM space_log s WHERE s.tenant_id = pl.tenant_id AND s.kind='download' AND s.detail->>'pledge' = pl.id) AS n, "
           "(SELECT array_agg(DISTINCT s.detail->>'fmt') FROM space_log s WHERE s.tenant_id = pl.tenant_id AND s.kind='download' AND s.detail->>'pledge' = pl.id) AS kinds, "
           "(SELECT max(s.at) FROM space_log s WHERE s.tenant_id = pl.tenant_id AND s.kind='download' AND s.detail->>'pledge' = pl.id) AS last "
           "FROM download_pledges pl LEFT JOIN tenant_users u ON u.id = pl.user_id LEFT JOIN tenants t ON t.id = pl.tenant_id "
           f"WHERE {' AND '.join(cond)} ORDER BY pl.at DESC LIMIT 2000")
    async with (db(realm_p) if realm_p else db(realm="lx")) as c:
        rows = await c.fetch(sql, *args)
    today = dt.datetime.now(KST).date()
    out = []
    for r in rows:
        tn = r["tname"]
        tn = json.loads(tn) if isinstance(tn, str) and tn.startswith("{") else tn
        tn = (tn or {}).get("ko") if isinstance(tn, dict) else (tn or r["tenant_id"])
        out.append({"id": r["id"], "org": _short_org(tn), "who": " · ".join(x for x in ((r["dept"] or "").strip(), " ".join(y for y in ((r["name"] or "").strip(), (r["position"] or "").strip()) if y)) if x) or "기관 담당자",
                    "at": r["at"].astimezone(KST).isoformat(timespec="minutes"), "purpose": r["purpose"], "until": r["until_date"].isoformat(),
                    "live": r["until_date"] >= today,
                    "kinds": [FMT_KO.get(k, "자료") for k in (r["kinds"] or []) if k],
                    "n": env(int(r["n"] or 0), "회", "recorded", "그 약속으로 받은 개인정보가 든 파일"),
                    "last": r["last"].astimezone(KST).isoformat(timespec="minutes") if r["last"] else None})
    return out


def _pledge_xlsx(rows: list[dict], lx: bool, d0, d1) -> bytes:
    wb = _book()
    head = (["기관"] if lx else []) + ["누가", "확인한 때", "사용 목적", "활용 기간(끝나는 날)", "지금", "받은 파일 종류", "받은 횟수", "마지막으로 받은 때"]
    data = []
    for r in rows:
        at = dt.datetime.fromisoformat(r["at"]).replace(tzinfo=None)
        last = dt.datetime.fromisoformat(r["last"]).replace(tzinfo=None) if r["last"] else None
        data.append(([r["org"]] if lx else []) + [r["who"], at, r["purpose"], dt.date.fromisoformat(r["until"]), "기간 안" if r["live"] else "기간 끝",
                                                  " · ".join(r["kinds"]) or "—", r["n"]["value"], last])
    o = 1 if lx else 0
    rng = f"{d0.isoformat() if d0 else '처음'} ~ {d1.isoformat() if d1 else '오늘'}"
    cap = cap_line("모든 기관" if lx else (rows[0]["org"] if rows else ""), "개인정보 이용 확인", rng, dt.datetime.now(KST).strftime("%Y-%m-%d"), "개인정보가 든 파일 받기 기록")
    _sheet(wb, "개인정보 이용 확인", head, data, num={o + 1: DATETIME, o + 3: DATE, o + 6: COUNT, o + 7: DATETIME},
           widths=([16] if lx else []) + [24, 18, 40, 16, 10, 30, 10, 18], cap=cap)
    _info(wb, [("기간", f"{rng}(확인한 날 기준)"), ("줄 수", len(rows)),
               ("규칙", "올린 대장의 필지별 내용(지번 · 대장 값)이 든 파일만 받을 때마다 사용 목적 · 활용 기간을 확인합니다."),
               ("만든 때", now_iso()[:16].replace("T", " "))])
    return _save(wb)


def _mgr(request: Request) -> Principal:
    p = _me(request)
    if p.role != "manager":
        raise ApiError("forbidden", "기관 관리자만 볼 수 있습니다")
    return p


@router.get("/spaces/me/pledges")
async def my_pledges(request: Request):
    p = _mgr(request)
    q = request.query_params
    d0, d1 = _day(q.get("from"), "from"), _day(q.get("to"), "to")
    return {"items": await _pledges(p, p.tenant_id, d0, d1), "as_of": now_iso()}


@router.get("/spaces/me/pledges.xlsx")
async def my_pledges_xlsx(request: Request):
    p = _mgr(request)
    q = request.query_params
    d0, d1 = _day(q.get("from"), "from"), _day(q.get("to"), "to")
    rows = await _pledges(p, p.tenant_id, d0, d1)
    blob = await run_in_threadpool(_pledge_xlsx, rows, False, d0, d1)
    org = await _org_short(p.tenant_id)
    return _xlsx_resp(blob, f"{org}_다운로드이력_{_today()}")


@router.get("/admin/pledges")
async def all_pledges(request: Request):
    require(principal(request), admin=True)
    q = request.query_params
    d0, d1 = _day(q.get("from"), "from"), _day(q.get("to"), "to")
    t = (q.get("tenant") or "").strip() or None
    from .spaces import _short_org
    async with db(realm="lx") as c:
        ts = await c.fetch("SELECT id, name FROM tenants WHERE id <> 'lx' ORDER BY id")
    orgs = []
    for r in ts:
        n = r["name"]
        n = json.loads(n) if isinstance(n, str) and n.startswith("{") else n
        orgs.append({"id": r["id"], "name": _short_org((n or {}).get("ko") if isinstance(n, dict) else (n or r["id"]))})
    return {"items": await _pledges(None, t, d0, d1), "orgs": orgs, "as_of": now_iso()}


@router.get("/admin/pledges.xlsx")
async def all_pledges_xlsx(request: Request):
    require(principal(request), admin=True)
    q = request.query_params
    d0, d1 = _day(q.get("from"), "from"), _day(q.get("to"), "to")
    t = (q.get("tenant") or "").strip() or None
    rows = await _pledges(None, t, d0, d1)
    blob = await run_in_threadpool(_pledge_xlsx, rows, True, d0, d1)
    return _xlsx_resp(blob, f"다운로드이력_모든기관_{_today()}")
