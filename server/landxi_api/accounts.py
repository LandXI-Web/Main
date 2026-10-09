"""계정 — 가입 신청 · 아이디 찾기 · 비밀번호 찾기(재설정 요청) · 승인하는 화면(LX 관리자 · 기관 관리자) (구현 2차 T5).

근거: 확인 대장 FR-4(계정을 관리자 화면에서 만들기·승인·비밀번호 재설정·잠금 — 확인) · D4-ⓑ(기관 사용자 가입 신청 + 승인 — 확인) ·
      원칙 49(입력 최소) · 원칙 72(LX 관리자는 다 보고 돕는다) · 원칙 77(아이디 = 메일 주소) · 10-01 사용자 "가입 신청과 아이디 비밀번호 찾기 등 기능".

로그인 없이 부르는 길(시도 제한 — 접속 주소마다 · 입력 검증):
  POST /accounts/signup         {site?, tenant_id?, name, login(메일), password, password2?, dept, contact?, consent:true} → 201 {ok, state:'pending'}
                                  contact = 연락처(선택 · 기관 사용자만 · 본인이 적는다 — 기관-8 ⓐ · 원칙 105) → 승인되면 그 계정에
                                  site 가 admin 이면 받지 않는다(LX 관리자 계정은 LX 관리자가 만든다) · app = LX 직원 신청 · gov = 기관 사용자 신청(tenant_id 필수)
  POST /accounts/find-id        {site?, tenant_id?, name} → {items:[가린 메일]} (예: te**@lx.or.kr) · 없으면 빈 목록
  POST /accounts/reset-request  {site?, tenant_id?, login} → {ok} — 계정이 있든 없든 같은 답(있는 메일을 알아내는 데 쓰이지 않게)
  POST /auth/password/change    {change_token, password} → 로그인과 같은 답(세션) — 임시 비밀번호로 들어온 사람이 새 비밀번호를 정한다
입구(site)는 바깥 주소면 공개 관문이 알린 x-lx-site 가 이기고, 이 PC 안이면 본문 site(없으면 tenant_id 유무로 app/gov).

승인하는 사람(관할 — 서버가 정본 · 원칙 39):
  LX 직원 신청 · LX 계정 재설정 = LX 관리자 · 기관 사용자 신청 = 그 기관 관리자(manager)만 승인 · 반려.
  LX 관리자는 기관 가입 승인에 관여하지 않는다(원칙 72 · 10-01 "LX 는 기관 가입 승인에는 관여하지 않지만 사용자 확인·통제·지원") —
  기관 신청은 보기만(서버도 결정 403 tenant_signup) · 기관 계정은 확인 · 잠금 · 지원(임시 비밀번호 · 재설정 요청 처리).
  기관 관리자는 자기 기관 것만 보이고 바꿀 수 있다. 내 계정에 걸린 일(재설정 · 잠금 · 역할 · 임시 비밀번호)은 스스로 하지 않는다(409 self_account).
  사용 중지된 계정(status disabled — 옛 아이디 정리 · 0015_mail_accounts.sql)은 목록에 '사용 중지'로 보이고 바꾸지 않는다(409 disabled_account).
  반려는 사유 필수(400 reason_required). 누가 승인·반려·재설정했는지는 요청 행(decided_*)과 audit_log(account.*) 두 곳에 남는다.

로그인한 관리자가 부르는 길:
  GET  /accounts/summary                        → {counts:{signup, reset}}
  GET  /accounts/requests?kind=signup|reset&state=pending|all
  POST /accounts/signup/{id}/decide             {decision: approve|reject, reason}
  POST /accounts/reset/{id}/decide              {decision: issue|reject, reason} → issue 면 {temp_password}(이 응답에서 한 번만)
  GET  /accounts/users?realm=lx|tenant&tenant_id=
  POST /accounts/users/{realm}/{user_id}/lock   {locked: bool, reason?}
  POST /accounts/users/{realm}/{user_id}/role   {role}
  POST /accounts/users/{realm}/{user_id}/temp-password → {temp_password}(한 번만)
  GET  /accounts/log                            → 계정 기록(누가 무엇을 · 최근 100)
  GET  /accounts/failures                       → 실패한 로그인(아이디 · 시각 · 입구 · 접속 주소 · 까닭 · 최근 200 — 비밀번호 값은 없다)
  GET  /accounts/senders?review=a,b&request=x,y → (LX 직원 · LX 관리자) 기관 요청을 보낸 사람 — 이름 · 부서 · 역할 · 기관 + 연락처(선택)
                                                  그 요청을 받은 LX 담당 직원과 LX 관리자만(다른 직원의 요청은 빈 칸 · 기관 계정 403 — 기관-8 ⓐ · 원칙 102)
  잠금: 관리자가 잠근 계정(status locked)과 비밀번호 5번 틀려 10분 잠긴 계정(lock_until · auth.py) 둘 다 lock {locked:false} 로 푼다.

저장 용량 할당 · 늘리기 요청 · 승인(제안 S-19 확인 · 원칙 66 · 91 '한도는 하드웨어에만' · 121 · LX 계정만 — 기관 계정은 범위 밖):
  할당 = 사람마다 정한 값(lx_users.storage_quota_gb) → 없으면 기본 할당(lx_settings storage.default_quota_gb — 계정 관리에서 LX 관리자가 정한다)
         → 없으면 '할당 없음'(지어내지 않는다). 쓴 양 = 내가 프로젝트장인 프로젝트의 저장 공간 합(projects.lead_storage — 한 곳 계산).
  할당을 넘어도 막지 않는다 — 90% 를 넘으면 내 정보 · 파일 올리는 자리에 한 줄 + 창으로 늘리기 요청을 권한다(막을지는 구현 확인 때 사용자에게 여쭘).
  POST /me/storage-request                      {want_gb(원하는 할당), why(이유 한 줄)} → 대기 중 요청은 한 사람에 하나
  GET  /accounts/requests?kind=storage          (LX 관리자) 늘리기 요청 — 가입 신청 · 재설정 요청과 같은 자리 · 같은 모양
  POST /accounts/storage/{id}/decide            {decision: approve|reject, reason} → 승인 = 할당이 원하는 값이 됨 · 반려 = 사유 필수 · 요청한 사람에게 알림(lx_notices)
  POST /accounts/users/lx/{user_id}/quota       {quota_gb | null} 사람마다 할당(null = 기본 할당을 따름)
  GET|PUT /accounts/storage-default             {quota_gb | null} 기본 할당(따로 정하지 않은 계정)
  내 계정의 요청 · 할당은 다른 관리자가 — 관리자 계정이 하나뿐이면 스스로(결재함과 같은 규칙 · approvals.solo_admin).

LX 부서 목록(제안 S-21 확인 · 10-01 사용자 "부서는 일단 LX 누리집 조직도에. 나중엔 사내 시스템에서 불러오는 작업을 할 예정"):
  처음 목록 = server/config/lx-departments.csv(누리집 조직도를 옮긴 파일 · 상위 · 부서 · 단위) — 읽는 곳은 한 군데(dept_source · DEPT_SOURCES · LX_DEPT_SOURCE)라
  사내 시스템 불러오기로 바꿔 끼운다. 고르는 칸의 이름 = '상위 › 부서'(예: 공간정보본부 › 플랫폼사업처 · 상위가 사장 · 부사장 · 감사면 부서만).
  GET  /accounts/depts                          {items:[이름]} — 로그인 없이도(가입 신청 창의 부서 고르기) · 관리자에게는 출처 · 단위 · 마지막 바꿈 · 목록에 없는 부서를 쓰는 계정
  POST /accounts/depts/parse                    (LX 관리자 · 엑셀 xlsx · CSV — 조직도 모양 또는 한 열) → 읽은 목록 미리 보기(저장 안 함)
  PUT  /accounts/depts                          {rows | names} 목록을 통째로 바꾼다 · POST /accounts/depts/reload 출처에서 다시 · /add {name, parent?} · /remove {label}
  목록에 없는 부서(지역본부 아래 지사 등)는 직접 적는다(막지 않는다) · 적은 이름이 목록의 부서와 같으면 목록 이름으로 맞춘다.
  이미 적힌 부서 이름은 바꾸지 않는다(내 정보에서 목록의 이름을 고르게 안내만).

내 정보(본인이 고친다 — 확인 17차 P-5 ⓐ · 원칙 105 · 121 · LX 직원 · LX 관리자 · LX 영업 계정. 기관 계정은 확인 범위 밖):
  GET   /me/profile                             → {login(고정), name, dept, contact, role_ko, org, changed_at, storage{quota_gb, used_bytes, projects}}
  PATCH /me/profile                             {name?, dept?, contact?} → 관리자 승인 없이 바로 바뀌고 감사 기록(account.profile — 바뀐 칸 전 · 후)에 남는다
  저장 용량 = 나에게 할당된 저장 용량(lx_users.storage_quota_gb — 할당 화면은 확인 전이라 없으면 null '할당 없음') + 지금 쓴 양
             (내가 프로젝트장인 프로젝트의 저장 공간 합 — projects.lead_storage · 프로젝트 한 장의 storage 와 같은 식).

임시 비밀번호: 사람이 불러 주기 쉬운 대문자·숫자 12자(헷갈리는 0·O·1·I·L 제외) · 하루 동안만 · 받은 사람은 로그인하면 새 비밀번호를 정해야 들어간다.
비밀번호 규칙(기본값 — 사용자 확인 대기): 10자 이상 · 영문과 숫자를 함께 · 메일 아이디를 그대로 넣지 않음.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import re
import secrets
from decimal import Decimal

from fastapi import APIRouter, File, Request, UploadFile
from starlette.concurrency import run_in_threadpool

from . import auth, config
from .deps import ApiError, Principal, db, pool, principal, redis, require
from .envelope import KST, env, now_iso

router = APIRouter()

MAIL = re.compile(r"^[a-z0-9._%+-]+@[a-z0-9-]+(\.[a-z0-9-]+)*\.[a-z]{2,}$")
ROLES = {"lx": ("admin", "staff", "sales"), "tenant": ("manager", "viewer")}
ROLE_KO = {("lx", "admin"): "LX 관리자", ("lx", "staff"): "LX 직원", ("lx", "sales"): "LX 영업",
           ("tenant", "manager"): "기관 관리자", ("tenant", "viewer"): "부서 사용자"}
NEW_ROLE = {"lx": "staff", "tenant": "viewer"}          # 승인으로 만든 계정의 첫 역할 — 바꾸기는 계정 목록에서
# 시도 제한(접속 주소마다) — (횟수, 초). 가입 신청은 한 기관 직원들이 같은 사무실 주소로 몰려도 막히지 않게 한 시간 20번
LIMITS = {"signup": (20, 3600), "find": (10, 600), "reset": (10, 600), "change": (10, 600)}
TEMP_ALPHA = "ABCDEFGHJKMNPQRSTUVWXYZ23456789"
TEMP_TTL = dt.timedelta(hours=24)
KEEP_REJECTED = dt.timedelta(days=30)
ACTION_KO = {
    "account.signup.request": "가입 신청", "account.signup.approve": "가입 승인", "account.signup.reject": "가입 거절",
    "account.reset.request": "비밀번호 재설정 요청", "account.reset.issue": "임시 비밀번호 발급", "account.reset.reject": "재설정 거절",
    "account.temp.issue": "임시 비밀번호 발급", "account.lock": "잠금", "account.unlock": "잠금 풀기", "account.role": "역할 변경",
    "account.password.change": "새 비밀번호 설정", "account.autolock": "자동 잠금(10분)", "account.profile": "내 정보 고침(본인)",
    "account.quota": "저장 용량 할당", "account.quota.default": "기본 할당 바꿈", "account.storage.request": "저장 용량 늘리기 요청",
    "account.storage.approve": "저장 용량 늘리기 승인", "account.storage.reject": "저장 용량 늘리기 반려", "account.depts": "부서 목록 바꿈",
}
FAIL_KO = {"password": "비밀번호 틀림", "unknown": "없는 아이디", "temp_locked": "잠긴 동안 시도", "locked": "잠긴 계정", "temp_expired": "임시 비밀번호 기간 지남",
           "disabled": "사용 중지된 계정"}
TENANT_SIGNUP_MSG = "기관 가입 신청은 그 기관 관리자가 승인합니다"
# LX 직원 가입 신청은 회사 메일만(구현 확인 2차 Q-2 ⓐ) — 기관 신청은 기관 메일이 제각각이라 제한하지 않는다(그 기관 관리자가 승인으로 거른다)
LX_MAIL_DOMAINS = tuple(d.strip().lower() for d in (config.get("LX_STAFF_MAIL_DOMAINS", "lx.or.kr") or "").split(",") if d.strip())
SITE_KO = {"app": "Land-XI", "admin": "LX 관리자", "gov": "기관"}


# ── 공용 ────────────────────────────────────────────────────────────────────────────────
def _iso(v):
    return v.astimezone(KST).isoformat(timespec="seconds") if v else None


def client_ip(request: Request) -> str:
    """접속 주소 — 바깥 요청은 공개 관문이 x-forwarded-for 를 새로 적어 넘긴다(게이트웨이는 127.0.0.1 에만 열려 있다)."""
    xff = request.headers.get("x-forwarded-for") or ""
    if xff:
        return xff.split(",")[0].strip()
    return request.client.host if request.client else ""


async def _limit(kind: str, request: Request, extra: str = ""):
    n, win = LIMITS[kind]
    k = f"lx:acct:{kind}:{client_ip(request)}{':' + extra if extra else ''}"
    try:
        r = await redis()
        c = await r.incr(k)
        if c == 1:
            await r.expire(k, win)
        if c <= n:
            return
        left = await r.ttl(k)
    except Exception:           # 제한 저장소가 멈춰도 신청·찾기는 막지 않는다(입력 검증은 그대로)
        return
    mins = max(1, round((left if left and left > 0 else win) / 60))
    raise ApiError("too_many_attempts", f"시도가 너무 많습니다. {mins}분 뒤 다시 해 주세요.", status=429)


def _site(request: Request, body: dict) -> tuple[str, str]:
    site = request.headers.get("x-lx-site") or body.get("site") or ("gov" if body.get("tenant_id") else "app")
    if site not in auth.SITE_REALM:
        raise ApiError("bad_request", "site 는 app | admin | gov")
    return site, auth.SITE_REALM[site]


def _josa(word: str, pair: str) -> str:
    """받침이 있으면 앞 글자('을' · '은'), 없으면 뒤 글자('를' · '는')."""
    c = ord(word[-1]) if word else 0
    return pair[0] if 0xAC00 <= c <= 0xD7A3 and (c - 0xAC00) % 28 else pair[1]


def _text(v, most: int, label: str, field: str) -> str:
    s = " ".join(str(v or "").split())
    if not s:
        raise ApiError("bad_request", f"{label}{_josa(label, '을를')} 적어 주세요", {"field": field})
    if len(s) > most:
        raise ApiError("bad_request", f"{label}{_josa(label, '은는')} {most}자까지입니다", {"field": field})
    return s


def _mail(v) -> str:
    s = str(v or "").strip().lower()
    if not s:
        raise ApiError("bad_request", "메일 주소를 적어 주세요", {"field": "login"})
    if len(s) > 254 or not MAIL.match(s):
        raise ApiError("bad_request", "메일 주소 형식을 확인하세요", {"field": "login"})
    return s


_CONTACT_BAD = re.compile(r"[\x00-\x1f\x7f<>]")


def _contact(v) -> str | None:
    """연락처(선택) — 전화 · 내선 등 한 줄 30자. 비우면 없음(None). 지어내거나 고쳐 적지 않는다(받은 그대로 · 공백만 정리)."""
    s = " ".join(_CONTACT_BAD.sub("", str(v or "")).split())
    if not s:
        return None
    if len(s) > 30:
        raise ApiError("bad_request", "연락처는 30자까지입니다", {"field": "contact"})
    return s


def check_password(pw: str, login: str = "") -> None:
    if not isinstance(pw, str) or len(pw) < 10:
        raise ApiError("bad_request", "비밀번호는 10자 이상입니다", {"field": "password"})
    if len(pw) > 128:
        raise ApiError("bad_request", "비밀번호는 128자까지입니다", {"field": "password"})
    if not (re.search(r"[A-Za-z]", pw) and re.search(r"\d", pw)):
        raise ApiError("bad_request", "영문과 숫자를 함께 넣어 주세요", {"field": "password"})
    local = (login or "").split("@")[0].lower()
    if len(local) >= 4 and local in pw.lower():
        raise ApiError("bad_request", "메일 아이디를 비밀번호에 넣지 마세요", {"field": "password"})


def mask(login: str) -> str:
    """te**@lx.or.kr — 앞 두 글자만 보이고 나머지는 가린다(짧으면 한 글자). 옛 아이디(메일 아님)도 같은 방식."""
    local, at, dom = (login or "").partition("@")
    keep = 2 if len(local) > 3 else 1
    return local[:keep] + "*" * max(2, len(local) - keep) + (at + dom if at else "")


def temp_password() -> str:
    while True:
        s = "".join(secrets.choice(TEMP_ALPHA) for _ in range(12))
        if re.search(r"[A-Z]", s) and re.search(r"\d", s):
            return f"{s[:4]}-{s[4:8]}-{s[8:]}"


async def _tenant_ok(conn, tid) -> dict:
    if not tid or not isinstance(tid, str):
        raise ApiError("bad_request", "기관을 고르세요", {"field": "tenant_id"})
    r = await conn.fetchrow("SELECT id, name FROM tenants WHERE id=$1 AND status='active' AND kind='user' AND id <> 'lx-demo'", tid)
    if not r:
        raise ApiError("bad_request", "기관을 고르세요", {"field": "tenant_id"})
    return r


def _tname(v) -> str:
    return ((v or {}).get("ko") or (v or {}).get("en") or "") if isinstance(v, dict) else (v or "")


async def _log(conn, actor: str, realm: str, action: str, subject: str, after: dict):
    await conn.execute("INSERT INTO audit_log(actor, realm, action, subject, after) VALUES ($1,$2,$3,$4,$5)", actor, realm, action, subject, after)


async def _purge(conn):
    """반려된 신청은 30일 뒤 지운다(동의 안내 '보관 기간'과 같은 값) · 끝난 신청의 비밀번호 해시는 남기지 않는다."""
    await conn.execute("DELETE FROM signup_requests WHERE state='rejected' AND decided_at < now() - $1::interval", KEEP_REJECTED)
    await conn.execute("UPDATE signup_requests SET pw_hash=NULL WHERE state <> 'pending' AND pw_hash IS NOT NULL")


# ── 로그인 없이 — 가입 신청 · 아이디 찾기 · 비밀번호 찾기 ─────────────────────────────────
@router.post("/accounts/signup", status_code=201)
async def signup(body: dict, request: Request):
    site, realm = _site(request, body)
    if site == "admin":
        raise ApiError("forbidden", "LX 관리자 계정은 LX 관리자가 만듭니다")
    await _limit("signup", request)
    name = _text(body.get("name"), 40, "이름", "name")
    login = _mail(body.get("login"))
    if realm == "lx" and LX_MAIL_DOMAINS and login.rpartition("@")[2] not in LX_MAIL_DOMAINS:
        raise ApiError("bad_request", f"LX 직원은 @{LX_MAIL_DOMAINS[0]} 메일로 신청합니다", {"field": "login"})
    pw = body.get("password") or ""
    check_password(pw, login)
    if body.get("password2") is not None and body.get("password2") != pw:
        raise ApiError("bad_request", "두 비밀번호가 서로 다릅니다", {"field": "password2"})
    dept = _text(body.get("dept"), 60, "부서", "dept")
    contact = _contact(body.get("contact")) if realm == "tenant" else None      # 기관 사용자만(LX 직원 연락처는 받지 않는다)
    if body.get("consent") is not True:
        raise ApiError("bad_request", "개인정보 수집·이용에 동의해야 신청할 수 있습니다", {"field": "consent"})
    pw_hash = await run_in_threadpool(auth.hash_password, pw)
    async with db(realm="lx") as conn:
        tid = None
        if realm == "tenant":
            tid = (await _tenant_ok(conn, body.get("tenant_id")))["id"]
            used = await conn.fetchval("SELECT 1 FROM tenant_users WHERE tenant_id=$1 AND lower(login)=$2", tid, login)
        else:
            dept = await dept_pick(conn, dept)          # LX 부서 목록에 있는 부서면 목록 이름으로(S-21) · 없으면 적은 그대로(지사 등 직접 적기)
            used = await conn.fetchval("SELECT 1 FROM lx_users WHERE lower(login)=$1", login)
        used = used or await conn.fetchval("SELECT 1 FROM signup_requests WHERE state='pending' AND realm=$1 AND coalesce(tenant_id,'')=coalesce($2,'') "
                                           "AND lower(login)=$3", realm, tid, login)
        if used:
            raise ApiError("conflict", "이미 가입했거나 신청한 메일 주소입니다", {"field": "login"}, 409)
        rid = "sr_" + secrets.token_hex(8)
        await conn.execute("INSERT INTO signup_requests(id, realm, tenant_id, login, name, dept, pw_hash, consent_at, contact) VALUES ($1,$2,$3,$4,$5,$6,$7,now(),$8)",
                           rid, realm, tid, login, name, dept, pw_hash, contact)
        await _log(conn, "public", "public", "account.signup.request", login,
                   {"request": rid, "realm": realm, "tenant_id": tid, "name": name, "dept": dept, "contact": bool(contact), "ip": client_ip(request)})
        await _purge(conn)
    return {"ok": True, "state": "pending", "at": now_iso()}


@router.post("/accounts/find-id")
async def find_id(body: dict, request: Request):
    site, realm = _site(request, body)
    await _limit("find", request)
    name = _text(body.get("name"), 40, "이름", "name")
    pl = await pool()
    if realm == "tenant":
        async with db(realm="lx") as conn:
            tid = (await _tenant_ok(conn, body.get("tenant_id")))["id"]
        rows = await pl.fetch("SELECT login FROM tenant_users WHERE tenant_id=$1 AND name=$2 AND status IN ('active','locked') ORDER BY login LIMIT 5", tid, name)
    else:
        rows = await pl.fetch("SELECT login FROM lx_users WHERE name=$1 AND status IN ('active','locked') ORDER BY login LIMIT 5", name)
    return {"items": [mask(r["login"]) for r in rows], "at": now_iso()}


@router.post("/accounts/reset-request")
async def reset_request(body: dict, request: Request):
    site, realm = _site(request, body)
    await _limit("reset", request)
    login = str(body.get("login") or "").strip()
    if not login or len(login) > 254:
        raise ApiError("bad_request", "메일 주소를 적어 주세요", {"field": "login"})
    async with db(realm="lx") as conn:
        tid = None
        if realm == "tenant":
            tid = (await _tenant_ok(conn, body.get("tenant_id")))["id"]
            u = await conn.fetchrow("SELECT id, login, name FROM tenant_users WHERE tenant_id=$1 AND lower(login)=lower($2) AND status IN ('active','locked')", tid, login)
        else:
            u = await conn.fetchrow("SELECT id, login, name FROM lx_users WHERE lower(login)=lower($1) AND status IN ('active','locked')", login)
        if u and not await conn.fetchval("SELECT 1 FROM reset_requests WHERE state='pending' AND realm=$1 AND user_id=$2", realm, u["id"]):
            rid = "rr_" + secrets.token_hex(8)
            await conn.execute("INSERT INTO reset_requests(id, realm, tenant_id, user_id, login) VALUES ($1,$2,$3,$4,$5)", rid, realm, tid, u["id"], u["login"])
            await _log(conn, "public", "public", "account.reset.request", u["login"],
                       {"request": rid, "realm": realm, "tenant_id": tid, "name": u["name"], "ip": client_ip(request)})
    return {"ok": True, "at": now_iso()}           # 있든 없든 같은 답


@router.post("/auth/password/change")
async def password_change(body: dict, request: Request):
    """임시 비밀번호로 들어온 사람이 새 비밀번호를 정한다 — 로그인이 준 바꾸기 표(15분 · 한 번)만 받는다. 성공하면 로그인과 같은 답(세션)."""
    tok = str(body.get("change_token") or "")
    if not tok.startswith("lxc_"):
        raise ApiError("unauthorized", "다시 로그인해 주세요")
    await _limit("change", request)
    r = await redis()
    key = auth.CHANGE_KEY + auth.token_hash(tok)
    raw = await r.get(key)
    if not raw:
        raise ApiError("unauthorized", "시간이 지났습니다. 임시 비밀번호로 다시 로그인해 주세요")
    c = json.loads(raw)
    pw = body.get("password") or ""
    check_password(pw, c.get("login") or "")
    if body.get("password2") is not None and body.get("password2") != pw:
        raise ApiError("bad_request", "두 비밀번호가 서로 다릅니다", {"field": "password2"})
    table = "lx_users" if c["realm"] == "lx" else "tenant_users"
    pl = await pool()
    u = await pl.fetchrow(f"SELECT id, pw_hash, role, status, name, must_change FROM {table} WHERE id=$1", c["user_id"])
    if not u or u["status"] != "active" or not u["must_change"]:
        await r.delete(key)
        raise ApiError("unauthorized", "다시 로그인해 주세요")
    try:
        same = await run_in_threadpool(auth.ph.verify, u["pw_hash"], pw)
    except Exception:
        same = False
    if same:
        raise ApiError("bad_request", "임시 비밀번호와 다른 비밀번호를 정하세요", {"field": "password"})
    h = await run_in_threadpool(auth.hash_password, pw)
    await pl.execute(f"UPDATE {table} SET pw_hash=$2, must_change=false, temp_pw_at=NULL WHERE id=$1", u["id"], h)
    await r.delete(key)
    async with db(realm="lx") as conn:
        await _log(conn, u["id"], c["realm"], "account.password.change", c.get("login") or "", {"realm": c["realm"], "tenant_id": c.get("tenant_id"), "name": u["name"]})
    return await auth.open_session(c["realm"], u, c.get("tenant_id"), c.get("site"), c.get("login") or "")


# ── 관리자 — 관할 ───────────────────────────────────────────────────────────────────────
def _who(request: Request) -> Principal:
    p = require(principal(request))
    if p.is_admin or (p.realm == "tenant" and p.role == "manager"):
        return p
    raise ApiError("forbidden", "계정을 관리할 권한이 없습니다")


def _mine(p: Principal, realm: str, user_id: str) -> bool:
    return p.realm == realm and p.user_id == user_id


def _in_scope(p: Principal, realm: str, tenant_id: str | None) -> bool:
    """LX 관리자 = 전부 · 기관 관리자 = 자기 기관 계정만(LX 계정 0)."""
    return p.is_admin or (realm == "tenant" and tenant_id == p.tenant_id)


def _can_decide_signup(p: Principal, realm: str) -> bool:
    """가입 신청을 승인 · 반려할 수 있나 — LX 직원 신청 = LX 관리자 · 기관 신청 = 그 기관 관리자만(LX 관리자는 보기만 · 원칙 72)."""
    return (realm == "lx" and p.is_admin) or (realm == "tenant" and p.realm == "tenant" and p.role == "manager")


async def _tenant_names(conn) -> dict:
    return {r["id"]: _tname(r["name"]) for r in await conn.fetch("SELECT id, name FROM tenants")}


async def _last_login(conn, ids: list[str]) -> dict:
    if not ids:
        return {}
    rows = await conn.fetch("SELECT actor, max(at) AS at FROM audit_log WHERE action='login' AND actor = ANY($1::text[]) GROUP BY actor", ids)
    return {r["actor"]: r["at"] for r in rows}


@router.get("/accounts/summary")
async def summary(request: Request):
    p = _who(request)
    pl = await pool()
    w = "" if p.is_admin else " AND realm='tenant' AND tenant_id=$1"
    a = [] if p.is_admin else [p.tenant_id]
    # 가입 신청 수 = 내가 승인 · 반려할 것(LX 관리자 = LX 직원 신청 · 기관 관리자 = 자기 기관 신청). LX 관리자가 보기만 하는 기관 신청은 따로(signup_view)
    s = await pl.fetchval("SELECT count(*) FROM signup_requests WHERE state='pending'" + (" AND realm='lx'" if p.is_admin else w), *a)
    sv = await pl.fetchval("SELECT count(*) FROM signup_requests WHERE state='pending' AND realm='tenant'") if p.is_admin else 0
    r = await pl.fetchval("SELECT count(*) FROM reset_requests WHERE state='pending'" + w, *a)
    st = 0
    if p.is_admin:                                     # 저장 용량 늘리기 요청(S-19 · LX 계정만)
        async with db(realm="lx") as conn:
            st = await conn.fetchval("SELECT count(*) FROM storage_requests WHERE state='pending'")
    return {"counts": {"signup": s, "reset": r, "signup_view": sv, "storage": st}, "at": now_iso()}


@router.get("/accounts/requests")
async def list_requests(request: Request, kind: str = "signup", state: str = "pending"):
    p = _who(request)
    if kind == "storage" and state in ("pending", "all"):
        return await _storage_requests(p, state)
    if kind not in ("signup", "reset") or state not in ("pending", "all"):
        raise ApiError("bad_request", "kind 는 signup | reset | storage · state 는 pending | all")
    pl = await pool()
    table = "signup_requests" if kind == "signup" else "reset_requests"
    q = f"SELECT * FROM {table} WHERE ($1::text = 'all' OR state = 'pending')"
    args = [state]
    if not p.is_admin:
        q += " AND realm='tenant' AND tenant_id=$2"
        args.append(p.tenant_id)
    rows = await pl.fetch(q + " ORDER BY created_at DESC LIMIT 200", *args)
    async with db(realm="lx") as conn:
        tn = await _tenant_names(conn)
        users = {}
        if kind == "reset":
            for realm_ in ("lx", "tenant"):
                ids = [r["user_id"] for r in rows if r["realm"] == realm_]
                if ids:
                    tbl = "lx_users" if realm_ == "lx" else "tenant_users"
                    for u in await conn.fetch(f"SELECT id, name, dept, role, status FROM {tbl} WHERE id = ANY($1::text[])", ids):
                        users[(realm_, u["id"])] = u
            last = await _last_login(conn, [r["user_id"] for r in rows])
    items = []
    for r in rows:
        it = {"id": r["id"], "kind": kind, "realm": r["realm"], "tenant_id": r["tenant_id"], "org": "LX" if r["realm"] == "lx" else tn.get(r["tenant_id"], ""),
              "login": r["login"], "state": r["state"], "reason": r["reason"], "created_at": _iso(r["created_at"]),
              "decided_name": r["decided_name"], "decided_at": _iso(r["decided_at"])}
        if kind == "signup":
            it.update({"name": r["name"], "dept": r["dept"], "consent_at": _iso(r["consent_at"]), "mine": False,
                       "can_decide": _can_decide_signup(p, r["realm"])})   # LX 관리자가 보는 기관 신청 = 보기만(원칙 72)
        else:
            u = users.get((r["realm"], r["user_id"]))
            it.update({"name": u["name"] if u else "", "dept": u["dept"] if u else None, "role_ko": ROLE_KO.get((r["realm"], u["role"])) if u else None,
                       "status": u["status"] if u else "gone", "last_login": _iso(last.get(r["user_id"])), "mine": _mine(p, r["realm"], r["user_id"])})
        items.append(it)
    return {"items": items, "at": now_iso()}


def _decision(body: dict, ok: tuple[str, ...]) -> tuple[str, str]:
    d = body.get("decision")
    if d not in ok:
        raise ApiError("bad_request", "decision 은 " + " | ".join(ok))
    reason = " ".join(str(body.get("reason") or "").split())[:200]
    if d == "reject" and not reason:
        raise ApiError("reason_required", "거절 사유를 적어 주세요", status=400)
    return d, reason


async def _drop_sessions(conn, realm: str, user_id: str):
    await conn.execute("DELETE FROM sessions WHERE realm=$1 AND user_id=$2", realm, user_id)


async def _forget_fails(realm: str, tenant_id: str | None, login: str):
    try:
        r = await redis()
        await r.delete(auth.fail_key(realm, tenant_id, login))
    except Exception:
        pass


@router.post("/accounts/signup/{rid}/decide")
async def decide_signup(rid: str, body: dict, request: Request):
    p = _who(request)
    d, reason = _decision(body, ("approve", "reject"))
    async with db(realm="lx") as conn:
        r = await conn.fetchrow("SELECT * FROM signup_requests WHERE id=$1 FOR UPDATE", rid)
        if not r or not _in_scope(p, r["realm"], r["tenant_id"]):
            raise ApiError("not_found", "신청이 없습니다")
        if not _can_decide_signup(p, r["realm"]):     # LX 관리자는 기관 가입 승인에 관여하지 않는다(보기 · 잠금 · 지원만 — 원칙 72)
            raise ApiError("tenant_signup", TENANT_SIGNUP_MSG, status=403)
        if r["state"] != "pending":
            raise ApiError("conflict", "이미 처리한 신청입니다", {"state": r["state"], "by": r["decided_name"]}, 409)
        uid = None
        if d == "approve":
            if r["realm"] == "lx":
                if await conn.fetchval("SELECT 1 FROM lx_users WHERE lower(login)=lower($1)", r["login"]):
                    raise ApiError("conflict", "이미 있는 아이디입니다", None, 409)
                uid = "u_" + secrets.token_hex(6)
                await conn.execute("INSERT INTO lx_users(id, login, pw_hash, role, status, name, dept) VALUES ($1,$2,$3,$4,'active',$5,$6)",
                                   uid, r["login"], r["pw_hash"], NEW_ROLE["lx"], r["name"], r["dept"])
            else:
                if await conn.fetchval("SELECT 1 FROM tenant_users WHERE tenant_id=$1 AND lower(login)=lower($2)", r["tenant_id"], r["login"]):
                    raise ApiError("conflict", "이미 있는 아이디입니다", None, 409)
                uid = "u_" + secrets.token_hex(6)
                await conn.execute("INSERT INTO tenant_users(id, tenant_id, login, pw_hash, role, status, name, dept, contact) VALUES ($1,$2,$3,$4,$5,'active',$6,$7,$8)",
                                   uid, r["tenant_id"], r["login"], r["pw_hash"], NEW_ROLE["tenant"], r["name"], r["dept"], r["contact"])
        await conn.execute("UPDATE signup_requests SET state=$2, reason=$3, user_id=$4, pw_hash=NULL, decided_by=$5, decided_realm=$6, decided_name=$7, decided_at=now() "
                           "WHERE id=$1", rid, "approved" if d == "approve" else "rejected", reason or None, uid, p.user_id, p.realm, p.name)
        await _log(conn, p.user_id, p.realm, f"account.signup.{d}", r["login"],
                   {"request": rid, "realm": r["realm"], "tenant_id": r["tenant_id"], "name": r["name"], "reason": reason or None,
                    "role": NEW_ROLE[r["realm"]] if d == "approve" else None})
    return {"ok": True, "state": "approved" if d == "approve" else "rejected", "at": now_iso()}


async def _issue_temp(conn, p: Principal, realm: str, u, tenant_id: str | None, action: str, extra: dict) -> str:
    tp = temp_password()
    h = await run_in_threadpool(auth.hash_password, tp)
    table = "lx_users" if realm == "lx" else "tenant_users"
    await conn.execute(f"UPDATE {table} SET pw_hash=$2, must_change=true, temp_pw_at=now() WHERE id=$1", u["id"], h)
    await _drop_sessions(conn, realm, u["id"])
    await conn.execute("UPDATE reset_requests SET state='issued', decided_by=$3, decided_realm=$4, decided_name=$5, decided_at=now() "
                       "WHERE state='pending' AND realm=$1 AND user_id=$2", realm, u["id"], p.user_id, p.realm, p.name)
    await _log(conn, p.user_id, p.realm, action, u["login"], {"realm": realm, "tenant_id": tenant_id, "name": u["name"], **extra})
    await _forget_fails(realm, tenant_id, u["login"])
    return tp


@router.post("/accounts/reset/{rid}/decide")
async def decide_reset(rid: str, body: dict, request: Request):
    p = _who(request)
    d, reason = _decision(body, ("issue", "reject"))
    async with db(realm="lx") as conn:
        r = await conn.fetchrow("SELECT * FROM reset_requests WHERE id=$1 FOR UPDATE", rid)
        if not r or not _in_scope(p, r["realm"], r["tenant_id"]):
            raise ApiError("not_found", "요청이 없습니다")
        if r["state"] != "pending":
            raise ApiError("conflict", "이미 처리한 요청입니다", {"state": r["state"], "by": r["decided_name"]}, 409)
        if _mine(p, r["realm"], r["user_id"]):
            raise ApiError("self_account", "내 계정의 요청은 다른 관리자가 처리합니다", status=409)
        if d == "reject":
            await conn.execute("UPDATE reset_requests SET state='rejected', reason=$2, decided_by=$3, decided_realm=$4, decided_name=$5, decided_at=now() WHERE id=$1",
                               rid, reason, p.user_id, p.realm, p.name)
            await _log(conn, p.user_id, p.realm, "account.reset.reject", r["login"], {"request": rid, "realm": r["realm"], "tenant_id": r["tenant_id"], "reason": reason})
            return {"ok": True, "state": "rejected", "at": now_iso()}
        table = "lx_users" if r["realm"] == "lx" else "tenant_users"
        u = await conn.fetchrow(f"SELECT id, login, name, status FROM {table} WHERE id=$1", r["user_id"])
        if not u:
            raise ApiError("not_found", "계정이 없습니다")
        if u["status"] == "disabled":
            raise ApiError("disabled_account", "사용 중지된 계정입니다", status=409)
        tp = await _issue_temp(conn, p, r["realm"], u, r["tenant_id"], "account.reset.issue", {"request": rid})
    return {"ok": True, "state": "issued", "temp_password": tp, "locked": u["status"] == "locked", "at": now_iso()}


# ── 관리자 — 계정 목록 · 잠금 · 역할 · 임시 비밀번호 · 기록 ─────────────────────────────────
@router.get("/accounts/users")
async def list_users(request: Request, realm: str | None = None, tenant_id: str | None = None):
    p = _who(request)
    out = []
    extra = {}
    async with db(realm="lx") as conn:
        tn = await _tenant_names(conn)
        if p.is_admin and realm in (None, "", "lx") and not tenant_id:
            from .approvals import solo_admin
            from .projects import storage_all, storage_default
            stor = await storage_all(conn)
            keys = await _dept_keys(conn)
            for u in await conn.fetch("SELECT id, login, role, status, name, dept, must_change, created_at, lock_until FROM lx_users "
                                      "ORDER BY (status = 'disabled'), login"):   # 사용 중지(옛 아이디)는 뒤로
                out.append({"realm": "lx", "tenant_id": None, "org": "LX", **_user(u, "lx"), "storage": _storage_out(stor.get(u["id"])),
                            "dept_listed": (_dkey(u["dept"]) in keys) if keys and (u["dept"] or "").strip() else None})
            extra = {"solo": await solo_admin(conn, p), "storage_default": env(await storage_default(conn), "GB", "recorded", "기본 할당(lx_settings)"),
                     "depts": env(len(set(keys.values())), "count", "recorded", "LX 부서 목록(lx_depts · 최상위 뺌)")}
        if realm in (None, "", "tenant"):
            tid = tenant_id if p.is_admin else p.tenant_id
            rows = await conn.fetch("SELECT id, tenant_id, login, role, status, name, dept, must_change, created_at, lock_until FROM tenant_users "
                                    "WHERE ($1::text IS NULL OR tenant_id=$1) ORDER BY tenant_id, (status = 'disabled'), login", tid)
            for u in rows:
                if u["tenant_id"] == "lx-demo" and not p.is_admin:
                    continue
                out.append({"realm": "tenant", "tenant_id": u["tenant_id"], "org": tn.get(u["tenant_id"], ""), **_user(u, "tenant")})
        last = await _last_login(conn, [x["id"] for x in out])
    for x in out:
        x["last_login"] = _iso(last.get(x["id"]))
        x["mine"] = _mine(p, x["realm"], x["id"])
    orgs = [{"id": k, "name": v} for k, v in tn.items() if k not in ("lx", "lx-demo")] if p.is_admin else [{"id": p.tenant_id, "name": tn.get(p.tenant_id, "")}]
    return {"items": out, "orgs": orgs, "roles": {k: [{"id": r, "label": ROLE_KO[(k, r)]} for r in v] for k, v in ROLES.items()}, **extra, "at": now_iso()}


def _user(u, realm: str) -> dict:
    held = bool(u["lock_until"] and u["lock_until"] > dt.datetime.now(dt.timezone.utc))      # 비밀번호 5번 틀려 잠시 잠김
    return {"id": u["id"], "login": u["login"], "name": u["name"], "dept": u["dept"], "role": u["role"], "role_ko": ROLE_KO.get((realm, u["role"]), u["role"]),
            "status": u["status"], "temp_locked": held, "lock_until": _iso(u["lock_until"]) if held else None,
            "must_change": bool(u["must_change"]), "created_at": _iso(u["created_at"])}


async def _target(conn, p: Principal, realm: str, uid: str):
    if realm not in ROLES:
        raise ApiError("not_found", "계정이 없습니다")
    if realm == "lx":
        u = await conn.fetchrow("SELECT id, login, role, status, name, NULL::text AS tenant_id FROM lx_users WHERE id=$1 FOR UPDATE", uid)
    else:
        u = await conn.fetchrow("SELECT id, login, role, status, name, tenant_id FROM tenant_users WHERE id=$1 FOR UPDATE", uid)
    if not u or not _in_scope(p, realm, u["tenant_id"]):
        raise ApiError("not_found", "계정이 없습니다")
    if _mine(p, realm, uid):
        raise ApiError("self_account", "내 계정은 다른 관리자가 바꿉니다", status=409)
    if u["status"] == "disabled":                      # 옛 아이디 — 메일 아이디로 옮긴 뒤 사용 중지(되살리지 않는다)
        raise ApiError("disabled_account", "사용 중지된 계정입니다", status=409)
    return u


@router.post("/accounts/users/{realm}/{uid}/lock")
async def lock_user(realm: str, uid: str, body: dict, request: Request):
    p = _who(request)
    locked = body.get("locked")
    if not isinstance(locked, bool):
        raise ApiError("bad_request", "locked 는 true | false")
    reason = " ".join(str(body.get("reason") or "").split())[:200] or None
    table = "lx_users" if realm == "lx" else "tenant_users"
    async with db(realm="lx") as conn:
        u = await _target(conn, p, realm, uid)
        if locked:
            await conn.execute(f"UPDATE {table} SET status='locked' WHERE id=$1", uid)
        else:                                            # 관리자 잠금 · 5번 틀림 10분 잠금 둘 다 푼다
            await conn.execute(f"UPDATE {table} SET status='active', lock_until=NULL WHERE id=$1", uid)
        if locked:
            await _drop_sessions(conn, realm, uid)
        await _log(conn, p.user_id, p.realm, "account.lock" if locked else "account.unlock", u["login"],
                   {"realm": realm, "tenant_id": u["tenant_id"], "name": u["name"], "reason": reason})
    if not locked:
        await _forget_fails(realm, u["tenant_id"], u["login"])
    return {"ok": True, "status": "locked" if locked else "active", "at": now_iso()}


@router.post("/accounts/users/{realm}/{uid}/role")
async def set_role(realm: str, uid: str, body: dict, request: Request):
    p = _who(request)
    role = body.get("role")
    if realm not in ROLES or role not in ROLES[realm]:
        raise ApiError("bad_request", "역할을 고르세요")
    table = "lx_users" if realm == "lx" else "tenant_users"
    async with db(realm="lx") as conn:
        u = await _target(conn, p, realm, uid)
        if u["role"] == role:
            return {"ok": True, "role": role, "at": now_iso()}
        await conn.execute(f"UPDATE {table} SET role=$2 WHERE id=$1", uid, role)
        await _drop_sessions(conn, realm, uid)          # 새 역할은 다시 로그인하면 적용 — 옛 역할 세션이 남지 않게
        await _log(conn, p.user_id, p.realm, "account.role", u["login"],
                   {"realm": realm, "tenant_id": u["tenant_id"], "name": u["name"], "from": ROLE_KO.get((realm, u["role"])), "to": ROLE_KO.get((realm, role))})
    return {"ok": True, "role": role, "at": now_iso()}


@router.post("/accounts/users/{realm}/{uid}/temp-password")
async def temp_pw(realm: str, uid: str, request: Request):
    p = _who(request)
    async with db(realm="lx") as conn:
        u = await _target(conn, p, realm, uid)
        tp = await _issue_temp(conn, p, realm, u, u["tenant_id"], "account.temp.issue", {})
    return {"ok": True, "temp_password": tp, "locked": u["status"] == "locked", "at": now_iso()}


@router.get("/accounts/log")
async def account_log(request: Request):
    p = _who(request)
    pl = await pool()
    if p.is_admin:
        rows = await pl.fetch("SELECT id, actor, realm, action, subject, after, at FROM audit_log WHERE action LIKE 'account.%' ORDER BY id DESC LIMIT 100")
    else:
        rows = await pl.fetch("SELECT id, actor, realm, action, subject, after, at FROM audit_log WHERE action LIKE 'account.%' "
                              "AND after->>'tenant_id' = $1 ORDER BY id DESC LIMIT 100", p.tenant_id)
    async with db(realm="lx") as conn:
        names = {r["id"]: r["name"] for r in await conn.fetch("SELECT id, name FROM lx_users")}
        names.update({r["id"]: r["name"] for r in await conn.fetch("SELECT id, name FROM tenant_users")})
        tn = await _tenant_names(conn)
    items = []
    for r in rows:
        a = r["after"] or {}
        who = "신청한 사람" if r["actor"] == "public" else "자동" if r["actor"] == "system" else names.get(r["actor"]) or "관리자"
        items.append({"at": _iso(r["at"]), "who": who, "action": r["action"], "action_ko": ACTION_KO.get(r["action"], r["action"]), "subject": r["subject"],
                      "name": a.get("name"), "org": "LX" if a.get("realm") == "lx" else tn.get(a.get("tenant_id"), ""), "reason": a.get("reason"),
                      "to": a.get("to")})
    return {"items": items, "at": now_iso()}


@router.get("/accounts/failures")
async def login_failures(request: Request):
    """실패한 로그인 — LX 관리자 = 전부 · 기관 관리자 = 자기 기관 입구의 실패만. 이 PC 안 요청은 접속 주소 대신 '이 PC'."""
    p = _who(request)
    pl = await pool()
    if p.is_admin:
        rows = await pl.fetch("SELECT at, realm, tenant_id, login, site, ip, reason FROM login_failures ORDER BY id DESC LIMIT 200")
    else:
        rows = await pl.fetch("SELECT at, realm, tenant_id, login, site, ip, reason FROM login_failures WHERE realm='tenant' AND tenant_id=$1 "
                              "ORDER BY id DESC LIMIT 200", p.tenant_id)
    async with db(realm="lx") as conn:
        tn = await _tenant_names(conn)
    items = [{"at": _iso(r["at"]), "login": r["login"], "org": "LX" if r["realm"] == "lx" else tn.get(r["tenant_id"], ""),
              "site_ko": SITE_KO.get(r["site"] or "", "이 PC"), "ip": "이 PC" if (r["ip"] or "") in ("127.0.0.1", "::1", "") else r["ip"],
              "reason": r["reason"], "reason_ko": FAIL_KO.get(r["reason"], r["reason"])} for r in rows]
    return {"items": items, "at": now_iso()}


# ── LX 요청함 — 기관 요청을 보낸 사람(기관-8 ⓐ · 원칙 102 · 63 · 72) ───────────────────────────────────
SENDER_MAX = 200


def _ids(v: str | None) -> list[str]:
    return [x for x in dict.fromkeys(s.strip() for s in str(v or "").split(",")) if re.fullmatch(r"[a-z]{2}_[0-9a-f]{8,24}", x)][:SENDER_MAX]


@router.get("/accounts/senders")
async def senders(request: Request, review: str | None = None, req: str | None = None):
    """기관 요청(검토 요청 · 분석 요청)을 보낸 사람 — 이름 · 부서 · 역할 · 기관 + 연락처(선택).
    보는 사람: LX 관리자 = 모두 · LX 직원 = 내가 받은 것(검토 요청 recipient · 분석 요청 lead_user = 그 서비스 담당)만 — 아니면 그 줄은 빈다.
    연락처는 그 사람이 스스로 적은 경우만(가입 신청 '연락처(선택)') · 기관 계정 · 영업에는 열지 않는다(403).
    주소 이름 request 는 요청 객체와 겹쳐 ?request= 를 req 로도 받는다."""
    p = require(principal(request))
    if not (p.realm == "lx" and p.role in ("admin", "staff")):
        raise ApiError("forbidden", "LX 직원 · LX 관리자만 봅니다")
    rv = _ids(review)
    rq = _ids(req or request.query_params.get("request"))
    out = {"reviews": {}, "requests": {}, "at": now_iso()}
    if not rv and not rq:
        return out
    async with db(realm="lx") as conn:
        rows = []
        if rv:
            rows += [("reviews", r) for r in await conn.fetch(
                "SELECT id, tenant_id, sender_id AS uid, recipient_id AS owner FROM feedback WHERE kind='review' AND id = ANY($1::text[])", rv)]
        if rq:
            rows += [("requests", r) for r in await conn.fetch(
                "SELECT id, tenant_id, requested_by AS uid, lead_user AS owner FROM analysis_requests WHERE id = ANY($1::text[])", rq)]
        mine = [(k, r) for k, r in rows if p.is_admin or (r["owner"] and r["owner"] == p.user_id)]
        uids = list({r["uid"] for _, r in mine if r["uid"]})
        users = {u["id"]: u for u in await conn.fetch(
            "SELECT id, tenant_id, name, dept, role, contact FROM tenant_users WHERE id = ANY($1::text[])", uids)} if uids else {}
        tn = await _tenant_names(conn)
    for k, r in mine:
        u = users.get(r["uid"])
        if not u or u["tenant_id"] != r["tenant_id"]:
            out[k][r["id"]] = {"name": None, "dept": None, "role_ko": None, "org": _short(tn.get(r["tenant_id"], "")), "contact": None}
            continue
        out[k][r["id"]] = {"name": u["name"] or None, "dept": (u["dept"] or "").strip() or None, "role_ko": ROLE_KO.get(("tenant", u["role"])),
                           "org": _short(tn.get(r["tenant_id"], "")),
                           "contact": ((u["contact"] or "").strip() or None) if (p.is_admin or r["owner"] == p.user_id) else None}
    return out


def _short(name: str) -> str:
    """기관 짧은 이름 — '전북특별자치도 남원시' → '남원시'(요청함 · 알림과 같은 말)."""
    w = str(name or "").split()
    return w[-1] if w else ""


# ── 내 정보(본인이 고친다 — 확인 17차 P-5 ⓐ · 원칙 105 · 121) ─────────────────────────────────
PROFILE_MAX = {"name": 40, "dept": 60, "contact": 30}
CONTACT_OK = re.compile(r"^[0-9가-힣+\-().\s#~/]{2,30}$")      # 내선 · 휴대전화 · 대표번호(예: 내선 1234 · 010-0000-0000) — 숫자가 하나는 있어야 한다


def _me_lx(request: Request) -> Principal:
    p = require(principal(request))
    if p.realm != "lx" or p.role not in ROLES["lx"]:
        raise ApiError("forbidden", "내 정보 고치기는 LX 계정에서 합니다")
    return p


def _storage_out(st: dict | None) -> dict | None:
    """저장 용량(projects.lead_storage · storage_all 한 출처) → 화면 값. 할당 없음이면 quota_gb · pct 가 null(지어내지 않는다)."""
    if st is None:
        return None
    return {"quota_gb": env(st["quota_gb"], "GB", "recorded", "할당(사람마다 정한 값 → 없으면 기본 할당 → 없으면 할당 없음)"),
            "quota_own": st["quota_own"],
            "used": env(st["bytes"], "bytes", "measured", "프로젝트장인 프로젝트의 저장 공간 합(학습데이터 파일 + 올린 파일)"),
            "projects": env(st["projects"], "count", "recorded", "프로젝트장인 프로젝트(진행 중 · 보관)"),
            "pct": env(st["pct"], "%", "measured", "할당 가운데 쓴 비율"), "warn": st["warn"]}


async def _profile(conn, p: Principal) -> dict:
    from .projects import lead_storage
    u = await conn.fetchrow("SELECT login, name, dept, contact, role FROM lx_users WHERE id=$1", p.user_id)
    if not u:
        raise ApiError("not_found", "계정이 없습니다")
    last = await conn.fetchval("SELECT max(at) FROM audit_log WHERE action='account.profile' AND actor=$1", p.user_id)
    st = _storage_out(await lead_storage(conn, p.user_id))
    pend = await conn.fetchrow("SELECT want_gb, why, created_at FROM storage_requests WHERE user_id=$1 AND state='pending'", p.user_id)
    done = await conn.fetchrow("SELECT state, want_gb, reason, decided_at FROM storage_requests WHERE user_id=$1 AND state <> 'pending' "
                               "ORDER BY decided_at DESC NULLS LAST LIMIT 1", p.user_id)
    st["pending"] = {"want_gb": env(float(pend["want_gb"]), "GB", "recorded", "원하는 할당"), "why": pend["why"], "at": _iso(pend["created_at"])} if pend else None
    st["last"] = {"state": done["state"], "want_gb": env(float(done["want_gb"]), "GB", "recorded", "원하는 할당"), "reason": done["reason"], "at": _iso(done["decided_at"])} if done else None
    listed = await dept_listed(conn, u["dept"])
    return {"login": u["login"], "name": u["name"] or "", "dept": u["dept"] or "", "contact": u["contact"] or "",
            "role_ko": ROLE_KO.get(("lx", u["role"]), "LX"), "org": "한국국토정보공사", "changed_at": _iso(last),
            "dept_listed": listed,
            "storage": st}


@router.get("/me/profile")
async def my_profile(request: Request):
    p = _me_lx(request)
    async with db(realm="lx") as conn:
        out = await _profile(conn, p)
    return {**out, "as_of": now_iso()}


@router.patch("/me/profile")
async def edit_profile(body: dict, request: Request):
    """이름 · 부서 · 연락처를 본인이 바로 고친다(관리자 승인 없음). 아이디(메일)는 고치지 않는다. 바뀐 칸만 감사 기록에 전 · 후로 남긴다."""
    p = _me_lx(request)
    if "login" in body:
        raise ApiError("bad_request", "아이디(메일)는 바꿀 수 없습니다", {"field": "login"})
    async with db(realm="lx") as conn:
        u = await conn.fetchrow("SELECT login, name, dept, contact FROM lx_users WHERE id=$1 FOR UPDATE", p.user_id)
        if not u:
            raise ApiError("not_found", "계정이 없습니다")
        new = {"name": u["name"], "dept": u["dept"], "contact": u["contact"]}
        if "name" in body:
            new["name"] = _text(body.get("name"), PROFILE_MAX["name"], "이름", "name")
        for k, label in (("dept", "부서"), ("contact", "연락처")):
            if k in body:
                s = " ".join(str(body.get(k) or "").split())
                if len(s) > PROFILE_MAX[k]:
                    raise ApiError("bad_request", f"{label}{_josa(label, '은는')} {PROFILE_MAX[k]}자까지입니다", {"field": k})
                if k == "contact" and s and not (CONTACT_OK.match(s) and re.search(r"\d", s)):
                    raise ApiError("bad_request", "연락처는 번호로 적어 주세요(예: 내선 1234)", {"field": k})
                if k == "dept" and s:            # 목록에 있는 부서면 목록 이름으로 맞춘다 · 없으면 적은 그대로(지사 등 — S-21)
                    s = await dept_pick(conn, s, keep=u["dept"])
                new[k] = s or None
        before = {k: u[k] for k in new if new[k] != u[k]}
        if before:
            await conn.execute("UPDATE lx_users SET name=$2, dept=$3, contact=$4 WHERE id=$1", p.user_id, new["name"], new["dept"], new["contact"])
            await conn.execute("INSERT INTO audit_log(actor, realm, action, subject, before, after) VALUES ($1,'lx','account.profile',$2,$3,$4)",
                               p.user_id, u["login"], before,
                               {"realm": "lx", "tenant_id": None, "name": new["name"], "fields": list(before), **{k: new[k] for k in before}})
        out = await _profile(conn, p)
    return {**out, "changed": list(before), "as_of": now_iso()}


# ── 저장 용량 할당 · 늘리기 요청 · 승인(제안 S-19 확인 · 원칙 66 · 91 · 121) ─────────────────────────────
QUOTA_MAX_GB = 100000          # 입력 확인 상한(100 TB) — 업무 값이 아니라 잘못 친 숫자를 막는 선
WHY_MAX = 120


def _gb(v, *, field: str = "quota_gb", allow_none: bool = False) -> float | None:
    """GB 값(소수 둘째 자리까지 · 10 MB 단위) — 비우면 None(allow_none) · 0 이하 · 숫자 아님 · 상한 넘음은 400."""
    if v is None or (isinstance(v, str) and not v.strip()):
        if allow_none:
            return None
        raise ApiError("bad_request", "용량을 GB 로 적어 주세요", {"field": field})
    try:
        x = float(str(v).replace(",", "").strip())
    except ValueError:
        raise ApiError("bad_request", "용량은 숫자(GB)로 적어 주세요", {"field": field})
    x = round(x, 2) if not math.isnan(x) else x
    if math.isnan(x) or x <= 0 or x > QUOTA_MAX_GB:
        raise ApiError("bad_request", f"용량은 0보다 크고 {QUOTA_MAX_GB:,} GB 까지입니다", {"field": field})
    return x


def gb_word(x) -> str:
    """100 → '100 GB' · 50.5 → '50.5 GB' · None → '할당 없음'"""
    if x is None:
        return "할당 없음"
    x = float(x)
    return f"{int(x):,} GB" if x == int(x) else f"{x:,.2f}".rstrip("0") + " GB"


async def _storage_requests(p: Principal, state: str) -> dict:
    if not p.is_admin:                                 # 기관 관리자 화면은 이 탭이 없다(LX 계정만)
        return {"items": [], "at": now_iso()}
    from .approvals import solo_admin
    from .projects import storage_all
    async with db(realm="lx") as conn:
        rows = await conn.fetch("SELECT s.*, u.name, u.dept, u.status FROM storage_requests s LEFT JOIN lx_users u ON u.id = s.user_id "
                                "WHERE ($1::text = 'all' OR s.state = 'pending') ORDER BY s.created_at DESC LIMIT 200", state)
        stor = await storage_all(conn)
        solo = await solo_admin(conn, p)
    items = []
    for r in rows:
        st = _storage_out(stor.get(r["user_id"]))
        mine = _mine(p, "lx", r["user_id"])
        items.append({"id": r["id"], "kind": "storage", "realm": "lx", "tenant_id": None, "org": "LX", "login": r["login"], "name": r["name"] or "",
                      "dept": r["dept"], "from_gb": env(float(r["from_gb"]) if r["from_gb"] is not None else None, "GB", "recorded", "요청할 때의 할당"),
                      "want_gb": env(float(r["want_gb"]), "GB", "recorded", "원하는 할당"),
                      "why": r["why"], "storage": st, "state": r["state"], "reason": r["reason"], "created_at": _iso(r["created_at"]),
                      "decided_name": r["decided_name"], "decided_at": _iso(r["decided_at"]), "mine": mine,
                      "can_decide": r["state"] == "pending" and (not mine or solo)})
    return {"items": items, "at": now_iso()}


@router.post("/me/storage-request", status_code=201)
async def storage_request(body: dict, request: Request):
    """저장 용량 늘리기 요청 — 원하는 할당(GB) · 이유 한 줄. 할당이 없으면(제한 없음) 요청할 것이 없다(409). 대기 중 요청은 하나."""
    p = _me_lx(request)
    want = _gb(body.get("want_gb"), field="want_gb")
    why = _text(body.get("why"), WHY_MAX, "이유", "why")
    from .projects import lead_storage
    async with db(realm="lx") as conn:
        u = await conn.fetchrow("SELECT login, name FROM lx_users WHERE id=$1", p.user_id)
        if not u:
            raise ApiError("not_found", "계정이 없습니다")
        st = await lead_storage(conn, p.user_id)
        if st["quota_gb"] is None:
            raise ApiError("conflict", "할당이 없어 늘리기 요청이 필요 없습니다", {"why": "no_quota"}, 409)
        if want <= st["quota_gb"]:
            raise ApiError("bad_request", f"지금 할당 {gb_word(st['quota_gb'])}보다 큰 값을 적어 주세요", {"field": "want_gb"})
        if await conn.fetchval("SELECT 1 FROM storage_requests WHERE user_id=$1 AND state='pending'", p.user_id):
            raise ApiError("conflict", "이미 보낸 요청이 있습니다. LX 관리자 확인을 기다려 주세요", {"why": "pending"}, 409)
        rid = "sq_" + secrets.token_hex(8)
        await conn.execute("INSERT INTO storage_requests(id, user_id, login, from_gb, want_gb, why) VALUES ($1,$2,$3,$4,$5,$6)",
                           rid, p.user_id, u["login"], Decimal(str(st["quota_gb"])), Decimal(str(want)), why)
        await _log(conn, p.user_id, "lx", "account.storage.request", u["login"],
                   {"request": rid, "realm": "lx", "tenant_id": None, "name": u["name"], "from": st["quota_gb"], "to": want, "reason": why})
        out = await _profile(conn, p)
    return {**out, "request": rid, "as_of": now_iso()}


@router.post("/accounts/storage/{rid}/decide")
async def decide_storage(rid: str, body: dict, request: Request):
    """늘리기 요청 승인(할당 = 원하는 값) · 반려(사유 필수) → 요청한 사람에게 알림 한 줄(대시보드 · 프로젝트 목록 · 내 정보)."""
    p = _who(request)
    if not p.is_admin:
        raise ApiError("forbidden", "저장 용량 요청은 LX 관리자가 처리합니다")
    d, reason = _decision(body, ("approve", "reject"))
    from .approvals import SOLO_NOTE, solo_admin
    from .projects import lead_storage
    async with db(realm="lx") as conn:
        r = await conn.fetchrow("SELECT * FROM storage_requests WHERE id=$1 FOR UPDATE", rid)
        if not r:
            raise ApiError("not_found", "요청이 없습니다")
        if r["state"] != "pending":
            raise ApiError("conflict", "이미 처리한 요청입니다", {"state": r["state"], "by": r["decided_name"]}, 409)
        solo = False
        if _mine(p, "lx", r["user_id"]):
            solo = await solo_admin(conn, p)
            if not solo:
                raise ApiError("self_account", "내 계정의 요청은 다른 관리자가 처리합니다", status=409)
        u = await conn.fetchrow("SELECT id, login, name, status, storage_quota_gb FROM lx_users WHERE id=$1 FOR UPDATE", r["user_id"])
        if not u:
            raise ApiError("not_found", "계정이 없습니다")
        if u["status"] == "disabled":
            raise ApiError("disabled_account", "사용 중지된 계정입니다", status=409)
        st = await lead_storage(conn, u["id"])
        want = float(r["want_gb"])
        if d == "approve":
            if st["quota_gb"] is not None and st["quota_gb"] >= want:
                raise ApiError("conflict", f"지금 할당({gb_word(st['quota_gb'])})이 이미 원하는 양 이상입니다. 거절하거나 그대로 두세요", status=409)
            await conn.execute("UPDATE lx_users SET storage_quota_gb=$2 WHERE id=$1", u["id"], Decimal(str(want)))
        await conn.execute("UPDATE storage_requests SET state=$2, reason=$3, decided_by=$4, decided_name=$5, decided_at=now() WHERE id=$1",
                           rid, "approved" if d == "approve" else "rejected", reason or None, p.user_id, p.name)
        if u["id"] != p.user_id:                       # 요청한 사람에게 알림(스스로 처리했으면 필요 없다)
            text = f"저장 용량 할당이 {gb_word(want)}로 늘었습니다" if d == "approve" else f"저장 용량 늘리기 요청({gb_word(want)})이 반려되었습니다"
            await conn.execute("INSERT INTO lx_notices(id, user_id, kind, project_id, text, note, by) VALUES ($1,$2,'account.storage',NULL,$3,$4,$5)",
                               "nt_" + secrets.token_hex(6), u["id"], text, reason or None, p.user_id)
        await _log(conn, p.user_id, "lx", f"account.storage.{d}", u["login"],
                   {"request": rid, "realm": "lx", "tenant_id": None, "name": u["name"], "from": st["quota_gb"], "to": want if d == "approve" else None,
                    "reason": reason or None, **({"single_admin": True, "note": SOLO_NOTE} if solo else {})})
    return {"ok": True, "state": "approved" if d == "approve" else "rejected", "at": now_iso()}


@router.post("/accounts/users/{realm}/{uid}/quota")
async def set_quota(realm: str, uid: str, body: dict, request: Request):
    """사람마다 저장 용량 할당 — LX 관리자 · LX 계정만. quota_gb 비움(null) = 기본 할당을 따름. 바꾼 전 · 후는 처리 기록에."""
    p = _who(request)
    if not p.is_admin or realm != "lx":
        raise ApiError("forbidden", "저장 용량 할당은 LX 관리자가 LX 계정에 정합니다")
    q = _gb(body.get("quota_gb"), allow_none=True)
    from .approvals import solo_admin
    from .projects import lead_storage
    async with db(realm="lx") as conn:
        u = await conn.fetchrow("SELECT id, login, name, status, storage_quota_gb FROM lx_users WHERE id=$1 FOR UPDATE", uid)
        if not u:
            raise ApiError("not_found", "계정이 없습니다")
        if u["status"] == "disabled":
            raise ApiError("disabled_account", "사용 중지된 계정입니다", status=409)
        if _mine(p, "lx", uid) and not await solo_admin(conn, p):
            raise ApiError("self_account", "내 계정은 다른 관리자가 바꿉니다", status=409)
        before = float(u["storage_quota_gb"]) if u["storage_quota_gb"] is not None else None
        if before != q:
            await conn.execute("UPDATE lx_users SET storage_quota_gb=$2 WHERE id=$1", uid, Decimal(str(q)) if q is not None else None)
            await _log(conn, p.user_id, "lx", "account.quota", u["login"],
                       {"realm": "lx", "tenant_id": None, "name": u["name"], "from": before, "to": q,
                        "reason": f"{_qword(before)} → {_qword(q)}"})
        st = _storage_out(await lead_storage(conn, uid))
    return {"ok": True, "storage": st, "at": now_iso()}


def _qword(x) -> str:
    """사람마다 할당 — 비어 있으면 '기본 할당 따름'."""
    return gb_word(x) if x is not None else "기본 할당 따름"


def _admin(request: Request) -> Principal:
    p = require(principal(request))
    if not p.is_admin:
        raise ApiError("forbidden", "LX 관리자만 바꿉니다")
    return p


@router.get("/accounts/storage-default")
async def get_storage_default(request: Request):
    _admin(request)
    from .projects import STORAGE_DEFAULT_KEY, storage_default
    async with db(realm="lx") as conn:
        q = await storage_default(conn)
        row = await conn.fetchrow("SELECT updated_at FROM lx_settings WHERE key=$1", STORAGE_DEFAULT_KEY)
    return {"quota_gb": env(q, "GB", "recorded", "기본 할당(따로 정하지 않은 LX 계정 · 없으면 할당 없음)"),
            "updated_at": _iso(row["updated_at"]) if row else None, "at": now_iso()}


@router.put("/accounts/storage-default")
async def put_storage_default(body: dict, request: Request):
    """기본 할당 — 따로 정하지 않은 LX 계정에 쓰는 값(설정 한 곳). 비우면 할당 없음(처음 값)."""
    p = _admin(request)
    q = _gb(body.get("quota_gb"), allow_none=True)
    from .projects import STORAGE_DEFAULT_KEY, storage_default
    async with db(realm="lx") as conn:
        before = await storage_default(conn)
        if before != q:
            await conn.execute("INSERT INTO lx_settings(key, value, updated_by, updated_at) VALUES ($1,$2,$3,now()) "
                               "ON CONFLICT (key) DO UPDATE SET value=EXCLUDED.value, updated_by=EXCLUDED.updated_by, updated_at=now()",
                               STORAGE_DEFAULT_KEY, q, p.user_id)
            await _log(conn, p.user_id, "lx", "account.quota.default", "",
                       {"realm": "lx", "tenant_id": None, "name": "기본 할당", "from": before, "to": q, "reason": f"{gb_word(before)} → {gb_word(q)}"})
    return {"quota_gb": env(q, "GB", "recorded", "기본 할당"), "at": now_iso()}


# ── LX 부서 목록(제안 S-21 확인 · 10-01 사용자 "부서는 일단 여기(LX 누리집 조직도)에 보자. 나중엔 사내 시스템에서 불러오는 작업을 할 예정") ─────────
# 고르는 칸에 보이는 이름(label) = '상위 › 부서'(예: 공간정보본부 › 플랫폼사업처) — 상위가 최상위(사장 · 부사장 · 감사)면 부서만(예: 전남광주지역본부).
# 최상위(사장 · 부사장 · 감사)는 조직도의 뼈대일 뿐 고르는 칸에는 내지 않는다. 목록에 없는 부서(지역본부 아래 지사 등)는 직접 적는다(막지 않는다).
# 계정에는 label 을 적는다(목록과 같은 말) · 직접 적은 이름이 목록의 어느 이름(부서 · 상위 부서 · label)과 같으면 label 로 맞춘다.
DEPT_MAX, DEPTS_MAX, DEPT_FILE_MB = 60, 3000, 2
DEPT_TOP = "최상위"
DEPT_SEP = " › "
DEPT_HEAD = {"부서", "부서명", "부서 이름", "부서이름", "조직", "조직명", "소속", "소속 부서", "소속부서", "dept", "department", "name"}
DEPT_SOURCE_KEY = "depts.source"                   # lx_settings — 지금 목록의 출처 한 줄(화면 '출처: …')
_DEPT_BAD = re.compile(r"[\x00-\x1f\x7f<>›]")


def _dkey(s) -> str:
    """같은 부서인지 견줄 때 — '›' · 공백 정리 · 대소문자 무시."""
    return " ".join(str(s or "").replace("›", " ").replace(">", " ").split()).lower()


def _dclean(s) -> str:
    return " ".join(_DEPT_BAD.sub(" ", str(s if s is not None else "")).split())


# ── 부서 목록을 읽어 오는 곳(한 군데 — 출처를 바꿔 끼운다) ──────────────────────────────────
# 지금 출처: csv = server/config/lx-departments.csv(LX 누리집 조직도를 옮긴 파일 · 열 상위 · 부서 · 단위 · '#' 줄은 메모) — 처음 목록.
#           관리자 올리기(엑셀 · CSV — POST /accounts/depts/parse → PUT /accounts/depts)는 그 위에 덮어쓴다.
# 사내 시스템에서 불러오려면: 아래 DEPT_SOURCES 에 '() -> (rows [{name, parent?, unit?}], 출처 한 줄)' 함수를 하나 더하고
#           server/.env 의 LX_DEPT_SOURCE 를 그 이름으로 바꾼다. 관리자 화면의 '출처에서 다시 불러오기'(POST /accounts/depts/reload)가 그 함수를 부른다.
DEPT_CSV = config.SERVER_ROOT / "config" / "lx-departments.csv"


def _dept_from_csv() -> tuple[list[dict], str]:
    raw = DEPT_CSV.read_bytes()
    rows = _dept_table(_csv_rows(raw))
    m = re.search(r"(\d{4}-\d{2}-\d{2})", raw.decode("utf-8-sig", "ignore"))
    return rows, "LX 누리집 조직도" + (f" · {m.group(1)}" if m else "")


DEPT_SOURCES = {"csv": _dept_from_csv}
DEPT_SOURCE = (config.get("LX_DEPT_SOURCE", "csv") or "csv").strip()


def dept_source() -> tuple[list[dict], str]:
    fn = DEPT_SOURCES.get(DEPT_SOURCE) or DEPT_SOURCES["csv"]
    return fn()


def _csv_rows(data: bytes) -> list[list]:
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ApiError("bad_request", "CSV 글자를 읽을 수 없습니다. UTF-8 로 저장한 뒤 다시 올려 주세요", {"field": "file", "why": "read"})
    return [r for r in csv.reader(io.StringIO(text)) if r and not str(r[0]).lstrip().startswith("#")]


def _dept_table(table: list[list]) -> list[dict]:
    """표 → [{name, parent, unit}]. 첫 줄에 '상위' · '부서' 머리글이 있으면 그 열을 쓰고(조직도 모양), 아니면 첫 열만(부서 이름 한 열)."""
    if not table:
        return []
    head = [_dclean(c) for c in table[0]]
    if "부서" in head and "상위" in head:
        i, j = head.index("부서"), head.index("상위")
        k = head.index("단위") if "단위" in head else None
        get = lambda r, n: (r[n] if n is not None and n < len(r) else None)   # noqa: E731
        return [{"name": get(r, i), "parent": get(r, j), "unit": get(r, k)} for r in table[1:]]
    body = table[1:] if head and head[0].lower() in DEPT_HEAD else table
    return [{"name": r[0] if r else None, "parent": None, "unit": None} for r in body]


def _dept_rows(raw: list[dict]) -> tuple[list[dict], dict]:
    """정리 — 빈 칸 · 겹침 · 60자 넘음은 빼고 센다 · label 을 붙인다(상위가 최상위가 아니면 '상위 › 부서')."""
    rows, seen, skip = [], set(), {"blank": 0, "dup": 0, "long": 0}
    for x in raw:
        name, parent, unit = _dclean(x.get("name")), _dclean(x.get("parent")) or None, _dclean(x.get("unit")) or None
        if not name:
            skip["blank"] += 1
        elif len(name) > DEPT_MAX:
            skip["long"] += 1
        elif (_dkey(parent), _dkey(name)) in seen:
            skip["dup"] += 1
        else:
            seen.add((_dkey(parent), _dkey(name)))
            rows.append({"name": name, "parent": parent, "unit": unit})
    units = {_dkey(r["name"]): r["unit"] for r in rows}
    out, labels = [], set()
    for r in rows:
        up = r["parent"] if r["parent"] and units.get(_dkey(r["parent"])) != DEPT_TOP else None
        label = f"{up}{DEPT_SEP}{r['name']}" if up else r["name"]
        if _dkey(label) in labels:
            skip["dup"] += 1
            continue
        labels.add(_dkey(label))
        out.append({**r, "label": label, "pick": r["unit"] != DEPT_TOP})
    if len(out) > DEPTS_MAX:
        raise ApiError("bad_request", f"부서는 {DEPTS_MAX:,}개까지입니다", {"field": "names"})
    return out, skip


async def _save_depts(conn, actor: str, rows: list[dict], how: str, note: str | None = None) -> None:
    before = int(await conn.fetchval("SELECT count(*) FROM lx_depts WHERE pick") or 0)
    await conn.execute("DELETE FROM lx_depts")
    if rows:
        await conn.executemany("INSERT INTO lx_depts(label, name, parent, unit, pick, pos, by) VALUES ($1,$2,$3,$4,$5,$6,$7)",
                               [(r["label"], r["name"], r["parent"], r["unit"], r["pick"], i, actor) for i, r in enumerate(rows)])
    if note is not None:
        await conn.execute("INSERT INTO lx_settings(key, value, updated_by, updated_at) VALUES ($1,$2,$3,now()) "
                           "ON CONFLICT (key) DO UPDATE SET value=EXCLUDED.value, updated_by=EXCLUDED.updated_by, updated_at=now()",
                           DEPT_SOURCE_KEY, {"note": note}, actor)
    n = sum(1 for r in rows if r["pick"])
    await _log(conn, actor, "lx", "account.depts", "",
               {"realm": "lx", "tenant_id": None, "name": "부서 목록", "from": before, "to": n, "reason": f"{how} · {before}개 → {n}개"})


async def _ensure_depts(conn) -> None:
    """처음 한 번 — 출처(지금은 누리집 조직도 CSV)에서 목록을 채운다. 출처 기록(lx_settings)이 있으면(관리자가 비웠어도) 다시 채우지 않는다."""
    if await conn.fetchval("SELECT 1 FROM lx_settings WHERE key=$1", DEPT_SOURCE_KEY):
        return
    try:
        raw, note = await run_in_threadpool(dept_source)
        rows, _ = _dept_rows(raw)
    except Exception:            # 출처를 못 읽으면 목록 없이(직접 적기) — 다음에 다시 해 본다
        return
    if not await conn.fetchval("INSERT INTO lx_settings(key, value, updated_by) VALUES ($1,$2,'system') ON CONFLICT DO NOTHING RETURNING key",
                               DEPT_SOURCE_KEY, {"note": note}):
        return                   # 다른 요청이 먼저 채웠다
    await _save_depts(conn, "system", rows, "처음 목록 — " + note)


async def dept_rows(conn) -> list[dict]:
    await _ensure_depts(conn)
    return [dict(r) for r in await conn.fetch("SELECT label, name, parent, unit, pick FROM lx_depts ORDER BY pos, label")]


async def dept_list(conn) -> list[str]:
    """고르는 칸에 내는 이름(label) — 최상위(사장 · 부사장 · 감사)는 빼고 조직도 순서대로."""
    return [r["label"] for r in await dept_rows(conn) if r["pick"]]


async def _dept_keys(conn) -> dict:
    """견줄 열쇠 → label — label · 부서 이름 · '상위 부서' 어느 것으로 적어도 같은 부서로 본다(부서 이름이 겹치면 label 만)."""
    rows = [r for r in await dept_rows(conn) if r["pick"]]
    by_name: dict[str, list[str]] = {}
    for r in rows:
        by_name.setdefault(_dkey(r["name"]), []).append(r["label"])
    keys = {}
    for r in rows:
        keys[_dkey(r["label"])] = r["label"]
        if r["parent"]:
            keys[_dkey(f"{r['parent']} {r['name']}")] = r["label"]
    for k, labels in by_name.items():
        if len(labels) == 1:
            keys.setdefault(k, labels[0])
    return keys


async def dept_listed(conn, dept: str | None) -> bool | None:
    """지금 적힌 부서가 목록에 있나 — 목록이 비었거나 부서가 비면 None(안내할 것 없음)."""
    if not (dept or "").strip():
        return None
    keys = await _dept_keys(conn)
    return (_dkey(dept) in keys) if keys else None


async def dept_pick(conn, dept: str | None, keep: str | None = None) -> str | None:
    """목록에 있는 부서면 목록의 이름(label)으로 맞춘다 · 없으면 적은 그대로(지사 등 직접 적기 — 막지 않는다). keep 은 예전 호환(쓰지 않음)."""
    if not dept:
        return dept
    keys = await _dept_keys(conn)
    return keys.get(_dkey(dept), dept)


def _read_table(name: str, data: bytes) -> list[list]:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext == "xlsx":
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            table = [list(row) for row in ws.iter_rows(values_only=True) if row and any(c is not None and str(c).strip() for c in row)]
            wb.close()
        except Exception:
            raise ApiError("bad_request", "엑셀 파일을 열 수 없습니다. 엑셀에서 xlsx 로 다시 저장한 뒤 올려 주세요", {"field": "file", "why": "read"})
        return [r for r in table if not str(r[0] or "").lstrip().startswith("#")]
    if ext == "csv":
        return _csv_rows(data)
    raise ApiError("bad_request", "엑셀(xlsx) 또는 CSV 파일을 올려 주세요", {"field": "file", "why": "type"})


def _rows_out(rows: list[dict]) -> list[dict]:
    return [{"label": r["label"], "name": r["name"], "parent": r["parent"], "unit": r["unit"]} for r in rows if r["pick"]]


@router.get("/accounts/depts")
async def get_depts(request: Request):
    """LX 부서 목록 — 로그인 없이도 이름만(가입 신청 창의 부서 고르기). LX 관리자에게는 출처 · 마지막 바꿈 · 단위 · 목록에 없는 부서를 쓰는 계정."""
    p = principal(request)
    async with db(realm="lx") as conn:
        rows = await dept_rows(conn)
        names = [r["label"] for r in rows if r["pick"]]
        out = {"items": names, "count": env(len(names), "count", "recorded", "LX 부서 목록(lx_depts · 최상위 뺌)")}
        if p and p.is_admin:
            src = await conn.fetchrow("SELECT value FROM lx_settings WHERE key=$1", DEPT_SOURCE_KEY)
            last = await conn.fetchrow("SELECT actor, at FROM audit_log WHERE action='account.depts' ORDER BY id DESC LIMIT 1")
            who = ("자동" if last["actor"] == "system" else await conn.fetchval("SELECT name FROM lx_users WHERE id=$1", last["actor"])) if last else None
            keys = await _dept_keys(conn)
            off = [{"name": u["name"], "login": u["login"], "dept": u["dept"]} for u in await conn.fetch(
                "SELECT name, login, dept FROM lx_users WHERE status <> 'disabled' AND coalesce(trim(dept), '') <> '' ORDER BY name")
                if keys and _dkey(u["dept"]) not in keys]
            out.update({"rows": _rows_out(rows), "source": ((src["value"] or {}).get("note") if src else None),
                        "source_kind": DEPT_SOURCE, "updated_at": _iso(last["at"]) if last else None, "updated_by": who, "unlisted": off})
    return {**out, "at": now_iso()}


@router.post("/accounts/depts/parse")
async def parse_depts(request: Request, file: UploadFile = File(...)):
    """엑셀(xlsx) · CSV 를 읽어 목록을 미리 보여 준다(저장하지 않음 — 확인한 뒤 PUT /accounts/depts).
    첫 줄이 '상위 · 부서 · 단위' 머리글이면 조직도 모양으로, 아니면 첫 열을 부서 이름 한 열로. '#' 로 시작하는 줄은 메모."""
    _admin(request)
    data = await file.read(DEPT_FILE_MB * 1024 * 1024 + 1)
    if len(data) > DEPT_FILE_MB * 1024 * 1024:
        raise ApiError("bad_request", f"부서 목록 파일은 {DEPT_FILE_MB}MB 까지입니다", {"field": "file", "why": "size"})
    table = await run_in_threadpool(_read_table, file.filename or "", data)
    rows, skip = _dept_rows(_dept_table(table))
    picks = [r for r in rows if r["pick"]]
    if not picks:
        raise ApiError("bad_request", "부서 이름을 찾지 못했습니다. 첫 열에 부서 이름을 한 줄에 하나씩 적어 주세요", {"field": "file", "why": "empty"})
    return {"rows": [{k: r[k] for k in ("name", "parent", "unit")} for r in rows], "labels": [r["label"] for r in picks],
            "count": env(len(picks), "count", "measured", "읽은 부서(최상위 뺌)"),
            "skipped": {k: env(v, "count", "measured", "읽지 않은 줄(빈 칸 · 겹침 · 60자 넘음)") for k, v in skip.items()}, "at": now_iso()}


@router.put("/accounts/depts")
async def put_depts(body: dict, request: Request):
    """부서 목록을 통째로 바꾼다(올린 파일을 미리 보고 확인한 뒤 · rows [{name, parent?, unit?}] 또는 names [이름]). 빈 목록이면 부서 칸은 직접 적기만.
    계정에 이미 적힌 부서 이름은 건드리지 않는다."""
    p = _admin(request)
    raw = body.get("rows")
    if raw is None and isinstance(body.get("names"), list):
        raw = [{"name": n} for n in body["names"]]
    if not isinstance(raw, list) or not all(isinstance(x, dict) for x in raw):
        raise ApiError("bad_request", "rows 는 [{name, parent, unit}] 목록입니다", {"field": "rows"})
    rows, _ = _dept_rows(raw)
    note = " ".join(str(body.get("source") or "").split())[:80] or f"관리자가 올린 파일 · {dt.datetime.now(KST):%Y-%m-%d}"
    async with db(realm="lx") as conn:
        await _ensure_depts(conn)
        await _save_depts(conn, p.user_id, rows, "목록 바꿈" if rows else "모두 비움", note=note if rows else "비움")
        names = await dept_list(conn)
    return {"items": names, "count": env(len(names), "count", "recorded", "LX 부서 목록(lx_depts · 최상위 뺌)"), "at": now_iso()}


@router.post("/accounts/depts/reload")
async def reload_depts(request: Request):
    """출처(지금 csv = 누리집 조직도 · 나중에 사내 시스템)에서 다시 불러와 목록을 바꾼다."""
    p = _admin(request)
    try:
        raw, note = await run_in_threadpool(dept_source)
    except ApiError:
        raise
    except Exception:
        raise ApiError("conflict", "출처에서 부서 목록을 읽지 못했습니다", status=409)
    rows, _ = _dept_rows(raw)
    async with db(realm="lx") as conn:
        await _ensure_depts(conn)
        await _save_depts(conn, p.user_id, rows, "출처에서 다시 불러옴 — " + note, note=note)
        names = await dept_list(conn)
    return {"items": names, "count": env(len(names), "count", "recorded", "LX 부서 목록(lx_depts · 최상위 뺌)"), "source": note, "at": now_iso()}


@router.post("/accounts/depts/{op}")
async def edit_dept(op: str, body: dict, request: Request):
    """하나 더하기(add {name, parent?}) · 빼기(remove {label}) — 더하면 목록 끝에. 겹치면 409 · 없는 이름 빼기는 404."""
    if op not in ("add", "remove"):
        raise ApiError("not_found", "없는 길입니다")
    p = _admin(request)
    s = _dclean(body.get("name") if op == "add" else body.get("label") or body.get("name"))
    up = _dclean(body.get("parent")) or None if op == "add" else None
    if not s:
        raise ApiError("bad_request", "부서 이름을 적어 주세요", {"field": "name"})
    if len(s) > DEPT_MAX or (up and len(up) > DEPT_MAX):
        raise ApiError("bad_request", f"부서 이름은 {DEPT_MAX}자까지입니다", {"field": "name"})
    async with db(realm="lx") as conn:
        rows = await dept_rows(conn)
        n0 = sum(1 for r in rows if r["pick"])
        if op == "add":
            label = f"{up}{DEPT_SEP}{s}" if up else s
            if any(_dkey(r["label"]) == _dkey(label) for r in rows):
                raise ApiError("conflict", "이미 목록에 있는 부서입니다", {"field": "name"}, 409)
            if len(rows) >= DEPTS_MAX:
                raise ApiError("bad_request", f"부서는 {DEPTS_MAX:,}개까지입니다", {"field": "name"})
            await conn.execute("INSERT INTO lx_depts(label, name, parent, unit, pick, pos, by) "
                               "VALUES ($1,$2,$3,NULL,true, coalesce((SELECT max(pos) + 1 FROM lx_depts), 0), $4)", label, s, up, p.user_id)
            how, n1 = f"더함 {label}", n0 + 1
        else:
            hit = next((r for r in rows if r["pick"] and _dkey(r["label"]) == _dkey(s)), None)
            if not hit:
                raise ApiError("not_found", "목록에 없는 부서입니다")
            await conn.execute("DELETE FROM lx_depts WHERE label=$1", hit["label"])
            how, n1 = f"뺌 {hit['label']}", n0 - 1
        await _log(conn, p.user_id, "lx", "account.depts", "",
                   {"realm": "lx", "tenant_id": None, "name": "부서 목록", "from": n0, "to": n1, "reason": f"{how} · {n0}개 → {n1}개"})
        names = await dept_list(conn)
    return {"items": names, "count": env(len(names), "count", "recorded", "LX 부서 목록(lx_depts · 최상위 뺌)"), "at": now_iso()}
