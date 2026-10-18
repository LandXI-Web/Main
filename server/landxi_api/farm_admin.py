"""영농관리 행정서비스 결정(10-18 원칙 291 · 확인 대장 10-18 10:47 now 답 바퀴 18) — 기관 세션 · 관할 · 받은 서비스 범위(공간과 같은 판정).

  [곤포사일리지 생산량 — 영농-q2 ⓑ '기관이 1개당 무게를 넣으면 생산량도']
  GET  /api/v1/spaces/me/settings/farm                     → {bale_kg(봉투 | null), updated{by, date}}
  PUT  /api/v1/spaces/me/settings/farm {bale_kg | null}    기관 관리자만 · 1~3,000 kg · 비우면 지움(그 뒤로는 개수만)
       생산량(kg) = 곤포사일리지 개수 × 1개당 무게 — 개수가 나오는 곳(통계 · 보고서 · 사료작물 견주기 · 엑셀)에 같은 셈.

  [내려받기 이용 약속 — 영농-q3 '개인정보가 있는건 확인까지 받고. 아닌건 그냥 받고']
  규칙(PERSONAL_RULE): 올린 대장의 필지별 내용(지번 · 대장 값)이 든 파일 = 개인정보가 든 파일.
    처음 한 번 이용 약속(사용 목적 · 활용 기간 — 서버에 누가 · 언제와 함께 기록) → 받을 때마다 확인 한 번(ok=1). 기간이 끝나면 다시 약속.
    그 밖의 파일(지도 파일 · 필지 엑셀 · 요약 · 통계 · 대장이 없는 보고서)은 바로 받는다.
  GET  /api/v1/spaces/me/pledge                            → {active{purpose, until, at} | null, rule}
  POST /api/v1/spaces/me/pledge {purpose, until, agree}    이용 약속 적기(사람마다 · 여러 번 · 마지막 것이 유효)

  [행정정보 주제별 목록 — 영농-행정정보목록 넣음]
  GET   /api/v1/spaces/me/ledgers                          주제(대장 종류)별 · 기준일 · 건수(대장 행 · 이은 필지) · 등록자
  PATCH /api/v1/spaces/me/ledgers/{import_id} {base_date}  기준일 적기(기관 관리자)

  [사료작물 재배 신청 대장 × 세 시점 — 영농-q1 ⓐ '필지마다 견주기만']
  GET  /api/v1/spaces/me/ledger/feedcrop?import_id=&growth=&harvest=&bale=&geom=   필지마다 시점별 AI 결과 있음 · 없음 → 맞음 · 다름(어느 시점이 없는지)
  GET  /api/v1/spaces/me/ledger/feedcrop.xlsx?…&ok=1                               같은 행 엑셀(개인정보가 든 파일 — 이용 약속 · 확인)
  경영체로 묶기 · 합계 넓이 · 이행 판정 같은 덧붙인 판단은 하지 않는다(사실 그대로 · 처분은 담당자).
어느 지역이든 같은 길 — 지역 이름 · 코드를 코드에 두지 않는다(관할 = 로그인한 기관).
"""
from __future__ import annotations

import datetime as dt
import io
import json
import re
import secrets
from urllib.parse import quote

from fastapi import APIRouter, Request
from fastapi.responses import Response
from starlette.concurrency import run_in_threadpool

from .deps import ApiError, Principal, db
from .envelope import KST, env, now_iso

router = APIRouter()
IMP_RE = re.compile(r"^imp_[0-9A-Z]{10,40}$")
CARD_RE = re.compile(r"^card-[a-z0-9-]{2,40}$")
PERSONAL_RULE = ("올린 대장의 필지별 내용(지번 · 대장 값)이 든 파일은 개인정보가 든 파일로 봅니다 — 처음 한 번 이용 약속(사용 목적 · 활용 기간)을 적고, "
                 "받을 때마다 한 번 더 확인합니다. 그 밖의 파일은 바로 받습니다.")
BALE_MAX = 3000.0
SLOTS = [("growth", "생육기", ("생육",)), ("harvest", "생산기", ("생산", "수확")), ("bale", "곤포사일리지", ("곤포",))]
NEG = {"비경작지", "휴경지", "휴경", "배경", "기타"}        # '찾지 못함'과 같은 뜻의 종류 — 재배 흔적으로 세지 않는다
SRC_FEED = "사료작물 재배 신청 대장 필지 × 시점별 서비스 결과(필지 번호로 맞춤)"


def _me(request: Request) -> Principal:
    from .spaces import _me as me
    return me(request)


def _manager(p: Principal):
    if p.role != "manager":
        raise ApiError("forbidden", "기관 관리자만 바꿀 수 있습니다")


def is_bale(name: str | None) -> bool:
    return "곤포" in str(name or "")


# ── 기관 설정 · 곤포사일리지 1개당 무게 ─────────────────────────────────────────
async def bale_kg(tenant: str) -> float | None:
    async with db(realm="lx") as c:
        if not await c.fetchval("SELECT to_regclass('tenant_settings') IS NOT NULL"):
            return None
        v = await c.fetchval("SELECT value FROM tenant_settings WHERE tenant_id=$1 AND key='bale_kg'", tenant)
    v = json.loads(v) if isinstance(v, str) else v
    try:
        x = float(v) if v is not None else None
    except (TypeError, ValueError):
        return None
    return x if x and x > 0 else None


def yield_env(n: int | None, kg: float | None):
    """생산량(kg) = 개수 × 1개당 무게 — 무게가 없거나 개수가 없으면 None(개수만 보인다)."""
    if not kg or n is None:
        return None
    return env(round(n * kg, 1), "kg", "estimate", "곤포사일리지 개수 × 기관이 넣은 1개당 무게", None)


@router.get("/spaces/me/settings/farm")
async def farm_settings(request: Request):
    p = _me(request)
    async with db(p) as c:
        r = await c.fetchrow("SELECT s.value, s.updated_at, u.name, u.dept FROM tenant_settings s LEFT JOIN tenant_users u ON u.id = s.updated_by "
                             "WHERE s.tenant_id=$1 AND s.key='bale_kg'", p.tenant_id)
    kg = await bale_kg(p.tenant_id)
    return {"bale_kg": env(kg, "kg", "recorded", "기관이 넣은 곤포사일리지 1개당 무게") if kg else None,
            "updated": {"by": " · ".join(x for x in ((r["dept"] or "").strip(), (r["name"] or "").strip()) if x) or None,
                        "date": r["updated_at"].astimezone(KST).date().isoformat()} if r else None,
            "can_edit": p.role == "manager", "as_of": now_iso()}


@router.put("/spaces/me/settings/farm")
async def farm_settings_put(request: Request, body: dict | None = None):
    p = _me(request)
    _manager(p)
    raw = (body or {}).get("bale_kg")
    kg = None
    if raw not in (None, ""):
        try:
            kg = round(float(str(raw).replace(",", "").strip()), 1)
        except ValueError:
            raise ApiError("bad_request", "무게는 숫자(kg)로 적어 주세요", {"field": "bale_kg"})
        if not (0 < kg <= BALE_MAX):
            raise ApiError("bad_request", f"무게는 0보다 크고 {BALE_MAX:,.0f} kg 까지입니다", {"field": "bale_kg"})
    async with db(p) as c:
        if kg is None:
            await c.execute("DELETE FROM tenant_settings WHERE tenant_id=$1 AND key='bale_kg'", p.tenant_id)
        else:
            await c.execute("INSERT INTO tenant_settings(tenant_id, key, value, updated_by, updated_at) VALUES ($1,'bale_kg',$2,$3,now()) "
                            "ON CONFLICT (tenant_id, key) DO UPDATE SET value=EXCLUDED.value, updated_by=EXCLUDED.updated_by, updated_at=now()",
                            p.tenant_id, kg, p.user_id)
        await c.execute("INSERT INTO space_log(tenant_id, kind, line, actor, realm, detail) VALUES ($1,'space',$2,$3,'tenant',$4)",
                        p.tenant_id, f"{p.name or '담당자'} — 곤포사일리지 1개당 무게 " + (f"{kg:,.1f}kg" if kg else "지움"), p.user_id, {"bale_kg": kg})
    return await farm_settings(request)


# ── 이용 약속(개인정보가 든 파일) ───────────────────────────────────────────────
def _pl_out(r) -> dict | None:
    if not r:
        return None
    return {"id": r["id"], "purpose": r["purpose"], "until": r["until_date"].isoformat(),
            "at": r["at"].astimezone(KST).isoformat(timespec="seconds")}


async def active_pledge(p: Principal) -> dict | None:
    async with db(p) as c:
        r = await c.fetchrow("SELECT id, purpose, until_date, at FROM download_pledges WHERE tenant_id=$1 AND user_id=$2 "
                             "AND until_date >= (now() AT TIME ZONE 'Asia/Seoul')::date ORDER BY at DESC LIMIT 1", p.tenant_id, p.user_id)
    return _pl_out(r)


async def need_personal(p: Principal, ok) -> dict:
    """개인정보가 든 파일 — 유효한 이용 약속이 있고 이번 받기를 확인(ok=1)했을 때만. 없으면 409(화면이 약속 창 · 확인 창을 띄운다)."""
    pl = await active_pledge(p)
    if not pl:
        raise ApiError("conflict", "개인정보가 든 파일입니다 — 이용 약속을 먼저 적어 주세요", {"need": "pledge", "rule": PERSONAL_RULE}, 409)
    if str(ok or "") not in ("1", "true"):
        raise ApiError("conflict", "개인정보가 든 파일입니다 — 받기 전에 확인해 주세요", {"need": "confirm", "pledge": pl}, 409)
    return pl


@router.get("/spaces/me/pledge")
async def pledge_get(request: Request):
    p = _me(request)
    return {"active": await active_pledge(p), "rule": PERSONAL_RULE, "as_of": now_iso()}


@router.post("/spaces/me/pledge")
async def pledge_post(request: Request, body: dict | None = None):
    p = _me(request)
    b = body or {}
    if not b.get("agree"):
        raise ApiError("bad_request", "약속에 동의해 주세요", {"field": "agree"})
    purpose = " ".join(str(b.get("purpose") or "").split())
    if len(purpose) < 2 or len(purpose) > 200:
        raise ApiError("bad_request", "사용 목적을 두 글자 이상 200자 안으로 적어 주세요", {"field": "purpose"})
    try:
        until = dt.date.fromisoformat(str(b.get("until") or "")[:10])
    except ValueError:
        raise ApiError("bad_request", "활용 기간(끝나는 날)을 골라 주세요", {"field": "until"})
    today = dt.datetime.now(KST).date()
    if until < today or until > today + dt.timedelta(days=366):
        raise ApiError("bad_request", "활용 기간은 오늘부터 1년 안으로 골라 주세요", {"field": "until"})
    pid = "pl_" + secrets.token_hex(8)
    async with db(p) as c:
        await c.execute("INSERT INTO download_pledges(id, tenant_id, user_id, purpose, until_date) VALUES ($1,$2,$3,$4,$5)",
                        pid, p.tenant_id, p.user_id, purpose, until)
        await c.execute("INSERT INTO space_log(tenant_id, kind, line, actor, realm, detail) VALUES ($1,'consent',$2,$3,'tenant',$4)",
                        p.tenant_id, f"{p.name or '담당자'} — 개인정보가 든 파일 이용 약속({until.isoformat()}까지)", p.user_id,
                        {"pledge": pid, "purpose": purpose, "until": until.isoformat()})
    return {"ok": True, "active": await active_pledge(p), "as_of": now_iso()}


# ── 행정정보 주제별 목록 ─────────────────────────────────────────────────────
@router.get("/spaces/me/ledgers")
async def ledgers(request: Request):
    from .ledger import KINDS
    p = _me(request)
    async with db(p) as c:
        rows = await c.fetch("SELECT i.id, i.kind, i.name, i.filename, i.rows, i.stats, i.state, i.base_date, i.created_at, i.confirmed_at, i.latest, "
                             "u.name AS uname, u.dept AS udept, u.position AS upos FROM ledger_imports i LEFT JOIN tenant_users u ON u.id = i.created_by "
                             "WHERE i.tenant_id=$1 AND i.state IN ('matched', 'matching') ORDER BY i.kind, coalesce(i.base_date, i.created_at::date) DESC, i.created_at DESC LIMIT 300",
                             p.tenant_id)
    groups: dict[str, dict] = {}
    for r in rows:
        st = r["stats"] or {}
        st = json.loads(st) if isinstance(st, str) else st
        at = r["confirmed_at"] or r["created_at"]
        g = groups.setdefault(r["kind"], {"kind": r["kind"], "label": KINDS.get(r["kind"]) or "대장", "items": []})
        g["items"].append({
            "import_id": r["id"], "name": re.sub(r"\.(xlsx|xls)\.csv$", r".\1", r["name"] or r["filename"] or "") or g["label"],
            "base_date": r["base_date"].isoformat() if r["base_date"] else None,
            "uploaded": at.astimezone(KST).date().isoformat() if at else None,
            "rows": env(int(r["rows"] or 0), "행", "recorded", "올린 대장 행"),
            "parcels": env(int(st.get("matched") or 0), "필지", "measured", "대장 × 연속지적(필지 번호)") if r["state"] == "matched" else None,
            "who": " · ".join(x for x in ((r["udept"] or "").strip(), " ".join(y for y in ((r["uname"] or "").strip(), (r["upos"] or "").strip()) if y)) if x) or "기관 담당자",
            "state": "이어 붙이는 중" if r["state"] == "matching" else None, "latest": bool(r["latest"])})
    order = list(KINDS)
    out = sorted(groups.values(), key=lambda g: order.index(g["kind"]) if g["kind"] in order else 99)
    for g in out:
        g["n"] = env(len(g["items"]), "건", "recorded", "올린 대장")
    return {"groups": out, "kinds": KINDS, "can_edit": p.role == "manager", "as_of": now_iso()}


@router.patch("/spaces/me/ledgers/{iid}")
async def ledger_base_date(iid: str, request: Request, body: dict | None = None):
    p = _me(request)
    _manager(p)
    if not IMP_RE.match(iid or ""):
        raise ApiError("not_found", "올린 대장이 없습니다")
    raw = str((body or {}).get("base_date") or "").strip()[:10]
    d = None
    if raw:
        try:
            d = dt.date.fromisoformat(raw)
        except ValueError:
            raise ApiError("bad_request", "기준일은 날짜로 골라 주세요", {"field": "base_date"})
    async with db(p) as c:
        n = await c.execute("UPDATE ledger_imports SET base_date=$3 WHERE id=$1 AND tenant_id=$2", iid, p.tenant_id, d)
    if not str(n).endswith(" 1"):
        raise ApiError("not_found", "올린 대장이 없습니다")
    return {"ok": True, "import_id": iid, "base_date": d.isoformat() if d else None}


# ── 사료작물 재배 신청 대장 × 세 시점 ─────────────────────────────────────────
async def _feed_import(p: Principal, import_id: str | None) -> dict | None:
    async with db(realm="lx") as c:
        if import_id:
            if not IMP_RE.match(import_id):
                raise ApiError("not_found", "올린 대장이 없습니다")
            r = await c.fetchrow("SELECT * FROM ledger_imports WHERE id=$1 AND tenant_id=$2 AND kind='feedcrop'", import_id, p.tenant_id)
            if not r:
                raise ApiError("not_found", "올린 사료작물 재배 신청 대장이 없습니다")
            return dict(r)
        r = await c.fetchrow("SELECT * FROM ledger_imports WHERE tenant_id=$1 AND kind='feedcrop' AND state='matched' "
                             "ORDER BY (created_by=$2) DESC, coalesce(base_date, created_at::date) DESC, created_at DESC LIMIT 1", p.tenant_id, p.user_id or "")
    return dict(r) if r else None


def default_picks(services: list[dict], given: dict[str, str | None]) -> dict[str, str | None]:
    """시점마다 서비스 — 고른 것 그대로 · 고르지 않은 시점은 서비스 이름의 말(생육 · 생산 · 곤포)로 · 생육기만은 그래도 없으면 첫 서비스(순수 함수)."""
    cards = {s["card"] for s in services}
    out: dict[str, str | None] = {}
    used: set[str] = set()
    for key, _label, words in SLOTS:
        v = given.get(key)
        if v == "-":                       # 일부러 비움
            out[key] = None
            continue
        if v and v in cards:
            out[key] = v
            used.add(v)
            continue
        hit = next((s["card"] for s in services if s["card"] not in used and any(w in str(s.get("name") or "") for w in words)), None)
        out[key] = hit
        if hit:
            used.add(hit)
    if not any(out.values()) and services and given.get("growth") != "-":
        out["growth"] = services[0]["card"]
    return out


def judge(slots: dict[str, dict | None]) -> tuple[str, list[str]]:
    """필지 하나 — 고른 시점마다 AI 결과 있음 · 없음 → ('ok' | 'diff', 없는 시점 이름). 고른 시점이 없으면 ('none', [])."""
    picked = [(k, v) for k, v in slots.items() if v is not None]
    if not picked:
        return "none", []
    miss = [dict((k, lb) for k, lb, _ in SLOTS)[k] for k, v in picked if not v.get("found")]
    return ("diff" if miss else "ok"), miss


async def feedcrop(p: Principal, import_id: str | None, given: dict[str, str | None], geom: bool = False) -> dict:
    from .gov_report import _addr_short, _ymd, _services, _svc
    from .spaces import class_ko
    services = [s for s in await _services(p) if s["parcel"]]
    picks = default_picks(services, given)
    slot_meta = [{"key": k, "label": lb, "card": picks.get(k), "name": next((s["name"] for s in services if s["card"] == picks.get(k)), None),
                  "shot": next((s.get("shot") for s in services if s["card"] == picks.get(k)), None)} for k, lb, _ in SLOTS]
    kg = await bale_kg(p.tenant_id)
    imp = await _feed_import(p, import_id)
    base = {"services": [{"card": s["card"], "name": s["name"], "shot": s.get("shot")} for s in services], "slots": slot_meta,
            "bale_kg": env(kg, "kg", "recorded", "기관이 넣은 곤포사일리지 1개당 무게") if kg else None, "rule": PERSONAL_RULE, "as_of": now_iso()}
    if not imp:
        return {**base, "ledger": None}
    async with db(realm="lx") as c:
        await c.execute("SET LOCAL jit = off")
        led: dict[str, dict] = {}
        for x in await c.fetch("SELECT pnu, payload FROM registry_snapshots WHERE import_id=$1 AND tenant_id=$2 AND pnu IS NOT NULL ORDER BY pnu, row_no",
                               imp["id"], p.tenant_id):
            pay = x["payload"] or {}
            pay = json.loads(pay) if isinstance(pay, str) else pay
            e = led.setdefault(x["pnu"], {"crops": [], "m2": None, "rows": 0})
            e["rows"] += 1
            cr = str(pay.get("use") or "").strip()
            if cr and cr not in e["crops"]:
                e["crops"].append(cr)
            try:
                a = float(str(pay.get("area") or "").replace(",", ""))
                e["m2"] = (e["m2"] or 0) + a
            except ValueError:
                pass
        loose = int(await c.fetchval("SELECT count(*) FROM registry_snapshots WHERE import_id=$1 AND tenant_id=$2 AND pnu IS NULL", imp["id"], p.tenant_id) or 0)
        pn = list(led)
        attrs = {a["pnu"]: dict(a) for a in await c.fetch("SELECT pnu, addr, emd, area_m2 FROM survey_parcels WHERE pnu = ANY($1::text[])", pn)} if pn else {}
        gj = {}
        if geom and pn:
            gj = {r["pnu"]: r["g"] for r in await c.fetch(
                "SELECT pnu, ST_AsGeoJSON(ST_SimplifyPreserveTopology(geom, 0.000004), 6) g FROM survey_parcels WHERE pnu = ANY($1::text[])", pn)}
    hits: dict[str, dict[str, dict]] = {}
    for sm in slot_meta:
        if not sm["card"] or not pn:
            continue
        try:
            _s, _g, _rows, _body, sets, codes, _only = await _svc(p, sm["card"])
        except ApiError:
            sm["card"] = None
            continue
        if not sets:
            continue
        sc = " AND left(d.pnu, 5) = ANY($3::text[])" if codes is not None else ""
        async with db(realm="lx") as c:
            await c.execute("SET LOCAL jit = off")
            got = await c.fetch("SELECT d.pnu, d.cls, d.cls_en, count(*) n, coalesce(sum(d.area_m2), 0) m2 FROM detections d "
                                f"WHERE d.job_id = ANY($1::text[]) AND d.edit_state <> 'deleted' AND d.pnu = ANY($2::text[]){sc} GROUP BY 1, 2, 3",
                                *([sets, pn, codes] if codes is not None else [sets, pn]))
        per: dict[str, dict] = {}
        for r in got:
            k = class_ko(r["cls"], r["cls_en"])
            if k in NEG:
                continue
            e = per.setdefault(r["pnu"], {"by": {}, "n": 0, "bale": 0})
            e["by"][k] = e["by"].get(k, 0.0) + float(r["m2"])
            e["n"] += int(r["n"])
            if is_bale(k):
                e["bale"] += int(r["n"])
        hits[sm["key"]] = per
    items = []
    for k in pn:
        a = attrs.get(k) or {}
        L = led[k]
        slots: dict[str, dict | None] = {}
        for sm in slot_meta:
            if not sm["card"]:
                slots[sm["key"]] = None
                continue
            e = (hits.get(sm["key"]) or {}).get(k)
            if sm["key"] == "bale":
                n = (e or {}).get("bale") or 0                 # 곤포사일리지 종류만 센다(다른 종류는 '있음'으로만)
                slots["bale"] = {"found": bool(e), "n": n, "kg": round(n * kg, 1) if kg and n else None}
            else:
                slots[sm["key"]] = {"found": bool(e), "ai": " · ".join(f"{x} {round(m):,}㎡" for x, m in sorted(e["by"].items(), key=lambda kv: -kv[1])) if e else None,
                                    "m2": round(sum(e["by"].values()), 1) if e else None}
        kk, miss = judge(slots)
        items.append({"pnu": k, "where": _addr_short(a.get("addr")) or k, "emd": a.get("emd") or "", "crop": " · ".join(L["crops"]) or "—",
                      "apply_m2": round(L["m2"], 1) if L["m2"] is not None else None,
                      "parcel_m2": round(float(a["area_m2"]), 1) if a.get("area_m2") else None,
                      "slots": slots, "k": kk, "miss": miss})
    items.sort(key=lambda x: ({"diff": 0, "ok": 1, "none": 2}[x["k"]], x["emd"], x["where"]))
    as_of = (imp["confirmed_at"] or imp["created_at"]).astimezone(KST).isoformat(timespec="seconds")
    cnt = {"ok": sum(1 for x in items if x["k"] == "ok"), "diff": sum(1 for x in items if x["k"] == "diff")}
    per_slot = []
    for sm in slot_meta:
        if not sm["card"]:
            continue
        f = sum(1 for x in items if (x["slots"].get(sm["key"]) or {}).get("found"))
        row = {"key": sm["key"], "label": sm["label"], "found": env(f, "필지", "inferred", SRC_FEED), "missing": env(len(items) - f, "필지", "inferred", SRC_FEED)}
        if sm["key"] == "bale":
            bn = sum((x["slots"].get("bale") or {}).get("n") or 0 for x in items)
            row["bales"] = env(bn, "개", "inferred", SRC_FEED)
            row["yield"] = yield_env(bn, kg)
        per_slot.append(row)
    # 읍면동 묶음(10-18 제안-견주기읍면동묶음 넣음) — 읍면동마다 신청 필지 · 맞음 · 다름 · 시점별 AI 결과 없음(현장 나갈 순서 · 다른 필지 많은 곳부터)
    emd_m: dict[str, dict] = {}
    for x in items:
        e = emd_m.setdefault(x["emd"] or "읍면동 밖", {"emd": x["emd"] or "읍면동 밖", "n": 0, "ok": 0, "diff": 0, "miss": {}})
        e["n"] += 1
        if x["k"] in ("ok", "diff"):
            e[x["k"]] += 1
        for m in x["miss"]:
            e["miss"][m] = e["miss"].get(m, 0) + 1
    by_emd = sorted(emd_m.values(), key=lambda e: (e["emd"] == "읍면동 밖", -e["diff"], -e["n"], e["emd"]))
    from .ledger import _uploader
    async with db(realm="lx") as c:
        up = await _uploader(c, imp)
    out = {**base,
           "ledger": {"import_id": imp["id"], "name": re.sub(r"\.(xlsx|xls)\.csv$", r".\1", imp["name"] or imp["filename"] or "") or "사료작물 재배 신청 대장",
                      "kind": "사료작물 재배 신청 대장", "kind_code": "feedcrop", "uploader": up,
                      "base_date": _ymd(imp["base_date"].isoformat()) if imp.get("base_date") else None, "date": _ymd(as_of),
                      "parcels": env(len(items), "필지", "measured", "올린 대장 × 연속지적(필지 번호)", as_of=as_of),
                      "loose": env(loose, "행", "measured", "필지와 잇지 못한 대장 행", as_of=as_of) if loose else None},
           "ok": env(cnt["ok"], "필지", "inferred", SRC_FEED), "diff": env(cnt["diff"], "필지", "inferred", SRC_FEED),
           "per_slot": per_slot,
           "by_emd": [{"emd": e["emd"], "n": env(e["n"], "필지", "measured", SRC_FEED), "ok": env(e["ok"], "필지", "inferred", SRC_FEED),
                       "diff": env(e["diff"], "필지", "inferred", SRC_FEED), "miss": {k: env(v, "필지", "inferred", SRC_FEED) for k, v in e["miss"].items()}} for e in by_emd],
           "items": [_pub(x) for x in items[:400]], "more": env(max(0, len(items) - 400), "필지", "recorded", "표에 싣지 않은 필지(엑셀에 모두)")}
    if geom:
        out["geojson"] = {"type": "FeatureCollection", "features": [
            {"type": "Feature", "id": x["pnu"], "geometry": json.loads(gj[x["pnu"]]),
             "properties": {"pnu": x["pnu"], "k": x["k"], "where": x["where"], "crop": x["crop"], "miss": " · ".join(x["miss"]), "emd": x["emd"] or "읍면동 밖"}}
            for x in items if gj.get(x["pnu"])]}
    out["_all"] = items
    return out


def _pub(x: dict) -> dict:
    """응답용 한 줄 — 넓이 · 무게는 봉투로(엑셀은 날 숫자 _all)."""
    m2 = lambda v, src: env(v, "m2", "measured", src) if v is not None else None   # noqa: E731
    sl = {}
    for k, v in x["slots"].items():
        if v is None:
            sl[k] = None
        elif k == "bale":
            sl[k] = {"found": v["found"], "n": v["n"], "kg": env(v["kg"], "kg", "estimate", "곤포사일리지 개수 × 기관이 넣은 1개당 무게") if v.get("kg") else None}
        else:
            sl[k] = {"found": v["found"], "ai": v["ai"], "m2": m2(v["m2"], SRC_FEED)}
    return {**x, "apply_m2": m2(x["apply_m2"], "대장 신청 면적"), "parcel_m2": m2(x["parcel_m2"], "연속지적 필지 면적"), "slots": sl}


def _picks_of(growth, harvest, bale) -> dict[str, str | None]:
    out = {}
    for k, v in (("growth", growth), ("harvest", harvest), ("bale", bale)):
        v = (v or "").strip()
        out[k] = "-" if v == "-" else (v if CARD_RE.match(v) else None)
    return out


@router.get("/spaces/me/ledger/feedcrop")
async def feedcrop_get(request: Request, import_id: str | None = None, growth: str | None = None, harvest: str | None = None,
                       bale: str | None = None, geom: int = 0):
    p = _me(request)
    out = await feedcrop(p, import_id, _picks_of(growth, harvest, bale), bool(geom))
    out.pop("_all", None)
    return out


@router.get("/spaces/me/ledger/feedcrop.xlsx")
async def feedcrop_xlsx(request: Request, import_id: str | None = None, growth: str | None = None, harvest: str | None = None,
                        bale: str | None = None, ok: str | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    p = _me(request)
    pl = await need_personal(p, ok)
    c = await feedcrop(p, import_id, _picks_of(growth, harvest, bale))
    if not c.get("ledger"):
        raise ApiError("not_found", "올린 사료작물 재배 신청 대장이 없습니다")
    slots = [s for s in c["slots"] if s["card"]]
    kg = (c.get("bale_kg") or {}).get("value")

    def build() -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "필지"
        head = ["필지 번호", "지번", "읍면동", "신청 작물", "신청 면적(㎡)", "필지 면적(㎡)"]
        for s in slots:
            if s["key"] == "bale":
                head += ["곤포사일리지 AI 결과", "곤포사일리지 개수"] + (["생산량(kg)"] if kg else [])
            else:
                head += [f"{s['label']} AI 결과", f"{s['label']} AI 면적(㎡)"]
        head += ["견준 결과", "AI 결과가 없는 시점"]
        ws.append(head)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F2F4F6")
        for x in c["_all"]:
            row = [x["pnu"], x["where"], x["emd"], x["crop"], x["apply_m2"], x["parcel_m2"]]
            for s in slots:
                v = x["slots"].get(s["key"]) or {}
                if s["key"] == "bale":
                    row += ["있음" if v.get("found") else "없음", v.get("n") or 0] + ([v.get("kg")] if kg else [])
                else:
                    row += [v.get("ai") or "없음", v.get("m2")]
            row += [{"ok": "맞음", "diff": "다름", "none": ""}[x["k"]], " · ".join(x["miss"])]
            ws.append(row)
        for r in ws.iter_rows(min_row=2, min_col=1, max_col=1):
            for cell in r:
                cell.number_format = "@"
        for i in range(len(head)):
            ws.column_dimensions[get_column_letter(i + 1)].width = 22 if i in (0, 1) else 16
        # 읍면동별 묶음(제안-견주기읍면동묶음) — 다른 필지 많은 곳부터 · 숫자는 숫자
        we = wb.create_sheet("읍면동별")
        lbs = [s["label"] for s in slots]
        we.append(["읍면동", "신청 필지", "맞음", "다름"] + [f"{lb} AI 결과 없음" for lb in lbs])
        for cell in we[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F2F4F6")
        for e in c["by_emd"]:
            we.append([e["emd"], e["n"]["value"], e["ok"]["value"], e["diff"]["value"]] + [(e["miss"].get(lb) or {}).get("value", 0) for lb in lbs])
        we.append(["합계", len(c["_all"]), c["ok"]["value"], c["diff"]["value"]] + [sum((e["miss"].get(lb) or {}).get("value", 0) for e in c["by_emd"]) for lb in lbs])
        for cell in we[we.max_row]:
            cell.font = Font(bold=True)
        we.freeze_panes = "A2"
        for i in range(4 + len(lbs)):
            we.column_dimensions[get_column_letter(i + 1)].width = 16 if i else 18
        info = wb.create_sheet("안내")
        L = c["ledger"]
        rows = [("올린 대장", L["name"]), ("기준일", L.get("base_date") or "—"), ("대장 필지", L["parcels"]["value"]),
                *[(f"{s['label']} 서비스", " · ".join(x for x in (s["name"], s.get("shot")) if x)) for s in slots],
                ("맞음(고른 시점 모두 AI 결과 있음)", c["ok"]["value"]), ("다름(AI 결과가 없는 시점이 있음)", c["diff"]["value"])]
        if kg:
            rows.append(("곤포사일리지 1개당 무게(kg · 기관이 넣은 값)", kg))
        rows += [("이용 약속", f"{pl['purpose']} · {pl['until']}까지"), ("만든 때", now_iso()),
                 ("안내", "AI 분석 결과는 참고자료입니다 — 처분은 담당자가 정합니다. 개인정보가 든 파일이니 약속한 목적 · 기간 안에서만 씁니다.")]
        for k, v in rows:
            info.append([k, v])
        info.column_dimensions["A"].width = 34
        info.column_dimensions["B"].width = 64
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    blob = await run_in_threadpool(build)
    name = f"사료작물_재배신청_AI결과_견주기_{dt.datetime.now(KST).strftime('%Y%m%d')}.xlsx"
    async with db(p) as cn:
        await cn.execute("INSERT INTO space_log(tenant_id, kind, card_id, line, actor, realm, detail) VALUES ($1,'download',$2,$3,$4,'tenant',$5)",
                         p.tenant_id, slots[0]["card"] if slots else None, f"{p.name or '담당자'} — 사료작물 재배 신청 대장과 AI 결과 내려받기(개인정보 · 이용 약속 확인)",
                         p.user_id, {"fmt": "feedcrop", "rows": len(c["_all"]), "import_id": c["ledger"]["import_id"], "pledge": pl["id"], "personal": True})
    return Response(blob, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"content-disposition": f"attachment; filename=\"feedcrop.xlsx\"; filename*=UTF-8''{quote(name)}", "cache-control": "no-store"})
